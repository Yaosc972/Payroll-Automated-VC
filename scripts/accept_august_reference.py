"""Independent original-cell reader and comparison for the August acceptance.

This module deliberately imports no platform parser, formula or comparator.
"""
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import re
import warnings

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

STAGES = ('入职满一个月', '入职满三个月', '入职满六个月', '转正')
GROUPS = (
    ('recruitment', '招聘负责人', (39, 41, 43, 52), 22, 23),
    ('recruitment', '协助招聘人', (46, 48, 50, 54), 24, 25),
    ('referral', '推荐人', (62, 64, 66, 68), 59, 58),
)
CURRENCIES = {'人民币': 'CNY', 'RMB': 'CNY', 'CNY': 'CNY', '美元': 'USD',
              'USD': 'USD', '欧元': 'EUR', 'EUR': 'EUR', '捷克郎': 'CZK',
              '捷克克朗': 'CZK', 'CZK': 'CZK'}


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def identifier(value):
    value = re.sub(r'\s+', '', text(value)).casefold()
    regional_staff = re.fullmatch(r'(?:海外|国内)[-—－_:：]*(zt\d+)', value)
    if regional_staff:
        value = regional_staff.group(1)
    return '' if value in {'', '-', '--', '无', '无工号', 'n/a', 'none'} else value


def date_string(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, (int, float)) and 1 <= value < 100000:
        return from_excel(value).strftime('%Y-%m-%d')
    value = text(value)
    match = re.fullmatch(r'(20\d\d)[/.-](\d{1,2})[/.-](\d{1,2})(?:[ T].*)?', value)
    if not match:
        return ''
    try:
        return date(*map(int, match.groups())).isoformat()
    except ValueError:
        return ''


def money(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(text(value).replace(',', ''))
    except InvalidOperation:
        return None
    return format(result.quantize(Decimal('.01'), rounding=ROUND_HALF_UP), '.2f') if result.is_finite() else None


def currency(value):
    raw = text(value)
    return CURRENCIES.get(raw.upper(), CURRENCIES.get(raw))


def period(value):
    raw = text(value)
    return int(raw) if re.fullmatch(r'20\d{2}(?:0[1-9]|1[0-2])', raw) else None


def identities(values):
    """Ordered strong identity evidence; names never become matching keys."""
    hire = date_string(values.get('入职日期') or values.get('候选人入职时间'))
    if not hire:
        return []
    return [(label, value, hire) for label, value in (
        ('employee', identifier(values.get('工号'))),
        ('document', identifier(values.get('证件号') or values.get('身份证'))),
        ('offer', identifier(values.get('海外offer流程编号'))),
    ) if value]


def read_reference(path, month=202608):
    """Read held-out workbook cached cells, never platform-normalized records."""
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always')
        wb = load_workbook(path, read_only=True, data_only=True)
        formula_wb = load_workbook(path, read_only=True, data_only=False)
        ws = wb['附件一-招聘奖金核算明细']
        formula_ws = formula_wb[ws.title]
        nodes, records, observations = [], [], Counter()
        for index, (cells, formulas) in enumerate(zip(
                ws.iter_rows(min_row=4, values_only=True),
                formula_ws.iter_rows(min_row=4, values_only=True)), start=4):
            if not any(cells):
                continue
            get = lambda column: cells[column - 1] if column <= len(cells) else None
            values = {label: get(column) for label, column in (
                ('姓名', 2), ('工号', 3), ('证件号', 4), ('入职日期', 5),
                ('工作地', 7), ('职级', 9), ('二级组织', 10), ('三级组织', 11),
                ('四级组织', 12), ('五级组织', 13), ('职位', 14), ('ABC类别', 21),
                ('招聘负责人工号', 22), ('协助招聘人工号', 24), ('招聘渠道', 26),
                ('海外offer流程编号', 27), ('招聘启动日期', 29), ('转正日期', 31),
                ('推荐人工号', 59))}
            values = {k: date_string(v) if k.endswith('日期') else text(v) for k, v in values.items()}
            if not any(values.values()):
                continue
            record_id = f'{ws.title}!{index}'
            records.append({'id': record_id, 'sourceRow': index, 'values': values})
            for kind, role, columns, id_column, name_column in GROUPS:
                raw_currency = get(56)
                note = text(get(72))
                if kind == 'referral' and '内推奖币种' in note and any(c in note for c in ('CZK', '捷克郎', '捷克克朗')):
                    raw_currency = 'CZK'
                for stage, column in zip(STAGES, columns):
                    raw_amount, raw_period = get(column), get(column + 1)
                    amount, scheduled = money(raw_amount), period(raw_period)
                    formula = formulas[column - 1] if column <= len(formulas) else None
                    missing_cache = raw_amount is None and isinstance(formula, str) and formula.startswith('=')
                    if missing_cache:
                        observations['formulaWithoutCachedAmount'] += 1
                    if amount == '0.00':
                        observations['zeroCells'] += 1
                        continue
                    if raw_amount is None and not missing_cache:
                        observations['emptyCells'] += 1
                        continue
                    if amount is None:
                        observations['unknownAmountCells'] += 1
                    elif scheduled != month:
                        observations['nonCurrentNonzeroCells'] += 1
                    nodes.append({'recordId': record_id, 'sourceRow': index, 'column': column,
                        'cell': f'{ws.title}!{get_column_letter(column)}{index}', 'kind': kind, 'role': role, 'stage': stage,
                        'amount': amount, 'month': scheduled, 'rawMonth': text(raw_period),
                        'currency': currency(raw_currency), 'recipientId': identifier(get(id_column)),
                        'recipientName': text(get(name_column)), 'formulaWithoutCache': missing_cache})
            if money(get(70)) not in (None, '0.00'):
                observations['anotherHalfNonzeroCells'] += 1
                nodes.append({'recordId': record_id, 'sourceRow': index, 'column': 70,
                    'cell': f'{ws.title}!BR{index}', 'kind': 'referral', 'role': '推荐人',
                    'stage': '另一半奖金', 'amount': money(get(70)), 'month': period(get(71)),
                    'rawMonth': text(get(71)), 'currency': currency(get(56)),
                    'recipientId': identifier(get(59)), 'recipientName': text(get(58))})
        sheets = [{'name': s.title, 'rows': s.max_row, 'columns': s.max_column} for s in wb]
        wb.close()
        formula_wb.close()
    current = [n for n in nodes if n['month'] == month and n['amount'] not in (None, '0.00')]
    return {'month': month, 'sheets': sheets, 'records': records, 'nodes': nodes,
            'currentNodeCount': len(current), 'currentRoles': dict(Counter(n['role'] for n in current)),
            'currentTotals': totals(current), 'observations': dict(observations),
            'warnings': sorted({str(w.message) for w in caught})}


def totals(nodes):
    result = defaultdict(Decimal)
    for node in nodes:
        if node.get('amount') is not None:
            result[node.get('currency') or 'UNKNOWN'] += Decimal(node['amount'])
    return {key: format(value, '.2f') for key, value in sorted(result.items())}


def match_records(reference_records, actual_records):
    indexes = defaultdict(list)
    for actual in actual_records:
        for key in identities(actual['values']):
            indexes[key].append(actual['id'])
    candidates = {}
    for ref in reference_records:
        possible = []
        for key in identities(ref['values']):
            matches = list(dict.fromkeys(indexes[key]))
            if matches:
                possible = matches
                break
        candidates[ref['id']] = possible
    uses = Counter(row_id for options in candidates.values() for row_id in options)
    return {ref_id: {'status': 'matched' if len(options) == 1 and uses[options[0]] == 1
                    else 'ambiguous' if options else 'unmatched', 'candidates': options}
            for ref_id, options in candidates.items()}


def compare(reference, actual):
    month = reference['month']
    matches = match_records(reference['records'], actual['records'])
    by_node = defaultdict(list)
    for node in actual['nodes']:
        by_node[(node['recordId'], node['kind'], node['role'], node['stage'])].append(node)
    results, used = [], set()
    current = [n for n in reference['nodes'] if n['month'] == month and n['amount'] not in (None, '0.00')]
    reference_keys = Counter((n['recordId'], n['kind'], n['role'], n['stage']) for n in current)
    for ref in current:
        match = matches[ref['recordId']]
        key = (match['candidates'][0], ref['kind'], ref['role'], ref['stage']) if match['status'] == 'matched' else None
        options = by_node.get(key, []) if key else []
        entry = {'reference': ref, 'matchStatus': match['status'], 'actual': None, 'dimensions': {}}
        if match['status'] != 'matched':
            entry['outcome'] = 'identity_' + match['status']
        elif reference_keys[(ref['recordId'], ref['kind'], ref['role'], ref['stage'])] != 1 or len(options) > 1:
            entry.update(outcome='node_ambiguous', actualCandidates=options)
        elif not options:
            entry['outcome'] = 'reference_only'
        else:
            node = options[0]
            used.add(node['id'])
            entry['actual'] = node
            for field in ('amount', 'month', 'currency', 'recipientId'):
                expected, found = ref.get(field), node.get(field)
                unknown = expected is None or (found is None and not (field == 'month' and node.get('monthKnown'))) or (
                    field == 'recipientId' and (not expected or not found))
                entry['dimensions'][field] = 'unknown' if unknown else 'equal' if expected == found else 'different'
            entry['eligibility'] = node.get('eligibility', 'unknown')
            states = entry['dimensions'].values()
            entry['outcome'] = 'different' if 'different' in states else 'unknown' if 'unknown' in states else 'equal'
        results.append(entry)
    extras = [n for n in actual['nodes'] if n['id'] not in used and n['month'] == month and n['amount'] not in (None, '0.00')]
    unknown_current = [n for n in actual['nodes'] if n['month'] == month and n['amount'] is None]
    summary = {'referenceNodes': len(current),
               'referenceUnknownCurrentAmounts': sum(n['month'] == month and n['amount'] is None for n in reference['nodes']),
               'outcomes': dict(Counter(r['outcome'] for r in results)),
               'dimensions': {field: dict(Counter(r['dimensions'].get(field, 'not_comparable') for r in results))
                              for field in ('amount', 'month', 'currency', 'recipientId')},
               'eligibility': dict(Counter(r.get('eligibility', 'not_comparable') for r in results)),
               'platformOnlyCurrentNodes': len(extras), 'unknownCurrentAmountNodes': len(unknown_current),
               'roster': dict(Counter(m['status'] for m in matches.values())),
               'referenceTotals': totals(current),
               'actualCurrentTotals': totals([n for n in actual['nodes'] if n['month'] == month])}
    covered = {m['candidates'][0] for m in matches.values() if m['status'] == 'matched'}
    return {'summary': summary, 'nodes': results, 'platformOnly': extras,
            'unknownCurrentAmounts': unknown_current, 'rosterMatches': matches,
            'platformOnlyRecords': [r for r in actual['records'] if r['id'] not in covered]}
