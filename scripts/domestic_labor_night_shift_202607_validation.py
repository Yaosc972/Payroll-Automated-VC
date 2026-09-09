"""用2026年7月华南、华东线下日考勤验证夜班补贴班次基线。"""
from __future__ import annotations

import argparse
import json
import math
import subprocess
import sys
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, Mapping, Sequence

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from bonus_platform.engine.domestic_labor.engines.yeban_butie import YeBanBuTieEngine
from bonus_platform.engine.domestic_labor.night_shift_config import (
    _normalize_shift_row,
    load_night_shift_config,
)


DEFAULT_SOUTH = Path(
    "/Users/zt27532/Documents/AI算薪/华南&华东7月份AI+线下补贴/"
    "线下-华南202607考勤补贴 (操作).xlsx"
)
DEFAULT_EAST = Path(
    "/Users/zt27532/Documents/AI算薪/华南&华东7月份AI+线下补贴/"
    "线下-华东7月华东枢纽外包费用.xlsx"
)
DEFAULT_OUTPUT = (
    PROJECT_ROOT
    / "outputs/01a061c0-14fb-7962-8f9a-6bbbefead92f/"
    "night_shift_202607_after_shift_master.json"
)
TOLERANCE = 0.011


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _number(value: Any) -> float:
    try:
        number = float(value)
        return 0.0 if math.isnan(number) else number
    except (TypeError, ValueError):
        return 0.0


def _day(value: Any) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return _text(value)[:10]


def _previous_config() -> Dict[str, Any]:
    """重建补充班次资料之前、已固化HD048之后的平台基线。"""
    raw = subprocess.check_output(
        [
            "git",
            "show",
            "HEAD:bonus_platform/engine/domestic_labor/data/night_shift_breaks.json",
        ],
        cwd=PROJECT_ROOT,
        text=True,
    )
    payload = json.loads(raw)
    payload["shifts"].append(
        {
            "shift_category": "华东班次",
            "shift_name": "理货入库18点班",
            "shift_code": "HD048",
            "shift_time": "18:00-27:00;",
            "regular_hours": 9,
            "break_periods": ["23:00-24:00", "06:00-06:30"],
            "note": "2026年7月线下核算数据验证",
        }
    )
    shifts = [_normalize_shift_row(row) for row in payload["shifts"]]
    return {
        "month": "202607",
        "revision": 0,
        "exists": False,
        "effective_shift_breaks": shifts,
        "shift_breaks": shifts,
        "jinjiang_exclusions": [],
        "jinjiang_list_confirmed": False,
    }


def _region_layout(label: str) -> Dict[str, Any]:
    if label == "east":
        return {
            "monthly_sheet": "月考勤",
            "monthly_start": 2,
            "monthly_columns": (2, 1, 3, 14),
            "daily_sheet": "日考勤",
            "daily_start": 2,
            "daily_columns": {
                "date": 1,
                "work_status": 3,
                "employee": 4,
                "name": 5,
                "area": None,
                "position": 13,
                "shift_name": 18,
                "shift": 19,
                "abnormal": 20,
                "abnormal_reason": 21,
                "schedule": 22,
                "start": 24,
                "end": 25,
                "evening_break": 75,
                "morning_break": 76,
                "amount": 78,
                "work_type": None,
            },
        }
    return {
        "monthly_sheet": "202607月考勤",
        "monthly_start": 3,
        "monthly_columns": (5, 6, 3, 9),
        "daily_sheet": "202607日考勤",
        "daily_start": 3,
        "daily_columns": {
            "date": 1,
            "work_status": 3,
            "employee": 4,
            "name": 5,
            "area": 7,
            "position": 14,
            "shift_name": 18,
            "shift": 19,
            "abnormal": 20,
            "abnormal_reason": 21,
            "schedule": 22,
            "start": 24,
            "end": 25,
            "evening_break": 81,
            "morning_break": 82,
            "amount": 84,
            "work_type": 97,
        },
    }


def _read_region(label: str, path: Path):
    layout = _region_layout(label)
    workbook = load_workbook(path, read_only=True, data_only=True)
    monthly = workbook[layout["monthly_sheet"]]
    employee_col, name_col, area_col, position_col = layout["monthly_columns"]
    people: Dict[str, Dict[str, Any]] = {}
    for row in monthly.iter_rows(min_row=layout["monthly_start"], values_only=True):
        employee_id = _text(row[employee_col])
        if employee_id and employee_id != "工号":
            people[employee_id] = {
                "工号": employee_id,
                "姓名": _text(row[name_col]),
                "工作地区": _text(row[area_col]),
                "岗位名称": _text(row[position_col]),
            }

    columns = layout["daily_columns"]
    daily = workbook[layout["daily_sheet"]]
    grouped = defaultdict(list)
    offline: Dict[tuple[str, str], Dict[str, Any]] = {}
    duplicate_keys = 0
    for row_number, row in enumerate(
        daily.iter_rows(min_row=layout["daily_start"], values_only=True),
        start=layout["daily_start"],
    ):
        employee_id = _text(row[columns["employee"]])
        attendance_date = _day(row[columns["date"]])
        if not employee_id or employee_id == "工号" or not attendance_date:
            continue
        area = (
            _text(row[columns["area"]])
            if columns["area"] is not None
            else _text(people.get(employee_id, {}).get("工作地区"))
        )
        position = _text(row[columns["position"]]) or _text(
            people.get(employee_id, {}).get("岗位名称")
        )
        attendance = {
            "工号": employee_id,
            "姓名": _text(row[columns["name"]]),
            "出勤日期": row[columns["date"]],
            "班次编号": _text(row[columns["shift"]]),
            "班次时间段": _text(row[columns["schedule"]]),
            "工作状态": _text(row[columns["work_status"]]),
            "工作地区": area,
            "岗位名称": position,
            "上班一": row[columns["start"]],
            "下班一": row[columns["end"]],
        }
        if columns["work_type"] is not None:
            attendance["计时"] = _text(row[columns["work_type"]])
        key = (employee_id, attendance_date)
        if key in offline:
            duplicate_keys += 1
        grouped[employee_id].append(attendance)
        offline[key] = {
            "employee_id": employee_id,
            "employee_name": attendance["姓名"],
            "attendance_date": attendance_date,
            "work_status": attendance["工作状态"],
            "work_area": area,
            "position": position,
            "amount": _number(row[columns["amount"]]),
            "shift_code": attendance["班次编号"],
            "source_shift_name": _text(row[columns["shift_name"]]),
            "shift_schedule": attendance["班次时间段"],
            "abnormal": _text(row[columns["abnormal"]]),
            "abnormal_reason": _text(row[columns["abnormal_reason"]]),
            "raw_start": attendance["上班一"],
            "raw_end": attendance["下班一"],
            "line_evening_break_hours": _number(row[columns["evening_break"]]),
            "line_morning_break_hours": _number(row[columns["morning_break"]]),
            "source_row": row_number,
        }
        people.setdefault(
            employee_id,
            {
                "工号": employee_id,
                "姓名": attendance["姓名"],
                "工作地区": area,
                "岗位名称": position,
            },
        )
    workbook.close()
    return people, grouped, offline, duplicate_keys


def _calculate(
    people: Mapping[str, Mapping[str, Any]],
    grouped: Mapping[str, Sequence[Mapping[str, Any]]],
    config: Mapping[str, Any],
):
    platform = {}
    statuses = Counter()
    reasons = Counter()
    engine = YeBanBuTieEngine()
    for employee_id, attendance in grouped.items():
        result = engine.calculate(dict(people[employee_id]), list(attendance), config=config)
        for row in result.details["daily_results"]:
            platform[(employee_id, _day(row.get("attendance_date")))] = row
            statuses[row.get("status")] += 1
            reasons[row.get("reason_code")] += 1
    return platform, statuses, reasons


def _money(value: float) -> float:
    return round(value + 1e-10, 2)


def _summarize(
    offline: Mapping[tuple[str, str], Mapping[str, Any]],
    platform: Mapping[tuple[str, str], Mapping[str, Any]],
    statuses: Counter,
    reasons: Counter,
) -> Dict[str, Any]:
    matched = set(offline) & set(platform)
    totals = Counter()
    by_shift = defaultdict(Counter)
    for key in matched:
        source = offline[key]
        result = platform[key]
        line_amount = _number(source.get("amount"))
        platform_amount = _number(result.get("amount"))
        difference = platform_amount - line_amount
        exact = abs(difference) <= TOLERANCE
        shift_code = _text(source.get("shift_code") or result.get("shift_code"))
        totals["matched_rows"] += 1
        totals["exact_rows"] += exact
        totals["line_paid_rows"] += line_amount > TOLERANCE
        totals["line_paid_exact_rows"] += line_amount > TOLERANCE and exact
        totals["platform_total"] += platform_amount
        totals["line_total"] += line_amount
        totals["absolute_difference"] += abs(difference)
        shift = by_shift[shift_code]
        shift["rows"] += 1
        shift["mismatch_rows"] += not exact
        shift["platform_total"] += platform_amount
        shift["line_total"] += line_amount
        shift["absolute_difference"] += abs(difference)

    mismatches = []
    for shift_code, values in by_shift.items():
        if not values["mismatch_rows"]:
            continue
        mismatches.append(
            {
                "shift_code": shift_code,
                "rows": values["rows"],
                "mismatch_rows": values["mismatch_rows"],
                "platform_yuan": _money(values["platform_total"]),
                "line_yuan": _money(values["line_total"]),
                "net_yuan": _money(values["platform_total"] - values["line_total"]),
                "absolute_yuan": _money(values["absolute_difference"]),
            }
        )
    mismatches.sort(
        key=lambda row: (-row["absolute_yuan"], -row["mismatch_rows"], row["shift_code"])
    )
    matched_rows = totals["matched_rows"]
    paid_rows = totals["line_paid_rows"]
    return {
        "platform_rows": len(platform),
        "matched_rows": matched_rows,
        "exact_rows": totals["exact_rows"],
        "exact_rate": round(totals["exact_rows"] / matched_rows, 6),
        "line_paid_rows": paid_rows,
        "line_paid_exact_rows": totals["line_paid_exact_rows"],
        "line_paid_exact_rate": round(totals["line_paid_exact_rows"] / paid_rows, 6),
        "platform_total_yuan": _money(totals["platform_total"]),
        "line_total_yuan": _money(totals["line_total"]),
        "net_difference_yuan": _money(totals["platform_total"] - totals["line_total"]),
        "absolute_difference_yuan": _money(totals["absolute_difference"]),
        "platform_only_rows": len(set(platform) - set(offline)),
        "line_only_rows": len(set(offline) - set(platform)),
        "status_counts": dict(statuses),
        "reason_counts": dict(reasons),
        "top_mismatch_shifts": mismatches[:30],
    }


def _overall(regions: Iterable[Mapping[str, Any]], mode: str) -> Dict[str, Any]:
    metrics = [region["modes"][mode] for region in regions]
    matched = sum(row["matched_rows"] for row in metrics)
    exact = sum(row["exact_rows"] for row in metrics)
    paid = sum(row["line_paid_rows"] for row in metrics)
    paid_exact = sum(row["line_paid_exact_rows"] for row in metrics)
    return {
        "line_rows": sum(region["source_rows"] for region in regions),
        "platform_rows": sum(row["platform_rows"] for row in metrics),
        "matched_rows": matched,
        "exact_rows": exact,
        "exact_rate": round(exact / matched, 6),
        "line_paid_rows": paid,
        "line_paid_exact_rows": paid_exact,
        "line_paid_exact_rate": round(paid_exact / paid, 6),
        "platform_total_yuan": _money(sum(row["platform_total_yuan"] for row in metrics)),
        "line_total_yuan": _money(sum(row["line_total_yuan"] for row in metrics)),
        "net_difference_yuan": _money(sum(row["net_difference_yuan"] for row in metrics)),
        "absolute_difference_yuan": _money(
            sum(row["absolute_difference_yuan"] for row in metrics)
        ),
        "platform_only_rows": sum(row["platform_only_rows"] for row in metrics),
        "line_only_rows": sum(row["line_only_rows"] for row in metrics),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--south", type=Path, default=DEFAULT_SOUTH)
    parser.add_argument("--east", type=Path, default=DEFAULT_EAST)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    configs = {
        "before": _previous_config(),
        "after": load_night_shift_config("202607", required=False),
    }
    report: Dict[str, Any] = {
        "month": "202607",
        "source_shift_file": "基础班次资料(2).xls",
        "baseline_counts": {
            mode: len(config["shift_breaks"]) for mode, config in configs.items()
        },
        "regions": {},
    }
    for label, path in (("south", args.south), ("east", args.east)):
        people, grouped, offline, duplicate_keys = _read_region(label, path)
        region: Dict[str, Any] = {
            "source_rows": len(offline),
            "source_duplicate_keys": duplicate_keys,
            "employees": len(grouped),
            "modes": {},
        }
        for mode, config in configs.items():
            platform, statuses, reasons = _calculate(people, grouped, config)
            region["modes"][mode] = _summarize(
                offline, platform, statuses, reasons
            )
        report["regions"][label] = region

    report["overall"] = {
        mode: _overall(report["regions"].values(), mode) for mode in configs
    }
    before = report["overall"]["before"]
    after = report["overall"]["after"]
    report["improvement"] = {
        "new_shift_codes": report["baseline_counts"]["after"]
        - report["baseline_counts"]["before"],
        "exact_rows_added": after["exact_rows"] - before["exact_rows"],
        "absolute_difference_reduced_yuan": _money(
            before["absolute_difference_yuan"] - after["absolute_difference_yuan"]
        ),
        "net_difference_change_yuan": _money(
            after["net_difference_yuan"] - before["net_difference_yuan"]
        ),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
