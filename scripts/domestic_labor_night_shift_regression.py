"""用三地区真实考勤回归国内劳务夜班补贴通用规则。"""
from __future__ import annotations

import argparse
import json
import math
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Sequence, Tuple

from openpyxl import load_workbook

from bonus_platform.engine.domestic_labor.engines.yeban_butie import YeBanBuTieEngine
from bonus_platform.engine.domestic_labor.night_shift_config import (
    baseline_shift_breaks_metadata,
    load_baseline_shift_breaks,
)


DEFAULT_DONGGUAN = Path(
    "/Users/zt27532/Documents/AI算薪/中国操作部5月考勤补贴-外包人员/202605日考勤数据 (操作).xlsx"
)
DEFAULT_HUADONG = Path(
    "/Users/zt27532/Documents/AI算薪/中国操作部5月考勤补贴-外包人员/5月华东枢纽外包费用.xlsx"
)
DEFAULT_JINJIANG = Path(
    "/Users/zt27532/Documents/AI算薪/中国操作部5月考勤补贴-外包人员/东南枢纽操作2026年5月夜班津贴.XLSX"
)
TOLERANCE = 0.011


def _number(value: Any) -> float:
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)) and not math.isnan(value):
        return float(value)
    try:
        return float(str(value).strip())
    except (TypeError, ValueError):
        return 0.0


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _date_text(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return _text(value)[:10]


def _load_break_map() -> Dict[str, List[str]]:
    return {
        row["shift_code"]: list(row.get("break_periods") or [])
        for row in load_baseline_shift_breaks()
    }


def _is_exact(actual: float, predicted: float) -> bool:
    return abs(actual - predicted) <= TOLERANCE


def _new_scope() -> Dict[str, Any]:
    return {
        "paid_records": 0,
        "calculated_records": 0,
        "exact_records": 0,
        "manual_or_pending_records": 0,
        "actual_total_on_calculated": 0.0,
        "predicted_total_on_calculated": 0.0,
        "absolute_error_total": 0.0,
        "standard_break_records": 0,
        "standard_break_exact_records": 0,
        "status_reasons": Counter(),
        "mismatch_reasons": Counter(),
        "mismatch_shifts": Counter(),
        "examples": [],
    }


def _record_result(
    scope: Dict[str, Any],
    *,
    row_number: int,
    attendance_date: str,
    shift_code: str,
    actual: float,
    predicted_result,
    source_break_hours: float | None,
) -> None:
    scope["paid_records"] += 1
    if predicted_result.status != "calculated" or predicted_result.amount is None:
        scope["manual_or_pending_records"] += 1
        scope["status_reasons"][predicted_result.reason_code] += 1
        return

    predicted = float(predicted_result.amount)
    rule_break_hours = predicted_result.break_minutes / 60
    scope["calculated_records"] += 1
    scope["actual_total_on_calculated"] += actual
    scope["predicted_total_on_calculated"] += predicted
    scope["absolute_error_total"] += abs(actual - predicted)
    exact = _is_exact(actual, predicted)
    if exact:
        scope["exact_records"] += 1

    standard_break = (
        source_break_hours is not None
        and abs(source_break_hours - rule_break_hours) <= TOLERANCE
    )
    if standard_break:
        scope["standard_break_records"] += 1
        if exact:
            scope["standard_break_exact_records"] += 1

    if exact:
        return
    mismatch_reason = "historical_break_differs_from_rule" if not standard_break else "formula_difference"
    scope["mismatch_reasons"][mismatch_reason] += 1
    scope["mismatch_shifts"][shift_code] += 1
    if len(scope["examples"]) < 15:
        scope["examples"].append(
            {
                "source_row": row_number,
                "date": attendance_date,
                "shift": shift_code,
                "actual": round(actual, 4),
                "predicted": round(predicted, 4),
                "source_break_hours": None if source_break_hours is None else round(source_break_hours, 4),
                "rule_break_hours": round(rule_break_hours, 4),
            }
        )


def _finalize_scope(scope: Dict[str, Any]) -> Dict[str, Any]:
    paid = scope["paid_records"]
    calculated = scope["calculated_records"]
    exact = scope["exact_records"]
    standard = scope["standard_break_records"]
    result = dict(scope)
    result.update(
        {
            "coverage_rate": round(calculated / paid, 6) if paid else None,
            "exact_rate_on_calculated": round(exact / calculated, 6) if calculated else None,
            "direct_match_rate_on_paid": round(exact / paid, 6) if paid else None,
            "standard_break_exact_rate": (
                round(scope["standard_break_exact_records"] / standard, 6) if standard else None
            ),
            "actual_total_on_calculated": round(scope["actual_total_on_calculated"], 4),
            "predicted_total_on_calculated": round(scope["predicted_total_on_calculated"], 4),
            "net_difference": round(
                scope["predicted_total_on_calculated"] - scope["actual_total_on_calculated"], 4
            ),
            "absolute_error_total": round(scope["absolute_error_total"], 4),
            "status_reasons": scope["status_reasons"].most_common(),
            "mismatch_reasons": scope["mismatch_reasons"].most_common(),
            "mismatch_shifts": scope["mismatch_shifts"].most_common(20),
        }
    )
    return result


def _attendance_row(row: Sequence[Any], *, employee_col: int, name_col: int) -> Dict[str, Any]:
    return {
        "工号": _text(row[employee_col]),
        "姓名": _text(row[name_col]),
        "出勤日期": _date_text(row[1]),
        "班次编号": _text(row[19]),
        "上班一": row[24],
        "下班一": row[25],
    }


def audit_dongguan(path: Path, break_map: Mapping[str, Sequence[str]]) -> Dict[str, Any]:
    engine = YeBanBuTieEngine()
    scope = _new_scope()
    special_paid_records = 0
    candidate_not_paid = 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["202605日考勤"]
    for row_number, row in enumerate(worksheet.iter_rows(min_row=3, values_only=True), start=3):
        if not _text(row[4]):
            continue
        attendance = _attendance_row(row, employee_col=4, name_col=5)
        shift_code = attendance["班次编号"]
        result = engine.calculate_day(attendance, break_map.get(shift_code, ()))
        ordinary_amount = _number(row[84])
        special_amount = _number(row[89])
        if abs(special_amount) > TOLERANCE:
            special_paid_records += 1
        if abs(ordinary_amount) > TOLERANCE:
            _record_result(
                scope,
                row_number=row_number,
                attendance_date=attendance["出勤日期"],
                shift_code=shift_code,
                actual=ordinary_amount,
                predicted_result=result,
                source_break_hours=_number(row[81]) + _number(row[82]),
            )
        elif result.status == "calculated" and (result.amount or 0) > TOLERANCE:
            candidate_not_paid += 1
    workbook.close()
    result = _finalize_scope(scope)
    result.update(
        {
            "scope_note": "仅验证线下普通夜班补贴非0的日明细；特殊夜班列及资格未发记录不计入准确率。",
            "special_paid_records_excluded": special_paid_records,
            "rule_candidates_without_offline_pay_excluded": candidate_not_paid,
        }
    )
    return result


def audit_huadong(path: Path, break_map: Mapping[str, Sequence[str]]) -> Dict[str, Any]:
    engine = YeBanBuTieEngine()
    scope = _new_scope()
    candidate_not_paid = 0
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["日考勤"]
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not _text(row[4]):
            continue
        attendance = _attendance_row(row, employee_col=4, name_col=5)
        shift_code = attendance["班次编号"]
        result = engine.calculate_day(attendance, break_map.get(shift_code, ()))
        actual_amount = _number(row[78])
        if abs(actual_amount) > TOLERANCE:
            _record_result(
                scope,
                row_number=row_number,
                attendance_date=attendance["出勤日期"],
                shift_code=shift_code,
                actual=actual_amount,
                predicted_result=result,
                source_break_hours=_number(row[75]) + _number(row[76]),
            )
        elif result.status == "calculated" and (result.amount or 0) > TOLERANCE:
            candidate_not_paid += 1
    workbook.close()
    result = _finalize_scope(scope)
    result.update(
        {
            "scope_note": "验证线下夜班补贴非0的日明细；历史人工休息时长与班次规则不一致时保留为差异。",
            "rule_candidates_without_offline_pay_excluded": candidate_not_paid,
        }
    )
    return result


def _iter_until_blanks(worksheet, start: int = 2, blank_limit: int = 200):
    blank_count = 0
    for row_number, row in enumerate(worksheet.iter_rows(min_row=start, values_only=True), start=start):
        if not any(value not in (None, "") for value in row[:16]):
            blank_count += 1
            if blank_count >= blank_limit:
                break
            continue
        blank_count = 0
        yield row_number, row


def audit_jinjiang(path: Path, break_map: Mapping[str, Sequence[str]]) -> Dict[str, Any]:
    engine = YeBanBuTieEngine()
    scope = _new_scope()
    workbook = load_workbook(path, read_only=True, data_only=True)

    raw_by_key: Dict[Tuple[str, str], Tuple[int, Sequence[Any]]] = {}
    raw_sheet = workbook["5月份考勤表"]
    for row_number, row in _iter_until_blanks(raw_sheet, start=2):
        employee_id = _text(row[3])
        attendance_date = _date_text(row[0])
        if employee_id and attendance_date:
            raw_by_key[(employee_id, attendance_date)] = (row_number, row)

    missing_raw_records = 0
    detail_sheet = workbook["夜班补贴明细"]
    for detail_row_number, detail_row in _iter_until_blanks(detail_sheet, start=2):
        employee_id = _text(detail_row[0])
        attendance_date = _date_text(detail_row[2])
        if not employee_id or not attendance_date:
            continue
        raw_item = raw_by_key.get((employee_id, attendance_date))
        if raw_item is None:
            missing_raw_records += 1
            continue
        raw_row_number, raw_row = raw_item
        attendance = {
            "工号": employee_id,
            "姓名": _text(raw_row[4]),
            "出勤日期": attendance_date,
            "班次编号": _text(raw_row[18]),
            "上班一": raw_row[23],
            "下班一": raw_row[24],
        }
        shift_code = attendance["班次编号"]
        result = engine.calculate_day(attendance, break_map.get(shift_code, ()))
        _record_result(
            scope,
            row_number=detail_row_number,
            attendance_date=attendance_date,
            shift_code=shift_code,
            actual=_number(detail_row[13]),
            predicted_result=result,
            source_break_hours=_number(detail_row[9]),
        )
    workbook.close()
    result = _finalize_scope(scope)
    result.update(
        {
            "scope_note": "仅验证已进入晋江夜班补贴明细的记录；计件、门禁等未进入明细的特殊资格记录不计入准确率。",
            "detail_records_missing_raw_attendance": missing_raw_records,
            "known_special_omitted_records_excluded": 577,
        }
    )
    return result


def _overall(regions: Mapping[str, Mapping[str, Any]]) -> Dict[str, Any]:
    keys = ["dongguan", "huadong", "jinjiang"]
    paid = sum(regions[key]["paid_records"] for key in keys)
    calculated = sum(regions[key]["calculated_records"] for key in keys)
    exact = sum(regions[key]["exact_records"] for key in keys)
    standard = sum(regions[key]["standard_break_records"] for key in keys)
    standard_exact = sum(regions[key]["standard_break_exact_records"] for key in keys)
    actual_total = sum(regions[key]["actual_total_on_calculated"] for key in keys)
    predicted_total = sum(regions[key]["predicted_total_on_calculated"] for key in keys)
    return {
        "paid_records": paid,
        "calculated_records": calculated,
        "exact_records": exact,
        "manual_or_pending_records": paid - calculated,
        "coverage_rate": round(calculated / paid, 6) if paid else None,
        "exact_rate_on_calculated": round(exact / calculated, 6) if calculated else None,
        "direct_match_rate_on_paid": round(exact / paid, 6) if paid else None,
        "standard_break_records": standard,
        "standard_break_exact_records": standard_exact,
        "standard_break_exact_rate": round(standard_exact / standard, 6) if standard else None,
        "actual_total_on_calculated": round(actual_total, 4),
        "predicted_total_on_calculated": round(predicted_total, 4),
        "net_difference": round(predicted_total - actual_total, 4),
    }


def _percent(value: Any) -> str:
    return "—" if value is None else f"{float(value):.2%}"


def _markdown_report(result: Mapping[str, Any]) -> str:
    lines = [
        "# 国内劳务夜班补贴真实数据回归",
        "",
        "## 结论指标",
        "",
        "| 地区 | 线下已发记录 | 自动计算 | 精确一致 | 自动覆盖率 | 自动计算内准确率 | 全量直接命中率 | 标准休息子集准确率 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|",
    ]
    labels = {"dongguan": "东莞", "huadong": "华东", "jinjiang": "晋江", "overall": "合计"}
    for key in ("dongguan", "huadong", "jinjiang", "overall"):
        item = result["regions"].get(key) if key != "overall" else result["overall"]
        lines.append(
            "| {label} | {paid} | {calculated} | {exact} | {coverage} | {accuracy} | {direct} | {standard} |".format(
                label=labels[key],
                paid=item["paid_records"],
                calculated=item["calculated_records"],
                exact=item["exact_records"],
                coverage=_percent(item["coverage_rate"]),
                accuracy=_percent(item["exact_rate_on_calculated"]),
                direct=_percent(item["direct_match_rate_on_paid"]),
                standard=_percent(item["standard_break_exact_rate"]),
            )
        )
    lines.extend(
        [
            "",
            "## 口径",
            "",
            "- 自动覆盖率 = 引擎可按已确认通用规则计算的记录 / 线下已发记录。",
            "- 自动计算内准确率 = 精确一致记录 / 引擎已计算记录。",
            "- 全量直接命中率 = 精确一致记录 / 线下已发记录，未确认或异常记录计为未覆盖。",
            "- 标准休息子集准确率仅统计线下休息扣减与班次规则表一致的记录，用来剥离历史人工调整影响。",
            "- 精确一致按金额绝对差异不超过0.011元判断，用于容纳Excel浮点误差。",
            "- 东莞特殊夜班列、线下未发但规则算出金额的资格记录，以及晋江已确认的577条特殊资格排除记录均不计入准确率。",
            "",
            "## 分地区差异",
            "",
        ]
    )
    for key in ("dongguan", "huadong", "jinjiang"):
        item = result["regions"][key]
        lines.extend([f"### {labels[key]}", "", f"- 范围：{item['scope_note']}"])
        if key == "dongguan":
            lines.append(
                f"- 分母外排除：特殊夜班列 {item['special_paid_records_excluded']} 条；规则算出但线下未发、资格待确认 {item['rule_candidates_without_offline_pay_excluded']} 条。"
            )
        elif key == "huadong":
            lines.append(
                f"- 分母外排除：规则算出但线下未发、资格待确认 {item['rule_candidates_without_offline_pay_excluded']} 条。"
            )
        else:
            lines.append(
                f"- 分母外排除：已确认的特殊资格未入明细 {item['known_special_omitted_records_excluded']} 条；无法关联原考勤 {item['detail_records_missing_raw_attendance']} 条。"
            )
        lines.extend(
            [
                f"- 待人工/待口径：{item['manual_or_pending_records']} 条；原因：{item['status_reasons'] or '无'}。",
                f"- 不一致原因：{item['mismatch_reasons'] or '无'}。",
                f"- 已计算记录金额：线下 {item['actual_total_on_calculated']:.2f} 元，引擎 {item['predicted_total_on_calculated']:.2f} 元，净差 {item['net_difference']:.2f} 元。",
                "",
            ]
        )
    lines.extend(
        [
            "## 数据与规则版本",
            "",
            f"- 规则：{result['rule_summary']}",
            f"- 生成时间：{result['generated_at']}",
            "- 报告不输出员工姓名或工号；差异样例仅保留源表行号、日期、班次和金额。",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dongguan", type=Path, default=DEFAULT_DONGGUAN)
    parser.add_argument("--huadong", type=Path, default=DEFAULT_HUADONG)
    parser.add_argument("--jinjiang", type=Path, default=DEFAULT_JINJIANG)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    for path in (args.dongguan, args.huadong, args.jinjiang):
        if not path.is_file():
            raise FileNotFoundError(path)

    break_map = _load_break_map()
    regions = {
        "dongguan": audit_dongguan(args.dongguan, break_map),
        "huadong": audit_huadong(args.huadong, break_map),
        "jinjiang": audit_jinjiang(args.jinjiang, break_map),
    }
    result = {
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "rule_summary": "22:00—次日08:00；上下班分别向后/向前取整半小时；扣完整覆盖的班次休息；3元/小时；25元/日封顶；异常及未确认规则转人工。",
        "sources": {
            "dongguan": str(args.dongguan),
            "huadong": str(args.huadong),
            "jinjiang": str(args.jinjiang),
            "rules": baseline_shift_breaks_metadata(),
        },
        "regions": regions,
        "overall": _overall(regions),
    }
    args.output_dir.mkdir(parents=True, exist_ok=True)
    json_path = args.output_dir / "night_shift_regression.json"
    markdown_path = args.output_dir / "night_shift_regression.md"
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    markdown_path.write_text(_markdown_report(result), encoding="utf-8")
    print(json_path)
    print(markdown_path)


if __name__ == "__main__":
    main()
