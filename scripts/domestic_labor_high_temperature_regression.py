"""用华南线下结果与华东考勤/测温依据回归高温补贴验证版。"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Dict

from openpyxl import load_workbook

# 允许直接执行 `python scripts/...py`，同时保持 `python -m scripts...` 可用。
PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from bonus_platform.engine.domestic_labor.engines.gaowen_butie import GaoWenBuTieEngine
from bonus_platform.engine.domestic_labor.parser import MultiFilePayrollDataLoader


DEFAULT_SOUTH = Path("/Users/zt27532/Downloads/华南.xlsx")
DEFAULT_EAST = Path("/Users/zt27532/Downloads/华东.xlsx")
DEFAULT_TEMPERATURE = Path(
    "/Users/zt27532/Downloads/2026年中国操作部高温津贴测温登记表 (1).xlsx"
)
TOLERANCE = 0.011


def _employee_id(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value or "").strip()


def _money(value: Any) -> float:
    return float(Decimal(str(value or 0)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _load_south_expected(path: Path) -> Dict[str, float]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["AI&线下核算比对版"]
        expected: Dict[str, float] = {}
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            employee_id = _employee_id(row[2] if len(row) > 2 else None)
            if not employee_id:
                continue
            expected[employee_id] = float(row[17] or 0)
        return expected
    finally:
        workbook.close()


def audit_south(south_path: Path, temperature_path: Path) -> Dict[str, Any]:
    expected = _load_south_expected(south_path)
    sheet_mapping = {
        "0:AI餐补": "ignore",
        "0:AI外宿补贴": "ignore",
        "1:仓库通知人": "ignore",
    }
    with MultiFilePayrollDataLoader(
        [str(south_path), str(temperature_path)],
        sheet_mapping=sheet_mapping,
    ) as loader:
        input_summary = loader.validate_inputs(["gaowen_butie"], "202607")
        monthly = {
            _employee_id(row.get("工号")): row
            for row in loader.monthly.rows
            if _employee_id(row.get("工号"))
        }
        daily_by_employee = loader.group_daily_by_employee()
        engine = GaoWenBuTieEngine(loader.temperature.rows if loader.temperature else [])

        matched = 0
        residuals = []
        platform_total = Decimal("0")
        for employee_id, offline_raw in expected.items():
            employee = monthly.get(employee_id)
            if employee is None:
                residuals.append({
                    "employee_id": employee_id,
                    "reason": "月考勤缺少该工号",
                    "offline_amount": _money(offline_raw),
                    "platform_amount": None,
                    "difference": None,
                })
                continue
            result = engine.calculate(employee, daily_by_employee.get(employee_id, []))
            offline_amount = _money(offline_raw)
            difference = _money(result.amount - offline_amount)
            platform_total += Decimal(str(result.amount))
            if abs(difference) <= TOLERANCE:
                matched += 1
                continue
            residuals.append({
                "employee_id": employee_id,
                "employee_name": str(employee.get("姓名") or ""),
                "department": str(employee.get("四级部门名称") or employee.get("三级部门名称") or ""),
                "position": str(employee.get("岗位名称") or ""),
                "measurement_site": result.details.get("测温网点", ""),
                "offline_raw_amount": offline_raw,
                "offline_amount": offline_amount,
                "platform_amount": result.amount,
                "difference": difference,
                "payable_days": result.details.get("计发天数", 0),
            })

        compared = len(expected)
        residuals.sort(key=lambda row: abs(row.get("difference") or 0), reverse=True)
        return {
            "scope": "华南2026年7月线下高温补贴逐人金额回归",
            "input_rows": {
                "monthly": input_summary["monthly_rows"],
                "daily": input_summary["daily_rows"],
                "temperature": input_summary["temperature_rows"],
                "offline_results": compared,
            },
            "compared_employees": compared,
            "matched_employees": matched,
            "mismatched_employees": compared - matched,
            "employee_accuracy": round(matched / compared, 6) if compared else None,
            "offline_raw_total": round(sum(expected.values()), 4),
            "offline_total_after_employee_rounding": _money(sum(_money(value) for value in expected.values())),
            "platform_total": _money(platform_total),
            "absolute_difference": _money(sum(abs(row.get("difference") or 0) for row in residuals)),
            "net_difference": _money(sum(row.get("difference") or 0 for row in residuals)),
            "residuals": residuals,
            "monthly_employees_not_in_offline_result": sorted(set(monthly) - set(expected)),
        }


def audit_east(east_path: Path) -> Dict[str, Any]:
    sheet_mapping = {
        "0:嘉善7月餐补": "ignore",
        "0:嘉善7月外宿补贴": "ignore",
    }
    with MultiFilePayrollDataLoader([str(east_path)], sheet_mapping=sheet_mapping) as loader:
        input_summary = loader.validate_inputs(["gaowen_butie"], "202607")
        daily_by_employee = loader.group_daily_by_employee()
        engine = GaoWenBuTieEngine(loader.temperature.rows if loader.temperature else [])
        area_summary = defaultdict(lambda: {
            "employees": 0,
            "paid_employees": 0,
            "payable_days": 0,
            "amount": Decimal("0"),
            "review_employees": 0,
        })
        for employee in loader.monthly.rows:
            employee_id = _employee_id(employee.get("工号"))
            result = engine.calculate(employee, daily_by_employee.get(employee_id, []))
            work_area = str(employee.get("工作地区") or "未填写地区")
            summary = area_summary[work_area]
            summary["employees"] += 1
            summary["amount"] += Decimal(str(result.amount))
            summary["payable_days"] += int(result.details.get("计发天数") or 0)
            if result.amount > 0:
                summary["paid_employees"] += 1
            if result.warnings:
                summary["review_employees"] += 1

        hot_dates = defaultdict(lambda: {"白班": set(), "夜班": set()})
        for (site, day, shift), temperature in engine.temperature_index.items():
            if temperature >= 33:
                hot_dates[site][shift].add(day.isoformat())

        return {
            "scope": "华东2026年7月高温补贴可计算性验证",
            "accuracy": None,
            "accuracy_note": "华东工作簿未提供逐人线下高温应发金额，不将推算结果冒充准确率。",
            "input_rows": {
                "monthly": input_summary["monthly_rows"],
                "daily": input_summary["daily_rows"],
                "temperature": input_summary["temperature_rows"],
            },
            "area_summary": {
                area: {
                    **summary,
                    "amount": _money(summary["amount"]),
                }
                for area, summary in area_summary.items()
            },
            "hot_dates_by_site_and_shift": {
                site: {shift: len(days) for shift, days in shifts.items()}
                for site, shifts in hot_dates.items()
            },
        }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--south", type=Path, default=DEFAULT_SOUTH)
    parser.add_argument("--east", type=Path, default=DEFAULT_EAST)
    parser.add_argument("--temperature", type=Path, default=DEFAULT_TEMPERATURE)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    report = {
        "south": audit_south(args.south, args.temperature),
        "east": audit_east(args.east),
    }
    payload = json.dumps(report, ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(payload + "\n", encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
