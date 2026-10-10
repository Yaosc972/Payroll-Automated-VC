"""Company OBS mode retains new confirmations through calculation and recovery."""
from io import BytesIO
import json
import shutil

from openpyxl import load_workbook
import pytest

from test_domestic_labor_temperature_sites import api, _workbook_bytes
from bonus_platform.engine.domestic_labor.engines.gaowen_butie import LIAOBU_SITE
from bonus_platform.engine.domestic_labor import persistent_storage


pytestmark = pytest.mark.usefixtures("bypass_domestic_labor_access_gate")


def test_company_upload_calculate_restore_and_export_new_confirmations(api, monkeypatch):
    module, client, directory = api
    monkeypatch.setenv("SIGMA_DOMESTIC_LABOR_STORAGE_BACKEND", "obs")
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_ENV", "skillhub-integration-test")
    objects = {}
    monkeypatch.setattr(persistent_storage, "obs_storage_configured", lambda: True)
    monkeypatch.setattr(persistent_storage, "obs_put_bytes", lambda key, data, **kwargs: objects.__setitem__(key, data))
    monkeypatch.setattr(persistent_storage, "obs_get_bytes", lambda key: objects.get(key))
    assert client.get("/api/domestic-labor/upload-mode").json() == {"mode": "server"}
    mapping = {"寮步新名称": LIAOBU_SITE}
    confirmation = {"temperature_exclusions": "none"}
    response = client.post("/api/domestic-labor/runs", data={
        "engines": "gaowen_butie", "attendance_month": "202607",
        "temperature_site_mapping": json.dumps(mapping),
        "supplemental_confirmation": json.dumps(confirmation),
    }, files={"file": ("合成核算资料.xlsx", _workbook_bytes("寮步新名称"))})
    assert response.status_code == 200, response.text
    run_id = response.json()["run_id"]
    assert response.json()["status"] == "已完成"
    run_dir = directory / run_id
    metadata = module.load_payroll_metadata(run_dir)
    assert metadata["temperatureSiteMapping"] == mapping
    assert metadata["supplementalConfirmation"] == confirmation
    assert metadata["results"][0]["gaowen_butie"] == 13.8
    assert persistent_storage._file_object_path(run_id, metadata["savedFileName"]) in objects
    shutil.rmtree(run_dir)
    restored = module.load_payroll_metadata(run_dir)
    assert restored["temperatureSiteMapping"] == mapping
    assert restored["results"] == metadata["results"]
    response = client.get(f"/api/domestic-labor/runs/{run_id}/export")
    assert response.status_code == 200, response.text
    filename = response.json()["file_name"]
    downloaded = client.get(f"/api/domestic-labor/runs/{run_id}/download/{filename}")
    assert downloaded.status_code == 200
    workbook = load_workbook(BytesIO(downloaded.content))
    assert "核算汇总" in workbook.sheetnames
    assert 13.8 in next(workbook["核算汇总"].iter_rows(min_row=2, max_row=2, values_only=True))
    shutil.rmtree(run_dir)
    recovered_download = client.get(f"/api/domestic-labor/runs/{run_id}/download/{filename}")
    assert recovered_download.status_code == 200
    assert recovered_download.content == downloaded.content
