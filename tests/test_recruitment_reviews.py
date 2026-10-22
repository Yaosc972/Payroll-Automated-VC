from copy import deepcopy
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from bonus_platform.engine.recruitment_drafts import DraftStore, DraftConflict
from bonus_platform.engine import recruitment_reviews as review
from test_recruitment_awards import person
from test_recruitment_drafts import book


def setup(tmp_path):
    store=DraftStore(tmp_path);v=person()['values'];v['体系外人员编号']='test'
    state=store.ingest(202607,[('资料.xlsx',book(list(v),[list(v.values())]))],0)
    rid=state['rows'][0]['id']
    return store,state,rid


def test_preview_atomic_apply_retry_undo_redo_and_restart(tmp_path):
    store,state,rid=setup(tmp_path)
    edits=[{'rowId':rid,'field':'职级','value':'P2-1','reason':'任职资料核对'}]
    preview=review.prepare_preview(store,202607,state['revision'],edits)
    assert store.read(202607)==state
    assert preview['people']==1
    assert next(i for i in preview['impact'] if i['kind']=='内推奖金')['after']==500
    saved=review.apply_preview(store,202607,state['revision'],preview['id'])
    assert review.apply_preview(store,202607,state['revision'],preview['id'])==saved
    assert saved['rows'][0]['calculation']['referral']['total']==500
    undone=store.history_action(202607,saved['revision'],'undo')
    assert undone['rows'][0]['calculation']['referral']['total']==400
    redone=DraftStore(tmp_path).history_action(202607,undone['revision'],'redo')
    assert redone['rows'][0]['calculation']['referral']['total']==500
    assert len(redone['edits'])==1 and len(redone['historyActions'])==2
    carry=store.carry_forward(202608,202607,0)
    with pytest.raises(ValueError): store.history_action(202608,carry['revision'],'undo')


def test_undo_restores_unresolved_source_conflict_instead_of_approving_old_value(tmp_path):
    store,state,rid=setup(tmp_path)
    value=person(职级='P2-1')['values'];value['体系外人员编号']='test'
    state=store.ingest(202607,[('另一来源.xlsx',book(list(value),[list(value.values())]))],state['revision'])
    assert '职级' in state['rows'][0]['unresolvedFields']
    state=store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':'职级','value':'P2-1','reason':'核对'}]})
    assert '职级' not in state['rows'][0]['unresolvedFields']
    state=store.history_action(202607,state['revision'],'undo')
    assert '职级' in state['rows'][0]['unresolvedFields']
    assert state['rows'][0]['calculation']['referral']['total'] is None


def test_stale_preview_and_duplicate_cells_cannot_apply(tmp_path):
    store,state,rid=setup(tmp_path);edit={'rowId':rid,'field':'职级','value':'P2-1','reason':'核对'}
    with pytest.raises(ValueError): review.prepare_preview(store,202607,state['revision'],[edit,edit])
    preview=review.prepare_preview(store,202607,state['revision'],[edit])
    newer=store.change(202607,state['revision'],{'edits':[dict(edit,value='P3-1')]})
    with pytest.raises(DraftConflict): review.apply_preview(store,202607,newer['revision'],preview['id'])
    assert store.read(202607)==newer


def test_preview_reports_recipient_change_even_when_amount_does_not_change(tmp_path):
    store,state,rid=setup(tmp_path)
    preview=review.prepare_preview(store,202607,state['revision'],[
        {'rowId':rid,'field':'招聘负责人姓名','value':'核实后的负责人','reason':'录用备注确认'}])
    assert preview['impact']==[]
    changes=[n for n in preview['nodeImpact'] if n['role']=='招聘负责人']
    assert changes and all(n['before']['amount']==n['after']['amount'] for n in changes)
    assert all(n['after']['recipientName']=='核实后的负责人' for n in changes)
    saved=review.apply_preview(store,202607,state['revision'],preview['id'])
    actual={n['id']:n for n in saved['rows'][0]['calculation']['recruitment']['nodes']}
    assert all(actual[n['nodeId']]['recipientName']==n['after']['recipientName'] for n in changes)


def test_preview_reports_date_and_month_effect_without_changing_total(tmp_path):
    store,state,rid=setup(tmp_path)
    preview=review.prepare_preview(store,202607,state['revision'],[
        {'rowId':rid,'field':'入职日期','value':'2026-07-06','reason':'核验实际任职日期'},
        {'rowId':rid,'field':'候选人入职时间','value':'2026-07-06','reason':'核验实际任职日期'},
        {'rowId':rid,'field':'招聘启动日期','value':'2026-05-25','reason':'保持已核实的相同招聘周期'}])
    assert preview['impact']==[]
    assert preview['currentImpact']
    assert any(n['before']['dueMonth']==202607 and n['after']['dueMonth']==202608 for n in preview['nodeImpact'])
    saved=review.apply_preview(store,202607,state['revision'],preview['id'])
    for change in preview['currentImpact']:
        kind='recruitment' if change['kind']=='招聘奖金' else 'referral'
        assert saved['rows'][0]['calculation'][kind]['currentAmount']==change['after']


def test_approval_identifier_matching_rejects_prefix_and_concatenated_numbers(tmp_path,monkeypatch):
    store,state,rid=setup(tmp_path)
    state=store.ingest(202607,[('编号依据.pdf',b'%PDF-test')],state['revision'])
    sid=state['sources'][-1]['id']
    monkeypatch.setattr(review,'extract_document',lambda *a:{'pageCount':1,'pages':[{'page':1,'text':'zt-check9\n记录 2026 0604 0020'}]})
    assert review.document(store,202607,sid)['matches']==[]
    monkeypatch.setattr(review,'extract_document',lambda *a:{'pageCount':1,'pages':[{'page':1,'text':'已核对工号 / zt-check，任职记录'}]})
    match=review.document(store,202607,sid)['matches']
    assert len(match)==1 and match[0]['rowId']==rid and match[0]['matchedBy']==['工号']


def test_workbook_round_trip_recognizes_cells_and_rejects_hidden_duplicate_formula_stale(tmp_path):
    store,state,rid=setup(tmp_path);data=review.editable_workbook(state)
    assert review.workbook_edits(state,data)==[]
    def exported(change):
        wb=load_workbook(BytesIO(data));change(wb.active);out=BytesIO();wb.save(out);return out.getvalue()
    column=4+review.GRID_FIELDS.index('职级')
    edits=review.workbook_edits(state,exported(lambda ws:setattr(ws.cell(2,column),'value','P2-1')))
    assert len(edits)==1 and edits[0]['rowId']==rid and edits[0]['field']=='职级'
    def duplicate(ws):
        ws.append([c.value for c in ws[2]]);ws.row_dimensions[3].hidden=True
    with pytest.raises(ValueError,match='重复'):review.workbook_edits(state,exported(duplicate))
    with pytest.raises(ValueError,match='公式'):review.workbook_edits(state,exported(lambda ws:setattr(ws.cell(2,column),'value','=1+1')))
    with pytest.raises(ValueError,match='过期'):review.workbook_edits({**state,'revision':state['revision']+1},data)


def test_approval_page_status_attachment_gates_and_netted_recovery_entries(tmp_path,monkeypatch):
    store,state,rid=setup(tmp_path)
    state=store.ingest(202607,[('审批.pdf',b'%PDF-test')],state['revision'])
    sid=state['sources'][-1]['id']
    monkeypatch.setattr(review,'document',lambda *args:{'pageCount':2,'mentionedAttachments':['最新名单.xlsx']})
    evidence={'sourceId':sid,'pageNumbers':[2],'status':'审批未结束','scopeChecked':True,'note':'核对第二页意见','attachmentIds':[]}
    edits=[{'rowId':rid,'field':'ABC类别','value':'B'}]
    with pytest.raises(ValueError,match='认可状态'):review.prepare_preview(store,202607,state['revision'],edits,evidence=evidence)
    evidence['status']='已认可'
    with pytest.raises(ValueError,match='尚未关联'):review.prepare_preview(store,202607,state['revision'],edits,evidence=evidence)
    pending=review.save_review(store,202607,state['revision'],evidence)
    assert pending['reviewEntries'][-1]['missingAttachments']==['最新名单.xlsx']
    state=store.ingest(202607,[('最新名单.xlsx',book(['体系外人员编号','姓名'],[]))],pending['revision'])
    evidence['attachmentIds']=[state['sources'][-1]['id']]
    arrangements=[{'rowId':rid,'recipientId':'a','recipientName':'甲','amount':-10,'currency':'欧元'}, {'rowId':rid,'recipientId':'b','recipientName':'乙','amount':10,'currency':'欧元'}]
    preview=review.prepare_preview(store,202607,state['revision'],edits,evidence=evidence,arrangements=arrangements)
    result=review.apply_preview(store,202607,state['revision'],preview['id'])
    assert result['edits'][-1]['evidencePages']==[2]
    assert len(result['adjustmentDrafts'])==2
    assert sum(r['amount'] for r in result['adjustmentDrafts'])==0
    assert all(not r['paymentReady'] and r['status']=='待核实际支付' for r in result['adjustmentDrafts'])
    assert result['rows'][0]['history']==state['rows'][0]['history']
    with pytest.raises(ValueError,match='页码'):review.save_review(store,202607,result['revision'],dict(evidence,pageNumbers=[3]))


def test_preview_and_history_api(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from bonus_platform.app import app
    from bonus_platform import recruitment_draft_api as api
    store,state,rid=setup(tmp_path);monkeypatch.setattr(api,'store',store)
    client=TestClient(app)
    p=client.post('/api/recruitment/drafts/202607/previews',json={'revision':state['revision'],'edits':[{'rowId':rid,'field':'职级','value':'P2-1','reason':'核对'}]})
    assert p.status_code==200
    response=client.post('/api/recruitment/drafts/202607/previews/'+p.json()['id']+'/apply',json={'revision':state['revision']})
    assert response.status_code==200 and response.json()['canUndo']
    response=client.post('/api/recruitment/drafts/202607/history',json={'revision':response.json()['revision'],'action':'undo'})
    assert response.status_code==200 and response.json()['canRedo']


def test_approved_cancellation_keeps_formula_and_history_and_undo_restores(tmp_path,monkeypatch):
    store,state,rid=setup(tmp_path)
    state=store.ingest(202607,[('审批.pdf',b'%PDF-test')],state['revision']);sid=state['sources'][-1]['id']
    monkeypatch.setattr(review,'document',lambda *args:{'pageCount':1,'mentionedAttachments':[]})
    evidence={'sourceId':sid,'pageNumbers':[1],'status':'已认可','scopeChecked':True,'note':'已核对原页明确取消招聘奖金'}
    preview=review.prepare_preview(store,202607,state['revision'],[],evidence=evidence,decisions=[{'rowId':rid,'kind':'recruitment','action':'exclude'}])
    assert preview['impact'][0]['before']==55 and preview['impact'][0]['after']==0
    state=review.apply_preview(store,202607,state['revision'],preview['id'])
    calculated=state['rows'][0]['calculation']['recruitment']
    assert calculated['total']==0 and calculated['formulaTotal']==55
    assert not calculated['nodes']
    assert state['awardDecisions'][0]['pageNumbers']==[1]
    state=store.history_action(202607,state['revision'],'undo')
    assert state['rows'][0]['calculation']['recruitment']['total']==55


def test_recipient_prefix_aliases_enrich_and_replacement_does_not_inherit_old_status(tmp_path):
    from test_recruitment_drafts import ROSTER
    store,state,rid=setup(tmp_path)
    state=store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':'招聘负责人工号','value':'海外-zt12345','reason':'编号对应'}]})
    state=store.ingest(202607,[('花名册.xlsx',book(ROSTER,[['领取人','zt12345','','2020-01-01','正式','','','EHR入职','中国','中国']]))],state['revision'])
    row=next(r for r in state['rows'] if r['id']==rid)
    assert row['values']['招聘负责人工号']=='zt12345'
    assert row['values']['招聘负责人入职日期']=='2020-01-01'
    state=store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':'招聘负责人工号','value':'外部-001','reason':'外部领取人待补身份'}]})
    row=next(r for r in state['rows'] if r['id']==rid)
    assert not row['values'].get('招聘负责人入职日期')
    assert not row['values'].get('招聘负责人人员状态')
    assert any(t['kind']=='missing_recipient' and rid in t['rowIds'] for t in state['todos'])


def test_reconciliation_excludes_reference_and_keeps_cancelled_nodes_visible(tmp_path,monkeypatch):
    from bonus_platform.engine.recruitment_reconcile import reconcile,export_reconciliation
    store,state,rid=setup(tmp_path)
    reference=person(职级='P3-1')['values']
    reference.update({'奖金标准':999,'招聘负责人入职1个月发放奖金':99,'协助招聘人入职1个月发放奖金':8})
    state=store.ingest(202607,[('核对台账.xlsx',book(list(reference),[list(reference.values())],title='附件一-招聘奖金核算明细'))],state['revision'])
    sid=state['sources'][-1]['id'];saved=deepcopy(state)
    result=reconcile(store,202607,sid)
    assert store.read(202607)==saved
    standard=next(d for d in result['moneyDifferences'] if d['kind']=='招聘奖金' and d['item']=='奖金标准')
    assert standard['reference']==999 and standard['calculated']==55
    helper=next(d for d in result['moneyDifferences'] if d['item'].startswith('协助招聘人'))
    assert helper['reference']==8 and helper['calculated']==0
    state=store.ingest(202607,[('审批.pdf',b'%PDF-test')],state['revision'])
    monkeypatch.setattr(review,'document',lambda *args:{'pageCount':1,'mentionedAttachments':[]})
    evidence={'sourceId':state['sources'][-1]['id'],'pageNumbers':[1],'status':'已认可','scopeChecked':True,'scopeComplete':False,'note':'核对后取消招聘奖金'}
    preview=review.prepare_preview(store,202607,state['revision'],[],evidence=evidence,decisions=[{'rowId':rid,'kind':'recruitment','action':'exclude'}])
    state=review.apply_preview(store,202607,state['revision'],preview['id'])
    result=reconcile(store,202607,sid)
    node=next(d for d in result['moneyDifferences'] if d['item']=='招聘负责人 · 入职满一个月')
    assert node['reference']==99 and node['calculated']==0
    assert '审批不发' in node['reason']
    assert any(t['title']=='审批剩余范围待核' for t in state['todos'])
    wb=load_workbook(BytesIO(export_reconciliation(result)))
    assert wb['金额差异'].max_row==result['summary']['moneyDifferences']+1
    assert len(wb.sheetnames)==6


def test_malformed_workbook_returns_useful_error_and_export_keeps_decision_proof(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from bonus_platform.app import app
    from bonus_platform import recruitment_draft_api as api
    store,state,rid=setup(tmp_path);monkeypatch.setattr(api,'store',store)
    client=TestClient(app)
    response=client.post('/api/recruitment/drafts/202607/worksheet/preview',data={'revision':state['revision']},files={'file':('bad.xlsx',b'broken')})
    assert response.status_code==400 and '完整表格' in response.json()['detail']
    state=store.ingest(202607,[('审批.pdf',b'%PDF-test')],state['revision']);sid=state['sources'][-1]['id']
    monkeypatch.setattr(review,'document',lambda *args:{'pageCount':1,'mentionedAttachments':[]})
    evidence={'sourceId':sid,'pageNumbers':[1],'status':'已认可','scopeChecked':True,'note':'明确不发'}
    preview=review.prepare_preview(store,202607,state['revision'],[],evidence=evidence,decisions=[{'rowId':rid,'kind':'recruitment','action':'exclude'}])
    state=review.apply_preview(store,202607,state['revision'],preview['id'])
    result=client.get('/api/recruitment/drafts/202607/calculation/export')
    wb=load_workbook(BytesIO(result.content));rows=list(wb['调整及依据'].values)
    assert rows[1][1]=='招聘奖金' and rows[1][5]=='审批.pdf' and rows[1][6]=='1' and rows[1][8]=='已采用'
    store.history_action(202607,state['revision'],'undo')
    result=client.get('/api/recruitment/drafts/202607/calculation/export')
    assert list(load_workbook(BytesIO(result.content))['调整及依据'].values)[1][8]=='已撤销'


def test_corrupt_approval_returns_reviewable_error_without_losing_sources(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from bonus_platform.app import app
    from bonus_platform import recruitment_draft_api as api
    store,state,rid=setup(tmp_path)
    state=store.ingest(202607,[('损坏审批.pdf',b'%PDF-broken')],state['revision'])
    sid=state['sources'][-1]['id'];monkeypatch.setattr(api,'store',store)
    response=TestClient(app).get('/api/recruitment/drafts/202607/documents/'+sid)
    assert response.status_code==400 and '审批原页' in response.json()['detail']
    assert store.read(202607)==state
