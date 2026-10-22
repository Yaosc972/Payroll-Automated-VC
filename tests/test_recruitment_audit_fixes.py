"""Regressions from the independent July audit; no July answers are inputs."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook

from bonus_platform.engine.recruitment_awards import award
from bonus_platform.engine.recruitment_employment import status_from_timeline, enrich
from bonus_platform.engine.recruitment_attention import project as attention, is_current
from bonus_platform.engine.recruitment_monthly import project, export
from bonus_platform.engine.recruitment_settlement import apply_settlement
from bonus_platform.engine import recruitment_reviews as reviews
from test_recruitment_awards import person
from test_recruitment_flow import history, state_with_rows
from test_recruitment_trials import independent
from test_recruitment_reviews import setup


def test_recipient_timeline_cannot_use_a_different_employment():
    row=person(推荐人入职日期='2026-07-01')
    old=history(lastWorkDate='2026-05-14');old.update(entryDate='2020-01-01',startDate='2020-01-01')
    row['employmentTimelines']={'推荐人':[old]}
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31)) is None
    current=history(regularizationDate='2026-07-15')
    row['employmentTimelines']['推荐人'].append(current)
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31),True)==''


def test_overseas_roster_not_overridden_by_domestic_hr_history():
    row=person(推荐人入职日期='2026-07-01')
    row['recipientSourceRegions']={'推荐人':'海外'}
    row['employmentTimelines']={'推荐人':[history(lastWorkDate='2026-07-15')]}
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31)) is None


def test_monthly_export_blank_adopts_unique_approved_fact_and_manual_clear_wins():
    row=person(工号='zt1',标签分类='国内',转正日期='')
    row.update(edits=[],sources=[],provenance={},conflicts=[],monthlyRosterFields=['转正日期'])
    enrich(row,{'zt1':[history(regularizationDate='2026-07-20')]},[])
    assert row['values']['转正日期']=='2026-07-20'
    assert row['conflicts']==[]
    assert row['provenance']['转正日期'][0]['value']=='2026-07-20'
    row['values']['转正日期']='';row['edits']=[{'field':'转正日期'}];row['conflicts']=[]
    enrich(row,{'zt1':[history(regularizationDate='2026-07-20')]},[])
    assert row['values']['转正日期']=='' and row['conflicts']==[]


def test_effective_probation_source_conflict_is_not_hidden_as_history():
    state,row,result=independent()
    row['values']['转正日期']=''
    row['unresolvedFields']=['转正日期']
    row['provenance']={'转正日期':[{'role':'employment_history','value':'2026-07-20'}]}
    for h in row['history']:
        for cell in h['cells']:cell.setdefault('label','')
    apply_settlement(state,row,result);row['calculation']=result
    assert is_current(row,'recruitment',202607,set())


def test_exit_conflict_affecting_this_month_cancellation_is_current():
    state,row,result=independent()
    row['values']['转正日期']='2026-09-27'
    row['unresolvedFields']=['最后工作日']
    row['provenance']={'最后工作日':[{'role':'employment_history','value':'2026-07-20'}]}
    for h in row['history']:
        for cell in h['cells']:cell.setdefault('label','')
    apply_settlement(state,row,result);row['calculation']=result
    assert is_current(row,'recruitment',202607,set())


def test_recruitment_task_does_not_hold_unrelated_referral_export():
    state=state_with_rows();row=state['rows'][0]
    state['todos']=[{'id':'cycle','kind':'calculation','title':'招聘奖金：招聘起止日期或标准周期待核','rowIds':[row['id']],'sourceIds':[]}]
    monthly=project(state,attention(state))
    referral=[n for n in monthly['nodes'] if n['kind']=='referral' and n['scope']=='current']
    assert referral and not any(n['review'] for n in referral)
    assert all(n['review'] for n in monthly['nodes'] if n['kind']=='recruitment' and n['scope']=='current')
    wb=load_workbook(BytesIO(export(state,monthly)))
    assert wb['附件三-内推奖金汇总'].cell(3,7).value > 0
    assert all(row[-3]==0 for row in list(wb['附件二-招聘奖金汇总表'].values)[2:])


def test_unknown_entitlement_cannot_report_zero_cancelled_balance():
    state,row,result=independent()
    row['values'].update(招聘渠道='未知渠道',最后工作日='2026-07-10')
    apply_settlement(state,row,result)
    assert result['recruitment']['total'] is None
    assert result['recruitment']['cancelledBalances'][0]['amount'] is None


@pytest.mark.parametrize('opening_text,scheduled',[('未转正，待发放',202607),('不享受招聘奖金',None),('暂缓发放',None),('',None)])
def test_only_explicit_pending_probation_is_carried_to_current_month(opening_text,scheduled):
    state,row,result=independent()
    row['values']['转正日期']='2026-06-27'
    row['history'][0]['cells'].append({'column':53,'value':opening_text,'label':'转正发放周期'})
    apply_settlement(state,row,result)
    node=result['recruitment']['nodes'][1]
    assert node.get('scheduledMonth')==scheduled
    if scheduled:
        assert node['amount']==255 and node['scheduleSource']['sourceId']=='june'
        assert node['dueMonth']==202606


def test_role_exclusion_does_not_reclassify_already_arranged_assistant():
    state,row,result=independent()
    row['values'].update(协助招聘人工号='zt-assist',协助招聘人姓名='协助',协助招聘人人员状态='正式',协助招聘人入职日期='2020-01-01')
    row['history'][0]['cells'] += [{'column':46,'value':13.5,'label':'协助首期金额'}, {'column':47,'value':202605,'label':'协助首期月份'}]
    row['awardDecisions']=[{'kind':'recruitment','role':'协助招聘人','action':'exclude','sourceId':'approval','reason':'仅取消未发协助'}]
    apply_settlement(state,row,result)
    locked=next(n for n in result['recruitment']['nodes'] if n['role']=='协助招聘人' and n.get('openingLocked'))
    assert locked['scheduledMonth']==202605 and locked['status']!='按规则不发'


def test_assistant_exclusion_keeps_primary_share_and_opening_history():
    state,row,result=independent()
    row['values'].update(协助招聘人工号='zt-assist',协助招聘人姓名='协助',协助招聘人人员状态='正式',协助招聘人入职日期='2020-01-01')
    apply_settlement(state,row,result)
    before=deepcopy(result['recruitment'])
    row['awardDecisions']=[{'kind':'recruitment','role':'协助招聘人','action':'exclude','sourceId':'approval','reason':'只取消协助资格'}]
    apply_settlement(state,row,result)
    after=result['recruitment']
    assert [(n['amount'],n.get('scheduledMonth')) for n in before['nodes'] if n['role']=='招聘负责人']==[(n['amount'],n.get('scheduledMonth')) for n in after['nodes'] if n['role']=='招聘负责人']
    assistant=[n for n in after['nodes'] if n['role']=='协助招聘人']
    assert assistant and all(n.get('scheduledMonth') is None for n in assistant)
    assert all(n['status']=='按规则不发' for n in assistant)
    assert after['total']==before['total']  # formula entitlement kept for audit


def test_role_decision_preview_apply_undo_and_invalid_role(tmp_path,monkeypatch):
    store,state,rid=setup(tmp_path)
    state=store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':f,'value':v,'reason':'原始资料'} for f,v in {'协助招聘人工号':'zt-a','协助招聘人姓名':'协助','协助招聘人人员状态':'正式','协助招聘人入职日期':'2020-01-01'}.items()]})
    state=store.ingest(202607,[('依据.pdf',b'%PDF-test')],state['revision']);sid=state['sources'][-1]['id']
    monkeypatch.setattr(reviews,'document',lambda *args:{'pageCount':1,'mentionedAttachments':[]})
    evidence={'sourceId':sid,'pageNumbers':[1],'status':'已认可','scopeChecked':True,'note':'仅取消协助资格'}
    decision={'rowId':rid,'kind':'recruitment','role':'协助招聘人','action':'exclude'}
    with pytest.raises(ValueError):
        reviews.prepare_preview(store,202607,state['revision'],[],evidence=evidence,decisions=[dict(decision,role='推荐人')])
    preview=reviews.prepare_preview(store,202607,state['revision'],[],evidence=evidence,decisions=[decision])
    assert preview['nodeImpact'] and all(n['role']=='协助招聘人' for n in preview['nodeImpact'])
    saved=reviews.apply_preview(store,202607,state['revision'],preview['id'])
    assert any(n['role']=='招聘负责人' and n.get('scheduledMonth',n['dueMonth'])==202607 for n in saved['rows'][0]['calculation']['recruitment']['nodes'])
    undone=store.history_action(202607,saved['revision'],'undo')
    assert any(n['role']=='协助招聘人' and n.get('scheduledMonth',n['dueMonth'])==202607 for n in undone['rows'][0]['calculation']['recruitment']['nodes'])
