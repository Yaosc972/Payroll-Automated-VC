"""工作表用途识别、确认及确认后的必需字段校验。"""
import json
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from bonus_platform.app import app
from bonus_platform.engine.domestic_labor.parser import MultiFilePayrollDataLoader as Loader

pytestmark = pytest.mark.usefixtures("bypass_domestic_labor_access_gate")


@pytest.mark.parametrize("name,headers,expected", [
    ("月考勤", ["工号", "考勤月份", "排班天数", "入宿时间", "退宿时间"], "monthly"),
    ("人员明细", ["工号", "考勤月份", "排班天数", "入宿时间"], "monthly"),
    ("Sheet1", ["工号", "出勤日期", "上班一"], "daily"),
    ("人员明细", ["工号", "入宿时间", "退宿时间"], "housing"),
    ("人员明细", ["工号", "考勤月份", "排班天数", "日期"], None),
    ("月考勤", ["工号", "入宿时间"], None),
    ("核对版", ["工号", "考勤月份", "排班天数"], None),
])
def test_sheet_evidence(name, headers, expected):
    assert Loader._sheet_type(name, headers) == expected


def workbook_bytes(headers):
    wb = Workbook()
    wb.active.title = "人员明细"
    wb.active.append(headers)
    values = {"工号": "TEST001", "姓名": "测试员工", "考勤月份": "202607", "排班天数": 22,
              "日期": "2026-07-01", "工作地区": "嘉善"}
    wb.active.append([values.get(header, "") for header in headers])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def test_confirmation_is_carried_into_background_calculation(monkeypatch, tmp_path):
    # 捕获后台实参，验证确认决定没有在校验后丢失。
    import importlib
    app_module = importlib.import_module("bonus_platform.app")
    runs_module = importlib.import_module("bonus_platform.engine.domestic_labor.runs")
    monkeypatch.setattr(runs_module, "DOMESTIC_LABOR_RUNS_DIR", tmp_path)
    monkeypatch.setattr(app_module, "DOMESTIC_LABOR_RUNS_DIR", tmp_path)
    captured = []
    monkeypatch.setattr(app_module, "_run_payroll_calculation", lambda *args, **kwargs: (captured.append(kwargs), {"status": "已完成"})[1])
    content = workbook_bytes(["工号", "考勤月份", "排班天数", "日期"])
    client = TestClient(app)
    def submit(mapping=None):
        return client.post("/api/domestic-labor/runs", files={"file": ("test.xlsx", content)},
                           data={"engines": "gangwei_butie", "attendance_month": "202607",
                                 "sheet_mapping": json.dumps(mapping or {})})
    response = submit()
    assert response.status_code == 409
    sheet = response.json()["detail"]["sheets"][0]
    assert set(sheet["candidates"]) == {"monthly", "daily"}
    assert sheet["row_count"] == 1
    assert submit({"0:人员明细": "housing"}).status_code == 400
    response = submit({"0:人员明细": "monthly"})
    assert response.status_code == 200
    assert response.json()["input_summary"]["monthly_rows"] == 1
    assert captured and captured[0]["sheet_mapping"] == {"0:人员明细": "monthly"}


def test_choice_cannot_bypass_missing_fields():
    response = TestClient(app).post(
        "/api/domestic-labor/runs",
        files={"file": ("test.xlsx", workbook_bytes(["姓名"]))},
        data={"engines": "gangwei_butie", "attendance_month": "202607",
              "sheet_mapping": '{"0:人员明细":"monthly"}'},
    )
    assert response.status_code == 400
    assert "工号" in response.json()["detail"]


def test_unrelated_sheets_do_not_require_confirmation(tmp_path):
    wb = Workbook()
    wb.active.title = "月考勤"
    wb.active.append(["工号", "考勤月份", "排班天数"])
    wb.active.append(["TEST001", "202607", 22])
    for name in ["7月证书补贴", "高温补贴费用明细", "夜班津贴", "五险费用2026.08"]:
        sheet = wb.create_sheet(name)
        sheet.append(["工号", "姓名", "金额"])
        sheet.append(["TEST001", "测试员工", 100])
    path = tmp_path / "mixed.xlsx"
    wb.save(path)
    with Loader([str(path)]) as loader:
        assert loader.validate_inputs(["gangwei_butie"], "202607")["monthly_rows"] == 1


def test_skipping_unrelated_sheets_still_requires_attendance(tmp_path):
    wb = Workbook()
    wb.active.title = "夜班津贴"
    wb.active.append(["工号", "姓名", "金额"])
    wb.active.append(["TEST001", "测试员工", 100])
    path = tmp_path / "result_only.xlsx"
    wb.save(path)
    with Loader([str(path)]) as loader:
        with pytest.raises(ValueError, match="未识别到月考勤"):
            loader.validate_inputs(["yeban_butie"], "202607")
