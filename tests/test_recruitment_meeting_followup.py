from copy import deepcopy
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from bonus_platform.engine.recruitment_awards import calculate_draft
from bonus_platform.engine.recruitment_detail_table import installment_period
from bonus_platform.engine.recruitment_monthly import project, export
from bonus_platform.engine.recruitment_worksheet import worksheet, export_results
from bonus_platform.engine import recruitment_drafts as drafts
from test_recruitment_flow import state_with_rows
from test_recruitment_reviews import setup


def test_no_payment_reason_stays_with_its_installment_and_amount_across_exports():
    state=state_with_rows();row=state['rows'][0]
    row['values'].update(招聘负责人人员状态='离职',招聘负责人最后工作日='2026-07-20',人工备注='业务原备注')
    calculate_draft(state)
    node=row['calculation']['recruitment']['nodes'][0]
    assert node['status']=='按规则不发' and node['amount']>0 and node['scheduledMonth'] is None
    monthly=project(state,{'current':[],'history':[]})
    entry=monthly['worksheet']['rows'][0]
    period=entry['cells']['recruitment.招聘负责人.0.period']
    assert period=='不发：招聘负责人：适用日期已不在职'
    assert entry['cells']['recruitment.招聘负责人.0.amount']==node['amount']
    assert entry['cells']['recruitmentAmount']==0
    assert entry['cells']['note']=='业务原备注'
    assert not monthly['recipientSummaries']['recruitment']['rows']
    assert row['calculation']['recruitment']['currentAmount']==0
    assert state['calculation']['totalsByCurrency']=={'欧元':200.0}

    table=monthly['worksheet'];wb=load_workbook(BytesIO(export(state,monthly)))
    ws=wb['附件一-招聘奖金核算明细']
    for key,expected in [('recruitment.招聘负责人.0.amount',node['amount']),('recruitment.招聘负责人.0.period',period)]:
        col=next(i for i,c in enumerate(table['columns'],1) if c['key']==key)
        assert ws.cell(3,col).value==expected
    complete=worksheet(state)
    assert complete['rows'][0]['cells']['recruitment.招聘负责人.0.scheduledMonth']==period
    wb=load_workbook(BytesIO(export_results(state)))
    ws=wb['核算总明细']
    for key,expected in [('recruitment.招聘负责人.0.amount',node['amount']),('recruitment.招聘负责人.0.scheduledMonth',period)]:
        col=next(i for i,c in enumerate(complete['columns'],1) if c['key']==key)
        assert ws.cell(4,col).value==expected
    assert all(r[2]!='招聘奖金' for r in list(wb['本月到期明细'].values)[1:])


def test_no_payment_reason_precedes_an_old_period_and_is_not_duplicated():
    node={'label':'转正','status':'按规则不发','scheduledMonth':202609,
          'scheduleText':'候选人：适用日期已不在职；招聘负责人：适用日期已不在职',
          'eligibilityReasons':['候选人：适用日期已不在职','招聘负责人：适用日期已不在职']}
    before=deepcopy(node)
    assert installment_period(node)=='不发：候选人：适用日期已不在职；招聘负责人：适用日期已不在职'
    assert node==before


def test_saved_revision_and_current_state_are_identical_independent_files(tmp_path):
    store,state,_=setup(tmp_path)
    folder=store._folder(state['month']);current=folder/'draft.json';revision=folder/'revisions'/f"{state['revision']}.json"
    assert current.read_bytes()==revision.read_bytes()
    assert current.stat().st_ino!=revision.stat().st_ino
    assert json.loads(current.read_text())['rows']==state['rows']
    assert not list(folder.rglob('*.tmp'))


def test_failed_file_replace_keeps_existing_current_state_and_cleans_temp(tmp_path,monkeypatch):
    store,state,rid=setup(tmp_path);folder=store._folder(state['month']);current=folder/'draft.json'
    before=current.read_bytes();replace=drafts.os.replace
    def fail_current(src,dst):
        if dst==current:raise OSError('isolated write failure')
        return replace(src,dst)
    monkeypatch.setattr(drafts.os,'replace',fail_current)
    with pytest.raises(OSError,match='isolated write failure'):
        store.change(state['month'],state['revision'],{'edits':[{'rowId':rid,'field':'人工备注','value':'隔离核验','reason':'核对'}]})
    assert current.read_bytes()==before
    assert not list(folder.rglob('*.tmp'))
