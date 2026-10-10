from datetime import date
from io import BytesIO
import json
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from bonus_platform.engine.domestic_labor.engines.gaowen_butie import GaoWenBuTieEngine, LIAOBU_SITE
from bonus_platform.engine.domestic_labor.temperature_sites import (
    prepare_temperature_sites, TemperatureSiteConfirmationRequired,
    validate_temperature_site_mapping,
)
from bonus_platform.engine.domestic_labor.templates import generate_template


pytestmark = pytest.mark.usefixtures("bypass_domestic_labor_access_gate")


def _loader(daily_site, temperature_site):
    return SimpleNamespace(
        daily=SimpleNamespace(rows=[{"测温网点": daily_site, "出勤日期": date(2026, 7, 1),
                                    "班次时间段": "19:00-28:00;", "正班时数": 8, "实际上班时数": 8}]),
        temperature=SimpleNamespace(rows=[{"测温网点": temperature_site, "班次日期": date(2026, 7, 1),
                                          "测温班次": "夜班", "测温温度": 34.5}]),
    )


def test_custom_site_shared_by_both_inputs_is_accepted_and_paid():
    loader = _loader("新仓 Ａ", "新仓A")
    audit = prepare_temperature_sites(loader)
    result = GaoWenBuTieEngine(loader.temperature.rows).calculate(
        {"工号": "T001", "姓名": "测试", "工作地区": "东莞", "岗位名称": "操作员"}, loader.daily.rows,
    )
    assert result.amount == 13.8
    assert loader.daily.rows[0]["原始测温网点"] == "新仓 Ａ"
    assert audit["temperature_site_matches"][0]["method"] == "format"


def test_unknown_alias_is_gated_without_mutating_rows():
    loader = _loader("", "寮步新名称")
    with pytest.raises(TemperatureSiteConfirmationRequired) as caught:
        prepare_temperature_sites(loader)
    assert caught.value.detail["records"] == [{"name": "寮步新名称", "sources": {"测温登记": 1}}]
    assert loader.temperature.rows[0]["测温网点"] == "寮步新名称"
    assert "原始测温网点" not in loader.temperature.rows[0]


def test_confirmation_is_batch_scoped_and_custom_aliases_can_converge():
    loader = _loader("新仓简称", "新仓全称")
    audit = prepare_temperature_sites(loader, {"新仓简称": "新仓全称", "新仓全称": "新仓全称"})
    assert loader.daily.rows[0]["测温网点"] == "新仓全称"
    assert audit["temperature_site_matches"][0]["method"] == "confirmed"
    with pytest.raises(TemperatureSiteConfirmationRequired):
        prepare_temperature_sites(_loader("新仓简称", "新仓全称"))


@pytest.mark.parametrize("mapping", [None, [], {"": "a"}, {"A": "B", "B": "A"}, {"Ａ": "x", "A": "y"}])
def test_invalid_mapping_is_rejected(mapping):
    with pytest.raises(ValueError):
        validate_temperature_site_mapping(mapping)


def test_mapping_to_an_unrelated_invented_target_is_rejected():
    with pytest.raises(ValueError, match="本次上传名称"):
        prepare_temperature_sites(_loader("", "寮步简称"), {"寮步简称": "凭空网点"})


def test_template_has_optional_site_dropdown_and_strict_shift_dropdown():
    wb = load_workbook(BytesIO(generate_template("gaowen_butie")))
    assert wb.sheetnames == ["测温登记", "不享有高温补贴名单", "填写说明"]
    for sheet_name in ("测温登记",):
        ws = wb[sheet_name]
        validations = list(ws.data_validations.dataValidation)
        site_column = list(next(ws.values)).index("测温网点") + 1
        site = next(item for item in validations if f"{ws.cell(1, site_column).column_letter}2" in item.sqref)
        assert site.type == "list" and site.showErrorMessage is False
        assert LIAOBU_SITE in site.formula1
        assert "候选" in ws.cell(1, site_column).comment.text
    shift = next(item for item in wb["测温登记"].data_validations.dataValidation if item.formula1 == '"白班,夜班"')
    assert shift.showErrorMessage is True and shift.errorStyle == "stop"


@pytest.fixture
def api(monkeypatch, tmp_path):
    import bonus_platform.app as app_module
    from bonus_platform.engine.domestic_labor import runs
    monkeypatch.setenv("SIGMA_DOMESTIC_LABOR_STORAGE_BACKEND", "local")
    monkeypatch.setattr(runs, "DOMESTIC_LABOR_RUNS_DIR", tmp_path)
    monkeypatch.setattr(app_module, "DOMESTIC_LABOR_RUNS_DIR", tmp_path)
    return app_module, TestClient(app_module.app), tmp_path


def _workbook_bytes(site, daily_site=""):
    wb = Workbook()
    monthly = wb.active
    monthly.title = "月考勤"
    monthly.append(["工号", "姓名", "工作地区", "岗位名称", "一级部门名称", "三级部门名称", "考勤月份"])
    monthly.append(["T001", "测试", "东莞", "操作员", "寮步区", "华南1号枢纽2", "202607"])
    daily = wb.create_sheet("日考勤")
    daily.append(["出勤日期", "工号", "姓名", "班次时间段", "正班时数", "实际上班时数", "测温网点"])
    daily.append([date(2026, 7, 1), "T001", "测试", "19:00-28:00;", 8, 8, daily_site])
    temperature = wb.create_sheet("测温登记")
    temperature.append(["班次日期", "测温班次", "测温网点", "测温温度"])
    temperature.append([date(2026, 7, 1), "夜班", site, 34.5])
    wb.create_sheet("不享有高温补贴名单").append(["工号", "姓名", "工作地区", "不享有原因", "生效日期", "结束日期"])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def test_multipart_name_gate_then_confirmed_payment_and_audit(api):
    app_module, client, directory = api
    data = {"engines": "gaowen_butie", "attendance_month": "202607", "supplemental_confirmation": json.dumps({"temperature_exclusions": "none"})}
    file = {"file": ("考勤.xlsx", _workbook_bytes("寮步新名称"))}
    response = client.post("/api/domestic-labor/runs", data=data, files=file)
    assert response.status_code == 409, response.text
    assert response.json()["detail"]["code"] == "temperature_site_confirmation_required"
    assert not list(directory.iterdir())
    response = client.post("/api/domestic-labor/runs", data={**data, "temperature_site_mapping": json.dumps({"寮步新名称": LIAOBU_SITE})}, files=file)
    assert response.status_code == 200, response.text
    metadata = app_module.load_payroll_metadata(app_module.get_payroll_run_dir(response.json()["run_id"]))
    assert metadata["results"][0]["gaowen_butie"] == 13.8
    assert metadata["temperatureSiteMapping"] == {"寮步新名称": LIAOBU_SITE}
    assert metadata["inputSummary"]["temperature_rows"] == 1
    audit = metadata["inputSummary"]["temperature_site_matches"]
    assert audit[0]["original"] == "寮步新名称" and audit[0]["method"] == "confirmed"


def test_uploaded_custom_site_outside_candidates_matches_attendance(api):
    app_module, client, directory = api
    response = client.post("/api/domestic-labor/runs", data={"engines": "gaowen_butie", "attendance_month": "202607", "supplemental_confirmation": json.dumps({"temperature_exclusions": "none"})},
                           files={"file": ("考勤.xlsx", _workbook_bytes("新网点A", "新网点 Ａ"))})
    assert response.status_code == 200, response.text
    metadata = app_module.load_payroll_metadata(directory / response.json()["run_id"])
    assert metadata["results"][0]["gaowen_butie"] == 13.8
    assert metadata["inputSummary"]["temperature_site_matches"][0]["original"] == "新网点 Ａ"


def test_direct_name_gate_reuses_same_files_and_clears_confirmation_after_retry(api, monkeypatch):
    app_module, client, directory = api
    monkeypatch.setattr(app_module, "domestic_labor_persistent_storage_enabled", lambda: True)
    content = _workbook_bytes("寮步新名称")
    run = app_module.create_payroll_run({"uploadMode": "direct", "fileNames": ["考勤.xlsx"],
        "expectedFiles": [{"filename": "考勤.xlsx", "originalFilename": "考勤.xlsx", "size": len(content)}]})
    path = directory / run["id"] / "考勤.xlsx"
    path.write_bytes(content)
    payload = {"engines": ["gaowen_butie"], "attendanceMonth": "202607", "supplementalConfirmation": {"temperature_exclusions": "none"}}
    url = f"/api/domestic-labor/runs/{run['id']}/direct-upload-complete"
    first = client.post(url, json=payload)
    assert first.status_code == 409, first.text
    assert app_module.load_payroll_metadata(path.parent)["status"] == "待确认"
    assert path.read_bytes() == content
    second = client.post(url, json={**payload, "temperatureSiteMapping": {"寮步新名称": LIAOBU_SITE}})
    assert second.status_code == 200, second.text
    metadata = app_module.load_payroll_metadata(path.parent)
    assert metadata["status"] == "已完成" and metadata["results"][0]["gaowen_butie"] == 13.8
    assert metadata["confirmationDetail"] is None and metadata["errorCode"] == ""
