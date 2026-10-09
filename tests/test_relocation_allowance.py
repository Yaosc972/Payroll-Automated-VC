from copy import deepcopy

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from bonus_platform.engine.china_employee_payroll.relocation import RelocationError, calculate_relocation
from bonus_platform.engine.china_employee_payroll import relocation_router as routes


@pytest.fixture
def source():
    return {
        "attendance": [{"StaffId": 1, "EmpRecordOID": "service-1", "AttendanceOrg": 10,
                        "Month": "2026-08-01", "TargetDays": 23, "AbsenteeismDuration": 0,
                        "Properties": {"AttendanceItem5": 0, "AttendanceItem6": 0, "AttendanceItem16": 0}}],
        "services": [{"objectId": "service-1", "userID": 1, "jobNumber": "zt00001", "entryDate": "2025-04-01",
                      "startDate": "2025-04-01", "translateProperties": {"OIdJobLevelText": "P2-1", "EmployTypeText": "内部员工"}}],
        "people": [{"userID": 1, "name": "测试员工"}],
        "organizations": [{"oId": 10, "startDate": "2025-01-01", "secondLevelOrganization": 2404320,
                           "thirdLevelOrganization": 3694315, "fourthLevelOrganization": None, "translateProperties": {
            "SecondLevelOrganizationText": "ABU航空事业部", "ThirdLevelOrganizationText": "航线部"}}],
    }


def row(source):
    return calculate_relocation("2026-08", source)["results"][0]


def test_cap_and_hour_conversion(source):
    assert row(source)["amount"] == 1000
    a = source["attendance"][0]
    a.update(TargetDays=20, AbsenteeismDuration=1)
    a["Properties"].update(AttendanceItem5=4, AttendanceItem6=8, AttendanceItem16=2)
    r = row(source)
    assert r["normalDays"] == 15.5
    assert r["amount"] == 775
    assert r["status"] == "已试算"
    a["TargetDays"] = 0
    assert row(source)["amount"] == 0


@pytest.mark.parametrize("change", ["cutoff", "grade", "excluded"])
def test_eligibility(source, change):
    service = source["services"][0]
    if change == "cutoff":
        service["entryDate"] = "2026-05-01"
    elif change == "grade":
        service["translateProperties"]["OIdJobLevelText"] = "P3-1"
    else:
        service["jobNumber"] = "zt18927"
    assert row(source)["amount"] == 0
    assert row(source)["status"] == "不发放"


def test_organization_scope_and_named_inclusion(source):
    unit = source["organizations"][0]
    org = source["organizations"][0]["translateProperties"]
    unit["secondLevelOrganization"] = 815717
    org["SecondLevelOrganizationText"] = "FBU仓储事业部"
    assert calculate_relocation("2026-08", source)["results"] == []
    source["services"][0]["jobNumber"] = "zt25631"
    assert row(source)["amount"] == 1000
    org.update(SecondLevelOrganizationText="ABU航空事业部", ThirdLevelOrganizationText="产品部", FourthLevelOrganizationText="产品二组")
    unit.update(secondLevelOrganization=2404320, thirdLevelOrganization=700388823, fourthLevelOrganization=999)
    source["services"][0]["jobNumber"] = "zt00001"
    assert calculate_relocation("2026-08", source)["results"] == []
    org["FourthLevelOrganizationText"] = "产品一组"
    unit["fourthLevelOrganization"] = 700388843
    assert row(source)["amount"] == 1000


def test_organization_rename_preserves_scope_and_displays_current_names(source):
    unit = source["organizations"][0]
    unit["translateProperties"].update(SecondLevelOrganizationText="新事业部名", ThirdLevelOrganizationText="新部门名")
    result = row(source)
    assert result["amount"] == 1000
    assert result["thirdOrg"] == "新部门名"
    assert result["thirdOrgId"] == 3694315
    assert result["attendanceOrgId"] == 10


@pytest.mark.parametrize("second,third,fourth", [
    (815717, 700566270, None), (3709467, 3713313, None),
    (815744, 3280413, 3654124), (2404320, 999, None)])
def test_same_names_do_not_grant_scope_to_other_units(source, second, third, fourth):
    source["organizations"][0].update(secondLevelOrganization=second, thirdLevelOrganization=third, fourthLevelOrganization=fourth)
    assert calculate_relocation("2026-08", source)["results"] == []


@pytest.mark.parametrize("second,third", [(3362203, 3713424), (3362202, 2957015)])
def test_functional_units_use_ids_and_parent_scope(source, second, third):
    unit = source["organizations"][0]
    unit.update(secondLevelOrganization=second, thirdLevelOrganization=third)
    unit["translateProperties"] = {}
    assert row(source)["amount"] == 1000
    assert row(source)["status"] == "已试算"
    unit["secondLevelOrganization"] = 815717
    assert calculate_relocation("2026-08", source)["results"] == []


def test_missing_ids_do_not_fall_back_to_names(source):
    del source["organizations"][0]["secondLevelOrganization"]
    result = calculate_relocation("2026-08", source)
    assert result["results"] == []
    assert result["summary"]["missingScopeCount"] == 1


def test_numeric_string_org_ids(source):
    source["organizations"][0].update(secondLevelOrganization="2404320", thirdLevelOrganization="3694315")
    assert row(source)["amount"] == 1000


@pytest.mark.parametrize("bad", [None, -1, "NaN", "Infinity", "unknown", True])
def test_missing_or_invalid_attendance_is_not_zero(source, bad):
    source["attendance"][0]["Properties"]["AttendanceItem5"] = bad
    assert row(source)["amount"] is None
    assert row(source)["status"] == "待确认"


def test_departure_is_display_only_and_new_organization_is_reviewed(source):
    later = deepcopy(source["services"][0])
    later.update(objectId="later", lastWorkDate="2026-08-20")
    source["services"].append(later)
    source["organizations"][0]["startDate"] = "2026-09-01"
    r = row(source)
    assert r["amount"] == 1000
    assert r["departureDate"] == "2026-08-20"
    assert "离职" not in r["reason"]
    assert "组织版本" in r["reason"]


def test_missing_scope_is_visible(source):
    source["services"] = []
    result = calculate_relocation("2026-08", source)
    assert result["summary"]["missingScopeCount"] == 1
    assert result["summary"]["amount"] == 0


def test_duplicate_and_wrong_month_fail(source):
    source["attendance"].append(deepcopy(source["attendance"][0]))
    with pytest.raises(RelocationError, match="多条记录"):
        row(source)
    source["attendance"].pop()
    source["attendance"][0]["Month"] = "2026-07-01"
    with pytest.raises(RelocationError, match="月份"):
        row(source)


@pytest.fixture
def client(monkeypatch, tmp_path, source):
    class FakeSource:
        def load(self, month, progress):
            progress("正在读取测试数据")
            return source
    monkeypatch.setattr(routes, "RUNS_DIR", tmp_path)
    monkeypatch.setattr(routes, "BeisenMonthlySource", FakeSource)
    monkeypatch.setattr(routes, "_active", False)
    app = FastAPI()
    app.include_router(routes.router)
    return TestClient(app)


BASE = "/api/china-employee-payroll/relocation-allowance/runs"


def test_start_read_history_and_export(client, source):
    source["people"][0]["name"] = "=HYPERLINK(1)"
    response = client.post(BASE, json={"month": "2026-08"})
    assert response.status_code == 202
    run_id = response.json()["runId"]
    assert "workerId" not in response.json()
    result = client.get(f"{BASE}/{run_id}").json()
    assert result["status"] == "completed"
    assert result["result"]["summary"]["amount"] == 1000
    history = client.get(BASE).json()["runs"]
    assert history[0]["runId"] == run_id
    assert "result" not in history[0]
    exported = client.get(f"{BASE}/{run_id}/export")
    assert exported.status_code == 200
    from io import BytesIO
    from openpyxl import load_workbook
    from urllib.parse import unquote
    wb = load_workbook(BytesIO(exported.content))
    assert wb.sheetnames == ["职场搬迁补贴", "考勤源数据"]
    ws = wb.worksheets[0]
    assert ws["A2"].value == "2026年8月"
    assert ws["C2"].value == "=HYPERLINK(1)" and ws["C2"].data_type == "s"
    assert ws["Q1"].value == "职场搬迁补贴"
    assert ws["Q2"].value == 1000
    assert ws.freeze_panes == "D2"
    assert "口径提醒" not in [c.value for c in ws[1]]
    assert not any("组织ID" in str(c.value) for c in ws[1])
    assert "2026 年 8 月 ABU职场搬迁补贴" in unquote(exported.headers["content-disposition"])
    assert wb.worksheets[1].max_row == 2
    assert "attendanceSource" not in result and "attendanceSource" not in history[0]


@pytest.mark.parametrize("month", ["2026-13", "2026-8", "1999-01", None])
def test_invalid_month(client, month):
    assert client.post(BASE, json={"month": month}).status_code == 400


def test_future_month_can_be_requested_and_empty_report_is_reported(client, monkeypatch):
    class EmptySource:
        def load(self, month, progress):
            assert month == "2099-10"
            raise RelocationError("所选月份没有北森月报，请确认月报已生成。")
    monkeypatch.setattr(routes, "BeisenMonthlySource", EmptySource)
    response = client.post(BASE, json={"month": "2099-10"})
    assert response.status_code == 202
    result = client.get(f"{BASE}/{response.json()['runId']}").json()
    assert result["status"] == "failed"
    assert "没有北森月报" in result["progress"]


def test_busy_rejected(client, monkeypatch):
    monkeypatch.setattr(routes, "_active", True)
    assert client.post(BASE, json={"month": "2026-08"}).status_code == 409


def test_failed_source_does_not_export(client, monkeypatch):
    class BadSource:
        def __init__(self):
            raise RelocationError("读取权限不足")
    monkeypatch.setattr(routes, "BeisenMonthlySource", BadSource)
    response = client.post(BASE, json={"month": "2026-08"})
    run_id = response.json()["runId"]
    result = client.get(f"{BASE}/{run_id}").json()
    assert result["status"] == "failed"
    assert result["progress"] == "读取权限不足"
    assert not routes._active
    assert client.get(f"{BASE}/{run_id}/export").status_code == 409


def test_other_instance_run_and_invalid_id(client):
    run_id = "a" * 32
    routes._save({"runId": run_id, "month": "2026-08", "status": "running", "workerId": "old", "createdAt": "2026-09-30", "progress": "读取中"})
    assert client.get(f"{BASE}/{run_id}").json()["status"] == "running"
    assert client.get(f"{BASE}/invalid").status_code == 404


def test_external_employee_and_missing_employment_type(source):
    s = source["services"][0]["translateProperties"]
    s["EmployTypeText"] = "实习生"
    assert row(source)["amount"] == 0
    assert row(source)["status"] == "不发放"
    del s["EmployTypeText"]
    assert row(source)["status"] == "待确认"


def test_invalid_entry_date_is_reviewed(source):
    source["services"][0]["entryDate"] = "bad-date"
    assert row(source)["status"] == "待确认"
    assert "无效" in row(source)["reason"]


def test_source_recovers_a_timed_out_page_without_duplicates(monkeypatch):
    from bonus_platform.engine.china_employee_payroll.relocation import BeisenMonthlySource, BeisenReadError
    monkeypatch.setenv('BEISEN_APP_KEY', 'test-key')
    monkeypatch.setenv('BEISEN_APP_SECRET', 'test-secret')
    monkeypatch.delenv('BEISEN_BASE_URL', raising=False)
    source = BeisenMonthlySource()
    calls = []
    def request(path, body, authenticate=False):
        if authenticate:
            source.token = 'test-token'
            return {}
        if 'Monthly/GetList' in path:
            index, size = body['PageIndex'], body['PageSize']
            calls.append((index, size))
            if (index, size) == (2, 100):
                raise BeisenReadError('timeout')
            first = (index - 1) * size
            return {'Data': {'Total': 250, 'DailyReportList': [{'StaffId': i + 1, 'AttendanceOrg': 10} for i in range(first, min(first + size, 250))]}}
        return {'data': []}
    monkeypatch.setattr(source, 'request', request)
    loaded = source.load('2026-08')
    assert len(loaded['attendance']) == 250
    assert len({r['StaffId'] for r in loaded['attendance']}) == 250
    assert (3, 50) in calls and (4, 50) in calls
    assert source.token == source.key == source.secret == ''
    assert source.client.is_closed


def test_incomplete_pagination_does_not_return_partial_data(monkeypatch):
    from bonus_platform.engine.china_employee_payroll.relocation import BeisenMonthlySource
    monkeypatch.setenv('BEISEN_APP_KEY', 'test-key')
    monkeypatch.setenv('BEISEN_APP_SECRET', 'test-secret')
    monkeypatch.delenv('BEISEN_BASE_URL', raising=False)
    source = BeisenMonthlySource()
    monkeypatch.setattr(source, 'request', lambda *args: {'Data': {'Total': 10, 'DailyReportList': []}})
    with pytest.raises(RelocationError, match='分页不完整'):
        source.load('2026-08')
    assert source.client.is_closed


@pytest.mark.parametrize('departure', ['2026-08-20', '2026-09-20'])
def test_departure_does_not_create_review(source, departure):
    source['services'][0]['lastWorkDate'] = departure
    assert row(source)['status'] == '已试算'
    assert row(source)['amount'] == 1000
    assert row(source)['departureDate'] == departure


def test_saved_batch_clears_obsolete_reviews_without_changing_amounts(client):
    run_id = 'b' * 32
    result = {'ruleVersion': '2026.09.30-01', 'summary': {'reviewCount': 2, 'amount': 2000},
              'results': [
                  {'status': '待确认', 'reason': '职能条线组织例外待确认；离职特殊口径待确认', 'amount': 1000},
                  {'status': '待确认', 'reason': '核算月后离职，发放口径待确认；组织版本晚于核算月', 'amount': 1000}]}
    routes._save({'runId': run_id, 'status': 'completed', 'result': result})
    updated = client.get(f'{BASE}/{run_id}').json()['result']
    assert updated['summary']['reviewCount'] == 1
    assert updated['summary']['amount'] == 2000
    assert updated['results'][0]['status'] == '已试算'
    assert updated['results'][1]['reason'] == '组织版本晚于核算月'
    assert updated['ruleVersion'] == '2026.09.30-01'


def test_historical_organization_uses_month_end_not_latest(monkeypatch):
    from bonus_platform.engine.china_employee_payroll.relocation import BeisenMonthlySource
    monkeypatch.setenv('BEISEN_APP_KEY', 'test-key')
    monkeypatch.setenv('BEISEN_APP_SECRET', 'test-secret')
    source = BeisenMonthlySource()
    current = {'oId': 10, 'startDate': '2026-09-04', 'stopDate': '9999-12-31', 'thirdLevelOrganization': 999}
    historic = {'oId': 10, 'startDate': '2025-01-01', 'stopDate': '2026-09-03', 'thirdLevelOrganization': 3694315}
    def request(path, body):
        assert path.endswith('GetSubOrganizations')
        assert body['queryDate'] == '2026-08-31'
        assert body['isWithSelf'] is False
        assert body['oId'] == 900109025
        return {'data': [current, historic, {'oId': 11, 'startDate': '2025-01-01', 'stopDate': '9999-12-31'}]}
    monkeypatch.setattr(source, 'request', request)
    try:
        assert source.load_organizations([10], '2026-08-31') == [historic]
    finally:
        source.client.close()


def test_historical_permission_failure_does_not_use_latest(monkeypatch):
    from bonus_platform.engine.china_employee_payroll.relocation import BeisenMonthlySource
    monkeypatch.setenv('BEISEN_APP_KEY', 'test-key')
    monkeypatch.setenv('BEISEN_APP_SECRET', 'test-secret')
    source = BeisenMonthlySource()
    def request(path, body):
        if path.endswith('GetByIds'):
            return {'data': [{'oId': 10, 'startDate': '2026-09-04', 'stopDate': '9999-12-31'}]}
        raise RelocationError('北森接口读取权限不足或授权已失效')
    monkeypatch.setattr(source, 'request', request)
    try:
        with pytest.raises(RelocationError, match='GetSubOrganizations'):
            source.load_organizations([10], '2026-08-31')
    finally:
        source.client.close()


def test_progress_counts_processed_rows_outside_scope(source):
    source["organizations"][0]["secondLevelOrganization"] = 815717
    events = []
    result = calculate_relocation("2026-08", source, events.append)
    assert result["results"] == []
    assert events[0]["completed"] == 0
    assert events[-1] == {"message": "正在生成明细：1/1", "stage": 3, "completed": 1, "total": 1, "unit": "条"}


def test_batch_exposes_real_progress_detail(client):
    response = client.post(BASE, json={"month": "2026-08"})
    run = client.get(f'{BASE}/{response.json()["runId"]}').json()
    assert run["status"] == "completed"
    assert run["progressDetail"] == {"stage": 3, "completed": 1, "total": 1, "unit": "条"}
    assert run["progressUpdatedAt"]


def test_export_snapshot_preserves_raw_units_and_scope(source):
    from bonus_platform.engine.china_employee_payroll.relocation_export import build_workbook
    from io import BytesIO
    from openpyxl import load_workbook
    source['attendance'][0]['Properties']['AttendanceItem6'] = 8
    result = calculate_relocation('2026-08', source)
    assert result['results'][0]['sickDays'] == 1
    snapshot = result.pop('attendanceSource')
    payload, name = build_workbook({'month': '2026-08', 'createdAt': '2026-10-06T11:56:00+08:00',
                                   'result': result, 'attendanceSource': snapshot})
    assert name == '2026 年 8 月 ABU职场搬迁补贴202610061156.xlsx'
    ws = load_workbook(BytesIO(payload)).worksheets[1]
    data = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
    assert data['病假小时数'] == 8
    assert set(data) == {'考勤月份', '工号', '姓名', '二级组织', '三级组织', '四级组织', '应出勤天数', '旷工天数', '事假小时数', '病假小时数', '工作日产假天数'}
    assert not ws.merged_cells.ranges
    assert ws['G2'].number_format == '0'
    assert ws.auto_filter.ref == 'A1:K2'
    source['organizations'][0]['secondLevelOrganization'] = 815717
    assert calculate_relocation('2026-08', source)['attendanceSource'] == []


def test_legacy_export_does_not_invent_source_snapshot(source):
    from bonus_platform.engine.china_employee_payroll.relocation_export import build_workbook
    from io import BytesIO
    from openpyxl import load_workbook
    result = calculate_relocation('2026-08', source)
    result.pop('attendanceSource')
    payload, _ = build_workbook({'month': '2026-08', 'createdAt': '2026-10-06T11:56:00+08:00', 'result': result})
    ws = load_workbook(BytesIO(payload)).worksheets[1]
    assert '未保存原始考勤快照' in ws['A2'].value
    assert ws.max_row == 2


def test_export_fractional_days_have_no_optional_decimal_point():
    from openpyxl import Workbook
    from bonus_platform.engine.china_employee_payroll.relocation_export import style_sheet
    ws = Workbook().active
    style_sheet(ws, ['天数'], [[0], [22.0], [0.5], [1.125]], [18])
    assert [ws.cell(r, 1).number_format for r in range(2, 6)] == ['0', '0', '0.0', '0.000']
    assert [ws.cell(r, 1).value for r in range(2, 6)] == [0, 22, 0.5, 1.125]
