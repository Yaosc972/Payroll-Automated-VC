from copy import deepcopy
from io import BytesIO
from openpyxl import load_workbook
from fastapi import FastAPI
from fastapi.testclient import TestClient
from bonus_platform.engine import recruitment_work_items as work
from bonus_platform.engine.recruitment_attention import project
from bonus_platform.engine.recruitment_beisen import fetch_employment, IDS, HISTORY
from bonus_platform.engine.recruitment_drafts import DraftStore
from bonus_platform.engine.recruitment_reviews import prepare_preview, apply_preview
from test_recruitment_attention import state_for, fact_task
from test_recruitment_awards import person
from test_recruitment_drafts import book
from test_recruitment_beisen import snapshot


def test_rehire_consistent_rule_is_automatic_while_real_gaps_remain_work_items():
    from bonus_platform.engine.recruitment_awards import ensure_calculated
    from bonus_platform.engine.recruitment_monthly import project as monthly_project
    ready=person(招聘渠道='二次入职',上次离职日期='2024-09-06');ready['id']='ready'
    conflict=person(招聘渠道='二次入职',上次离职日期='2026-01-01');conflict['id']='conflict'
    missing=person(招聘渠道='二次入职');missing['id']='missing'
    state=state_for([ready,conflict,missing],[])
    original_values=[deepcopy(r['values']) for r in state['rows']]
    state['calculation']['engineVersion']='4.8'
    state['todos'].append({**fact_task(ready,'calculation','招聘奖金：二次入职专项条款执行依据待核'),'id':'old-rule-task'})
    ensure_calculated(state)
    monthly=monthly_project(state,project(state))
    assert not any('ready' in item['rowIds'] for item in monthly['workItems'])
    assert {rid for item in monthly['workItems'] for rid in item['rowIds']}=={'missing'}
    assert next(item for item in monthly['workItems'] if item['rowIds']==['missing'])['mode']=='facts'
    assert any('二次入职' in item['title'] and '40%' in item['detail'] for item in monthly['automaticResults'])
    assert [r['values'] for r in state['rows']]==original_values and state['revision']==1


def test_missing_offer_work_is_grouped_by_region_not_job_family():
    rows=[]
    for index,job in enumerate(['操作类','操作类','操作管理类','操作支持类']):
        row=person(hire='2026-07-01',职类=job,招聘启动日期='');row['id']=str(index)
        row['scopeStatus']='操作类范围待核' if index<2 else '已纳入底稿'
        row['originKinds']=['roster'];row['offerMatched']=False
        rows.append(row)
    state=state_for(rows,[fact_task(r,'missing_offer','待补录用') for r in rows]);before=deepcopy(state)
    items=work.project(state,project(state))
    assert not any(i['key']=='operator-scope' for i in items)
    intake=next(i for i in items if i['key']=='intake:海外')
    assert intake['rowIds']==['0','1','2','3']
    assert intake['mode']=='intake'
    assert state==before  # Neither grouping nor opening is an approval.


def test_unknown_currency_requests_a_country_instead_of_an_unrelated_standard_type():
    row=person(工作地='非中国',二级组织='',三级组织='',四级组织='',五级组织='',奖金地区类型='',推荐人工号='',推荐人姓名='')
    state=state_for([row],[])
    items=work.project(state,project(state))
    item=next(i for i in items if i['key']=='country')
    assert item['title']=='核对实际工作国家/地区' and not item['ruleMaintenance']
    data=work.detail(state,item)
    assert [f['field'] for f in data['fields']]==['工作地']
    field=data['fields'][0]
    assert field['label']=='实际工作国家/地区'
    assert '德国' in field['placeholder'] and '币种' in field['help']
    assert '非中国' in data['people'][0]['explanation']
    row['values']['工作地']='德国'
    from bonus_platform.engine.recruitment_awards import calculate_draft
    calculate_draft(state)
    assert row['calculation']['recruitment']['currency']=='欧元'
    assert not any(i['key']=='country' for i in work.project(state,project(state)))


def test_recipient_fields_are_shared_across_roles_but_not_distinct_employments():
    first=person();first['id']='a';first['values'].update(招聘负责人工号='zt-shared',招聘负责人入职日期='',招聘负责人人员状态='',招聘负责人姓名='同一领取人')
    second=deepcopy(first);second['id']='b';second['values'].update(招聘负责人工号='zt-other',招聘负责人入职日期='2020-01-01',协助招聘人工号='zt-shared',协助招聘人入职日期='',协助招聘人人员状态='',协助招聘人姓名='同一领取人')
    state=state_for([first,second],[]);item=next(i for i in work.project(state,project(state)) if i['key']=='recipient')
    fields=work.detail(state,item)['fields']
    combined=next(f for f in fields if len(f['bindings'])==2)
    assert {b['field'] for b in combined['bindings']}=={'招聘负责人人员状态','协助招聘人人员状态'}


def test_missing_start_date_does_not_reask_known_recruitment_relationships():
    row=person(hire='2026-07-01',招聘启动日期='')
    row['offerMatched']=False
    state=state_for([row],[])
    item=next(i for i in work.project(state,project(state)) if i['mode']=='intake')
    fields={f['field'] for f in work.detail(state,item)['fields']}
    assert '招聘启动日期' in fields
    assert not fields & {'核算范围','招聘渠道','ABC类别','招聘负责人工号'}


def test_unmatched_roster_exposes_real_gaps_and_checked_offer_sources():
    row=person(招聘渠道='',ABC类别='',招聘启动日期='',招聘负责人工号='',招聘负责人姓名='',推荐人工号='',推荐人姓名='')
    row.update(originKinds=['roster'],offerMatched=False,sources=[])
    state=state_for([row],[fact_task(row,'missing_offer','海外人员待补Offer')])
    state['sources']=[{'id':'offer','name':'海外录用.xlsx','roles':['overseas_offer'],'active':True},
                      {'id':'unused','name':'未采用.xlsx','roles':['overseas_offer'],'active':False}]
    before=deepcopy(state)
    item=next(i for i in work.project(state,project(state)) if i['mode']=='intake')
    data=work.detail(state,item)
    assert {f['field'] for f in data['fields']}=={'招聘渠道','ABC类别','招聘启动日期','招聘负责人工号'}
    assert data['people'][0]['offerStatus']=='未匹配到本次入职的录用记录'
    assert set(data['people'][0]['missingFields'])=={'招聘渠道','ABC类别','招聘启动日期','招聘负责人工号'}
    assert data['checkedSources']==[{'id':'offer','name':'海外录用.xlsx'}]
    assert state==before


def test_unmatched_offer_does_not_hide_actual_employment_difference():
    row=person();row.update(originKinds=['roster'],offerMatched=False)
    state=state_for([row],[fact_task(row,'employment_conflict','入职日有差异')])
    assert any(t['kind']=='employment' for t in project(state)['current'])


def test_complete_roster_facts_do_not_require_offer_or_scope_confirmation():
    row=person();row.update(originKinds=['roster'],offerMatched=False)
    state=state_for([row],[fact_task(row,'missing_offer','尚未匹配Offer')])
    assert not work.project(state,project(state))


def test_missing_offer_known_recruiter_name_is_not_reasked_as_required_id():
    row=person(招聘负责人工号='',招聘启动日期='');row.update(originKinds=['roster'],offerMatched=False)
    state=state_for([row],[fact_task(row,'missing_offer','尚未匹配Offer')])
    item=next(i for i in work.project(state,project(state)) if i['mode']=='intake')
    assert {f['field'] for f in work.detail(state,item)['fields']}=={'招聘启动日期'}


def test_known_recruiter_name_and_complete_facts_do_not_leave_an_empty_intake():
    row=person(招聘负责人工号='');row.update(originKinds=['roster'],offerMatched=False)
    state=state_for([row],[fact_task(row,'missing_offer','尚未匹配Offer'),
                          fact_task(row,'missing_fields','补充核算所需信息：招聘负责人工号')])
    assert not work.project(state,project(state))


def test_supplement_roundtrip_matches_original_employment_and_clears_real_gaps(tmp_path):
    store=DraftStore(tmp_path)
    state=store.ingest(202607,[('花名册.xlsx',book(['姓名','工号','入职日期','人员状态','人员来源','职级','二级组织','offer流程编号'],
        [['甲','zt1','2026-07-01','正式','OEHR入职','P1-1','德国','OF1'],['乙','zt2','2026-07-15','正式','OEHR入职','P1-1','德国','OF2']]))],0)
    before=deepcopy(state);ids={r['id'] for r in state['rows']}
    item=next(i for i in work.project(state,project(state)) if i['mode']=='intake')
    data=work.detail(state,item)
    template=load_workbook(BytesIO(work.supplement_workbook(state,data)))
    assert template.active.max_row==2 and template.active.title=='导入_月度数据'
    assert load_workbook(BytesIO(work.supplement_workbook(state,data,True))).active.max_row==3
    assert state==before
    headers=[c.value for c in template.active[1]]
    assert '核算范围' not in headers and '职级' not in headers
    for f,value in {'ABC类别':'C','招聘渠道':'招聘网站','招聘启动日期':'2026-05-20','招聘负责人工号':'zt-owner'}.items():
        template.active.cell(2,headers.index(f)+1,value)
    output=BytesIO();template.save(output)
    result=store.ingest(202607,[('录用信息补充表.xlsx',output.getvalue())],state['revision'])
    assert len(result['rows'])==2 and {r['id'] for r in result['rows']}==ids
    row=next(r for r in result['rows'] if r['values']['工号']=='zt1')
    assert row['calculation']['recruitment']['total']==55
    assert not any(t['kind'] in {'intake_scope','field:招聘渠道','standard','cycle','field:招聘负责人工号'} and row['id'] in t['rowIds'] for t in project(result)['current'])
    assert not row['values'].get('核算范围')


def test_targeted_lookup_does_not_scan_offers_or_return_credentials():
    class Client:
        def __init__(self):self.paths=[]
        def post(self,path,body):
            self.paths.append(path)
            if path==IDS:return {'data':[{'code':200,'userId':1}]}
            assert path==HISTORY
            return {'data':[{'jobNumber':'zt1','approvalStatus':4,'userID':1,'secret':'not retained'}]}
    client=Client();result=fetch_employment(['ZT1','zt1'],client)
    assert client.paths==[IDS,HISTORY]
    assert result['employmentOnly'] and result['scopeRecipientJobNumbers']==['zt1']
    assert 'secret' not in result['recipientHistories'][0]


def test_employment_gap_snapshot_preserves_original_offer_source(tmp_path):
    store=DraftStore(tmp_path)
    state=store.save_beisen_snapshot(202607,0,snapshot())
    source=state['sources'][-1]['id']
    supplemental={'version':1,'employmentOnly':True,'queriedAt':'2026-09-22T10:00:00','scopeJobNumbers':[],
                  'scopeRecipientJobNumbers':['zt9'],'candidateHistories':[],'recipientHistories':[],
                  'offers':[],'requirements':[],'forms':[]}
    state=store.save_beisen_snapshot(202607,state['revision'],supplemental)
    assert next(s for s in state['sources'] if s['id']==source)['active']
    assert state['sources'][-1]['employmentOnly']


def test_explicit_batch_scope_edit_is_preserved_and_stale_rejected(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    store=DraftStore(tmp_path);monkeypatch.setattr(api,'store',store)
    state=store.ingest(202607,[('roster.xlsx',book(['姓名','工号','入职日期','人员状态','人员来源','职类'],
          [['甲','zt1','2026-07-01','正式','EHR入职','操作类'],['乙','zt2','2026-07-02','正式','EHR入职','操作类']]))],0)
    items=work.project(state,project(state));item=next(i for i in items if i['mode']=='intake')
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    assert client.get(f"/api/recruitment/drafts/202607/work-items/{item['id']}?revision=0").status_code==409
    assert client.get(f"/api/recruitment/drafts/202607/work-items/{item['id']}/supplement.xlsx?revision=0").status_code==409
    data=client.get(f"/api/recruitment/drafts/202607/work-items/{item['id']}?revision={state['revision']}").json()
    assert len(data['people'])==2
    download=client.get(f"/api/recruitment/drafts/202607/work-items/{item['id']}/supplement.xlsx?revision={state['revision']}")
    assert download.status_code==200
    assert load_workbook(BytesIO(download.content)).active.max_row==2
    edits=[{'rowId':r['id'],'field':'核算范围','value':'不纳入核算','reason':'测试明确排除两条重复任职'} for r in state['rows']]
    preview=prepare_preview(store,202607,state['revision'],edits)
    first_id=next(r['id'] for r in state['rows'] if r['values']['工号']=='zt1')
    assert {i['rowId'] for i in preview['currentImpact']}=={first_id}
    assert store.read(202607)['revision']==state['revision']
    result=apply_preview(store,202607,state['revision'],preview['id'])
    assert len(result['operations'])==1 and len(result['edits'])==2
    assert all(r['values']['核算范围']=='不纳入核算' for r in result['rows'])
    assert not any(i['key']=='operator-scope' for i in work.project(result,project(result)))
