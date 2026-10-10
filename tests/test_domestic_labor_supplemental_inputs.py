from datetime import date
from io import BytesIO
import json

import pytest
from openpyxl import load_workbook

from test_domestic_labor_temperature_sites import api, _workbook_bytes
from bonus_platform.engine.domestic_labor.engines.gaowen_butie import LIAOBU_SITE, JIASHAN_SITE
from bonus_platform.engine.domestic_labor.exporter import ExcelExporter
from bonus_platform.engine.domestic_labor.templates import generate_template

pytestmark = pytest.mark.usefixtures("bypass_domestic_labor_access_gate")


def _supplemented_book(region="东莞", exclusions=None, housing=None):
    site = LIAOBU_SITE if region == "东莞" else JIASHAN_SITE
    wb = load_workbook(BytesIO(_workbook_bytes(site)))
    wb["月考勤"]["C2"] = region
    for day in (2, 3):
        wb["日考勤"].append([date(2026, 7, day), "T001", "测试", "19:00-28:00;", 8, 8, ""])
        wb["测温登记"].append([date(2026, 7, day), "夜班", site, 34.5])
    if exclusions is None:
        exclusions = [["T001", "测试", region, "办公地点有空调", date(2026, 7, 2), date(2026, 7, 2)]]
    for record in exclusions:
        wb["不享有高温补贴名单"].append(record)
    if housing is not None:
        ws = wb.create_sheet("住宿名单")
        ws.append(["工号", "姓名", "入住时间", "退宿时间"])
        for record in housing:
            ws.append(record)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def _post(client, content, subject="gaowen_butie", confirmation=None):
    data = {"engines": subject, "attendance_month": "202607"}
    if confirmation is not None:
        data["supplemental_confirmation"] = json.dumps(confirmation)
    return client.post("/api/domestic-labor/runs", data=data, files={"file": ("核算输入.xlsx", content)})


@pytest.mark.parametrize("region,expected", [("东莞", 27.6), ("嘉善", 18.4)])
def test_period_exclusion_applies_only_to_inclusive_dates_and_is_auditable(api, region, expected):
    module, client, directory = api
    response = _post(client, _supplemented_book(region))
    assert response.status_code == 200, response.text
    metadata = module.load_payroll_metadata(directory / response.json()["run_id"])
    result = metadata["results"][0]
    assert result["gaowen_butie"] == expected
    detail = result["subject_details"]["gaowen_butie"]["details"]
    assert [row["reason_code"] for row in detail["daily_results"]] == ["calculated", "supplemental_period_excluded", "calculated"]
    day = detail["daily_results"][1]
    assert day["exclusion_reason"] == "办公地点有空调"
    assert day["exclusion_start"] == day["exclusion_end"] == "2026-07-02"
    assert "2026-07-02" in ExcelExporter._high_temperature_process(result, day)


def test_open_ended_exclusion_and_name_collision_match_by_id(api):
    module, client, directory = api
    content = _supplemented_book(exclusions=[["T001", "测试", "东莞", "办公地点有空调", date(2026, 7, 2), None]])
    wb = load_workbook(BytesIO(content))
    monthly = list(wb["月考勤"].values)[1]
    wb["月考勤"].append(["T002", *monthly[1:]])
    for row in list(wb["日考勤"].values)[1:]:
        clone = list(row); clone[1] = "T002"; wb["日考勤"].append(clone)
    output = BytesIO(); wb.save(output)
    response = _post(client, output.getvalue())
    assert response.status_code == 200, response.text
    rows = module.load_payroll_metadata(directory / response.json()["run_id"])["results"]
    assert {row["employee_id"]: row["gaowen_butie"] for row in rows} == {"T001": 13.8, "T002": 41.4}


@pytest.mark.parametrize("record,match", [
    (["MISSING", "测试", "东莞", "办公地点有空调", "2026-07-01", None], "工号未匹配"),
    (["T001", "另一姓名", "东莞", "办公地点有空调", "2026-07-01", None], "姓名"),
    (["T001", "测试", "嘉善", "办公地点有空调", "2026-07-01", None], "工作地区"),
    (["T001", "测试", "东莞", "出库发运组过渡性高温补贴", "2026-07-01", None], "已确认规则"),
    (["T001", "测试", "东莞", "固定排除名单", "2026-07-01", None], "不符合"),
    (["T001", "测试", "东莞", "办公地点有空调", None, None], "日期无效"),
    (["T001", "测试", "东莞", "办公地点有空调", "2026-07-02", "2026-07-01"], "早于"),
])
def test_invalid_exclusion_is_rejected_before_calculation(api, record, match):
    _, client, directory = api
    response = _post(client, _supplemented_book(exclusions=[record]))
    assert response.status_code == 400 and match in response.text
    assert not list(directory.iterdir())


def test_overlapping_exclusions_are_rejected(api):
    _, client, _ = api
    row = ["T001", "测试", "东莞", "办公地点有空调", "2026-07-01", "2026-07-02"]
    response = _post(client, _supplemented_book(exclusions=[row, [*row[:4], "2026-07-02", "2026-07-03"]]))
    assert response.status_code == 400 and "重叠" in response.text


def test_empty_exclusions_require_confirmation_but_missing_sheet_is_not_empty(api):
    _, client, _ = api
    content = _supplemented_book(exclusions=[])
    response = _post(client, content)
    assert response.status_code == 409
    assert response.json()["detail"]["sheets"][0]["key"] == "temperature_exclusions"
    assert _post(client, content, confirmation={"temperature_exclusions": "none"}).status_code == 200
    wb = load_workbook(BytesIO(content)); del wb["不享有高温补贴名单"]
    output = BytesIO(); wb.save(output)
    response = _post(client, output.getvalue(), confirmation={"temperature_exclusions": "none"})
    assert response.status_code == 400 and "缺少不享有" in response.text


def test_direct_empty_confirmation_reuses_uploaded_files(api, monkeypatch):
    module, client, directory = api
    monkeypatch.setattr(module, "domestic_labor_persistent_storage_enabled", lambda: True)
    content = _supplemented_book(exclusions=[])
    run = module.create_payroll_run({"uploadMode": "direct", "fileNames": ["考勤.xlsx"], "expectedFiles": [{"filename": "考勤.xlsx", "size": len(content)}]})
    path = directory / run["id"] / "考勤.xlsx"; path.write_bytes(content)
    url = f"/api/domestic-labor/runs/{run['id']}/direct-upload-complete"
    payload = {"engines": ["gaowen_butie"], "attendanceMonth": "202607"}
    first = client.post(url, json=payload)
    assert first.status_code == 409 and module.load_payroll_metadata(path.parent)["status"] == "待确认"
    response = client.post(url, json={**payload, "supplementalConfirmation": {"temperature_exclusions": "none"}})
    assert response.status_code == 200 and path.read_bytes() == content


@pytest.mark.parametrize("records,match", [
    ([["MISSING", "测试", "2026-06-01", None]], "工号未匹配"),
    ([["T001", "测试", "2026-07-02", "2026-07-01"]], "早于"),
    ([["T001", "测试", "2026-06-01", None], ["T001", "测试", "2026-07-01", None]], "重叠"),
])
def test_housing_invalid_ids_and_dates_are_not_silently_ignored(api, records, match):
    _, client, _ = api
    response = _post(client, _supplemented_book(housing=records), subject="waisu_butie")
    assert response.status_code == 400 and match in response.text


def test_housing_cross_month_multiple_stays_and_empty_confirmation(api):
    module, client, directory = api
    records = [["T001", "测试", "2026-06-20", "2026-07-01"], ["T001", "测试", "2026-07-02", None]]
    response = _post(client, _supplemented_book(housing=records), subject="waisu_butie")
    assert response.status_code == 200, response.text
    summary = response.json()["input_summary"]
    assert summary["housing_rows"] == 2
    content = _supplemented_book(housing=[])
    assert _post(client, content, subject="waisu_butie").status_code == 409
    response = _post(client, content, subject="waisu_butie", confirmation={"housing": "none"})
    assert response.status_code == 200
    assert response.json()["input_summary"]["supplemental_confirmation"] == {"housing": "none"}


def test_measurement_invalid_date_and_non_finite_temperature_are_rejected(api):
    _, client, _ = api
    for value, column in [("bad-date", "A2"), ("nan", "D2")]:
        wb = load_workbook(BytesIO(_supplemented_book())); wb["测温登记"][column] = value
        output = BytesIO(); wb.save(output)
        assert _post(client, output.getvalue()).status_code == 400


def test_downloaded_templates_have_requested_sheets_and_month_metadata(api):
    _, client, _ = api
    for subject, sheets in [("gaowen_butie", ["测温登记", "不享有高温补贴名单", "填写说明"]), ("waisu_butie", ["住宿名单", "填写说明"])]:
        response = client.get(f"/api/domestic-labor/templates/{subject}/download?attendance_month=2026-07")
        wb = load_workbook(BytesIO(response.content))
        assert wb.sheetnames == sheets and wb["填写说明"]["B3"].value == "202607"
        assert all(cell.value is None for sheet in wb if sheet.title != "填写说明" for row in sheet.iter_rows(min_row=2) for cell in row)
