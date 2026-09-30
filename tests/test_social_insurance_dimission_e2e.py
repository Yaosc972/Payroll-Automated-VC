"""Real connector/rules/HTTP/persistence/export; only Beisen responses are synthetic."""
from io import BytesIO
import socket
import subprocess
import time
import zipfile
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bonus_platform.app import app
from bonus_platform.engine.social_insurance.template_schemas import TEMPLATE_SCHEMAS
from tests.test_social_insurance_report_package import _template_bytes


def test_pending_dimission_sync_review_and_export(tmp_path, monkeypatch):
    for name in ("RUNS", "BASELINES", "SNAPSHOTS", "RELEASES", "TEMPLATE_LIBRARY"):
        monkeypatch.setenv(f"SIGMA_SOCIAL_INSURANCE_{name}_DIR", str(tmp_path / name))
    monkeypatch.setenv("SIGMA_SOCIAL_INSURANCE_STORAGE_BACKEND", "local")
    monkeypatch.setenv("SIGMA_LABOR_AUTH_REQUIRED", "0")
    monkeypatch.delenv("SIGMA_SOCIAL_INSURANCE_TEMPLATE_FILE", raising=False)
    monkeypatch.delenv("VERCEL", raising=False)
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]
    root = Path(__file__).resolve().parents[1]
    process = subprocess.Popen(["node", str(root / "social_insurance_connector/test/fixtures/dimission-server.mjs"), str(port)],
                               stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    monkeypatch.setenv("SIGMA_SOCIAL_INSURANCE_CONNECTOR_URL", f"http://127.0.0.1:{port}")
    monkeypatch.setenv("SIGMA_SOCIAL_INSURANCE_CONNECTOR_TOKEN", "fixture")
    try:
        for _ in range(100):
            try:
                with socket.create_connection(("127.0.0.1", port), timeout=.1):
                    break
            except OSError:
                time.sleep(.05)
        client = TestClient(app)
        response = client.post("/api/social-insurance/runs/sync", json={
            "periodStart":"2026-08-16", "periodEnd":"2026-09-15",
            "confirmationDate":"2026-09-16", "subject":"SZ001", "forceRefresh":True})
        assert response.status_code == 200, response.text
        run = response.json()
        employees = {e["report"]["姓名"]:e for e in run["employees"]}
        assert employees["审批中离职测试"]["status"] == "needs_review"
        for name in ("正常入职测试", "重新入职测试", "撤回离职测试", "15日前非自愿测试"):
            assert employees[name]["status"] == "ready"
        for name in ("15日前自愿测试", "BP指定购买测试"):
            assert employees[name]["status"] == "excluded"
        prefix = f'/api/social-insurance/runs/{run["id"]}'
        assert client.post(prefix + "/confirm").status_code == 409
        assert client.post(prefix + "/generate-package").status_code == 409
        pending = employees["审批中离职测试"]
        result = client.patch(prefix + f'/employees/{pending["id"]}', json={
            "decision":"exclude", "reviewNote":"测试：业务确认本月不购买", "confirmed":True})
        assert result.status_code == 200, result.text
        # BP exceptions use the existing explicit business override, not an
        # invented upstream field or an automatic assumption about BP intent.
        bp_employee = employees["BP指定购买测试"]
        result = client.patch(prefix + f'/employees/{bp_employee["id"]}', json={
            "decision":"include", "reviewNote":"测试：BP明确指定本月购买", "confirmed":True})
        assert result.status_code == 200, result.text
        assert client.post(prefix + "/confirm").status_code == 200
        route = "shenzhen-social-medical"
        result = client.post(prefix + "/template", data={"route":route},
            files={"file":("深圳测试模板.xlsx", _template_bytes(route), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert result.status_code == 200, result.text
        result = client.post(prefix + "/generate-package")
        assert result.status_code == 200, result.text
        downloaded = client.get(prefix + "/package/download")
        assert downloaded.status_code == 200
        with zipfile.ZipFile(BytesIO(downloaded.content)) as archive:
            filename = next(n for n in archive.namelist() if n.startswith("报盘文件/") and n.endswith(".xlsx"))
            workbook = load_workbook(BytesIO(archive.read(filename)), read_only=True)
            sheet = workbook[TEMPLATE_SCHEMAS[route]["sheet"]]
            names = [row[1] for row in sheet.iter_rows(min_row=TEMPLATE_SCHEMAS[route]["dataStartRow"],values_only=True)]
            assert set(names) == {"正常入职测试", "重新入职测试", "撤回离职测试", "15日前非自愿测试", "BP指定购买测试"}
            workbook.close()
    finally:
        process.terminate()
        process.communicate(timeout=5)
