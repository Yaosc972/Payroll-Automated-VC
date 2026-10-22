from io import BytesIO

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bonus_platform.engine import recruitment_reviews as reviews
from bonus_platform.engine.recruitment_drafts import DraftStore
from bonus_platform.engine.recruitment_worksheet import export_results
from test_recruitment_assignments import setup


def test_manual_batch_keeps_audit_and_undo_without_note_or_source(tmp_path):
    store,state=setup(tmp_path)
    edits=[{'rowId':row['id'],'field':'职级','value':'P2-1'} for row in state['rows']]
    preview=reviews.prepare_preview(store,202607,state['revision'],edits)
    assert store.read(202607)==state
    saved=reviews.apply_preview(store,202607,state['revision'],preview['id'])
    for old,row,edit in zip(state['rows'],saved['rows'],saved['edits']):
        assert row['values']['职级']=='P2-1'
        assert row['original']==old['original']
        assert row['provenance']==old['provenance']
        assert edit['before']==old['values']['职级']
        assert edit['reason']=='手工修改' and edit['at']
        assert edit['evidenceId']=='' and edit['evidenceRefs']==[]
    undone=store.history_action(202607,saved['revision'],'undo')
    assert [r['values']['职级'] for r in undone['rows']]==[r['values']['职级'] for r in state['rows']]
    redone=DraftStore(tmp_path).history_action(202607,undone['revision'],'redo')
    assert all(r['values']['职级']=='P2-1' for r in redone['rows'])
    assert all(e['active'] for e in redone['edits'])


def test_source_adoption_records_real_location_but_manual_value_does_not(tmp_path):
    store,state=setup(tmp_path);row=state['rows'][0];sid=state['sources'][0]['id']
    edit={'rowId':row['id'],'field':'职级','value':row['values']['职级'],'evidenceId':sid,
          'evidenceRefs':[{'sourceId':sid,'sheet':'伪造位置','row':999}]}
    saved=store.change(202607,state['revision'],{'edits':[edit]})
    actual=saved['edits'][-1]
    assert actual['reason']=='采用原始资料'
    assert actual['evidenceRefs']==row['provenance']['职级']
    assert actual['evidenceRefs'][0]['row']==2
    workbook=load_workbook(BytesIO(export_results(saved)),read_only=True)
    try:
        audit=list(workbook['调整及依据'].values)
        assert audit[0][-1]=='来源位置'
        assert '第 2 行' in audit[1][-1] and '999' not in audit[1][-1]
    finally:workbook.close()
    # Merely linking a file does not make an arbitrary new value a source value.
    saved=store.change(202607,saved['revision'],{'edits':[{**edit,'value':'P4-1'}]})
    assert saved['edits'][-1]['evidenceRefs']==[]
    assert saved['edits'][-1]['reason']=='手工修改'
    saved=store.change(202607,saved['revision'],{'edits':[{'rowId':row['id'],'field':'职级','value':'P2-1','reason':'线下已核实'}]})
    assert saved['edits'][-1]['reason']=='线下已核实'
    assert saved['edits'][-1]['evidenceId']==''


def test_optional_source_still_must_be_valid(tmp_path):
    store,state=setup(tmp_path)
    edit={'rowId':state['rows'][0]['id'],'field':'职级','value':'P2-1','evidenceId':'missing'}
    with pytest.raises(ValueError,match='依据文件不存在'):
        store.change(202607,state['revision'],{'edits':[edit]})
    state['sources'][0]['comparisonOnly']=True
    store._write(store._folder(202607)/'draft.json',state)
    edit['evidenceId']=state['sources'][0]['id']
    with pytest.raises(ValueError,match='不能作为修改依据'):
        store.change(202607,state['revision'],{'edits':[edit]})


def test_excel_return_without_basis_previews_and_saves(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    store,state=setup(tmp_path);monkeypatch.setattr(api,'store',store)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    wb=load_workbook(BytesIO(reviews.editable_workbook(state)))
    ws=wb.active;column=[c.value for c in ws[1]].index('职级')+1
    for index in range(2,4):ws.cell(index,column,'P2-1')
    data=BytesIO();wb.save(data);wb.close()
    response=client.post('/api/recruitment/drafts/202607/worksheet/preview',
                         data={'revision':state['revision']},files={'file':('回传.xlsx',data.getvalue())})
    assert response.status_code==200,response.text
    preview=response.json();assert len(preview['changes'])==2
    assert store.read(202607)==state
    response=client.post('/api/recruitment/drafts/202607/previews/'+preview['id']+'/apply',json={'revision':state['revision']})
    assert response.status_code==200,response.text
    saved=store.read(202607)
    assert all(e['reason']=='回传表格核对更正' and not e['evidenceId'] for e in saved['edits'])
    assert all(r['values']['职级']=='P2-1' for r in saved['rows'])


@pytest.mark.parametrize('kind',['arrangements','decisions'])
def test_special_changes_still_require_approval(tmp_path,kind):
    store,state=setup(tmp_path);rid=state['rows'][0]['id']
    payload={'arrangements':[{'rowId':rid,'recipientId':'zt-old','recipientName':'原负责人','amount':10,'currency':'人民币'}]} if kind=='arrangements' else {'decisions':[{'rowId':rid,'kind':'recruitment','action':'exclude'}]}
    with pytest.raises(ValueError,match='审批'):
        reviews.prepare_preview(store,202607,state['revision'],[],**payload)
    assert store.read(202607)==state
