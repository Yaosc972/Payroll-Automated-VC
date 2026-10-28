from copy import deepcopy
from io import BytesIO
import json

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bonus_platform.engine.recruitment_drafts import DraftStore
from bonus_platform.engine import recruitment_external_results as external
from test_recruitment_awards import person
from test_recruitment_drafts import book


FIELDS=['姓名','工号','入职日期','转正日期','最后工作日','标签分类','工作地','ABC类别','招聘负责人工号']


def opening(store, rows=None):
    rows=rows or [['历史甲','outside-01','2026-06-01','','','国内','中国大陆','B','zt-owner']]
    return store.ingest(202609,[('8月期初.xlsx',book(FIELDS,rows,title='附件一-招聘奖金核算明细'))],0)


def upload(store, state, rows, *, fields=FIELDS, filename='BP累计人员.xlsx', title='BP维护的全年人员资料'):
    return store.ingest(state['month'],[(filename,book(fields,rows,title=title))],state['revision'],purpose='external-personnel')


def people(state):
    return {r['values']['姓名']:r for r in state['rows']}


def test_bp_original_workbook_reads_personnel_detail_not_payee_summary_or_award_amounts(tmp_path):
    from bonus_platform.engine.recruitment_draft_sources import read_source
    fields=FIELDS+['候选人入职时间','奖金对应币种']
    wb=load_workbook(BytesIO(book(fields,[
        ['历史甲','outside-01','2026-06-01','2026-09-02','','国内','中国大陆','A','zt-other','2026-06-01','人民币'],
        ['本月新增','outside-02','2026-09-03','','','国内','中国大陆','C','zt-owner','2026-09-03','人民币'],
    ],title='附件一-招聘奖金核算明细')))
    detail=wb.active;detail.insert_rows(1);detail.cell(1,2,'花名册');detail.insert_rows(3)
    detail['AM3']='入职1个月发放奖金';detail['AN3']='入职1个月发放周期'
    detail['AM4']=99999;detail['AN4']=202609
    detail['AM5']=88888;detail['AN5']=202609
    summary=wb.create_sheet('附件二-招聘奖金汇总表')
    summary.append(['工号','姓名','币种',202608]);summary.append(['zt-owner','领取人甲','美元/欧元',77777])
    data=BytesIO();wb.save(data);content=data.getvalue()
    parsed=read_source(tmp_path/'BP原表.xlsx','bp-original',content=content,personnel_only=True)
    assert [(s['role'],s['recordCount']) for s in parsed['sheets']]==[('cumulative',2),('unknown',0)]
    assert all(not r['historyCells'] for r in parsed['records'])
    store=DraftStore(tmp_path/'drafts');before=opening(store);previous=deepcopy(before['rows'][0])
    after=store.ingest(202609,[('BP原表.xlsx',content)],before['revision'],purpose='external-personnel')
    rows=people(after)
    assert set(rows)=={'历史甲','本月新增'}
    assert rows['历史甲']['values']['转正日期']=='2026-09-02'
    assert rows['历史甲']['values']['ABC类别']==previous['values']['ABC类别']
    assert rows['历史甲']['history']==previous['history']
    assert not rows['本月新增']['history'] and not after.get('externalResults')
    assert after['sources'][-1]['cumulativeSummary']['notApplied']==[]
    assert after['sources'][-1]['cumulativeSummary']['newHires']==1


def test_cumulative_adds_only_current_hires_and_observes_historical_facts(tmp_path):
    store=DraftStore(tmp_path);state=opening(store);old=deepcopy(state['rows'][0])
    state=upload(store,state,[
        ['历史甲','outside-01','2026-06-01','2026-09-01','2026-09-20','海外','英国','A','zt-other'],
        ['本月新增','outside-02','2026-09-03','','','国内','中国大陆','C','zt-owner'],
        ['过往未纳入','outside-03','2026-05-01','','','国内','中国大陆','C','zt-owner'],
        ['未来新人','outside-04','2026-10-01','','','国内','中国大陆','C','zt-owner'],
    ])
    rows=people(state);assert set(rows)=={'历史甲','本月新增'}
    row=rows['历史甲'];assert row['id']==old['id']
    assert row['values']['转正日期']=='2026-09-01' and row['values']['最后工作日']=='2026-09-20'
    for field in ('ABC类别','工作地','标签分类','招聘负责人工号','入职日期'):
        assert row['values'][field]==old['values'][field]
        assert field not in row['unresolvedFields']
    assert row['provenance']['ABC类别'][-1]['value']=='A'
    assert row['provenance']['ABC类别'][-1]['sheet']=='BP维护的全年人员资料'
    assert row['provenance']['ABC类别'][-1]['row']==2
    assert row['provenance']['ABC类别'][-1]['column']==8
    summary=state['sources'][-1]['cumulativeSummary']
    assert summary['newHires']==1 and summary['updatedPeople']==1
    assert {d['code'] for d in summary['notApplied']}=={'past_hire_not_carried','future_hire'}


def test_prior_future_and_blank_monthly_dates_never_replace_history(tmp_path):
    store=DraftStore(tmp_path)
    state=opening(store,[['历史甲','outside-01','2026-06-01','2026-08-15','2026-08-20','国内','中国大陆','B','zt-owner']])
    state=upload(store,state,[['历史甲','outside-01','2026-06-01','','','国内','中国大陆','B','zt-owner']])
    state=upload(store,state,[['历史甲','outside-01','2026-06-01','2026-07-15','2026-10-20','国内','中国大陆','B','zt-owner']],filename='新增观察.xlsx')
    row=state['rows'][0]
    assert row['values']['转正日期']=='2026-08-15' and row['values']['最后工作日']=='2026-08-20'
    assert not {'转正日期','最后工作日'} & set(row['unresolvedFields'])


def test_monthly_date_conflicts_remain_visible_and_manual_oa_edits_win(tmp_path):
    store=DraftStore(tmp_path);state=opening(store)
    state=upload(store,state,[['历史甲','outside-01','2026-06-01','2026-09-01','','国内','中国大陆','B','zt-owner']])
    state=upload(store,state,[['历史甲','outside-01','2026-06-01','2026-09-02','','国内','中国大陆','B','zt-owner']],filename='第二份累计.xlsx')
    row=state['rows'][0];assert row['values']['转正日期']=='2026-09-01'
    assert '转正日期' in row['unresolvedFields']
    state=store.change(202609,state['revision'],{'edits':[{'rowId':row['id'],'field':'转正日期','value':'2026-09-03','reason':'已核对审批原件'}]})
    edits=deepcopy(state['edits'])
    state=upload(store,state,[['历史甲','outside-01','2026-06-01','2026-09-04','','国内','中国大陆','B','zt-owner']],filename='第三份累计.xlsx')
    row=state['rows'][0]
    assert row['values']['转正日期']=='2026-09-03' and state['edits']==edits
    assert '转正日期' not in row['unresolvedFields']
    assert {p['value'] for p in row['provenance']['转正日期']}=={'2026-09-01','2026-09-02','2026-09-04'}


def test_two_employments_require_exact_date_and_conflicting_identifiers_are_skipped(tmp_path):
    store=DraftStore(tmp_path)
    state=opening(store,[['同名甲','outside-01','2025-01-01','','','国内','中国大陆','B','zt-owner'],
                         ['同名甲','outside-01','2026-06-01','','','国内','中国大陆','B','zt-owner']])
    state=upload(store,state,[['同名甲','outside-01','','2026-09-01','','国内','中国大陆','B','zt-owner'],
                              ['同名甲','outside-01','2026-06-01','2026-09-02','','国内','中国大陆','B','zt-owner']])
    assert len(state['rows'])==2
    dates={r['values']['入职日期']:r['values'].get('转正日期','') for r in state['rows']}
    assert dates=={'2025-01-01':'','2026-06-01':'2026-09-02'}
    assert state['sources'][-1]['cumulativeSummary']['notApplied'][0]['code']=='missing_hire_date'
    separate=DraftStore(tmp_path/'ids')
    fields=['姓名','工号','证件号','入职日期','转正日期']
    state=separate.ingest(202609,[('台账.xlsx',book(fields,[['甲','outside-a','cert-a','2026-06-01',''],['乙','outside-b','cert-b','2026-06-01','']],title='附件一-招聘奖金核算明细'))],0)
    state=upload(separate,state,[['甲','outside-a','cert-b','2026-06-01','2026-09-05']],fields=fields)
    assert len(state['rows'])==2 and not any(r['values'].get('转正日期') for r in state['rows'])
    assert state['sources'][-1]['cumulativeSummary']['notApplied'][0]['code']=='ambiguous_identity'
    assert not any(r.get('identityConflict') for r in state['rows'])


def test_duplicate_reopen_and_carry_do_not_expand_previous_cumulative_file(tmp_path):
    store=DraftStore(tmp_path);content=book(FIELDS,[
        ['本月新增','outside-02','2026-09-03','','','国内','中国大陆','C','zt-owner'],
        ['下月人员','outside-03','2026-10-01','','','国内','中国大陆','C','zt-owner']],title='BP累计')
    state=store.ingest(202609,[('累计.xlsx',content)],0,purpose='external-personnel')
    rid=state['rows'][0]['id'];revision=state['revision']
    repeated=DraftStore(tmp_path).ingest(202609,[('相同内容.xlsx',content)],revision,purpose='external-personnel')
    assert repeated['revision']==revision and len(repeated['sources'])==1
    assert repeated['rows'][0]['id']==rid
    carried=store.carry_forward(202610,202609,0)
    assert len(carried['rows'])==1 and carried['rows'][0]['id']==rid
    assert carried['rows'][0]['values']['ABC类别']=='C'
    assert carried['sources'][0]['forMonth']==202609
    refreshed=store.refresh(202610,carried['revision'])
    assert len(refreshed['rows'])==1
    current=store.ingest(202610,[('再次确认本月采用.xlsx',content)],refreshed['revision'],purpose='external-personnel')
    assert set(people(current))=={'本月新增','下月人员'}


def test_no_identifier_new_hire_is_retained_but_same_name_is_not_merged(tmp_path):
    store=DraftStore(tmp_path);state=store.read(202609)
    fields=['姓名','入职日期','ABC类别']
    state=upload(store,state,[['无工号甲','2026-09-01','C']],fields=fields)
    rid=state['rows'][0]['id']
    reopened=store.refresh(202609,state['revision']);assert reopened['rows'][0]['id']==rid
    state=upload(store,reopened,[['无工号甲','2026-09-01','B']],fields=fields)
    assert len(state['rows'])==1 and state['rows'][0]['values']['ABC类别']=='C'
    assert state['sources'][-1]['cumulativeSummary']['notApplied'][0]['code']=='ambiguous_identity'


@pytest.mark.parametrize('first_id,second_id',[('', 'outside-1'),('outside-1','')])
def test_missing_identifier_on_either_version_does_not_duplicate_calculable_employment(tmp_path,first_id,second_id):
    store=DraftStore(tmp_path)
    values=person(hire='2026-09-01',姓名='同名新人',工号=first_id,推荐人工号='')['values']
    fields=list(values)
    state=upload(store,store.read(202609),[list(values.values())],fields=fields)
    rid=state['rows'][0]['id'];total=state['rows'][0]['calculation']['recruitment']['total']
    assert total>0
    values['工号']=second_id
    state=upload(store,state,[list(values.values())],fields=fields,filename='补齐人员资料.xlsx')
    assert len(state['rows'])==1 and state['rows'][0]['id']==rid
    assert state['rows'][0]['values'].get('工号','')==first_id
    assert sum(r['calculation']['recruitment']['total'] for r in state['rows'])==total
    assert state['sources'][-1]['cumulativeSummary']['newHires']==0
    assert state['sources'][-1]['cumulativeSummary']['notApplied'][0]['code']=='ambiguous_identity'


def test_same_name_and_hire_date_with_distinct_confirmed_ids_remain_separate(tmp_path):
    store=DraftStore(tmp_path)
    values=person(hire='2026-09-01',姓名='同名新人',工号='outside-1',推荐人工号='')['values']
    fields=list(values)
    state=upload(store,store.read(202609),[list(values.values())],fields=fields)
    first=state['rows'][0]
    values['工号']='outside-2'
    state=upload(store,state,[list(values.values())],fields=fields,filename='另一位同名新人.xlsx')
    assert len(state['rows'])==2
    assert {r['values']['工号'] for r in state['rows']}=={'outside-1','outside-2'}
    assert state['sources'][-1]['cumulativeSummary']['newHires']==1
    assert state['sources'][-1]['cumulativeSummary']['notApplied']==[]
    assert sum(r['calculation']['recruitment']['total'] for r in state['rows'])==2*first['calculation']['recruitment']['total']


def test_new_employment_sharing_identifier_keeps_old_hire_and_requires_review(tmp_path):
    store=DraftStore(tmp_path);state=opening(store);old_id=state['rows'][0]['id']
    state=upload(store,state,[['历史甲','outside-01','2026-09-05','','','国内','中国大陆','C','zt-owner']])
    assert len(state['rows'])==2
    old=next(r for r in state['rows'] if r['id']==old_id)
    new=next(r for r in state['rows'] if r['id']!=old_id)
    assert old['values']['入职日期']=='2026-06-01'
    assert new['employmentCandidates'][0]['rowId']==old_id
    assert new['rule']['version']=='unresolved'


def test_duplicate_source_rows_count_people_once_and_preserve_conflicts_on_refresh(tmp_path):
    store=DraftStore(tmp_path)
    state=upload(store,store.read(202609),[
        ['新人','outside-1','2026-09-05','','','国内','中国大陆','B','zt-owner'],
        ['新人','outside-1','2026-09-05','','','国内','中国大陆','C','zt-owner']])
    assert len(state['rows'])==1
    assert state['sources'][0]['cumulativeSummary']['newHires']==1
    assert 'ABC类别' in state['rows'][0]['unresolvedFields']
    state=store.refresh(202609,state['revision'])
    assert len(state['rows'])==1 and 'ABC类别' in state['rows'][0]['unresolvedFields']


@pytest.mark.parametrize('reverse',[False,True])
def test_reused_offer_and_hire_date_never_merge_different_people(tmp_path,reverse):
    store=DraftStore(tmp_path);state=opening(store)
    fields=['姓名','海外offer流程编号','入职日期','转正日期']
    rows=[['Ashley Siddon','offer-1','2026-09-29',''],['Callum Halford','offer-1','2026-09-29',''],
          ['正常新人','offer-2','2026-09-20','']]
    state=upload(store,state,rows[::-1] if reverse else rows,fields=fields)
    assert set(people(state))=={'历史甲','正常新人'}
    summary=state['sources'][-1]['cumulativeSummary']
    assert summary['newHires']==1
    assert {x['name'] for x in summary['notApplied']}=={'Ashley Siddon','Callum Halford'}
    assert all(x['code']=='ambiguous_identity' and x['row'] for x in summary['notApplied'])
    refreshed=store.refresh(202609,state['revision'])
    assert set(people(refreshed))=={'历史甲','正常新人'}


def test_existing_person_is_not_updated_from_a_different_name_with_same_identifier(tmp_path):
    store=DraftStore(tmp_path);state=opening(store);before=deepcopy(state['rows'][0])
    state=upload(store,state,[['另一人','outside-01','2026-06-01','2026-09-02','','国内','中国大陆','C','zt-owner']])
    assert len(state['rows'])==1
    assert state['rows'][0]['values']==before['values'] and state['rows'][0]['history']==before['history']
    assert state['sources'][-1]['cumulativeSummary']['notApplied'][0]['code']=='ambiguous_identity'


def test_cumulative_matching_ignores_name_case_and_repeated_spaces(tmp_path):
    store=DraftStore(tmp_path)
    state=opening(store,[['Ashley Siddon','outside-01','2026-06-01','','','国内','中国大陆','C','zt-owner']])
    state=upload(store,state,[['ASHLEY  SIDDON','outside-01','2026-06-01','2026-09-02','','国内','中国大陆','C','zt-owner']])
    assert len(state['rows'])==1 and state['rows'][0]['values']['转正日期']=='2026-09-02'
    assert state['sources'][-1]['cumulativeSummary']['notApplied']==[]


def test_current_month_opening_does_not_accept_new_recruitment_facts(tmp_path):
    store=DraftStore(tmp_path)
    state=opening(store,[['已有甲','outside-1','2026-09-01','','','国内','中国大陆','B','zt-owner']])
    state=upload(store,state,[['已有甲','outside-1','2026-09-01','2026-09-15','','国内','英国','A','zt-other']])
    row=state['rows'][0]
    assert row['values']['ABC类别']=='B' and row['values']['招聘负责人工号']=='zt-owner'
    assert row['values']['工作地']=='中国大陆' and row['values']['转正日期']=='2026-09-15'
    assert not {'ABC类别','招聘负责人工号','工作地'} & set(row['unresolvedFields'])


def test_lost_or_changed_cumulative_identity_proof_is_reported(tmp_path):
    store=DraftStore(tmp_path)
    state=upload(store,store.read(202609),[['新人','outside-1','2026-09-05','','','国内','中国大陆','C','zt-owner']])
    row_id=state['rows'][0]['id'];source=state['sources'][0]
    original=store._folder(202609)/'sources'/source['storedName']
    original.write_bytes(book(FIELDS,[['新人','outside-1','2026-10-05','','','国内','中国大陆','C','zt-owner']],title='BP维护的全年人员资料'))
    reread=store.read(202609)
    assert reread['rows'][0]['id']==row_id
    assert reread['rows'][0]['unavailableSourceIds']==[source['id']]
    assert reread['rows'][0]['rule']['version']=='unresolved'


def test_financial_cells_never_become_opening_or_external_results(tmp_path):
    store=DraftStore(tmp_path);state=store.read(202609)
    fields=FIELDS+['招聘人入职1月奖金','招聘人入职1月发放周期','本月核定金额']
    state=upload(store,state,[['本月新增','outside-02','2026-09-03','','','国内','中国大陆','C','zt-owner',99999,202609,88888]],fields=fields,title='附件一-招聘奖金核算明细')
    row=state['rows'][0]
    assert row['history']==[] and not state.get('baseline') and not state.get('externalResults')
    assert not any(key in row['values'] for key in fields[-3:])
    assert state['sources'][0]['roles']==['cumulative']
    assert not any(n['amount'] in {99999,88888} for a in row['calculation'].values() if isinstance(a,dict) for n in a.get('nodes',[]))


@pytest.mark.parametrize('initial,purpose',[('', 'external-personnel'),('external-personnel','')])
def test_same_content_cannot_change_import_purpose(tmp_path,initial,purpose):
    store=DraftStore(tmp_path)
    content=book(FIELDS,[['本月新增','outside-02','2026-09-03','','','国内','中国大陆','C','zt-owner']],title='导入_月度数据')
    state=store.ingest(202609,[('资料.xlsx',content)],0,purpose=initial)
    with pytest.raises(ValueError,match='其他用途'):
        store.ingest(202609,[('改名.xlsx',content)],state['revision'],purpose=purpose)
    assert store.read(202609)['revision']==state['revision']
    assert store.read(202609)['sources'][0].get('purpose','')==initial


def test_external_results_and_cumulative_imports_reject_cross_purpose_content(tmp_path):
    content=book(['姓名','体系外人员编号','入职日期','被招聘人','奖金项目','领取人','本月核定金额','币种'],
                 [['候选甲','outside-1','2026-09-01','候选甲','招聘负责人奖金','领取甲',100,'CNY']])
    store=DraftStore(tmp_path/'facts')
    state=store.ingest(202609,[('人员.xlsx',content)],0,purpose='external-personnel')
    with pytest.raises(ValueError,match='其他用途'):
        external.prepare_import(store,202609,state['revision'],[('金额.xlsx',content)])
    result_store=DraftStore(tmp_path/'results')
    preview=external.prepare_import(result_store,202609,0,[('金额.xlsx',content)])
    state=result_store.read(202609)
    with pytest.raises(ValueError,match='其他用途'):
        result_store.ingest(202609,[('人员.xlsx',content)],state['revision'],purpose='external-personnel')
    repeated=external.prepare_import(result_store,202609,state['revision'],[('再传金额.xlsx',content)])
    assert repeated['token']==preview['token']


def test_cached_parsed_purpose_is_validated_before_reuse(tmp_path):
    store=DraftStore(tmp_path);state=upload(store,store.read(202609),[
        ['本月新增','outside-02','2026-09-03','','','国内','中国大陆','C','zt-owner'],
        ['未来人员','outside-03','2026-10-01','','','国内','中国大陆','C','zt-owner']])
    source=state['sources'][0];path=store._folder(202609)/'parsed'/f"{source['id']}.json"
    parsed=json.loads(path.read_text());parsed['sourcePurpose']=''
    for row in parsed['records']:row['role']='supplement'
    path.write_text(json.dumps(parsed))
    reopened=store.read(202609)
    assert set(people(reopened))=={'本月新增'}
    assert json.loads(path.read_text())['sourcePurpose']=='external-personnel'


def test_upload_api_template_and_invalid_purpose(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    store=DraftStore(tmp_path);monkeypatch.setattr(api,'store',store)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    template=client.get('/api/recruitment/drafts/202609/external-personnel/template')
    assert template.status_code==200
    workbook=load_workbook(BytesIO(template.content));sheet=workbook['体系外累计人员']
    fields=[cell.value for cell in sheet[1]];values={'姓名':'接口新人','体系外人员编号':'outside-1','入职日期':'2026-09-01','工作地':'中国大陆'}
    sheet.append([values.get(field,'') for field in fields]);output=BytesIO();workbook.save(output)
    response=client.post('/api/recruitment/drafts/202609/sources',data={'revision':0,'purpose':'external-personnel'},files=[('files',('202608招聘奖金台账.xlsx',output.getvalue()))])
    assert response.status_code==200,response.text
    data=response.json();assert data['month']==202609 and data['rows'][0]['values']['姓名']=='接口新人'
    assert data['sources'][0]['purpose']=='external-personnel'
    assert data['sources'][0]['cumulativeSummary']['newHires']==1
    bad=client.post('/api/recruitment/drafts/202609/sources',data={'revision':data['revision'],'purpose':'unknown'},files=[('files',('资料.xlsx',output.getvalue()))])
    assert bad.status_code==400


def test_generic_supplement_keeps_explicit_historical_and_future_rows(tmp_path):
    store=DraftStore(tmp_path)
    state=store.ingest(202609,[('补充.xlsx',book(['体系外人员编号','姓名','入职日期'],[
        ['old','历史','2026-05-01'],['future','未来','2026-10-01']]))],0)
    assert set(people(state))=={'历史','未来'}
