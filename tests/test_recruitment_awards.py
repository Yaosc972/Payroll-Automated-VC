from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bonus_platform.engine.recruitment_awards import CATALOG, award, calculate_draft, lookup, cycle, channel
from bonus_platform.engine.recruitment_drafts import DraftStore
from test_recruitment_drafts import book


def person(hire='2026-07-01', **values):
    v={'姓名':'核验人员','工号':'zt-check','入职日期':hire,'候选人入职时间':hire,'人员状态':'正式',
       '标签分类':'海外','工作地':'德国','二级组织':'德国','奖金地区类型':'海外发达国家','ABC类别':'C','职级':'P1-1',
       '招聘渠道':'招聘网站','招聘启动日期':'2026-05-20','招聘负责人工号':'zt-owner','招聘负责人姓名':'招聘负责人',
       '招聘负责人入职日期':'2020-01-01','招聘负责人人员状态':'正式','推荐人工号':'zt-ref','推荐人姓名':'推荐人',
       '推荐人入职日期':'2020-01-01','推荐人人员状态':'正式','推荐人职级':'P2-1','推荐人职位':'产品运营',
       '直线经理工号':'zt-manager','推荐人转正日期':'2020-07-01'}
    v.update(values)
    return {'id':'employee-span','values':v,'rule':{'version':'legacy' if hire<'2026-07-01' else '2026-07','specialScope':''},
            'unresolvedFields':[],'scopeStatus':'已纳入底稿','history':[]}


@pytest.mark.parametrize('category,level,developed,developing,ratios',[
    ('A','P4-1',2100,1100,[.3,.3,.4,0]),('A','M3-2',1300,700,[.3,.3,.4,0]),
    ('B','P3-1',1100,600,[.3,.3,.4,0]),('B','M2-1',600,300,[.3,.7,0,0]),
    ('C','P3-2',900,400,[.3,.7,0,0]),('C','P2-1',500,200,[.5,.5,0,0]),('C','P1-1',400,100,[.5,.5,0,0]),
])
def test_complete_overseas_referral_standards_and_merged_installments(category,level,developed,developing,ratios):
    for region,expected in [('海外发达国家',developed),('海外发展中国家',developing)]:
        r=person(ABC类别=category,职级=level,奖金地区类型=region)
        standard=lookup(r,'referral')
        assert standard['amount']==expected
        assert standard['ratios']==ratios
        assert award(r,'referral',202607)['total']==expected


@pytest.mark.parametrize('kind,region,category,grade,amount,days,ratios',[
    ('recruitment','国内','A','P3-1',6000,90,[.2,0,0,.8]),
    ('recruitment','国内','B','P2-1',1500,60,[.2,0,0,.8]),
    ('recruitment','国内','C','P1-1',150,42,[.5,0,0,.5]),
    ('recruitment','国内','C','P2-1',300,49,[.5,0,0,.5]),
    ('recruitment','国内','C','M3-1',600,80,[.5,0,0,.5]),
    ('recruitment','海外','A','P3-1',1250,120,[.3,.3,.4,0]),
    ('recruitment','海外','A','M4-1',1650,120,[.3,.3,.4,0]),
    ('recruitment','海外','B','P1-1',450,80,[.3,.7,0,0]),
    ('recruitment','海外','B','P3-1',650,80,[.3,.7,0,0]),
    ('recruitment','海外','C','P1-1',55,42,[.5,.5,0,0]),
    ('recruitment','海外','C','M2-1',110,49,[.5,.5,0,0]),
    ('recruitment','海外','C','P3-1',180,80,[.5,.5,0,0]),
    ('referral','国内','A','P4-1',10000,None,[.3,0,0,.7]),
    ('referral','国内','A','M3-1',6000,None,[.3,0,0,.7]),
    ('referral','国内','B','P3-1',5000,None,[.3,0,0,.7]),
    ('referral','国内','B','M2-1',2000,None,[.3,0,0,.7]),
    ('referral','国内','C','P3-1',4000,None,[.3,0,0,.7]),
    ('referral','国内','C','P2-1',1500,None,[.5,0,0,.5]),
    ('referral','国内','C','P1-1',600,None,[.5,0,0,.5]),
])
def test_remaining_new_standard_tables_have_complete_amount_cycle_and_installments(kind,region,category,grade,amount,days,ratios):
    row=person(标签分类=region,ABC类别=category,职级=grade,奖金地区类型='国内发展中国家' if region=='国内' else '海外发达国家')
    standard=lookup(row,kind)
    assert standard['amount']==amount
    assert standard['ratios']==ratios
    if days is not None: assert standard['cycle']==days


def test_old_and_new_in_same_month_are_isolated_including_old_special_scope():
    old=person('2026-06-30');new=person('2026-07-01')
    assert award(old,'referral',202608)['standard']==150
    assert award(new,'referral',202608)['standard']==400
    old['values']['职级']='P1-3';new['values']['职级']='P1-3'
    old['rule']['specialScope']='FBU德国';new['rule']['specialScope']='FBU德国'
    assert lookup(old,'referral')['scope']=='FBU德国'
    assert lookup(new,'referral')['scope']=='集团统一'
    assert '旧特殊规则适用日期及依据待核' not in award(old,'referral',202607)['issues']


def test_explicit_probation_is_a_result_not_a_missing_employment_task():
    row=person(推荐人人员状态='试用',推荐人转正日期='')
    node=next(n for n in award(row,'referral',202607)['nodes'] if n['dueMonth']==202607)
    assert node['status']=='按规则不发'
    assert node['scheduledMonth'] is None
    assert not node['issues']
    assert '推荐人：适用日期尚未转正' in node['eligibilityReasons']
    row['values']['推荐人人员状态']=''
    assert any('正式在职状态待核' in i for n in award(row,'referral',202607)['nodes'] for i in n['issues'])


@pytest.mark.parametrize('identifier',['体系外海外-Jin','海外-GL001669'])
def test_offline_external_recipient_does_not_require_internal_employment(identifier):
    row=person(招聘负责人工号=identifier,招聘负责人入职日期='',招聘负责人人员状态='')
    row['provenance']={'招聘负责人工号':[{'sourceId':'offline','role':'ledger','value':identifier}]}
    result=award(row,'recruitment',202607)
    assert not result['nodes'][0]['issues']
    assert result['nodes'][0]['employmentBasis']['sourceId']=='offline'
    row['values']['招聘负责人最后工作日']='2026-07-02'
    assert award(row,'recruitment',202607)['nodes'][0]['status']=='按规则不发'
    row['values']['招聘负责人最后工作日']=''
    row['provenance']={}
    assert '招聘负责人：任职日期待核' in award(row,'recruitment',202607)['nodes'][0]['issues']


def test_zero_exclusion_does_not_require_an_approval_even_when_holidays_exist():
    v=person(标签分类='国内',招聘启动日期='2026-05-01',周期剔除天数=0)['values']
    result,issues=cycle(v,42,'2026-07')
    assert result['exclusionDays']==0
    assert '周期剔除依据待核' not in issues


@pytest.mark.parametrize('previous,hire,factor',[
    ('2025-07-01','2026-07-01',0),('2025-06-30','2026-07-01',.4),
    ('2027-02-28','2028-02-28',0),('2027-02-28','2028-02-29',.4),
    ('2028-02-29','2029-02-28',0),('2028-02-29','2029-03-01',.4),
])
def test_rehire_calendar_month_boundary(previous,hire,factor):
    r=person(hire,招聘渠道='二次入职',上次离职日期=previous)
    assert channel(r['values'],'2026-07')[0]==factor
    assert ('二次入职专项条款执行依据待核' in channel(r['values'],'2026-07')[1]) == (factor==0)
    if factor==0:assert award(r,'recruitment',202607)['total'] is None
    assert award(r,'referral',202607)['total']==0


def test_complete_rehire_over_twelve_months_explains_consistent_rule_without_an_approval():
    row=person(招聘渠道='二次入职',上次离职日期='2024-09-06',标签分类='国内',
               管培生已批名单资格='是',周期豁免依据='已确认的管培生名单',招聘负责人最后工作日='2026-06-12')
    before=deepcopy(row)
    result=award(row,'recruitment',202607)
    assert result['issues']==[]
    assert result['channelFactor']==.4 and result['total']==60
    assert result['currentAmount']==0  # A channel conclusion does not approve a recipient's eligibility.
    assert result['nodes'][0]['status']=='按规则不发'
    basis=result['rehireBasis']
    assert basis['previousExitDate']=='2024-09-06' and basis['hireDate']=='2026-07-01'
    assert basis['twelveMonthDate']=='2025-09-06'
    assert '超过12个月' in basis['description'] and '40%' in basis['description']
    assert '第2、3页' in basis['source'] and '第8页' in basis['source']
    assert row==before and not row['values'].get('二次入职专项口径依据')
    assert award(row,'referral',202607)['total']==0


@pytest.mark.parametrize('previous',['','无法辨认','2026-07-01','2026-07-02'])
def test_missing_or_invalid_rehire_dates_are_not_automatically_cleared(previous):
    result=award(person(招聘渠道='二次入职',上次离职日期=previous),'recruitment',202607)
    assert result['total'] is None
    assert '二次入职的上一段离职日期待核' in result['issues']
    assert not result.get('rehireBasis')


def test_rehire_date_conflict_is_not_presented_as_a_confirmed_automatic_result():
    row=person(招聘渠道='二次入职',上次离职日期='2024-09-06')
    row['unresolvedFields']=['上次离职日期']
    result=award(row,'recruitment',202607)
    assert result['total'] is None
    assert '核算字段来源冲突：上次离职日期' in result['issues']
    assert not result.get('rehireBasis')


def test_rehire_existing_confirmation_and_legacy_rule_are_preserved():
    values=person(招聘渠道='二次入职',二次入职专项口径依据='已批准的适用条款',二次入职间隔结论='超过十二个月')['values']
    assert channel(values,'2026-07')[:2]==(.4,[])
    values['二次入职间隔结论']='不超过十二个月'
    assert channel(values,'2026-07')[:2]==(0,[])
    assert channel(person(招聘渠道='二次入职')['values'],'legacy')[:2]==(.4,[])
    row=person(招聘渠道='二次入职',上次离职日期='2024-09-06');row['rule']['version']='unresolved'
    result=award(row,'recruitment',202607)
    assert result['total'] is None and '本次任职与规则版本待核' in result['issues']
    assert not result.get('rehireBasis')
    confirmed=person(招聘渠道='二次入职',上次离职日期='2026-01-01',二次入职专项口径依据='已批准采用专项不发条款')
    assert award(confirmed,'recruitment',202607)['total']==0


def test_missing_rehire_date_is_unknown_and_referral_ineligibility_is_independent():
    r=person(招聘渠道='二次入职')
    assert award(r,'recruitment',202607)['total'] is None
    assert award(r,'referral',202607)['total']==0


def test_unknown_channel_never_defaults_to_zero_or_full_bonus():
    r=person(招聘渠道='其他',其他渠道='尚未认可的新来源')
    assert award(r,'recruitment',202607)['total'] is None
    r['values']['其他渠道']='招聘广告'
    assert channel(r['values'],'2026-07')[0]==1


@pytest.mark.parametrize('grade',['P4-1','M4-1','P6-2','M9-1'])
def test_high_referrer_grades_are_ineligible_in_both_versions_and_special_scope(grade):
    for hire in ('2026-06-30','2026-07-01'):
        r=person(hire,推荐人职级=grade);r['rule']['specialScope']='FBU德国'
        assert award(r,'referral',202607)['total']==0


def test_domestic_graduate_exemption_requires_evidence_and_retains_channel_share():
    r=person(标签分类='国内',奖金地区类型='国内发展中国家',二级组织='中国大陆',工作地='深圳',招聘渠道='校招',
             招聘启动日期='2026-01-01',周期豁免依据='已核对批准名单',应届校招='是',延后入职原因='公司集训',原计划入职日期='2026-06-16')
    a=award(r,'recruitment',202607)
    assert a['cycle']['exempt'] and a['cycle']['factor']==1 and a['channelFactor']==.4
    assert a['total']==60
    r['values']['原计划入职日期']='2026-06-15'
    assert not award(r,'recruitment',202607)['cycle']['exempt']
    r['values'].pop('原计划入职日期');r['values']['管培生已批名单资格']='是'
    assert award(r,'recruitment',202607)['cycle']['exempt']
    r['values'].pop('周期豁免依据')
    assert not award(r,'recruitment',202607)['cycle']['exempt']


def test_negative_exclusion_is_not_a_cycle_exemption_and_holidays_are_not_stacked():
    r=person(周期剔除天数=-90)
    assert award(r,'recruitment',202607)['total'] is None
    v=person(标签分类='国内',招聘启动日期='2026-05-01',周期剔除天数=10,周期剔除依据='已批采用总天数')['values']
    c,issues=cycle(v,42,'2026-07')
    assert c['exclusionDays']==10 and c['holidayDays']>0
    assert c['actualDays']==61-10


def test_new_overseas_c_recruitment_has_both_halves_and_rounding_balances():
    r=person(协助招聘人工号='zt-helper',协助招聘人入职日期='2020-01-01',协助招聘人人员状态='正式')
    a=award(r,'recruitment',202607)
    assert len(a['nodes'])==4
    assert {n['label'] for n in a['nodes']}=={'入职满一个月','入职满三个月'}
    assert round(sum(n['amount'] for n in a['nodes']),2)==a['total']==55
    assert sum(n['amount'] for n in a['nodes'] if n['role']=='协助招聘人')==16.5


def test_legacy_each_installment_rounds_and_exposes_difference_separately():
    row=person('2026-04-27',招聘启动日期='2026-04-23',协助招聘人工号='zt-helper')
    a=award(row,'recruitment',202607)
    assert a['total']==a['formulaTotal']==71.5
    assert [n['amount'] for n in a['nodes'] if n['role']=='招聘负责人']==[25.03,25.03]
    assert [n['amount'] for n in a['nodes'] if n['role']=='协助招聘人']==[10.73,10.73]
    assert a['scheduledTotal']==71.52
    assert a['roundingDifference']==.02
    owner=next(r for r in a['recipientTotals'] if r['role']=='招聘负责人')
    assert owner=={'role':'招聘负责人','allocatedAmount':50.05,'scheduledAmount':50.06,'roundingDifference':.01}
    assert all(n['paymentStatus']=='未核实' for n in a['nodes'])


def test_new_rule_keeps_existing_final_installment_remainder():
    row=person('2026-07-01',招聘启动日期='2026-06-27',协助招聘人工号='zt-helper')
    a=award(row,'recruitment',202609)
    assert [n['amount'] for n in a['nodes'] if n['role']=='招聘负责人']==[25.03,25.02]
    assert a['scheduledTotal']==a['total']==71.5
    assert a['roundingDifference']==0


def test_recipient_waiting_to_join_is_not_eligible_and_last_day_is_checked():
    r=person(推荐人人员状态='待入职',推荐人入职日期='2026-08-01')
    a=award(r,'referral',202607)
    assert any(n['status']=='按规则不发' and any('推荐人' in i for i in n['eligibilityReasons']) for n in a['nodes'] if n['dueMonth']==202607)
    r=person(招聘负责人最后工作日='2026-07-29')
    a=award(r,'recruitment',202607)
    assert a['nodes'][0]['status']=='按规则不发'
    assert a['nodes'][0]['scheduledMonth'] is None
    assert not a['nodes'][0]['paymentReady']


def test_core_conflict_blocks_amount_but_nonessential_display_field_does_not():
    r=person();r['unresolvedFields']=['姓名']
    assert award(r,'referral',202607)['total']==400
    r['unresolvedFields']=['ABC类别']
    assert award(r,'referral',202607)['total'] is None


def test_organization_priority_used_for_both_currency_and_referral():
    r=person(二级组织='墨西哥',工作地='德国',奖金地区类型='海外发展中国家')
    for kind in ('recruitment','referral'): assert award(r,kind,202607)['currency']=='美元'
    assert award(r,'referral',202607)['standard']==100


def test_save_recalculate_and_export_share_same_values_and_never_mark_history_paid(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    from bonus_platform.app import app
    from fastapi.testclient import TestClient
    store=DraftStore(tmp_path);monkeypatch.setattr(api,'store',store)
    values=person()['values'];values['体系外人员编号']='outside'
    draft=store.ingest(202607,[('补充.xlsx',book(list(values),[list(values.values())]))],0)
    rid=draft['rows'][0]['id']
    for role,kind in [('招聘负责人','recruitment'),('推荐人','referral')]:
        assert draft['rows'][0]['values'][role+'入职日期']=='2020-01-01'
        assert not any('任职日期待核' in issue for node in draft['rows'][0]['calculation'][kind]['nodes'] for issue in node['issues'])
    assert draft['rows'][0]['calculation']['referral']['total']==400
    draft=store.change(202607,draft['revision'],{'edits':[{'rowId':rid,'field':'入职日期','value':'2026-06-30','reason':'真实入职记录'},
        {'rowId':rid,'field':'候选人入职时间','value':'2026-06-30','reason':'真实入职记录'}]})
    saved=DraftStore(tmp_path).read(202607)
    assert saved['rows'][0]['calculation']['referral']['total']==150
    client=TestClient(app)
    response=client.get('/api/recruitment/drafts/202607/calculation/export')
    assert response.status_code==200
    wb=load_workbook(BytesIO(response.content),data_only=True)
    assert wb.sheetnames==['核算总明细','本月到期明细','领取人分币种汇总','未结与异常','调整及依据']
    columns={c.value:c.column for c in wb['核算总明细'][3]}
    assert wb['核算总明细'].cell(4,columns['内推试算总额']).value==150
    assert not saved['paymentReady']
    assert saved['calculation']['draftRevision']==saved['revision']


def test_currency_totals_never_add_different_currencies():
    a=person();b=person(二级组织='美国',工作地='美国');b['id']='another'
    state={'rows':[a,b],'month':202607,'revision':1,'todos':[],'summary':{}}
    result=calculate_draft(state)
    assert set(result['calculation']['totalsByCurrency'])=={'欧元','美元'}
    assert result['calculation']['totalsByCurrency']['欧元']==result['calculation']['totalsByCurrency']['美元']


def test_real_source_category_aliases_do_not_create_false_conflicts(tmp_path):
    values=person(ABC类别='C1类',职类='操作管理类')['values'];values['体系外人员编号']='old'
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('历史.xlsx',book(list(values),[list(values.values())],title='附件一-招聘奖金核算明细'))],0)
    newer=dict(values,ABC类别='C')
    draft=store.ingest(202607,[('补充.xlsx',book(list(newer),[list(newer.values())]))],draft['revision'])
    row=draft['rows'][0]
    assert 'ABC类别' not in row['unresolvedFields']
    assert row['scopeStatus']=='已纳入底稿'
    assert row['calculation']['referral']['total']==400
    assert row['provenance']['ABC类别'][0]['rawValue']=='C1类'


def test_domestic_non_headcount_keeps_original_without_false_c_category_conflict(tmp_path):
    values=person(ABC类别='C',标签分类='国内',二级组织='中国大陆',奖金地区类型='国内发展中国家')['values'];values['体系外人员编号']='domestic'
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('历史.xlsx',book(list(values),[list(values.values())],title='附件一-招聘奖金核算明细'))],0)
    newer=dict(values,ABC类别='不占编')
    draft=store.ingest(202607,[('补充.xlsx',book(list(newer),[list(newer.values())]))],draft['revision'])
    row=draft['rows'][0]
    assert row['values']['原始ABC类别']=='不占编'
    assert 'ABC类别' not in row['unresolvedFields']
    assert row['calculation']['recruitment']['standard']==150


def test_known_recruiting_site_aliases_do_not_block_the_real_source_amount(tmp_path):
    v=person()['values'];v['体系外人员编号']='channel'
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('历史.xlsx',book(list(v),[list(v.values())],title='附件一-招聘奖金核算明细'))],0)
    latest=dict(v,招聘渠道='Boss直聘')
    draft=store.ingest(202607,[('录用.xlsx',book(list(latest),[list(latest.values())]))],draft['revision'])
    row=draft['rows'][0]
    assert '招聘渠道' not in row['unresolvedFields']
    assert row['calculation']['recruitment']['total']==55
    assert row['provenance']['招聘渠道'][-1]['rawValue']=='Boss直聘'
    conflict=dict(v,招聘渠道='内推')
    draft=store.ingest(202607,[('新资料.xlsx',book(list(conflict),[list(conflict.values())]))],draft['revision'])
    assert '招聘渠道' in draft['rows'][0]['unresolvedFields']
    assert draft['rows'][0]['calculation']['recruitment']['total'] is None
