from io import BytesIO
from datetime import datetime
import zipfile

import pytest
from openpyxl import Workbook

from bonus_platform.engine.recruitment_drafts import DraftStore, DraftConflict, read_source


def book(headers, rows, title='Sheet1', broken_dimension=False):
    wb = Workbook(); ws = wb.active; ws.title = title
    ws.append(headers)
    for row in rows: ws.append(row)
    b = BytesIO(); wb.save(b)
    if not broken_dimension: return b.getvalue()
    out = BytesIO()
    with zipfile.ZipFile(b) as src, zipfile.ZipFile(out, 'w') as dest:
        for name in src.namelist():
            data = src.read(name)
            if name == 'xl/worksheets/sheet1.xml':
                import re
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', data)
            dest.writestr(name, data)
    return out.getvalue()


ROSTER = ['姓名', '工号', 'offer流程编号', '入职日期', '人员状态', '实际转正日期', '离职日期', '人员来源', '办公地址', '二级部门']
OFFER = ['offer编码', '招聘需求编码', '编制ABC分类', '招聘负责人工号', '招聘渠道', '其他渠道', '状态']


def test_ledger_multirow_headers_preserve_referral_hierarchy(tmp_path):
    wb=Workbook();ws=wb.active;ws.title='附件一-招聘奖金核算明细'
    ws.append(['花名册','','','直接上级','','间接上级1',''])
    ws.append(['姓名','工号','入职日期','','','','','推荐人是否直接/间接上级','推荐人工号'])
    ws.append(['','','','姓名','工号','姓名','工号','','','入职1个月发放周期'])
    ws.append(['候选人','zt1',datetime(2026,1,1),'经理','zt2','总监','zt3','是','zt3',202602])
    path=tmp_path/'ledger.xlsx';wb.save(path)
    rec=read_source(path,'opening')['records'][0]
    assert rec['values']['姓名']=='候选人' and rec['values']['工号']=='zt1'
    assert rec['values']['推荐人是否直接/间接上级']=='是'
    assert rec['values']['直接上级工号']=='zt2'
    assert rec['values']['间接上级1工号']=='zt3'


def test_combined_roster_dates_and_missing_offers_remain_in_draft(tmp_path):
    store = DraftStore(tmp_path)
    data = book(ROSTER, [
        ['甲', 'zt1', '', datetime(2026,7,1), '正式', datetime(2026,7,20), '', 'EHR入职', '深圳', '中国'],
        ['乙', 'zt2', 'OF2', datetime(2026,7,2), '离职', '', datetime(2026,7,30), 'OEHR入职', '德国', '德国'],
    ], broken_dimension=True)
    draft = store.ingest(202607, [('海外花名册.xlsx', data)], 0)
    assert len(draft['rows']) == 2
    rows = {r['values']['工号']: r for r in draft['rows']}
    assert rows['zt1']['values']['标签分类'] == '国内'
    assert rows['zt1']['values']['转正日期'] == '2026-07-20'
    assert rows['zt2']['values']['最后工作日'] == '2026-07-30'
    assert rows['zt1']['rule']['version'] == '2026-07'
    assert not any('缺少EHR' in t['title'] for t in draft['todos'])
    assert sum(len(t['rowIds']) for t in draft['todos'] if t['kind']=='missing_offer') == 2
    assert not draft.get('paymentReady')


def test_employment_dates_separate_and_conflicts_do_not_pick_rule(tmp_path):
    store = DraftStore(tmp_path)
    draft = store.ingest(202607, [('补充.xlsx', book(
        ['体系外人员编号','姓名','工号','入职日期','候选人入职时间'],
        [['a','甲','zt1','2026-06-30','2026-06-30'],['b','甲','zt1','2026-07-01','2026-07-01'],['c','乙','zt2','2026-06-30','2026-07-01']]))],0)
    assert len(draft['rows']) == 3
    assert {r['rule']['version'] for r in draft['rows'] if r['values']['工号']=='zt1'} == {'legacy','2026-07'}
    assert next(r for r in draft['rows'] if r['values']['工号']=='zt2')['rule']['version'] == 'unresolved'


def test_saved_edit_survives_duplicate_upload_restart_and_new_source(tmp_path):
    store = DraftStore(tmp_path)
    data=book(ROSTER,[['甲','zt1','OF1','2026-07-01','正式','','','OEHR入职','德国','德国']])
    draft=store.ingest(202607,[('roster.xlsx',data)],0)
    row_id=draft['rows'][0]['id']
    draft=store.change(202607,draft['revision'],{'edits':[{'rowId':row_id,'field':'招聘负责人工号','value':'zt-correct','reason':'已核对原件'}]})
    rev=draft['revision']
    draft=DraftStore(tmp_path).ingest(202607,[('另一个文件名.xlsx',data)],rev)
    assert draft['revision']==rev
    draft=store.ingest(202607,[('offer.xlsx',book(OFFER,[['OF1','R1','C','zt-original','招聘网站','','已入职']]))],rev)
    row=draft['rows'][0]
    assert row['values']['招聘负责人工号']=='zt-correct'
    assert row['original']['招聘负责人工号']=='zt-original'
    assert row['edits'][0]['reason']=='已核对原件'
    assert row['id']==row_id
    assert not any(t['kind']=='missing_offer' for t in draft['todos'])
    with pytest.raises(DraftConflict):
        store.change(202607,rev,{'edits':[{'rowId':row_id,'field':'姓名','value':'stale','reason':'test'}]})
    assert store.read(202607)['rows'][0]['values']['姓名']=='甲'


def test_common_recipient_issue_groups_and_resolves_once(tmp_path):
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('补充.xlsx',book(['体系外人员编号','姓名','入职日期','招聘负责人工号'],[['a','甲','2026-07-01','zt-owner'],['b','乙','2026-07-02','zt-owner']]))],0)
    groups=[t for t in draft['todos'] if t['kind']=='missing_recipient']
    assert len(groups)==1 and len(groups[0]['rowIds'])==2
    draft=store.ingest(202607,[('roster.xlsx',book(ROSTER,[['负责人','zt-owner','','2020-01-01','正式','','','EHR入职','中国','中国']]))],draft['revision'])
    assert not any(t['kind']=='missing_recipient' for t in draft['todos'])
    assert all(r['values']['招聘负责人入职日期']=='2020-01-01' for r in draft['rows'])


def test_parser_upgrade_recovers_recipient_dates_without_losing_saved_edits(tmp_path):
    import json
    from bonus_platform.engine.recruitment_draft_sources import PARSER_VERSION
    store=DraftStore(tmp_path)
    fields=['体系外人员编号','姓名','入职日期','招聘负责人入职日期','协助招聘人入职日期','推荐人入职日期']
    content=book(fields,[['a','甲','2026-07-01',datetime(2020,1,1),datetime(2021,1,1),datetime(2022,1,1)]])
    draft=store.ingest(202607,[('补充.xlsx',content)],0)
    rid=draft['rows'][0]['id']
    draft=store.change(202607,draft['revision'],{'edits':[{'rowId':rid,'field':'招聘负责人入职日期','value':'2019-12-01','reason':'核对原件'}]})
    folder=tmp_path/'202607';sid=draft['sources'][0]['id']
    # Emulate a saved source parsed before recipient entry dates were supported.
    draft['sources'][0].pop('parserVersion')
    store._write(folder/'draft.json',draft)
    parsed=json.loads((folder/'parsed'/f'{sid}.json').read_text())
    parsed.pop('parserVersion',None)
    for field in fields[3:]: parsed['records'][0]['values'].pop(field)
    store._write(folder/'parsed'/f'{sid}.json',parsed)
    refreshed=DraftStore(tmp_path).ingest(202607,[('再次上传.xlsx',content)],draft['revision'])
    row=refreshed['rows'][0]
    assert len(refreshed['sources'])==1
    assert refreshed['sources'][0]['parserVersion']==PARSER_VERSION
    assert row['id']==rid
    assert row['values']['招聘负责人入职日期']=='2019-12-01'
    assert row['values']['协助招聘人入职日期']=='2021-01-01'
    assert row['values']['推荐人入职日期']=='2022-01-01'
    assert row['original']['招聘负责人入职日期']=='2020-01-01'
    assert row['edits'][0]['reason']=='核对原件'
    assert store.read(202607)==refreshed


def test_missing_identity_retained_and_no_payment_inferred_from_history(tmp_path):
    store=DraftStore(tmp_path)
    data=book(['姓名','工号','入职日期','招聘人入职1月奖金','备注'],[['无工号','','2026-06-01',300,'人工提前安排']],title='附件一-招聘奖金核算明细')
    draft=store.ingest(202607,[('6月台账.xlsx',data)],0)
    row=draft['rows'][0]
    assert row['values']['姓名']=='无工号'
    assert row['history'][0]['cells']
    assert row['history'][0]['paymentStatus']=='unverified'
    assert not draft['baseline']
    draft=store.change(202607,draft['revision'],{'baseline':draft['sources'][0]['id']})
    assert draft['baseline']==draft['sources'][0]['id']
    assert draft['rows'][0]['history'][0]['paymentStatus']=='unverified'


def test_multiple_sheets_and_unknown_file_do_not_discard_good_data(tmp_path):
    wb=Workbook(); wb.active.title='说明'; wb.active.append(['请看花名册'])
    ws=wb.create_sheet('合并资料'); ws.append(ROSTER); ws.append(['甲','zt1','','2026-07-01','正式','','','EHR入职','中国','中国'])
    b=BytesIO(); wb.save(b)
    draft=DraftStore(tmp_path).ingest(202607,[('package.xlsx',b.getvalue()),('审批.pdf',b'%PDF-1.4\n')],0)
    assert len(draft['rows'])==1
    assert len(draft['sources'])==2
    assert any(t['kind']=='attachment_review' for t in draft['todos'])


def test_draft_api_recovers_after_reload_and_rejects_stale_edits(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from bonus_platform.app import app
    from bonus_platform import recruitment_draft_api as api
    monkeypatch.setattr(api,'store',DraftStore(tmp_path))
    client=TestClient(app)
    data=book(ROSTER,[['甲','zt1','','2026-07-01','正式','','','EHR入职','中国','中国']])
    result=client.post('/api/recruitment/drafts/202607/sources',data={'revision':0},files=[('files',('roster.xlsx',data))])
    assert result.status_code==200,result.text
    draft=result.json(); row=draft['rows'][0]
    assert 'provenance' not in row
    detail=client.get(f"/api/recruitment/drafts/202607/rows/{row['id']}").json()['row']
    assert detail['provenance']['入职日期'][0]['column']==4
    monkeypatch.setattr(api,'store',DraftStore(tmp_path))
    assert client.get('/api/recruitment/drafts/202607').json()['revision']==draft['revision']
    resp=client.post('/api/recruitment/drafts/202607/changes',json={'revision':0,'edits':[]})
    assert resp.status_code==409
    export=client.get('/api/recruitment/drafts/202607/export')
    assert export.status_code==200
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(export.content))
    assert wb.sheetnames==['工作底稿','集中待办','人工更正','历史原记录','资料清单']
    assert client.get('/api/recruitment/drafts/202613').status_code==400


def test_recipient_conflicts_and_migration_dates_require_review(tmp_path):
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[
        ('补充.xlsx',book(['体系外人员编号','姓名','入职日期','招聘负责人工号'],[['a','甲','2026-07-01','zt-owner']])),
        # Migration is an explicit supplement, not a roster-discovered hire.
        ('迁移补充.xlsx',book(['姓名','工号','入职日期','人员来源','二级组织'],[['迁移人员','zt-moved','2026-07-01','人员迁移','中国']],title='导入_月度数据')),
        ('one.xlsx',book(ROSTER,[['负责人','zt-owner','','2020-01-01','正式','','','EHR入职','中国','中国'],['迁移人员','zt-moved','','2026-07-01','正式','','','人员迁移','中国','中国']])),
        ('two.xlsx',book(ROSTER,[['负责人','zt-owner','','2020-01-01','离职','','2026-07-05','EHR入职','中国','中国']])),
    ],0)
    moved=next(r for r in draft['rows'] if r['values'].get('工号')=='zt-moved')
    assert moved['rule']['version']=='unresolved'
    assert moved['values']['标签分类']=='国内'
    assert any(t['kind']=='recipient_conflict' for t in draft['todos'])


def test_edit_hire_date_is_durable_across_source_refresh_and_new_month(tmp_path):
    store=DraftStore(tmp_path)
    raw=book(ROSTER,[['甲','zt1','OF1','2026-07-01','正式','','','OEHR入职','德国','德国']])
    draft=store.ingest(202607,[('roster.xlsx',raw)],0)
    row=draft['rows'][0]
    draft=store.change(202607,draft['revision'],{'edits':[{'rowId':row['id'],'field':'入职日期','value':'2026-06-30','reason':'实际入职记录'}]})
    assert draft['rows'][0]['rule']['version']=='legacy'
    draft=store.ingest(202607,[('offer.xlsx',book(OFFER,[['OF1','R1','C','zt-owner','招聘网站','','已入职']]))],draft['revision'])
    assert len(draft['rows'])==1 and draft['rows'][0]['rule']['version']=='legacy'
    next_month=store.carry_forward(202608,202607,0)
    assert next_month['rows'][0]['id']==row['id']
    assert next_month['rows'][0]['values']['入职日期']=='2026-06-30'
    assert next_month['rows'][0]['rule']['version']=='legacy'
    assert next_month['parentMonth']==202607


def test_invalid_batch_changes_are_atomic_and_unknown_ops_rejected(tmp_path):
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('补充.xlsx',book(['体系外人员编号','姓名','入职日期'],[['a','甲','2026-07-01']]))],0)
    rid=draft['rows'][0]['id']; rev=draft['revision']
    with pytest.raises(ValueError):
        store.change(202607,rev,{'edits':[{'rowId':rid,'field':'姓名','value':'wrong','reason':'test'}, {'rowId':rid,'field':'入职日期','value':'tomorrow','reason':'test'}]})
    assert store.read(202607)['rows'][0]['values']['姓名']=='甲'
    assert store.read(202607)['revision']==rev
    with pytest.raises(ValueError): store.change(202607,rev,{'rules':'free-edit'})


def test_missing_rule_facts_and_transfers_have_specific_tasks(tmp_path):
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('补充.xlsx',book(['体系外人员编号','姓名','入职日期','招聘渠道'],[['a','甲','2026-07-01','招聘网站'],['b','乙','2026-07-02','跨体系调动']]))],0)
    first=next(r for r in draft['rows'] if r['values']['姓名']=='甲')
    second=next(r for r in draft['rows'] if r['values']['姓名']=='乙')
    assert any(t['kind']=='missing_fields' and first['id'] in t['rowIds'] for t in draft['todos'])
    assert not any(t['kind']=='missing_fields' and second['id'] in t['rowIds'] for t in draft['todos'])
    assert second['scopeStatus']=='调动记录，保留核对'


def test_same_employee_number_does_not_bind_another_hire_offer(tmp_path):
    store=DraftStore(tmp_path)
    ledger=book(['姓名','工号','入职日期'],[['旧任职','zt1','2025-01-01']],title='附件一-招聘奖金核算明细')
    draft=store.ingest(202607,[('历史.xlsx',ledger),('roster.xlsx',book(ROSTER,[['新任职','zt1','OFNEW','2026-07-01','正式','','','OEHR入职','德国','德国']])),('offer.xlsx',book(OFFER,[['OFNEW','RNEW','A','zt-owner','招聘网站','','已入职']]))],0)
    old=next(r for r in draft['rows'] if r['values'].get('姓名')=='旧任职')
    assert not old['offerMatched']
    assert not old['values'].get('ABC类别')


def test_duplicate_ledger_rows_keep_both_sources_and_raise_review(tmp_path):
    data=book(['姓名','工号','入职日期','招聘人入职1月奖金'],[['甲','zt1','2026-06-01',300],['甲','zt1','2026-06-01',300]],title='附件一-招聘奖金核算明细')
    draft=DraftStore(tmp_path).ingest(202607,[('历史.xlsx',data)],0)
    assert len(draft['rows'])==1
    assert len(draft['rows'][0]['history'])==2
    assert any(t['kind']=='duplicate_source_row' for t in draft['todos'])


@pytest.mark.parametrize('reverse', [False, True])
def test_rehire_offers_stay_with_their_employment_not_a_later_roster_row(tmp_path, reverse):
    ledger=book(['姓名','工号','入职日期','最后工作日','招聘启动日期','招聘负责人工号'],[
        ['甲','zt1','2025-08-29','2026-01-06','2025-08-11','zt-old'],
        ['甲','zt1','2026-03-16','','2026-03-12','zt-new'],
    ],title='附件一-招聘奖金核算明细')
    offers=[['OLD','R1','C','zt-old','招聘网站','','已入职','zt1','2025-08-11'],
            ['NEW','R2','C','zt-new','招聘网站','','已入职','zt1','2026-03-12']]
    if reverse: offers.reverse()
    draft=DraftStore(tmp_path).ingest(202607,[('历史.xlsx',ledger),
        ('人员.xlsx',book(ROSTER,[['甲','zt1','','2026-07-01','正式','','','EHR入职','中国大陆','国内']])),
        ('Offer.xlsx',book(OFFER+['工号','招聘启动日期'],offers))],0)
    rows={r['values']['入职日期']:r for r in draft['rows']}
    assert len(rows)==3
    for start,owner in [('2025-08-29','zt-old'),('2026-03-16','zt-new')]:
        assert rows[start]['values']['招聘负责人工号']==owner
        assert rows[start]['offerMatched']
        assert not rows[start]['conflicts']
        assert len(rows[start]['history'])==1
    extra=rows['2026-07-01']
    assert not extra['offerMatched']
    assert not extra['values'].get('招聘负责人工号')
    assert not extra['conflicts']
    assert not any(t['kind']=='employment_conflict' and rows['2025-08-29']['id'] in t['rowIds'] for t in draft['todos'])


def test_dated_offers_separate_rehires_but_same_employment_conflict_is_kept(tmp_path):
    draft=DraftStore(tmp_path).ingest(202607,[
        ('历史.xlsx',book(['姓名','工号','入职日期','招聘负责人工号'],[
            ['甲','zt1','2025-08-29','zt-old']],title='附件一-招聘奖金核算明细')),
        ('本月新增.xlsx',book(['姓名','工号','入职日期'],[['甲','zt1','2026-07-01']],title='导入_月度数据')),
        ('Offer.xlsx',book(OFFER+['工号','入职日期'],[
            ['OLD','R1','C','zt-old','招聘网站','','已入职','zt1','2025-08-29'],
            ['NEW','R2','C','zt-new','招聘网站','','已入职','zt1','2026-07-01'],
            ['NEW','R2','C','zt-other','招聘网站','','已入职','zt1','2026-07-01']]))],0)
    old,new=sorted(draft['rows'],key=lambda r:r['values']['入职日期'])
    assert old['values']['招聘负责人工号']=='zt-old'
    assert not old['conflicts']
    conflicts=[c for c in new['conflicts'] if c['field']=='招聘负责人工号']
    assert len(conflicts)==1
    assert set(conflicts[0]['values'])=={'zt-new','zt-other'}


def test_confirmed_employment_continuation_does_not_reopen_date_review(tmp_path):
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[
        ('历史.xlsx',book(['姓名','工号','入职日期'],[['甲','zt1','2026-03-16']],title='附件一-招聘奖金核算明细')),
        ('人员.xlsx',book(ROSTER,[['甲','zt1','','2026-07-01','正式','','','EHR入职','中国大陆','国内']]))],0)
    old=next(r for r in draft['rows'] if r['values']['入职日期']=='2026-03-16')
    assert any(t['kind']=='employment_conflict' and old['id'] in t['rowIds'] for t in draft['todos'])
    updated=store.change(202607,draft['revision'],{'edits':[{
        'rowId':old['id'],'field':'任职说明','value':'任职延续：七月实习转内部，三月招聘任职持续有效。','reason':'已核验任职历史'}]})
    assert not any(t['kind']=='employment_conflict' and old['id'] in t['rowIds'] for t in updated['todos'])
    assert next(r for r in updated['rows'] if r['id']==old['id'])['values']['入职日期']=='2026-03-16'


@pytest.mark.parametrize('confirmed,intern_count,expected',[(True,1,True),(False,1,False),(True,2,False)])
def test_intern_offer_matches_only_unique_verified_continuation(tmp_path,confirmed,intern_count,expected):
    store=DraftStore(tmp_path)
    headers=['工号','姓名','EmploymentForm','ApplicantPosition','招聘负责人工号',
             'LookupPrefix_RecruitmentRequirement_extzhiweifenlei_109025_933701443']
    offers=[['zt1','甲','实习生日薪','HRBP实习生','zt-old','C']]*intern_count
    offers.append(['zt1','甲','合同用工','高级专员','zt-new','C'])
    state=store.ingest(202607,[
        ('历史.xlsx',book(['姓名','工号','入职日期','职位','招聘负责人工号','转正日期'],
                         [['甲','zt1','2024-08-22','实习生','zt-old','2026-01-01']],title='附件一-招聘奖金核算明细')),
        ('花名册.xlsx',book(ROSTER,[['甲','zt1','','2025-07-01','正式','2026-01-01','','EHR入职','中国大陆','国内']])),
        ('录用.xlsx',book(headers,offers))],0)
    rid=state['rows'][0]['id']
    if confirmed:
        state=store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':'任职说明',
            'value':'任职延续：已核验实习转正式，沿用原招聘日期','reason':'已核验任职历史'}]})
    row=state['rows'][0]
    assert row['offerMatched'] is expected
    assert row['values']['招聘负责人工号']=='zt-old'
    assert row['values']['入职日期']=='2024-08-22'
    assert not any(c['field']=='招聘负责人工号' for c in row['conflicts'])
    assert any(t['kind']=='offer_employment' for t in state['todos']) is not expected
    if expected:
        assert row['employmentContext']['hrHireDates']==['2025-07-01']
        assert row['employmentContext']['recruitmentHireDate']=='2024-08-22'
        assert row['employmentContext']['probationDate']=='2026-01-01'
        assert row['offerMatchBasis']['sources'][0]['row']==2
        restored=store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':'任职说明','value':'','reason':'撤回任职核验'}]})
        assert not restored['rows'][0]['offerMatched']


def test_linked_continuation_resolves_both_records_without_changing_original_hire(tmp_path):
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[
        ('历史.xlsx',book(['姓名','工号','入职日期'],[['甲','zt1','2026-03-16']],title='附件一-招聘奖金核算明细')),
        ('人员.xlsx',book(ROSTER,[['甲','zt1','','2026-07-01','正式','','','EHR入职','中国大陆','国内']]))],0)
    old=next(r for r in draft['rows'] if r['values']['入职日期']=='2026-03-16')
    new=next(r for r in draft['rows'] if r['values']['入职日期']=='2026-07-01')
    updated=store.change(202607,draft['revision'],{'edits':[
        {'rowId':new['id'],'field':'任职说明','value':'任职延续：沿用原招聘任职'+old['id'],'reason':'已核验实习转内部'},
        {'rowId':new['id'],'field':'核算范围','value':'不纳入核算','reason':'已核验实习转内部'}]})
    assert not any(t['kind']=='employment_conflict' and old['id'] in t['rowIds'] for t in updated['todos'])
    restored=store.change(202607,updated['revision'],{'edits':[
        {'rowId':new['id'],'field':'核算范围','value':'纳入核算','reason':'撤回此前任职延续处理'}]})
    assert any(t['kind']=='employment_conflict' and old['id'] in t['rowIds'] for t in restored['todos'])


@pytest.mark.parametrize('status,entry,last,expected',[
    ('已废弃','2026-06-01','2026-07-15',False),
    ('已驳回','2026-06-01','2026-07-15',False),
    ('已撤回','2026-06-01','2026-07-15',False),
    ('审批中','2025-06-01','2026-07-15',False),
    ('审批中','2026-06-01','2026-08-15',False),
    ('审批中','2026-06-01','2026-07-15',True),
])
def test_only_applicable_pending_resignation_needs_review(tmp_path,status,entry,last,expected):
    draft=DraftStore(tmp_path).ingest(202607,[
        ('历史.xlsx',book(['姓名','工号','入职日期'],[['甲','zt1','2026-06-01']],title='附件一-招聘奖金核算明细')),
        ('离职.xlsx',book(['工号','入职日期','最后工作日','离职类型','审批状态'],[['zt1',entry,last,'主动离职',status]]))],0)
    assert any(t['kind']=='resignation_review' for t in draft['todos']) is expected
    assert not draft['rows'][0]['values'].get('最后工作日')


def test_carry_does_not_treat_previous_snapshot_future_hires_as_current_coverage(tmp_path):
    store=DraftStore(tmp_path)
    raw=book(ROSTER,[['甲','zt1','OF1','2026-07-01','正式','','','OEHR入职','德国','德国'],['预入职','zt2','OF2','2026-08-01','待入职','','','OEHR入职','德国','德国']])
    draft=store.ingest(202607,[('roster.xlsx',raw)],0)
    august=store.carry_forward(202608,202607,0)
    assert len(august['rows'])==len(draft['rows'])==1
    assert any(t['kind']=='coverage' for t in august['todos'])


def test_concurrent_edits_do_not_lose_a_saved_revision(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('补充.xlsx',book(['体系外人员编号','姓名','入职日期'],[['a','甲','2026-07-01']]))],0)
    rid=draft['rows'][0]['id']
    def edit(value):
        try:
            return store.change(202607,draft['revision'],{'edits':[{'rowId':rid,'field':'人工备注','value':value,'reason':'并发核对'}]})['revision']
        except DraftConflict: return 'conflict'
    with ThreadPoolExecutor(max_workers=2) as executor: result=list(executor.map(edit,['甲修改','乙修改']))
    assert sorted(map(str,result))==['2','conflict']
    assert len(store.read(202607)['edits'])==1


def test_bad_workbook_does_not_discard_other_uploaded_sources(tmp_path):
    draft=DraftStore(tmp_path).ingest(202607,[('好资料.xlsx',book(['体系外人员编号','姓名'],[['a','甲']])),('损坏.xlsx',b'not a workbook')],0)
    assert len(draft['rows'])==1 and len(draft['sources'])==2
    assert any(t['kind']=='source_error' for t in draft['todos'])


def test_category_aliases_and_confirmed_domestic_scope_survive_overseas_organization(tmp_path):
    store=DraftStore(tmp_path)
    data=book(['体系外人员编号','姓名','工号','入职日期','ABC类别','二级组织','工作地','标签分类'],[['a','甲','zt1','2026-07-01','C类','德国团队','中国大陆','国内']])
    draft=store.ingest(202607,[('补充.xlsx',data),('roster.xlsx',book(ROSTER,[['甲','zt1','OF1','2026-07-01','正式','','','EHR入职','中国大陆','德国团队']])),('offer.xlsx',book(OFFER,[['OF1','R1','C','zt-owner','招聘网站','','已入职']]))],0)
    row=draft['rows'][0]
    assert row['values']['标签分类']=='国内'
    assert row['values']['奖金地区类型']=='国内发展中国家'
    assert row['calculation']['recruitment']['currency']=='人民币'
    assert row['calculation']['referral']['currency']=='人民币'
    assert row['values']['ABC类别']=='C'
    assert 'ABC类别' not in row['unresolvedFields']
    assert row['provenance']['ABC类别'][0]['rawValue']=='C类'


def test_identity_edit_cannot_silently_merge_two_saved_records(tmp_path):
    store=DraftStore(tmp_path)
    draft=store.ingest(202607,[('补充.xlsx',book(['体系外人员编号','姓名','工号','入职日期'],[['a','甲','zt1','2026-07-01'],['b','乙','zt2','2026-07-01']]))],0)
    rid=next(r['id'] for r in draft['rows'] if r['values']['工号']=='zt2')
    with pytest.raises(ValueError,match='同一任职'):
        store.change(202607,draft['revision'],{'edits':[{'rowId':rid,'field':'工号','value':'zt1','reason':'核对身份'}]})
    assert len(store.read(202607)['rows'])==2
    assert store.read(202607)['revision']==draft['revision']
