from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from test_recruitment_assignments import setup
from bonus_platform.engine import recruitment_special as special
from bonus_platform.engine.recruitment_awards import calculate_draft
from bonus_platform.engine.recruitment_attention import project as attention
from bonus_platform.engine.recruitment_monthly import project, export
from bonus_platform.engine.recruitment_drafts import DraftConflict


def fixture(tmp_path, monkeypatch):
    store, state = setup(tmp_path)
    import json
    source_path=store._folder(202607)/'parsed'/(state['sources'][0]['id']+'.json')
    parsed=json.loads(source_path.read_text())
    for record in parsed['records']:record['values']['奖金地区类型']='国内发展中国家'
    store._write(source_path,parsed)
    state=store._save(store._folder(202607),store._rebuild(store._folder(202607),state))
    from bonus_platform.engine.recruitment_draft_sources import PARSER_VERSION
    state['sources'].append({'id':'approval','name':'批准原件.pdf','active':True,'attachment':True,'parserVersion':PARSER_VERSION})
    store._write(store._folder(202607)/'parsed'/'approval.json',{'records':[]})
    store._write(store._folder(202607)/'draft.json', state)
    monkeypatch.setattr(special, 'evidence_check', lambda *a, **k: {
        'sourceId':'approval','status':'已认可','scopeChecked':True,'pageNumbers':[1],
        'note':'批准第1条','attachmentIds':[]})
    row = state['rows'][0]
    payload = {'revision':state['revision'],'type':'adjustment','rowId':row['id'],
        'kind':'recruitment','role':'招聘负责人','nodeId':row['calculation']['recruitment']['nodes'][0]['id'],
        'recipientId':'zt-new','recipientName':'新领取人','amount':25.5,'currency':'人民币',
        'month':202607,'lineRef':'第1条','evidence':{'sourceId':'approval'}}
    return store, state, payload


def apply(store, state, payload):
    p=special.preview(store, state['month'], payload)
    return p,special.apply_preview(store, state['month'], state['revision'], p['id'])


def test_adjustment_preview_export_idempotency_and_undo(tmp_path, monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch)
    p=special.preview(store,202607,payload)
    assert store.read(202607)==state
    assert any(x['after']-x['before']==25.5 for x in p['currentImpact'])
    saved=special.apply_preview(store,202607,state['revision'],p['id'])
    assert special.apply_preview(store,202607,state['revision'],p['id'])==saved
    monthly=project(saved,attention(saved))
    recipient=next(x for x in monthly['recipientSummaries']['recruitment']['rows'] if x['name']=='新领取人')
    assert recipient['supplement']==25.5 and recipient['confirmed']==25.5
    wb=load_workbook(BytesIO(export(saved,monthly)))
    assert len(wb.sheetnames)==3
    assert any('补发' in str(c.value) for row in wb.worksheets[1] for c in row)
    with pytest.raises(ValueError,match='重复'):
        special.preview(store,202607,{**payload,'revision':saved['revision']})
    undone=store.history_action(202607,saved['revision'],'undo')
    assert not any(n.get('adjustmentId') for r in undone['rows'] for n in r['calculation']['recruitment']['nodes'])


def test_adjustment_rechecks_changed_facts_and_inactive_evidence(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch)
    _,saved=apply(store,state,payload)
    changed=deepcopy(saved);changed['rows'][0]['values']['ABC类别']='A';calculate_draft(changed)
    node=next(n for n in changed['rows'][0]['calculation']['recruitment']['nodes'] if n.get('adjustmentId'))
    assert node['issues'] and node['status']=='安排需重新核对'
    changed=deepcopy(saved);next(s for s in changed['sources'] if s['id']=='approval')['active']=False;calculate_draft(changed)
    node=next(n for n in changed['rows'][0]['calculation']['recruitment']['nodes'] if n.get('adjustmentId'))
    assert node['issues'] and node.get('scheduledMonth') is None


def test_cutoff_defaults_month_scope_and_rollback(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch)
    p={'revision':state['revision'],'type':'cutoff','kind':'referral','mode':'due','month':202607,
       'sourceId':state['sources'][0]['id'],'reason':'核算口径第3条'}
    _,saved=apply(store,state,p)
    assert special.cutoff(saved,'referral')['mode']=='due'
    future=deepcopy(saved);future['month']=202608
    assert special.cutoff(future,'referral')['mode']=='month_end'
    undone=store.history_action(202607,saved['revision'],'undo')
    assert special.cutoff(undone,'referral')['mode']=='month_end'


def test_exception_is_node_condition_and_month_scoped(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch)
    row=state['rows'][0];row['values']['招聘负责人最后工作日']='2026-01-01'
    state.setdefault('edits',[]).append({'id':'fact','rowId':row['id'],'field':'招聘负责人最后工作日','value':'2026-01-01','active':True,'reason':'原始离职材料','at':'2026-07-01'})
    store._write(store._folder(202607)/'draft.json',calculate_draft(state))
    p={k:v for k,v in payload.items() if k in ('revision','rowId','kind','nodeId','month','evidence','lineRef')}
    p.update(type='exception',conditions=['recipient_not_employed'])
    _,saved=apply(store,state,p)
    node=next(n for n in saved['rows'][0]['calculation']['recruitment']['nodes'] if n['id']==p['nodeId'])
    assert node['scheduledMonth']==202607 and node['exceptionId']
    future=deepcopy(saved);future['month']=202608;calculate_draft(future)
    assert not any(n.get('exceptionId') for n in future['rows'][0]['calculation']['recruitment']['nodes'])


def test_stale_and_bad_inputs_rejected(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch)
    with pytest.raises(DraftConflict):special.preview(store,202607,{**payload,'revision':-1})
    for patch in ({'amount':float('nan')},{'month':202606},{'lineRef':''},{'kind':'invalid'}):
        with pytest.raises(ValueError):special.preview(store,202607,{**payload,**patch})


def test_future_deduction_carries_once_and_revocation_is_undoable(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch)
    # The source fixture is a parsed empty approval; carry-forward needs its stored name.
    state['sources'][-1]['storedName']='approval.pdf'
    store._write(store._folder(202607)/'draft.json',state)
    _,saved=apply(store,state,{**payload,'amount':-12.75,'month':202608})
    july=project(saved,attention(saved))
    assert not any(n.get('adjustmentId') and n['scope']=='current' for n in july['nodes'])
    august=store.carry_forward(202608,202607,0)
    monthly=project(august,attention(august))
    recipient=next(r for r in monthly['recipientSummaries']['recruitment']['rows'] if r['name']=='新领取人')
    assert recipient['deduction']==-12.75 and recipient['confirmed']==-12.75
    record=special.context(august)['records'][0]
    p=special.preview(store,202608,{'revision':august['revision'],'type':'revoke','recordId':record['id']})
    revoked=special.apply_preview(store,202608,august['revision'],p['id'])
    assert not any(n.get('adjustmentId') for r in revoked['rows'] for n in r['calculation']['recruitment']['nodes'])
    restored=store.history_action(202608,revoked['revision'],'undo')
    assert sum(n['amount'] for r in restored['rows'] for n in r['calculation']['recruitment']['nodes'] if n.get('adjustmentId'))==-12.75
    september=store.carry_forward(202609,202608,0)
    assert not any(n.get('adjustmentId') and n['scope']=='current' for n in project(september,attention(september))['nodes'])


def test_held_adjustment_stays_visible_as_pending_not_confirmed(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch);_,saved=apply(store,state,payload)
    next(s for s in saved['sources'] if s['id']=='approval')['active']=False;calculate_draft(saved)
    monthly=project(saved,attention(saved))
    recipient=next(r for r in monthly['recipientSummaries']['recruitment']['rows'] if r['name']=='新领取人')
    assert recipient['confirmed']==0 and recipient['pending']==25.5
    assert monthly['worksheet']['rows'][0]['review']
    assert not saved['ruleValidation']['passed']


def test_cutoff_changes_boundary_person_only_and_referrer_exception_is_specific():
    from test_recruitment_awards import person
    from bonus_platform.engine.recruitment_awards import calculate_row
    r=person(hire='2026-06-01',最后工作日='2026-07-10')
    stable=person(hire='2026-06-01');stable['id']='stable'
    state={'month':202607,'sources':[{'id':'basis','active':True}], 'operations':[], 'rows':[r,stable]}
    normal=calculate_row(state,r)['referral']
    assert any(n.get('eligibilityReasons') for n in normal['nodes'])
    old_stable=calculate_row(state,stable)['referral']
    state['operations']=[{'id':'policy','active':True}]
    state['cutoffPolicies']=[{'id':'policy','operationId':'policy','month':202607,'kind':'referral','mode':'due','sourceId':'basis'}]
    revised=calculate_row(state,r)['referral']
    assert not any(n.get('eligibilityReasons') for n in revised['nodes'])
    assert calculate_row(state,stable)['referral']['currentAmount']==old_stable['currentAmount']
    # A probation exemption must not also waive the candidate's employment condition.
    r['values'].update(推荐人转正日期='2026-09-01',推荐人人员状态='试用')
    state['cutoffPolicies']=[];award=calculate_row(state,r)['referral'];node=award['nodes'][0]
    state['eligibilityExceptions']=[{'id':'ex','operationId':'policy','rowId':r['id'],'kind':'referral','month':202607,
        'nodeId':node['id'],'conditions':['recipient_not_regularized'],'fingerprint':special.fingerprint(award),
        'sourceId':'basis','reason':'仅批准推荐人试用期例外'}]
    result=calculate_row(state,r)['referral']['nodes'][0]
    assert any(x.startswith('候选人') for x in result['eligibilityReasons'])
    assert not any('未转正' in x for x in result['eligibilityReasons'])
    assert result.get('scheduledMonth') is None


def test_review_withdrawal_and_no_early_exception(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch);_,saved=apply(store,state,payload)
    saved['reviewEntries'].append({'sourceId':'approval','status':'已驳回','scopeChecked':True})
    calculate_draft(saved)
    assert any(n.get('adjustmentId') and n['issues'] for n in saved['rows'][0]['calculation']['recruitment']['nodes'])
    with pytest.raises(ValueError,match='资格障碍'):
        special.preview(store,202607,{**payload,'revision':saved['revision'],'type':'exception','conditions':['candidate_not_employed']})


def test_real_pdf_api_preview_apply_retry_export_and_revoke(tmp_path,monkeypatch):
    import fitz
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from bonus_platform import recruitment_draft_api as api
    store,state=setup(tmp_path)
    doc=fitz.open();page=doc.new_page();page.insert_text((60,60),'APPROVED: recruitment supplement 25.50 CNY; no attachments. This document is synthetic evidence for isolated test acceptance only.')
    state=store.ingest(202607,[('approved.pdf',doc.tobytes())],state['revision'])
    monkeypatch.setattr(api,'store',store)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    url='/api/recruitment/drafts/202607'
    payload={'revision':state['revision'],'type':'adjustment','rowId':state['rows'][0]['id'],'kind':'recruitment',
        'role':'招聘负责人','recipientId':'qa-payee','recipientName':'验收领取人','amount':25.5,'currency':'人民币',
        'month':202607,'lineRef':'1','evidence':{'sourceId':state['sources'][-1]['id'],'pageNumbers':[1],
         'scopeChecked':True,'note':'approved test entry','status':'已认可'}}
    bad=deepcopy(payload);bad['evidence']['scopeChecked']=False
    assert client.post(url+'/special/preview',json=bad).status_code==400
    response=client.post(url+'/special/preview',json=payload);assert response.status_code==200,response.text
    token=response.json()['id'];path=url+'/special/previews/'+token+'/apply'
    response=client.post(path,json={'revision':state['revision']});assert response.status_code==200,response.text
    saved=response.json();retry=client.post(path,json={'revision':state['revision']})
    assert retry.status_code==200 and retry.json()['revision']==saved['revision']
    recipient=next(r for r in saved['monthly']['recipientSummaries']['recruitment']['rows'] if r['name']=='验收领取人')
    assert recipient['confirmed']==25.5 and recipient['supplement']==25.5
    export_response=client.get(url+'/monthly/export',params={'revision':saved['revision']});assert export_response.status_code==200
    wb=load_workbook(BytesIO(export_response.content));assert len(wb.worksheets)==3
    row=next(row for row in wb.worksheets[1].iter_rows(values_only=True) if '验收领取人' in row)
    headers=list(next(wb.worksheets[1].iter_rows(min_row=2,max_row=2,values_only=True)))
    assert row[headers.index('补发金额')]==25.5 and row[headers.index('合计发放')]==25.5
    detail=wb.worksheets[0];headers=[c.value for c in detail[2]]
    assert detail.cell(3,headers.index('招聘奖金补发')+1).value==25.5
    records=client.get(url+'/special').json();record=records['records'][0]
    p=client.post(url+'/special/preview',json={'revision':saved['revision'],'type':'revoke','recordId':record['id']}).json()
    response=client.post(url+'/special/previews/'+p['id']+'/apply',json={'revision':saved['revision']})
    assert response.status_code==200
    assert not any(r['name']=='验收领取人' for r in response.json()['monthly']['recipientSummaries']['recruitment']['rows'])


def test_new_exception_requires_rechecking_an_existing_supplement(tmp_path,monkeypatch):
    store,state,payload=fixture(tmp_path,monkeypatch)
    row=state['rows'][0];row['values']['招聘负责人最后工作日']='2026-01-01'
    state.setdefault('edits',[]).append({'id':'fact','rowId':row['id'],'field':'招聘负责人最后工作日',
        'value':'2026-01-01','active':True,'reason':'离职事实','at':'2026-07-01'})
    store._write(store._folder(202607)/'draft.json',calculate_draft(state))
    _,saved=apply(store,state,payload)
    p={'type':'exception','revision':saved['revision'],'rowId':row['id'],'kind':'recruitment',
        'nodeId':payload['nodeId'],'month':202607,'evidence':payload['evidence'],'lineRef':'批准资格第2条',
        'conditions':['recipient_not_employed']}
    _,updated=apply(store,saved,p)
    node=next(n for n in updated['rows'][0]['calculation']['recruitment']['nodes'] if n.get('adjustmentId'))
    assert node['issues'] and node['scheduledMonth'] is None
