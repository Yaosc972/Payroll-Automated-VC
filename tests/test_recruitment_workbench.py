from copy import deepcopy
from io import BytesIO

from openpyxl import load_workbook

from test_recruitment_awards import person
from test_recruitment_reviews import setup
from test_recruitment_drafts import book
from bonus_platform.engine.recruitment_awards import calculate_draft


def test_note_suggestion_uses_same_row_hierarchy_then_normal_rule(tmp_path):
    from bonus_platform.engine.recruitment_actions import referral_suggestion
    from bonus_platform.engine.recruitment_drafts import DraftStore
    from bonus_platform.engine.recruitment_reviews import prepare_preview,apply_preview
    values=person(招聘渠道='内推',推荐人工号='',推荐人姓名='',人工备注='内推人为甲经理，业务确认不需要发内推奖金',
                  间接上级2姓名='甲经理',间接上级2工号='zt-boss')['values']
    store=DraftStore(tmp_path)
    state=store.ingest(202607,[('台账.xlsx',book(list(values),[list(values.values())],title='附件一-招聘奖金核算明细'))],0)
    row=state['rows'][0];before=deepcopy(row['calculation']['recruitment'])
    suggestion=referral_suggestion(row,state)
    assert suggestion and [e['value'] for e in suggestion['edits']]==['甲经理','zt-boss']
    assert not row['values'].get('推荐人工号')
    assert row['calculation']['referral']['total'] is None
    preview=prepare_preview(store,202607,state['revision'],suggestion['edits'])
    updated=apply_preview(store,202607,state['revision'],preview['id'])
    changed=updated['rows'][0]
    assert changed['calculation']['referral']['total']==0
    assert '推荐人与候选人属于上下级汇报关系' in changed['calculation']['referral']['reasons']
    assert changed['calculation']['recruitment']==before
    assert referral_suggestion(changed,updated) is None
    # Arbitrary notes, negative/conditional wording, comparison-only files,
    # conflicting manager IDs and unrelated source rows cannot supply facts.
    for note in ['领导说不发内推','内推人不是甲经理','如果内推人为甲经理，不发内推']:
        other=deepcopy(row);other['values']['人工备注']=note
        assert referral_suggestion(other,state) is None
    other=deepcopy(row);other['provenance']['间接上级2工号'][0]['row']+=1
    assert referral_suggestion(other,state) is None
    other=deepcopy(row);other['unresolvedFields']=['间接上级2工号']
    assert referral_suggestion(other,state) is None
    comparison=deepcopy(state);comparison['sources'][0]['comparisonOnly']=True
    assert referral_suggestion(row,comparison) is None


def test_complete_result_is_one_employment_with_all_recipient_installments(tmp_path):
    from bonus_platform.engine.recruitment_worksheet import worksheet, export_results
    store, state, rid = setup(tmp_path)
    data = worksheet(state)
    assert len(data['rows']) == 1
    cells = data['rows'][0]['cells']
    assert cells['姓名'] and cells['招聘启动日期']
    assert cells['recruitment.standard'] == 55
    assert cells['recruitment.total'] == 55
    assert cells['recruitment.招聘负责人.0.amount'] == 27.5
    assert cells['referral.total'] == 400
    wb = load_workbook(BytesIO(export_results(state)), data_only=True)
    assert len(wb.sheetnames) == 5
    ws = wb['核算总明细']
    assert ws.max_row == 4
    assert ws.max_column > 80
    assert ws.freeze_panes == 'D4'
    assert not any(isinstance(c.value, str) and c.value.startswith('{') for row in ws for c in row)
    assert '招聘负责人' in ' '.join(str(c.value) for c in ws[2])


def test_task_opens_the_actual_conflicting_values_and_source_locations(tmp_path):
    from bonus_platform.engine.recruitment_actions import task_detail
    from bonus_platform.engine.recruitment_attention import project
    store, state, rid = setup(tmp_path)
    value = person(招聘渠道='校招')['values']; value['体系外人员编号'] = 'test'
    state = store.ingest(202607, [('另一资料.xlsx', book(list(value), [list(value.values())]))], state['revision'])
    group = next(t for scope in ('current', 'history') for t in project(state)[scope] if t['kind'] == 'field:招聘渠道')
    data = task_detail(state, group['id'])
    entry = data['rows'][0]
    field = next(f for f in entry['fields'] if f['field'] == '招聘渠道')
    assert {v['value'] for v in field['options']} >= {'招聘网站', '校招'}
    assert all(o['sources'][0]['file'] and o['sources'][0]['row'] for o in field['options'])
    assert data['instruction'] and entry['impact']
    saved = store.change(202607, state['revision'], {'edits': [{'rowId': rid, 'field': '招聘渠道', 'value': '招聘网站', 'reason': '核对原资料后采用'}]})
    assert not any(t['kind'] == group['kind'] for scope in ('current', 'history') for t in project(saved)[scope])


def test_reference_matching_without_employee_number_requires_multiple_facts():
    from bonus_platform.engine.recruitment_reconcile import match_reference
    row = {'id': 'a', 'values': {'姓名': '候选人甲', '入职日期': '2026-07-01', '职位': '专员', '二级组织': '国内业务', '招聘启动日期': '2026-06-01'}}
    assert match_reference(row['values'], [row]) == row
    other = deepcopy(row); other['id'] = 'b'
    assert match_reference(row['values'], [row, other]) is None
    assert match_reference({'姓名': '候选人甲', '入职日期': '2026-07-01'}, [row]) is None


def test_employment_review_shows_unmatched_roster_date_without_adopting_it(tmp_path):
    from bonus_platform.engine.recruitment_drafts import DraftStore
    from bonus_platform.engine.recruitment_actions import task_detail
    from bonus_platform.engine.recruitment_attention import project
    from test_recruitment_drafts import ROSTER
    state=DraftStore(tmp_path).ingest(202607,[
        ('历史.xlsx',book(['姓名','工号','入职日期','候选人入职时间'],[['甲','zt1','2024-08-22','2024-08-22']],title='附件一-招聘奖金核算明细')),
        ('花名册.xlsx',book(ROSTER,[['甲','zt1','','2025-07-01','正式','','','EHR入职','中国大陆','国内']]))],0)
    task=next(t for scope in ('current','history') for t in project(state)[scope] if t['kind']=='employment')
    review=task_detail(state,task['id'])['rows'][0]['employmentReview']
    assert review['issues']==['台账与花名册入职日不同']
    assert [(o['label'],o['value']) for o in review['observations']]==[('台账入职日','2024-08-22'),('花名册入职日','2025-07-01')]
    assert review['observations'][1]['sources'][0]['file']=='花名册.xlsx'
    assert review['observations'][1]['sources'][0]['row']==2
    assert state['rows'][0]['values']['入职日期']=='2024-08-22'
    assert state['rows'][0]['values']['候选人入职时间']=='2024-08-22'


def test_employment_review_distinguishes_duplicate_rows_from_date_conflict(tmp_path):
    from bonus_platform.engine.recruitment_drafts import DraftStore
    from bonus_platform.engine.recruitment_actions import task_detail
    from bonus_platform.engine.recruitment_attention import project
    state=DraftStore(tmp_path).ingest(202607,[('台账.xlsx',book(['姓名','工号','入职日期'],[
        ['甲','zt1','2026-07-01'],['甲','zt1','2026-07-01']],title='附件一-招聘奖金核算明细'))],0)
    task=next(t for scope in ('current','history') for t in project(state)[scope] if t['kind']=='employment')
    review=task_detail(state,task['id'])['rows'][0]['employmentReview']
    assert review['issues']==['同一台账有重复任职行']
    assert {o['sources'][0]['row'] for o in review['observations'] if o['label']=='台账记录'}=={2,3}


def test_ledger_reader_keeps_upper_header_monthly_adjustment_columns(tmp_path):
    from openpyxl import Workbook
    from bonus_platform.engine.recruitment_draft_sources import read_source
    wb = Workbook(); ws = wb.active; ws.title = '附件一-招聘奖金核算明细'
    ws.append(['', '', '', '', '应发奖金', '招聘负责人应扣减'])
    ws.append(['姓名', '工号', '入职日期', '招聘负责人入职奖金明细', '', ''])
    ws.append(['', '', '', '入职1个月发放奖金', '', 202607])
    ws.append(['甲', 'zt1', '2026-07-01', 30, 60, 10])
    path = tmp_path / '台账.xlsx'; wb.save(path)
    result = read_source(path, 'source')
    rec = result['records'][0]
    assert any(c['column'] == 5 and c['value'] == 60 for c in rec['historyCells'])
    assert any(c['column'] == 6 and '招聘负责人应扣减' in c['label'] for c in rec['historyCells'])


def test_old_offer_for_another_recruitment_does_not_contaminate_current_facts(tmp_path):
    from bonus_platform.engine.recruitment_drafts import DraftStore
    headers=['工号','姓名','招聘启动日期','招聘渠道','招聘负责人工号','LookupPrefix_RecruitmentRequirement_extzhiweifenlei_109025_933701443']
    ledger=book(['工号','姓名','入职日期','招聘启动日期','招聘渠道','招聘负责人工号','ABC类别'],[['zt1','甲','2026-06-11','2026-05-29','二次入职','zt2','C']],title='附件一-招聘奖金核算明细')
    offers=book(headers,[['zt1','甲','2026-05-29','二次入职','zt2','C'],['zt1','甲','2024-02-23','招聘网站','zt3','B']])
    state=DraftStore(tmp_path).ingest(202607,[('台账.xlsx',ledger),('录用资料.xlsx',offers)],0)
    row=state['rows'][0]
    assert not {'招聘渠道','招聘负责人工号','ABC类别','招聘启动日期'} & set(row['unresolvedFields'])
    assert {r['row'] for r in row['sources'] if r['sheet']=='Sheet1'}=={2}


def test_two_offers_in_same_recruitment_keep_real_conflict(tmp_path):
    from bonus_platform.engine.recruitment_drafts import DraftStore
    headers=['工号','姓名','招聘启动日期','招聘渠道','招聘负责人工号','LookupPrefix_RecruitmentRequirement_extzhiweifenlei_109025_933701443']
    ledger=book(['工号','姓名','入职日期','招聘启动日期'],[['zt1','甲','2026-07-01','2026-05-29']],title='导入_月度数据')
    offers=book(headers,[['zt1','甲','2026-05-29','招聘网站','zt2','C'],['zt1','甲','2026-05-29','招聘网站','zt3','C']])
    state=DraftStore(tmp_path).ingest(202607,[('台账.xlsx',ledger),('录用资料.xlsx',offers)],0)
    assert '招聘负责人工号' in state['rows'][0]['unresolvedFields']


def test_monthly_comparison_counts_amounts_dates_and_missing_ledger_nodes(tmp_path):
    from bonus_platform.engine.recruitment_reconcile import monthly_comparison
    store,state,rid=setup(tmp_path)
    row=state['rows'][0]
    source=state['sources'][0]
    rec={'values':row['values'],'ref':{'sheet':'台账','row':2},'historyCells':[
        {'column':39,'label':'招聘负责人入职奖金明细 / 入职1个月发放奖金','value':99},
        {'column':40,'label':'招聘负责人入职奖金明细 / 入职1个月发放周期','value':202607},
        {'column':41,'label':'招聘负责人入职奖金明细 / 入职3个月发放奖金','value':999},
        {'column':42,'label':'招聘负责人入职奖金明细 / 入职3个月发放周期','value':202606},
    ]}
    before=deepcopy(state)
    result=monthly_comparison(state,[rec],source)
    assert state==before
    assert result['summary']['referenceNodes']==1
    assert result['summary']['differentAmounts']==1
    assert result['summary']['draftOnlyNodes']>0
    assert result['differences'][0]['calculated']==27.5
    assert all(d['reference']!=999 for d in result['differences'])
    assert not any(n['paymentReady'] for a in row['calculation'].values() if isinstance(a,dict) for n in a.get('nodes',[]))


def test_monthly_difference_explains_approved_exemption_without_missing_day_error(tmp_path):
    from bonus_platform.engine.recruitment_reconcile import monthly_comparison
    store,state,rid=setup(tmp_path);row=state['rows'][0]
    row['values'].update(标签分类='国内',奖金地区类型='国内发展中国家',管培生已批名单资格='是',周期豁免依据='认可名单')
    calculate_draft(state)
    rec={'values':row['values'],'ref':{'sheet':'台账','row':2},'historyCells':[
        {'column':39,'label':'招聘负责人入职奖金明细 / 入职1个月发放奖金','value':99},
        {'column':40,'label':'招聘负责人入职奖金明细 / 入职1个月发放周期','value':202607}]}
    result=monthly_comparison(state,[rec],state['sources'][0])
    assert '周期豁免' in result['differences'][0]['reason']


def test_person_eligibility_task_has_actionable_fields(tmp_path):
    from bonus_platform.engine.recruitment_actions import task_fields
    store,state,rid=setup(tmp_path)
    row=state['rows'][0]
    row['calculation']['recruitment']['nodes'][0]['issues']=['招聘负责人：适用日期在职状态待核']
    assert '招聘负责人人员状态' in task_fields(row,{'kind':'recipient'},202607)


def test_scope_decision_is_explicit_and_does_not_fake_payment(tmp_path):
    import pytest
    store,state,rid=setup(tmp_path)
    before=deepcopy(state['rows'][0]['history'])
    with pytest.raises(ValueError):
        store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':'核算范围','value':'随便','reason':'核对范围'}]})
    state=store.change(202607,state['revision'],{'edits':[{'rowId':rid,'field':'核算范围','value':'不纳入核算','reason':'已核对本次任职范围'}]})
    assert state['rows'][0]['history']==before
    assert state['rows'][0]['calculation']['recruitment']['total']==0
    assert state['rows'][0]['calculation']['recruitment']['formulaTotal']==55
    assert state['paymentReady'] is False


def test_monthly_reference_currency_is_independent_and_czech_referral_is_separate(tmp_path):
    from bonus_platform.engine.recruitment_reconcile import monthly_comparison,reference_currency
    store,state,rid=setup(tmp_path);row=state['rows'][0]
    rec={'values':{**row['values'],'标准币种':'人民币'},'ref':{'sheet':'台账','row':2},'historyCells':[
        {'column':39,'label':'招聘负责人入职奖金明细 / 入职1个月发放奖金','value':27.5},
        {'column':40,'label':'招聘负责人入职奖金明细 / 入职1个月发放周期','value':202607},
        {'column':70,'label':'内推奖明细 / 另一半奖金','value':'内推奖币种：捷克郎'}]}
    assert reference_currency(rec,'recruitment')=='人民币'
    assert reference_currency(rec,'referral')=='捷克克朗'
    result=monthly_comparison(state,[rec],state['sources'][0]);item=result['differences'][0]
    assert item['reference']==item['calculated']==27.5
    assert item['category']=='币种不同'
    assert item['currency']=='人民币' and item['calculatedCurrency']=='欧元'
    assert result['summary'].get('equalNodes',0)==0
    assert next(t for t in result['totalsByCurrency'] if t['currency']=='人民币')['reference']==27.5


def test_confirmed_domestic_scope_uses_domestic_standards_in_both_rule_versions(tmp_path):
    from bonus_platform.engine.recruitment_drafts import DraftStore
    for hire in ('2026-06-01', '2026-07-01'):
        for country in ('德国', '捷克'):
            values = person(hire, 标签分类='国内', 工作地='中国大陆', 二级组织='FBU仓储事业部',
                            三级组织='欧洲区', 四级组织=country+'区', 特殊地区规则='FBU'+country,
                            奖金地区类型='海外发达国家', 体系外人员编号='domestic')['values']
            store = DraftStore(tmp_path / (hire+country))
            state = store.ingest(202607, [('人员.xlsx', book(list(values), [list(values.values())]))], 0)
            row = state['rows'][0]
            assert row['values']['标签分类'] == '国内'
            assert row['values']['奖金地区类型'] == '国内发展中国家'
            assert row['rule']['specialScope'] == ''
            assert row['calculation']['recruitment']['standard'] == 150
            assert row['calculation']['referral']['standard'] == 600
            for kind in ('recruitment', 'referral'):
                assert row['calculation'][kind]['currency'] == '人民币'
                assert all(n['currency'] == '人民币' for n in row['calculation'][kind]['nodes'])
            assert row['original']['特殊地区规则'] == 'FBU'+country
            assert row['original']['奖金地区类型'] == '海外发达国家'
            revision = state['revision']
            state = store.change(202607, revision, {'edits':[{'rowId':row['id'], 'field':'人工备注', 'value':'核对保留', 'reason':'保存地区依据'}]})
            assert state['rows'][0]['id'] == row['id']
            assert state['rows'][0]['calculation']['referral']['standard'] == 600


def test_domestic_priority_does_not_expand_to_unconfirmed_or_overseas_people(tmp_path):
    from bonus_platform.engine.recruitment_drafts import DraftStore
    for label, location in [('海外','中国大陆'), ('','中国大陆'), ('国内','德国'), ('国内','中国香港')]:
        values = person(标签分类=label, 工作地=location, 二级组织='德国区', 体系外人员编号='scope')['values']
        state = DraftStore(tmp_path / (label+location)).ingest(202607, [('人员.xlsx',book(list(values),[list(values.values())]))],0)
        row = state['rows'][0]
        assert row['values']['标签分类'] == '海外'
        assert row['calculation']['recruitment']['currency'] == '欧元'
        assert row['calculation']['referral']['standard'] == 400


def test_legacy_template_entry_also_ignores_overseas_special_scope_for_domestic_people():
    from bonus_platform.engine.calculator import _referral_scope, _normalized_referral_region
    from bonus_platform.engine.models import ImportRow
    from bonus_platform.engine.recruitment_import_validation import _normalized_special_region_rule
    values = {'标签分类':'国内', '工作地':'中国大陆', '二级组织':'FBU仓储事业部', '四级组织':'捷克区',
              '推荐人工号':'zt-ref', '特殊地区规则':'FBU捷克'}
    assert _normalized_special_region_rule(values) == ''
    row = ImportRow(source_row=2, values=values)
    assert _referral_scope(row) == '集团统一'
    assert _normalized_referral_region(row) == '国内发展中国家'
    assert _referral_scope(ImportRow(source_row=2, values={**values,'标签分类':'海外','工作地':'捷克'})) == 'FBU捷克'
