from copy import deepcopy
from io import BytesIO
import json
from datetime import datetime
import zipfile
import pytest
from openpyxl import Workbook, load_workbook
from bonus_platform.engine.recruitment_drafts import DraftStore, DraftConflict
from bonus_platform.engine import recruitment_external_results as external
from bonus_platform.engine.recruitment_attention import project as attention
from bonus_platform.engine.recruitment_monthly import project as monthly, export
from bonus_platform.engine.recruitment_settlement import capture_confirmed_installments
from test_recruitment_awards import person
from test_recruitment_drafts import book as facts_book

HEADERS=['被招聘人','奖金项目','领取人','本月核定金额','币种','人员编号','入职日期','领取人编号','核算月份','发放状态','金额类型']

def book(rows, headers=HEADERS, title='体系外核算结果'):
    wb=Workbook();wb.properties.created=datetime(2026,1,1);wb.properties.modified=datetime(2026,1,1)
    ws=wb.active;ws.title=title;ws.append(headers)
    for row in rows:ws.append(row)
    output=BytesIO();wb.save(output);stable=BytesIO()
    with zipfile.ZipFile(output) as source,zipfile.ZipFile(stable,'w') as target:
        for filename in source.namelist():
            info=zipfile.ZipInfo(filename,date_time=(2026,1,1,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED
            target.writestr(info,source.read(filename))
    return stable.getvalue()

def upload(store, rows, revision=0):
    return external.prepare_import(store,202608,revision,[('体系外结果.xlsx',book(rows))])

def adopt(store, preview):
    return external.apply_import(store,202608,preview['revision'],preview['token'])[0]

def view(state):return monthly(state,attention(state))

def test_external_result_is_final_amount_without_internal_facts(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['体系外甲','招聘负责人奖金','体系外乙',123.45,'CNY']])
    assert preview['readyCount']==1
    state=adopt(store,preview);result=view(state)
    assert result['totals']==[{'currency':'人民币','confirmed':123.45,'pending':0.0,'count':1}]
    assert not attention(state)['current']
    assert not result['coverage']['syncRequired']
    assert result['nodes'][0]['externalResultId']
    assert capture_confirmed_installments(state)==[]
    workbook=load_workbook(BytesIO(export(state,result)),data_only=True)
    assert len(workbook.sheetnames)==3

def test_repeated_upload_and_apply_do_not_add_money_or_revision(tmp_path):
    store=DraftStore(tmp_path);data=[['甲','招聘负责人奖金','乙',100,'CNY']]
    preview=upload(store,data);state=adopt(store,preview);revision=state['revision']
    repeated=upload(store,data,revision)
    assert repeated['unchangedCount']==1 and repeated['readyCount']==0
    assert repeated['revision']==revision
    again,_=external.apply_import(store,202608,revision,preview['token'])
    assert again['revision']==revision and len(again['externalResults'])==1
    assert view(again)['totals'][0]['confirmed']==100

def test_correction_replaces_recipient_currency_and_all_components_undo_restores(tmp_path):
    store=DraftStore(tmp_path)
    first=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY'],['甲','招聘负责人奖金','乙',20,'CNY','','','','','','补发']])
    state=adopt(store,first)
    corrected=upload(store,[['甲','招聘负责人奖金','丙',80,'EUR']],state['revision'])
    assert corrected['rows'][0]['change']=='replacement'
    state=adopt(store,corrected)
    assert [(n['recipientName'],n['currency'],n['amount']) for n in view(state)['nodes']]==[('丙','欧元',80)]
    state=store.history_action(202608,state['revision'],'undo')
    assert sorted(n['amount'] for n in view(state)['nodes'])==[20,100]
    state=store.history_action(202608,state['revision'],'redo')
    assert [n['amount'] for n in view(state)['nodes']]==[80]

def test_invalid_latest_source_does_not_restore_old_amount_or_automatic_result(tmp_path):
    store=DraftStore(tmp_path)
    state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]))
    state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',80,'CNY']],state['revision']))
    source=next(s for s in state['sources'] if s['id']==state['externalResults'][-1]['sourceId'])
    source['active']=False
    store._write(store._folder(202608)/'draft.json',state)
    state=store.refresh(202608,state['revision'])
    assert not any(n['amount'] for n in view(state)['nodes'] if n['scope']=='current')
    assert any('失效' in t['title'] for t in attention(state)['current'])

def test_blank_zero_hold_and_deduction_are_distinct_and_errors_survive_reopen(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['缺金额','招聘负责人奖金','乙',None,'CNY'],['不发','招聘负责人奖金','乙',0,'CNY'],
                          ['暂缓','招聘负责人奖金','乙',90,'CNY','','','','','暂缓'],
                          ['扣回','招聘负责人奖金','乙',-15,'CNY']])
    assert preview['blockedCount']==1 and preview['readyCount']==3
    state=adopt(store,preview)
    assert view(state)['totals'][0]['confirmed']==-15
    assert any(n['scope']=='held' and n['amount']==90 for n in view(state)['nodes'])
    pending=external.context(store,state)['pendingImports']
    assert len(pending)==1 and pending[0]['blockedCount']==1
    reopened=external.repreview(store,202608,state['revision'],preview['token'],[])
    assert reopened['blockedCount']==1 and reopened['appliedCount']==3

def test_carry_maps_people_but_requires_new_month_amount(tmp_path):
    store=DraftStore(tmp_path)
    state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]))
    carried=store.carry_forward(202609,202608,0)
    assert len(carried['externalSubjects'])==1
    assert not any(n['scope']=='current' and n['amount'] for n in view(carried)['nodes'])
    assert len([t for t in attention(carried)['current'] if t['kind']=='external_results_missing'])==1

def test_ambiguous_employment_requires_selection_and_other_rows_can_adopt(tmp_path):
    store=DraftStore(tmp_path);state=store.read(202608)
    state=store.ingest(202608,[('任职.xlsx',facts_book(['体系外人员编号','姓名','入职日期'],[['same','同名','2025-01-01'],['same','同名','2026-07-01']]))],0)
    second=next(r['id'] for r in state['rows'] if r['values']['入职日期']=='2026-07-01')
    preview=upload(store,[['同名','招聘负责人奖金','乙',100,'CNY'],['新外部','内推奖金','丙',20,'USD']],state['revision'])
    blocked=next(r for r in preview['rows'] if r['status']=='blocked')
    assert len(blocked['candidates'])==2
    state=adopt(store,preview)
    resolved=external.repreview(store,202608,state['revision'],preview['token'],[{'itemId':blocked['id'],'rowId':second}])
    assert resolved['readyCount']==1

def test_mixed_roles_keep_real_internal_role_gap_without_tainting_external_node(tmp_path):
    store=DraftStore(tmp_path);state=store.read(202608)
    row=person(hire='2026-07-02',姓名='甲',协助招聘人工号='zt-assistant',协助招聘人姓名='内部乙')
    fields=['体系外人员编号']+list(row['values'])
    state=store.ingest(202608,[('人员.xlsx',facts_book(fields,[['mixed']+list(row['values'].values())]))],0)
    preview=upload(store,[['甲','招聘负责人奖金','外部乙',100,'CNY']],state['revision'])
    state=adopt(store,preview);result=view(state)
    assert any(n['role']=='协助招聘人' and n['review'] for n in result['nodes'] if n['scope']=='current')
    assert all(not n['review'] and not n['reasons'] for n in result['nodes'] if n.get('externalResultId'))
    assert attention(state)['current']

def test_accepted_month_and_stale_revision_reject_adoption(tmp_path):
    store=DraftStore(tmp_path);preview=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']])
    with pytest.raises(DraftConflict):external.apply_import(store,202608,0,preview['token'])
    store._write(store._folder(202608)/'accepted.json',{'revision':preview['revision'],'state':store.read(202608)})
    with pytest.raises(ValueError,match='确认'):adopt(store,preview)


def test_all_component_sources_are_verified_and_latest_slot_is_invalid_as_a_group(tmp_path):
    store=DraftStore(tmp_path)
    first=book([['甲','招聘负责人奖金','乙',100,'CNY']])
    extra=book([['甲','招聘负责人奖金','乙',20,'CNY','','','','','','补发']])
    preview=external.prepare_import(store,202608,0,[('常规.xlsx',first),('补发.xlsx',extra)])
    state=adopt(store,preview)
    second=next(s for s in state['sources'] if s['name']=='补发.xlsx')
    state=store.change(202608,state['revision'],{'source':{'id':second['id'],'active':False}})
    assert not external.effective(state)['valid']
    assert state['rows'][0]['calculation']['recruitment']['currentAmount'] is None
    assert not any(n['scope']=='current' and n['amount'] for n in view(state)['nodes'])


def test_resolutions_cannot_bind_another_person(tmp_path):
    store=DraftStore(tmp_path)
    state=store.ingest(202608,[('人员.xlsx',facts_book(['体系外人员编号','姓名','入职日期'],[['a','同名','2025-01-01'],['a','同名','2026-07-01'],['b','其他人','2026-07-01']]))],0)
    other=next(r['id'] for r in state['rows'] if r['values']['姓名']=='其他人')
    preview=upload(store,[['同名','招聘负责人奖金','乙',100,'CNY']],state['revision'])
    item=preview['rows'][0]
    resolved=external.repreview(store,202608,preview['revision'],preview['token'],[{'itemId':item['id'],'rowId':other}])
    assert resolved['blockedCount']==1 and '对应' in resolved['rows'][0]['issues'][0]


def test_no_payment_does_not_require_a_payee_and_pending_is_one_action(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['甲','招聘负责人奖金','',0,'CNY'],['乙','招聘负责人奖金','',None,'CNY']])
    assert preview['readyCount']==1 and preview['blockedCount']==1 and preview['peopleCount']==2
    state=adopt(store,preview)
    tasks=[w for w in view(state)['workItems'] if w['key']=='external_results']
    assert len(tasks)==1 and tasks[0]['mode']=='externalResults'


def test_changed_original_becomes_invalid_and_restoring_bytes_restores_the_same_version(tmp_path):
    store=DraftStore(tmp_path);preview=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]);state=adopt(store,preview)
    source=state['sources'][0];path=store._folder(202608)/'sources'/source['storedName'];original=path.read_bytes()
    path.write_bytes(b'changed')
    changed=store.read(202608)
    assert len(external.effective(changed)['invalid'])==1
    assert not any(n['scope']=='current' and n['amount'] for n in view(changed)['nodes'])
    path.write_bytes(original)
    restored=store.read(202608)
    assert view(restored)['totals'][0]['confirmed']==100


def test_reupload_restores_missing_or_changed_source_and_does_not_double_count(tmp_path):
    store=DraftStore(tmp_path);rows=[['甲','招聘负责人奖金','乙',100,'CNY']]
    content=book(rows)
    def reupload(revision=0):
        return external.prepare_import(store,202608,revision,[('体系外结果.xlsx',content)])
    state=adopt(store,reupload());revision=state['revision']
    source=state['sources'][0];path=store._folder(202608)/'sources'/source['storedName']
    path.unlink()
    preview=reupload(revision)
    assert path.is_file() and preview['revision']==revision+1 and preview['unchangedCount']==1
    state=store.read(202608)
    assert view(state)['totals'][0]['confirmed']==100 and len(state['externalResults'])==1
    state=store.change(202608,state['revision'],{'source':{'id':source['id'],'active':False}})
    preview=reupload(state['revision'])
    assert preview['unchangedCount']==1 and view(store.read(202608))['totals'][0]['confirmed']==100


def test_duplicate_slot_with_blank_amount_returns_diagnostics_instead_of_error(tmp_path):
    store=DraftStore(tmp_path);state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]))
    preview=upload(store,[['甲','招聘负责人奖金','乙',None,'CNY'],['甲','招聘负责人奖金','乙',200,'CNY']],state['revision'])
    assert preview['blockedCount']==2
    assert any('重复' in issue for row in preview['rows'] for issue in row['issues'])


def test_public_api_projection_accepts_external_labels(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    store=DraftStore(tmp_path);state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]))
    monkeypatch.setattr(api,'store',store)
    payload=api.public(state)
    assert payload['externalResults']['summary']['totalsByCurrency']=={'人民币':100}
    assert payload['monthly']['nodes'][0]['externalResultId']


def test_new_file_resolves_only_the_matching_old_pending_slot(tmp_path):
    store=DraftStore(tmp_path)
    old=upload(store,[['甲','招聘负责人奖金','乙',None,'CNY'],['丙','招聘负责人奖金','丁',None,'CNY']])
    corrected=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']],old['revision'])
    state=adopt(store,corrected)
    pending=external.context(store,state)['pendingImports']
    assert len(pending)==1 and pending[0]['blockedCount']==1
    reopened=external.repreview(store,202608,state['revision'],old['token'],[])
    assert next(r for r in reopened['rows'] if r['candidateName']=='甲')['status']=='superseded'
    final=upload(store,[['丙','招聘负责人奖金','丁',200,'CNY']],state['revision'])
    state=adopt(store,final)
    assert not external.context(store,state)['pendingImports']
    assert not any(w['key']=='external_results' for w in view(state)['workItems'])


def test_number_and_name_conflict_is_not_silently_bound_to_wrong_person(tmp_path):
    store=DraftStore(tmp_path)
    state=store.ingest(202608,[('人员.xlsx',facts_book(['体系外人员编号','姓名','入职日期'],[['ID-A','甲','2026-07-01'],['ID-B','乙','2026-07-01']]))],0)
    preview=upload(store,[['甲','招聘负责人奖金','丙',100,'CNY','ID-B']],state['revision'])
    assert preview['blockedCount']==1
    assert '姓名与人员编号' in preview['rows'][0]['issues'][0]


def test_source_restoration_is_per_file_even_in_a_different_batch(tmp_path):
    store=DraftStore(tmp_path);regular=book([['甲','招聘负责人奖金','乙',100,'CNY']])
    preview=external.prepare_import(store,202608,0,[('常规.xlsx',regular)]);state=adopt(store,preview)
    source=state['sources'][0];state=store.change(202608,state['revision'],{'source':{'id':source['id'],'active':False}})
    supplement=book([['甲','招聘负责人奖金','乙',20,'CNY','','','','','','补发']])
    preview=external.prepare_import(store,202608,state['revision'],[('常规.xlsx',regular),('补发.xlsx',supplement)])
    assert preview['readyCount']==2
    assert view(adopt(store,preview))['totals'][0]['confirmed']==120


def test_undo_then_reupload_same_file_can_really_adopt_again(tmp_path):
    store=DraftStore(tmp_path);rows=[['甲','招聘负责人奖金','乙',100,'CNY']]
    state=adopt(store,upload(store,rows));state=store.history_action(202608,state['revision'],'undo')
    preview=upload(store,rows,state['revision'])
    assert preview['readyCount']==1
    state,count=external.apply_import(store,202608,preview['revision'],preview['token'])
    assert count==1 and view(state)['totals'][0]['confirmed']==100
    assert len(external.effective(state)['valid'])==1


def test_unrecognized_status_is_blocked_and_stop_payment_never_pays(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY','','','','','停发'],['丙','招聘负责人奖金','丁',100,'CNY','','','','','状态写错']])
    assert preview['blockedCount']==2
    assert preview['rows'][0]['outcome']=='exclude'
    assert any('状态无法识别' in i for i in preview['rows'][1]['issues'])


def test_adding_optional_person_number_reuses_the_same_subject_and_slot(tmp_path):
    store=DraftStore(tmp_path);state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]))
    preview=upload(store,[['甲','招聘负责人奖金','乙',80,'CNY','EXT-甲','2026-07-01']],state['revision'])
    assert preview['rows'][0]['change']=='replacement'
    state=adopt(store,preview)
    assert len(state['externalSubjects'])==1 and state['externalSubjects'][0]['candidateId']=='EXT-甲'
    assert state['externalSubjects'][0]['hireDate']=='2026-07-01'
    assert view(state)['totals'][0]['confirmed']==80
    conflicting=upload(store,[['甲','招聘负责人奖金','乙',90,'CNY','DIFFERENT']],state['revision'])
    assert conflicting['blockedCount']==1


def test_external_then_ordinary_special_adjustment_is_blocked(tmp_path):
    from bonus_platform.engine import recruitment_special as special
    store=DraftStore(tmp_path);state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]))
    with pytest.raises(ValueError,match='体系外'):
        special.preview(store,202608,{'revision':state['revision'],'type':'adjustment','basisType':'manual',
            'rowId':state['rows'][0]['id'],'kind':'recruitment','role':'招聘负责人','month':202608,
            'recipientId':'external-payee','recipientName':'乙','amount':20,'currency':'人民币'})


def test_existing_special_adjustment_blocks_only_its_external_role(tmp_path):
    from bonus_platform.engine import recruitment_special as special
    store=DraftStore(tmp_path);row=person(hire='2026-07-02',姓名='甲')
    fields=['体系外人员编号']+list(row['values'])
    state=store.ingest(202608,[('人员.xlsx',facts_book(fields,[['mixed']+list(row['values'].values())]))],0)
    rid=state['rows'][0]['id']
    planned=special.preview(store,202608,{'revision':state['revision'],'type':'adjustment','basisType':'manual',
        'rowId':rid,'kind':'recruitment','role':'招聘负责人','month':202608,
        'recipientId':'external-payee','recipientName':'乙','amount':20,'currency':'欧元'})
    state=special.apply_preview(store,202608,state['revision'],planned['id'])
    preview=upload(store,[['甲','招聘负责人奖金','乙',100,'EUR'],['甲','内推奖金','丙',50,'EUR']],state['revision'])
    assert preview['blockedCount']==1 and preview['readyCount']==1
    assert '补发或扣回' in next(r for r in preview['rows'] if r['role']=='招聘负责人')['issues'][0]


def test_known_recipient_number_must_match_name_but_unknown_external_number_is_allowed(tmp_path):
    store=DraftStore(tmp_path)
    state=store.ingest(202608,[('人员.xlsx',facts_book(['体系外人员编号','姓名','工号','入职日期','招聘负责人工号','招聘负责人姓名'],
        [['candidate','候选甲','zt-candidate','2026-07-01','zt-known','实际领取人']]))],0)
    preview=upload(store,[['体系外乙','招聘负责人奖金','错误姓名',100,'CNY','','','zt-known'],
                          ['体系外丙','招聘负责人奖金','新外部领取人',100,'CNY','','','GL999']],state['revision'])
    assert preview['blockedCount']==1 and preview['readyCount']==1
    assert any('领取人姓名与编号' in i for i in preview['rows'][0]['issues'])
    reused=upload(store,[['体系外丁','招聘负责人奖金','实际领取人',100,'CNY']],preview['revision'])
    assert reused['rows'][0]['recipientId']=='zt-known' and reused['readyCount']==1


def test_same_batch_optional_number_is_one_identity_before_duplicate_slot_check(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY'],['甲','招聘负责人奖金','乙',80,'CNY','EXT-甲']])
    assert preview['blockedCount']==2 and preview['readyCount']==0
    assert len({r['rowId'] for r in preview['rows']})==1
    assert any('重复' in i for r in preview['rows'] for i in r['issues'])


def test_same_name_multiple_explicit_ids_block_only_the_unidentified_batch_row(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY','EXT-1'],['甲','招聘负责人奖金','乙',80,'CNY','EXT-2'],['甲','招聘负责人奖金','乙',90,'CNY']])
    assert preview['readyCount']==2 and preview['blockedCount']==1
    assert '编号' in preview['rows'][-1]['issues'][0]


def test_correction_closes_old_multifile_pending_when_unchanged_source_is_first(tmp_path):
    store=DraftStore(tmp_path)
    unchanged=book([['甲','招聘负责人奖金','乙',20,'CNY','','','','','','补发']])
    missing=book([['甲','招聘负责人奖金','乙',None,'CNY']])
    old=external.prepare_import(store,202608,0,[('补发.xlsx',unchanged),('常规缺项.xlsx',missing)])
    corrected=book([['甲','招聘负责人奖金','乙',100,'CNY']])
    new=external.prepare_import(store,202608,old['revision'],[('补发.xlsx',unchanged),('常规更正.xlsx',corrected)])
    state=adopt(store,new)
    assert view(state)['totals'][0]['confirmed']==120
    assert not external.context(store,state)['pendingImports']
    assert not any(w['key']=='external_results' for w in view(state)['workItems'])


def test_same_batch_optional_hire_date_uses_one_employment_and_blocks_duplicate_slot(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY'],['甲','招聘负责人奖金','乙',80,'CNY','','2026-07-01']])
    assert preview['blockedCount']==2 and len({r['rowId'] for r in preview['rows']})==1
    ambiguous=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY','','2025-01-01'],
                            ['甲','招聘负责人奖金','乙',80,'CNY','','2026-01-01'],
                            ['甲','招聘负责人奖金','乙',90,'CNY']],preview['revision'])
    assert ambiguous['readyCount']==2 and ambiguous['blockedCount']==1
    assert '多段任职' in ambiguous['rows'][-1]['issues'][0]


def test_same_role_components_share_one_identified_payee_summary_without_number(tmp_path):
    store=DraftStore(tmp_path)
    preview=upload(store,[['甲','招聘负责人奖金','乙',100,'CNY'],['甲','招聘负责人奖金','乙',20,'CNY','','','','','','补发']])
    state=adopt(store,preview);summaries=view(state)['recipientSummaries']['recruitment']['rows']
    assert len(summaries)==1 and summaries[0]['name']=='乙' and summaries[0]['confirmed']==120
    assert len({n['recipientKey'] for n in view(state)['nodes']})==1


def test_confirmed_month_rejects_upload_without_changing_any_saved_file(tmp_path,monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from bonus_platform import recruitment_draft_api as api
    store=DraftStore(tmp_path);state=adopt(store,upload(store,[['甲','招聘负责人奖金','乙',100,'CNY']]))
    folder=store._folder(202608)
    store._write(folder/'accepted.json',{'revision':state['revision'],'state':state})
    original={str(p.relative_to(folder)):p.read_bytes() for p in folder.rglob('*') if p.is_file()}
    monkeypatch.setattr(api,'store',store);app=FastAPI();app.include_router(api.router)
    client=TestClient(app)
    response=client.post('/api/recruitment/drafts/202608/external-results/import',
        data={'revision':state['revision']},files=[('files',('更正.xlsx',book([['乙','招聘负责人奖金','丙',200,'CNY']])))])
    assert response.status_code==400 and '已确认' in response.json()['detail']
    assert {str(p.relative_to(folder)):p.read_bytes() for p in folder.rglob('*') if p.is_file()}==original
    saved=store.read(202608)
    assert saved['revision']==state['revision'] and saved['sources']==state['sources'] and saved['externalImports']==state['externalImports']
