"""Durable payroll isolation, leases, auth and storage failure contracts."""
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from io import BytesIO
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import time
from uuid import uuid4
from urllib.parse import unquote
from zipfile import ZipFile, ZIP_STORED

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

import bonus_platform.app as app_module
import bonus_platform.auth as auth
from bonus_platform.engine.china_employee_payroll import activities, durable, worker
from bonus_platform.engine.china_employee_payroll.durable_store import Store, LeaseLost
from bonus_platform.engine.china_employee_payroll.meal_allowance import ParsedAttendance, REQUIRED_COLUMNS
from test_china_payroll_activities import wx_file
from bonus_platform.engine import company_run_storage

BASE = "/api/china-employee-payroll/activities"
_PRIVATE_PUT = durable.storage.put_labor_supabase_private_object
_PRIVATE_GET = durable.storage.get_labor_supabase_private_object
_SIGNED_DOWNLOAD = durable.storage.create_labor_supabase_signed_download


@pytest.fixture
def environment(tmp_path, monkeypatch):
    monkeypatch.setenv("SIGMA_CHINA_PAYROLL_BACKEND", "sqlite")
    monkeypatch.setenv("SIGMA_CHINA_PAYROLL_SQLITE_PATH", str(tmp_path / "activities.sqlite"))
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_ENV", "local")
    monkeypatch.delenv("VERCEL", raising=False)
    monkeypatch.delenv("VERCEL_ENV", raising=False)
    current = {"user": {"id": "feishu-1", "name": "创建人", "status": "active"}, "modules": [{"id": "employee", "enabled": True, "canEnter": True}]}
    monkeypatch.setattr(auth, "current_user_from_request", lambda request: current)
    monkeypatch.setattr(activities, "current_user_from_request", lambda request: current)
    monkeypatch.setattr(app_module, "current_user_from_request", lambda request: current)
    monkeypatch.setattr(app_module, "_labor_current_user_from_request", lambda request: current)
    store = Store()
    store.initialize_local()
    objects = tmp_path / "objects"
    objects.mkdir()
    def location(key):
        return objects / hashlib.sha256(key.encode()).hexdigest()
    def put(key, content, **kwargs):
        location(key).write_bytes(content)
    def get(key, **kwargs):
        path = location(key)
        return path.read_bytes() if path.exists() else None
    monkeypatch.setattr(durable.storage, "put_labor_supabase_private_object", put)
    monkeypatch.setattr(durable.storage, "get_labor_supabase_private_object", get)
    monkeypatch.setattr(durable.storage, "create_labor_supabase_signed_upload_for_object", lambda key, **kwargs: {"signedUrl": "https://example.invalid/upload", "objectKey": key})
    monkeypatch.setattr(durable.storage, "create_labor_supabase_signed_download", lambda key, **kwargs: {"signedUrl": "https://example.invalid/download"})
    return store, TestClient(app_module.app), current, objects


def create(client, subject="meal-allowance"):
    response = client.post(BASE, json={"subject": subject, "month": "2026-09", "requestId": uuid4().hex, "creator": {"name": "伪造", "userId": "fake"}})
    assert response.status_code == 201, response.text
    return response.json()


def fake_source(monkeypatch):
    row = dict.fromkeys(REQUIRED_COLUMNS, "")
    row.update({"员工": "测试", "工号": "zt1", "人员状态": "在职", "二级组织": "HQU技术部", "三级组织": "技术部", "考勤日期": date(2026, 9, 1), "首打卡(含补签)": datetime(2026, 9, 1, 9, 10), "末打卡(含补签)": datetime(2026, 9, 1, 21, 5), "当前班次": "深圳灵活打卡8:30-9:10", "日期类型": "工作日"})
    monkeypatch.setattr(worker.BeisenMealSource, "__init__", lambda self: None)
    monkeypatch.setattr(worker.BeisenMealSource, "load", lambda self, month, progress: ParsedAttendance(rows=[row], headers=REQUIRED_COLUMNS, files=[], missing_columns=[]))


def test_concurrent_enqueue_restart_and_stale_worker_fencing(environment):
    store, client, _, _ = environment
    activity = create(client)
    with ThreadPoolExecutor(8) as pool:
        accepted = list(pool.map(lambda _: Store().enqueue(activity["activityId"]), range(16)))
    assert len({item["runId"] for item in accepted}) == 1
    # An independent Python process reopens the queued record from disk.
    code = "from bonus_platform.engine.china_employee_payroll.durable_store import Store; import json,sys; print(json.dumps(Store().get(sys.argv[1])))"
    recovered = json.loads(subprocess.check_output([sys.executable, "-c", code, activity["activityId"]]))
    assert recovered["runId"] == accepted[0]["runId"]
    job, old = store.claim(lease_seconds=-1)
    same, new = Store().claim()
    assert job["runId"] == same["runId"] and old != new
    with pytest.raises(LeaseLost):
        store.update(job["runId"], old, {"status": "completed"}, terminal=True)
    assert store.claim() is None
    store.update(job["runId"], new, {"status": "completed", "progress": "done"}, terminal=True)
    assert Store().get(activity["activityId"])["status"] == "completed"


def test_exhausted_leases_fail_without_infinite_running(environment):
    store, client, _, _ = environment
    record = create(client)
    store.enqueue(record["activityId"])
    for _ in range(3):
        assert store.claim(lease_seconds=-1)
    assert store.claim() is None
    assert store.get(record["activityId"])["status"] == "failed"
    assert len(store.enqueue(record["activityId"])["attempts"]) == 2


def test_group_worker_snapshot_export_and_activity_isolation(environment, monkeypatch):
    store, client, _, objects = environment
    fake_source(monkeypatch)
    one, two = create(client), create(client)
    response = client.post(f"{BASE}/{one['activityId']}/start")
    assert response.status_code == 200 and response.json()["status"] == "running"
    assert worker.process_one(store)
    recovered = Store().get(one["activityId"])
    assert recovered["status"] == "completed" and recovered["summary"]["totalAmount"] == 20
    assert Store().get(two["activityId"])["status"] == "draft"
    job = store.job(recovered["runId"])
    response = client.get(f"/api/china-employee-payroll/meal-allowance/runs/{job['runId']}")
    assert response.status_code == 200 and response.json()["summary"]["totalAmount"] == 20
    assert not {"snapshot", "inputs", "creator", "objectKey"} & response.json().keys()
    book = load_workbook(BytesIO(durable.read_file(job["export"])), read_only=True)
    assert book.sheetnames == ["餐补核算结果", "北森考勤源数据"]
    book.close()
    assert client.get(f"/api/china-employee-payroll/meal-allowance/{job['runId']}/export").json()["filename"].startswith("2026年9月")
    assert client.get(f"/api/china-employee-payroll/relocation-allowance/runs/{job['runId']}").status_code == 404
    assert client.post(f"{BASE}/{one['activityId']}/start").json()["runId"] == job["runId"]
    assert len(list(objects.iterdir())) == 3


def test_wx_signed_upload_verified_before_calculation(environment):
    store, client, _, _ = environment
    activity = create(client, "wx-meal-allowance")
    raw = wx_file()
    response = client.post(f"{BASE}/{activity['activityId']}/uploads", json={"files": [{"name": "WX九月.xlsx", "size": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}]})
    assert response.status_code == 200, response.text
    key = store.get(activity["activityId"])["inputs"][0]["objectKey"]
    durable.storage.put_labor_supabase_private_object(key, raw)
    run = client.post(f"{BASE}/{activity['activityId']}/start").json()
    assert run["status"] == "running"
    worker.process_one(store)
    completed = store.get(activity["activityId"])
    assert completed["summary"]["totalAmount"] == 20
    assert completed["status"] == "completed"
    assert client.post(f"{BASE}/{activity['activityId']}/uploads", json={"files": [{"name": "x.xlsx", "size": 1, "sha256": "a"*64}]}).status_code == 409


def test_bad_input_or_storage_failure_never_publish_completed(environment, monkeypatch):
    store, client, _, _ = environment
    activity = create(client, "wx-meal-allowance")
    store.uploads(activity["activityId"], [{"name": "bad.xlsx", "size": 5, "sha256": "a"*64, "objectKey": "missing"}])
    run = store.enqueue(activity["activityId"])
    worker.process_one(store)
    assert store.get(activity["activityId"])["status"] == "failed"
    assert client.get(f"/api/china-employee-payroll/meal-allowance/{run['runId']}/export").status_code == 409
    fake_source(monkeypatch)
    other = create(client)
    store.enqueue(other["activityId"])
    def fail(*args, **kwargs):
        raise RuntimeError("secret-database-password")
    monkeypatch.setattr(durable.storage, "put_labor_supabase_private_object", fail)
    worker.process_one(store)
    failed = store.get(other["activityId"])
    assert failed["status"] == "failed" and "secret" not in failed["progress"]
    assert "export" not in store.job(failed["runId"])


@pytest.mark.parametrize("path,method", [(BASE, "get"), (BASE, "post"), (BASE + "/" + "a"*32 + "/start", "post"), (BASE + "/" + "a"*32 + "/uploads", "post"), ("/api/china-employee-payroll/meal-allowance/runs/test", "get"), ("/api/china-employee-payroll/meal-allowance/test/export", "get"), ("/api/china-employee-payroll/relocation-allowance/runs/test/export", "get")])
def test_all_api_access_uses_existing_employee_module_permission(environment, monkeypatch, path, method):
    _, client, current, _ = environment
    current["modules"][0]["canEnter"] = False
    assert getattr(client, method)(path).status_code == 403
    monkeypatch.setattr(auth, "current_user_from_request", lambda request: None)
    assert getattr(client, method)(path).status_code == 401


def test_creator_fallback_and_idempotency_are_server_owned(environment):
    store, client, current, _ = environment
    item = create(client)
    assert item["creator"] == {"name": "创建人", "userId": "feishu-1", "avatarUrl": activities.DEFAULT_AVATAR}
    current["user"]["id"] = "feishu-2"
    assert client.get(f"{BASE}/{item['activityId']}").json()["creator"]["userId"] == "feishu-1"
    response = client.post(BASE, json={"subject": item["subject"], "month": item["month"], "name": item["name"], "requestId": item["activityId"]})
    assert response.status_code == 409


def test_vercel_fails_closed_without_durable_backend(environment, monkeypatch):
    _, client, _, _ = environment
    monkeypatch.setenv("VERCEL", "1")
    with pytest.raises(RuntimeError):
        Store()
    monkeypatch.delenv("SIGMA_CHINA_PAYROLL_BACKEND")
    assert client.get(BASE).status_code == 503


def test_expired_lease_cannot_heartbeat(environment):
    store, client, _, _ = environment
    item = create(client)
    store.enqueue(item["activityId"])
    job, token = store.claim(-1)
    with pytest.raises(LeaseLost):
        store.update(job["runId"], token)


def test_database_namespace_matches_existing_storage_normalization(environment, monkeypatch):
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_ENV", "  local\n")
    monkeypatch.setenv("SIGMA_CHINA_PAYROLL_BACKEND", " SQLite\n")
    assert durable.enabled()
    assert Store().environment == durable.storage.labor_persistent_environment() == "local"


def test_company_obs_private_adapter_publishes_and_exports(environment, monkeypatch):
    store, client, _, _ = environment
    fake_source(monkeypatch)
    for name, value in {"SIGMA_LABOR_STORAGE_BACKEND": "obs", "OBS_ENDPOINT": "obs.example.invalid",
                        "OBS_BUCKET": "test", "OBS_ACCESS_KEY": "test", "OBS_SECRET_KEY": "test"}.items():
        monkeypatch.setenv(name, value)
    objects = {}
    monkeypatch.setattr(durable.storage, "put_labor_supabase_private_object", _PRIVATE_PUT)
    monkeypatch.setattr(durable.storage, "get_labor_supabase_private_object", _PRIVATE_GET)
    monkeypatch.setattr(durable.storage, "create_labor_supabase_signed_download", _SIGNED_DOWNLOAD)
    monkeypatch.setattr(durable.storage, "obs_put_bytes", lambda key, data, **kw: objects.__setitem__(key, data))
    monkeypatch.setattr(durable.storage, "obs_get_bytes", lambda key: objects.get(key))
    monkeypatch.setattr(durable.storage, "obs_signed_download_for_key", lambda key, **kw: "https://obs.example.invalid/" + key)
    item = create(client)
    run = store.enqueue(item["activityId"])
    worker.process_one(store)
    assert store.get(item["activityId"])["status"] == "completed"
    assert durable.get_run(run["runId"])["summary"]["totalAmount"] == 20
    response = client.get(f"/api/china-employee-payroll/meal-allowance/{run['runId']}/export")
    assert response.status_code == 200
    assert response.content == durable.read_file(store.job(run["runId"])["export"])
    assert "no-store" in response.headers["cache-control"]
    assert len(objects) == 3


@pytest.mark.parametrize("backend", ["obs", "s3"])
def test_company_server_upload_worker_export_and_retry(environment, monkeypatch, backend):
    store, client, _, _ = environment
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_BACKEND", backend)
    # Company browser transfers must not issue any private signed URLs.
    def no_signing(*args, **kwargs):
        pytest.fail("Company file transfers must go through the authenticated API")
    monkeypatch.setattr(durable.storage, "create_labor_supabase_signed_upload_for_object", no_signing)
    monkeypatch.setattr(durable.storage, "create_labor_supabase_signed_download", no_signing)
    assert client.get(BASE + "/config/upload").json() == {"directUpload": False}
    activity = create(client, "wx-meal-allowance")
    raw = wx_file()
    path = f"{BASE}/{activity['activityId']}/start"
    response = client.post(path, files={"attendance_files": ("WX九月.xlsx", raw, durable.XLSX)})
    assert response.status_code == 200, response.text
    run = response.json()
    manifest = store.get(activity["activityId"])["inputs"][0]
    assert manifest["size"] == len(raw)
    assert manifest["sha256"] == hashlib.sha256(raw).hexdigest()
    assert durable.read_file(manifest) == raw
    assert client.post(path, files={"attendance_files": ("重复.xlsx", b"different")}).json()["runId"] == run["runId"]
    assert worker.process_one(Store())
    assert Store().get(activity["activityId"])["summary"]["totalAmount"] == 20
    export_path = f"/api/china-employee-payroll/meal-allowance/{run['runId']}/export"
    exported = client.get(export_path)
    assert exported.status_code == 200
    assert exported.headers["content-disposition"].startswith("attachment; filename*=UTF-8''")
    filename = unquote(exported.headers["content-disposition"].split("UTF-8''")[1])
    assert filename.startswith("2026年9月") and filename.endswith(".xlsx")
    assert exported.content == durable.read_file(Store().job(run["runId"])["export"])
    book = load_workbook(BytesIO(exported.content), read_only=True)
    assert book.sheetnames == ["餐补核算结果", "WX技术部考勤源"]
    book.close()
    assert client.post(f"{BASE}/{activity['activityId']}/uploads", json={"files": []}).status_code == 409
    # The API still verifies the stored export before returning bytes.
    key = Store().job(run["runId"])["export"]["objectKey"]
    durable.storage.put_labor_supabase_private_object(key, b"corrupted")
    assert client.get(export_path).status_code == 409


def test_company_server_upload_validation_and_storage_failure(environment, monkeypatch):
    store, client, _, _ = environment
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_BACKEND", "obs")
    activity = create(client, "wx-meal-allowance")
    path = f"{BASE}/{activity['activityId']}/start"
    for name, content in [("wrong.txt", b"data"), ("../bad.xlsx", b"data"), ("empty.xlsx", b"")]:
        assert client.post(path, files={"attendance_files": (name, content)}).status_code == 400
    monkeypatch.setattr(durable, "MAX_FILE_BYTES", 5)
    assert client.post(path, files={"attendance_files": ("large.xlsx", b"123456")}).status_code == 400
    monkeypatch.setattr(durable, "MAX_TOTAL_BYTES", 5)
    assert client.post(path, files=[("attendance_files", ("a.xlsx", b"123")), ("attendance_files", ("b.xlsx", b"123"))]).status_code == 400
    store.uploads(activity["activityId"], [{"objectKey": "previous-input"}])
    def fail(*args, **kwargs):
        raise RuntimeError("private-storage-credential")
    monkeypatch.setattr(durable.storage, "put_labor_supabase_private_object", fail)
    response = client.post(path, files={"attendance_files": ("ok.xlsx", b"123")})
    assert response.status_code == 503 and "credential" not in response.text
    unchanged = Store().get(activity["activityId"])
    assert unchanged["status"] == "draft" and unchanged["runId"] == ""
    assert unchanged["inputs"] == [{"objectKey": "previous-input"}]


def test_company_server_file_endpoints_keep_module_permissions(environment, monkeypatch):
    store, client, current, _ = environment
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_BACKEND", "obs")
    item = create(client, "wx-meal-allowance")
    raw = wx_file()
    run = client.post(f"{BASE}/{item['activityId']}/start", files={"attendance_files": ("WX.xlsx", raw)}).json()
    worker.process_one(store)
    current["modules"][0]["canEnter"] = False
    assert client.post(f"{BASE}/{item['activityId']}/start", files={"attendance_files": ("WX.xlsx", raw)}).status_code == 403
    assert client.get(f"/api/china-employee-payroll/meal-allowance/{run['runId']}/export").status_code == 403
    monkeypatch.setattr(auth, "current_user_from_request", lambda request: None)
    assert client.post(f"{BASE}/{item['activityId']}/start", files={"attendance_files": ("WX.xlsx", raw)}).status_code == 401
    assert client.get(f"/api/china-employee-payroll/meal-allowance/{run['runId']}/export").status_code == 401


def test_company_server_upload_exceeds_vercel_body_limit(environment, monkeypatch):
    store, client, _, _ = environment
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_BACKEND", "obs")
    item = create(client, "wx-meal-allowance")
    buffer = BytesIO(wx_file())
    with ZipFile(buffer, "a", compression=ZIP_STORED) as archive:
        archive.writestr("large-file-fixture.bin", b"0" * (5 * 1024 * 1024))
    raw = buffer.getvalue()
    assert len(raw) > 5 * 1024 * 1024
    response = client.post(f"{BASE}/{item['activityId']}/start", files={"attendance_files": ("大文件.xlsx", raw)})
    assert response.status_code == 200, response.text
    assert worker.process_one(store)
    assert Store().get(item["activityId"])["summary"]["totalAmount"] == 20


def test_company_server_relocation_export_and_storage_error(environment, monkeypatch):
    store, client, _, _ = environment
    monkeypatch.setenv("SIGMA_LABOR_STORAGE_BACKEND", "obs")
    item = create(client, "relocation-allowance")
    store.enqueue(item["activityId"])
    job, token = store.claim()
    raw, filename = b"saved-relocation-export", "2026年9月 ABU职场搬迁补贴.xlsx"
    manifest = durable.put(item, filename, raw, "exports")
    store.update(job["runId"], token, {"status": "completed", "export": manifest}, terminal=True)
    path = f"/api/china-employee-payroll/relocation-allowance/runs/{job['runId']}/export"
    response = client.get(path)
    assert response.status_code == 200 and response.content == raw
    assert unquote(response.headers["content-disposition"].split("UTF-8''")[1]) == filename
    assert client.get(f"/api/china-employee-payroll/meal-allowance/{job['runId']}/export").status_code == 404
    def fail(*args, **kwargs):
        raise RuntimeError("private-storage-credential")
    monkeypatch.setattr(durable.storage, "get_labor_supabase_private_object", fail)
    response = client.get(path)
    assert response.status_code == 503 and "credential" not in response.text


def test_company_obs_legacy_history_survives_empty_instance_cache(environment, monkeypatch, tmp_path):
    store, client, _, _ = environment
    monkeypatch.setattr(app_module, "CHINA_EMPLOYEE_PAYROLL_RUNS_DIR", tmp_path / "empty-instance")
    for name, value in {"SIGMA_CHINA_EMPLOYEE_STORAGE_BACKEND": "obs", "OBS_ENDPOINT": "obs.example.invalid",
                        "OBS_BUCKET": "test", "OBS_ACCESS_KEY": "test", "OBS_SECRET_KEY": "test"}.items():
        monkeypatch.setenv(name, value)
    run_id = "china_employee_payroll_20260901_120000_123456"
    prefix = "china-employee-meal/local/" + run_id + "/"
    summary = {"dateStart": "2026-09-01", "sourceType": "wx", "totalAmount": 20, "employeeCount": 1}
    objects = {prefix + "metadata.json": json.dumps({"runId": run_id, "summary": summary, "createdAt": "2026-09-01", "sourceType": "wx"}).encode(),
               prefix + "result.json": json.dumps({"summary": summary, "results": []}).encode()}
    monkeypatch.setattr(company_run_storage, "obs_list_objects", lambda key: [{"key": k} for k in objects if k.startswith(key)])
    monkeypatch.setattr(company_run_storage, "obs_get_bytes", lambda key: objects.get(key))
    assert not store.records()
    listed = client.get(BASE).json()
    assert listed["total"] == 1
    legacy = listed["activities"][0]
    assert legacy["status"] == "completed" and legacy["month"] == "2026-09" and legacy["amount"] == 20
    assert client.get(f"{BASE}/{legacy['activityId']}").json() == legacy
    assert client.post(f"{BASE}/{legacy['activityId']}/start").json()["runId"] == run_id
    result = client.get(f"/api/china-employee-payroll/meal-allowance/runs/{run_id}")
    assert result.status_code == 200 and result.json()["summary"] == summary
    assert (tmp_path / "empty-instance" / run_id / "result.json").exists()
    assert not store.records()
