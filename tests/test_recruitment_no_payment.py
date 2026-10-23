from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from bonus_platform.engine.recruitment_awards import calculate_draft, calculate_row
from bonus_platform.engine.recruitment_no_payment import project as display, installments
from bonus_platform.engine.recruitment_monthly import project as monthly, export
from bonus_platform.engine.recruitment_worksheet import worksheet, export_results
from bonus_platform.engine.recruitment_settlement import capture_confirmed_installments
from test_recruitment_flow import state_with_rows
from test_recruitment_calculation_fixes import confirmed_state


def cancel(state, kind='recruitment', basis='manual'):
    row=state['rows'][0]
    row.setdefault('awardDecisions',[]).append({'kind':kind,'action':'exclude',
        'basisType':basis,'reason':'已核对本次奖金不发放'})
    calculate_draft(state)
    return row


def test_early_departure_keeps_known_amounts_only_in_display():
    state=state_with_rows();row=state['rows'][0]
    expected=[(n['label'],n['amount']) for n in row['calculation']['recruitment']['nodes']]
    row['values'].update(人员状态='离职',最后工作日='2026-07-20')
    calculate_draft(state);before=deepcopy(state)
    projected=display(state,row)['recruitment']
    assert [(n['label'],n['amount']) for n in projected['nonPaymentInstallments']]==expected
    assert all(n['scheduleText']=='入职未满一个月离职' for n in projected['nonPaymentInstallments'])
    assert not projected['nodes'] and projected['total']==0 and projected['currentAmount']==0
    assert not capture_confirmed_installments(state)
    assert state==before
    assert display(state,row)['recruitment']==projected


@pytest.mark.parametrize('kind',['recruitment','referral'])
@pytest.mark.parametrize('basis',['manual','approval'])
def test_whole_cancellation_preserves_amounts_and_undo_removes_display_only_installments(kind,basis):
    state=state_with_rows();row=state['rows'][0]
    expected=[(n['role'],n['label'],n['amount']) for n in row['calculation'][kind]['nodes']]
    cancel(state,kind,basis);before=deepcopy(state)
    projected=display(state,row)[kind]
    assert [(n['role'],n['label'],n['amount']) for n in projected['nonPaymentInstallments']]==expected
    assert all('已核对本次奖金不发放' in n['scheduleText'] and n['scheduledMonth'] is None
               for n in projected['nonPaymentInstallments'])
    assert projected['nodes']==[] and projected['total']==0
    assert state==before
    row['awardDecisions']=[];calculate_draft(state)
    assert 'nonPaymentInstallments' not in display(state,row)[kind]
    assert [(n['role'],n['label'],n['amount']) for n in row['calculation'][kind]['nodes']]==expected


@pytest.mark.parametrize('values,kind',[
    ({'招聘渠道':'其他'},'recruitment'),
    ({'推荐人工号':'','推荐人姓名':''},'referral'),
    ({'职级':'','招聘启动日期':''},'recruitment'),
    ({'招聘启动日期':''},'recruitment'),
])
def test_zero_or_unknown_basis_never_gains_a_positive_display_amount(values,kind):
    state=state_with_rows();state['rows'][0]['values'].update(values)
    row=cancel(state,kind);before=deepcopy(state)
    assert not display(state,row)[kind].get('nonPaymentInstallments')
    assert state==before


def test_unresolved_identity_does_not_gain_display_installments():
    state=state_with_rows();row=state['rows'][0]
    row['unresolvedFields']=['入职日期'];cancel(state)
    assert row['calculation']['recruitment']['total'] is None
    assert not display(state,row)['recruitment'].get('nonPaymentInstallments')


def test_historical_arrangements_and_recoveries_remain_unchanged():
    state,row=confirmed_state();state['month']=202609
    row['values']['招聘启动日期']='2026-06-27'
    row['calculation']=calculate_row(state,row)
    original=deepcopy(row['calculation']['recruitment']['nodes'])
    row['awardDecisions']=[{'kind':'recruitment','action':'exclude','basisType':'manual','reason':'本次不发'}]
    row['calculation']=calculate_row(state,row);before=deepcopy(state)
    projected=display(state,row)['recruitment']
    actual=projected['nodes']
    historical=next(n for n in actual if n['openingLocked'])
    assert (historical['amount'],historical['scheduledMonth'])==(27.5,202607)
    retained=projected['nonPaymentInstallments']
    assert len(retained)==1 and retained[0]['label']=='入职满三个月'
    assert retained[0]['amount']==next(n['amount'] for n in original if n['label']=='入职满三个月')
    assert len({(n['role'],n['label']) for n in installments(projected)})==len(installments(projected))
    assert projected['recoveries']==row['calculation']['recruitment']['recoveries']
    assert state==before


def test_external_roles_are_not_replaced_by_automatic_display_amounts():
    state=state_with_rows();row=cancel(state)
    row['externalOwnership']=[{'kind':'recruitment','role':'招聘负责人'}]
    before=deepcopy(state)
    assert not display(state,row)['recruitment'].get('nonPaymentInstallments')
    assert state==before
    row['externalOnly']=True
    assert display(state,row) is row['calculation']


def test_whole_cancelled_award_is_visible_in_due_month_and_exports_match_api(tmp_path,monkeypatch):
    import bonus_platform.recruitment_draft_api as api
    state=state_with_rows();state['month']=202609
    row=cancel(state);before=deepcopy(state)
    projected=display(state,row)['recruitment']
    expected=next(n for n in projected['nonPaymentInstallments'] if n['label']=='入职满三个月')
    monthly_view=monthly(state,{'current':[],'history':[]})
    entry=monthly_view['worksheet']['rows'][0]
    key='recruitment.招聘负责人.1'
    period='不发：手工取消招聘奖金：已核对本次奖金不发放'
    assert entry['current'] and not entry['newHire']
    assert entry['cells'][key+'.amount']==expected['amount']
    assert entry['cells'][key+'.period']==period
    assert entry['cells']['recruitmentAmount']==0
    assert not monthly_view['recipientSummaries']['recruitment']['rows']
    wb=load_workbook(BytesIO(export(state,monthly_view)))
    table=monthly_view['worksheet'];ws=wb['附件一-招聘奖金核算明细']
    for field,value in [(key+'.amount',expected['amount']),(key+'.period',period)]:
        column=next(i for i,c in enumerate(table['columns'],1) if c['key']==field)
        assert ws.cell(3,column).value==value
    complete=worksheet(state)
    assert complete['rows'][0]['cells'][key+'.amount']==expected['amount']
    assert complete['rows'][0]['cells'][key+'.scheduledMonth']==period
    wb=load_workbook(BytesIO(export_results(state)));ws=wb['核算总明细']
    for field,value in [(key+'.amount',expected['amount']),(key+'.scheduledMonth',period)]:
        column=next(i for i,c in enumerate(complete['columns'],1) if c['key']==field)
        assert ws.cell(4,column).value==value
    assert all(line[2]!='招聘奖金' for line in list(wb['本月到期明细'].values)[1:])
    monkeypatch.setattr(api.store,'read',lambda month:deepcopy(state))
    response=api.row_detail(state['month'],row['id'])
    detail=response['row']['calculation']['recruitment']
    shown=next(n for n in detail['nonPaymentInstallments'] if n['label']=='入职满三个月')
    assert shown['amount']==expected['amount'] and shown['periodText']==period
    assert detail['nodes']==[]
    assert state==before
