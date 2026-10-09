"""Contract checks for API-sourced meal batches and month-scoped uploads."""
import json
import sys
from datetime import date, datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import bonus_platform.app as app_module
from bonus_platform.engine.china_employee_payroll import meal_runs
from bonus_platform.engine.china_employee_payroll.meal_allowance import ParsedAttendance, REQUIRED_COLUMNS, WX_REQUIRED_COLUMNS
from bonus_platform.engine.china_employee_payroll.relocation import RelocationError

BASE = "/api/china-employee-payroll/meal-allowance"


@pytest.fixture(autouse=True)
def isolate_runs(tmp_path, monkeypatch):
    folder = tmp_path / "runs"
    folder.mkdir()
    monkeypatch.setattr(app_module, "CHINA_EMPLOYEE_PAYROLL_RUNS_DIR", folder)
    monkeypatch.setattr(meal_runs, "_active", False)


@pytest.fixture
def client():
    return TestClient(app_module.app)


def source_rows():
    row = dict.fromkeys(REQUIRED_COLUMNS, "")
    row.update({"员工": "=HYPERLINK(\"https://invalid\",\"张三\")", "工号": "zt10001", "人员状态": "在职",
                "二级组织": "HQU技术部", "三级组织": "技术部", "考勤日期": date(2026, 9, 1),
                "首打卡(含补签)": datetime(2026, 9, 1, 9, 10), "末打卡(含补签)": datetime(2026, 9, 1, 21, 5),
                "当前班次": "深圳灵活打卡8:30-9:10", "日期类型": "工作日", "备注": "=1+1"})
    return [row, {**row, "工号": "zt10002", "员工": "不满足打卡员工", "末打卡(含补签)": datetime(2026, 9, 1, 18, 5)}]


def fake_source(monkeypatch, loader=None):
    def load(self, month, progress):
        if loader:
            return loader(month, progress)
        progress({"stage": 0, "completed": 3, "total": 3, "unit": "个组织", "message": "已匹配范围"})
        progress({"stage": 1, "completed": 2, "total": 2, "unit": "条", "message": "日考勤已读取"})
        return ParsedAttendance(rows=source_rows(), headers=REQUIRED_COLUMNS, files=[], missing_columns=[], source_label="集团技术部 · 北森考勤")
    monkeypatch.setitem(sys.modules, "bonus_platform.engine.china_employee_payroll.meal_beisen", SimpleNamespace(BeisenMealSource=type("FakeSource", (), {"load": load})))


def test_beisen_job_persists_source_and_returns_compact_result(client, monkeypatch):
    observed = []

    def load(month, progress):
        assert month == "2026-09"
        progress({"stage": 1, "completed": 1, "total": 2, "unit": "条", "message": "正在读取日考勤"})
        folder = next(app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR.iterdir())
        observed.append(client.get(f"{BASE}/runs/{folder.name}").json())
        return ParsedAttendance(rows=source_rows(), headers=REQUIRED_COLUMNS, files=[], missing_columns=[])

    fake_source(monkeypatch, load)
    response = client.post(f"{BASE}/beisen/runs", json={"month": "2026-09"})
    assert response.status_code == 202
    assert response.json()["status"] == "running"
    run_id = response.json()["runId"]
    assert observed[0]["status"] == "running"
    assert observed[0]["progressDetail"] == {"stage": 1, "completed": 1, "total": 2, "unit": "条"}
    assert "results" not in observed[0] and "workerId" not in observed[0]
    result = client.get(f"{BASE}/runs/{run_id}").json()
    assert result["status"] == "completed"
    assert result["sourceType"] == "hr" and result["sourceOrigin"] == "beisen"
    assert result["month"] == "2026-09"
    assert result["summary"]["totalAmount"] == 20
    assert all("daily" not in row for row in result["results"])
    assert "dailyRows" not in result and "data-source" not in result
    snapshot = json.loads((app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR / run_id / "data-source.json").read_text())
    assert snapshot["columns"] == REQUIRED_COLUMNS
    assert len(snapshot["rows"]) == 2
    assert snapshot["rows"][0]["首打卡(含补签)"] == "2026-09-01T09:10:00"
    listed = client.get(f"{BASE}/runs").json()["runs"][0]
    assert listed["status"] == "completed" and listed["monthLabel"] == "2026年9月"


def test_beisen_export_has_named_source_fields_and_safe_text(client, monkeypatch):
    fake_source(monkeypatch)
    run_id = client.post(f"{BASE}/beisen/runs", json={"month": "2026-09"}).json()["runId"]
    response = client.get(f"{BASE}/{run_id}/export")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content), data_only=False)
    assert workbook.sheetnames == ["餐补核算结果", "北森考勤源数据"]
    result, source = workbook.worksheets
    assert result["A1"].value == "核算月份" and result["A2"].value == "2026年9月"
    assert result["H2"].value == 1 and result["H2"].number_format == "0"
    assert result["I2"].value == 20 and result["I2"].number_format == "#,##0.00"
    assert result["A3"].value == "合计" and result["I3"].value == 20
    assert result["C2"].data_type == "s"
    assert [cell.value for cell in source[1]] == REQUIRED_COLUMNS
    assert source["A2"].data_type == "s"
    assert source["M2"].value == "=1+1" and source["M2"].data_type == "s"
    assert source["I2"].value == "2026-09-01T09:10:00"
    assert source.max_row == 3  # Retain non-payable scoped records for reconciliation.
    assert source.freeze_panes == "C2" and source.auto_filter.ref == "A1:M3"
    assert source["A1"].fill.fgColor.rgb == "00173E80"
    assert not list((app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR / run_id).glob(".export-*"))


def test_export_rejects_truncated_source_snapshot(client, monkeypatch):
    fake_source(monkeypatch)
    run_id = client.post(f"{BASE}/beisen/runs", json={"month": "2026-09"}).json()["runId"]
    path = app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR / run_id / "data-source.json"
    snapshot = json.loads(path.read_text())
    snapshot["rows"].pop()
    meal_runs.save_json(path, snapshot)
    response = client.get(f"{BASE}/{run_id}/export")
    assert response.status_code == 409
    assert "不完整" in response.json()["detail"]


@pytest.mark.parametrize("month", [None, "2026-13", "2026-9", "1999-09", "2026-09-01"])
def test_beisen_rejects_invalid_month(client, month):
    assert client.post(f"{BASE}/beisen/runs", json={"month": month}).status_code == 400
    assert not list(app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR.iterdir())


def test_beisen_active_guard_and_other_instance_does_not_fail_job(client, monkeypatch):
    monkeypatch.setattr(meal_runs, "_active", True)
    assert client.post(f"{BASE}/beisen/runs", json={"month": "2026-09"}).status_code == 409
    run_id = meal_runs.new_run_id()
    folder = app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR / run_id
    folder.mkdir()
    meal_runs.save_json(folder / "metadata.json", {"runId": run_id, "status": "running", "workerId": "old-process", "month": "2026-09"})
    response = client.get(f"{BASE}/runs/{run_id}").json()
    assert response["status"] == "running"
    assert "results" not in response
    assert client.get(f"{BASE}/{run_id}/export").status_code == 409
    assert json.loads((folder / "metadata.json").read_text())["status"] == "running"


@pytest.mark.parametrize("failure", [RuntimeError("token=secret-access-key"), RelocationError("北森日考勤字段不完整")])
def test_beisen_failure_is_safe_and_releases_active_guard(client, monkeypatch, failure):
    def load(month, progress):
        raise failure
    fake_source(monkeypatch, load)
    run_id = client.post(f"{BASE}/beisen/runs", json={"month": "2026-09"}).json()["runId"]
    response = client.get(f"{BASE}/runs/{run_id}")
    assert response.json()["status"] == "failed"
    assert "secret-access-key" not in response.text
    assert not meal_runs._active
    assert not (app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR / run_id / "result.json").exists()
    fake_source(monkeypatch)
    assert client.post(f"{BASE}/beisen/runs", json={"month": "2026-09"}).status_code == 202


def wx_file(rows=None, omit=None):
    columns = [key for key in WX_REQUIRED_COLUMNS if key != omit]
    row = {"姓名": "WX员工", "工号": "WX001", "部门": "研发组", "组织架构": "WX-PBU技术部-研发组", "人员类型": "正式",
           "员工状态": "在职", "日期": datetime(2026, 9, 1), "班次": "09:10-18:10", "上班 1 打卡时间": "09:10", "下班 1 打卡时间": "21:10"}
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(columns)
    sheet.append(columns)
    for item in rows if rows is not None else [row]:
        sheet.append([{**row, **item}.get(key) for key in columns])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def upload_wx(client, content, month="2026-09"):
    return client.post(BASE, data={"source_type": "wx", "month": month}, files={"attendance_files": ("wx.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})


@pytest.mark.parametrize("content,month,reason", [
    (wx_file(), "2026-08", "核算月份"),
    (wx_file([{}, {}]), "2026-09", "重复"),
    (wx_file(omit="下班 1 打卡时间"), "2026-09", "缺少必需字段"),
    (wx_file([{"日期": "不正确的日期"}]), "2026-09", "无效的考勤日期"),
    (wx_file([{}, {"日期": datetime(2026, 8, 1)}]), "2026-09", "核算月份"),
])
def test_wx_upload_rejects_invalid_or_mismatched_attendance(client, content, month, reason):
    response = upload_wx(client, content, month)
    assert response.status_code == 400
    assert reason in response.json()["detail"]
    assert not list(app_module.CHINA_EMPLOYEE_PAYROLL_RUNS_DIR.iterdir())


def test_wx_monthly_upload_exports_clean_snapshot(client):
    response = upload_wx(client, wx_file())
    assert response.status_code == 200
    result = response.json()
    assert result["month"] == "2026-09" and result["sourceType"] == "wx"
    assert result["summary"]["totalAmount"] == 20
    exported = client.get(f"{BASE}/{result['runId']}/export")
    source = load_workbook(BytesIO(exported.content)).worksheets[1]
    assert source["A1"].value == "员工" and source["B1"].value == "工号"
    assert source["J2"].value == "21:10"  # Preserve source punch precision and type.
    assert source["A2"].value == "WX员工"
