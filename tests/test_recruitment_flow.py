from copy import deepcopy
from datetime import datetime
from io import BytesIO
import json
import threading
import time

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from bonus_platform.engine.recruitment_awards import award, calculate_draft
from bonus_platform.engine.recruitment_drafts import DraftStore
from bonus_platform.engine.recruitment_employment import enrich, status_from_timeline
from bonus_platform.engine.recruitment_monthly import project, export
from bonus_platform.engine.recruitment_beisen import parsed_snapshot
from test_recruitment_awards import person
from test_recruitment_beisen import snapshot
from test_recruitment_drafts import book


def history(**kw):
    return dict(userID=1,jobNumber='zt1',lUOffer='offer-current',entryDate='2026-07-01',startDate='2026-07-01',
                stopDate='9999-12-31',approvalStatus=4,serviceType=0,queriedAt='2026-09-16T12:00:00',
                source={'sourceId':'api','sheet':'北森同步','row':1},**kw)


def test_snapshot_preserves_approved_timelines_without_creating_hires():
    payload=snapshot();payload['recipientHistories']=[history()]
    records=parsed_snapshot(payload,'proof')['records']
    assert len([r for r in records if r['role']=='employment_history'])==3
    payload['recipientHistories'][0]['approvalStatus']=2
    assert len([r for r in parsed_snapshot(payload,'proof')['records'] if r['role']=='employment_history'])==2


def test_previous_exit_unique_and_manual_override_protected():
    row=person(工号='zt1');row.update(edits=[],sources=[],provenance={},conflicts=[])
    old=history();old.update(entryDate='2025-01-01',lastWorkDate='2026-01-19',lUOffer='old')
    enrich(row,{'zt1':[old,history()]},[])
    assert row['values']['上次离职日期']=='2026-01-19'
    assert row['provenance']['上次离职日期'][0]['sourceId']=='api'
    row['edits']=[{'field':'上次离职日期'}];row['values']['上次离职日期']='2026-01-20'
    enrich(row,{'zt1':[old,history()]},[])
    assert row['values']['上次离职日期']=='2026-01-20'
    row['edits']=[]
    enrich(row,{'zt1':[old,history()]},[])
    assert row['conflicts'][0]['field']=='上次离职日期'


def test_timeline_uses_node_date_and_does_not_infer_formal_from_null():
    row=person(推荐人入职日期='2026-07-01');row['employmentTimelines']={'推荐人':[history(lastWorkDate='2026-09-10')]}
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31))==''
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31),True)=='转正日期待核对'
    node=next(n for n in award(row,'referral',202607)['nodes'] if n['dueMonth']==202607)
    assert node['status']=='本月到期' and not node['issues']
    assert not node.get('eligibilityReasons') and not node['paymentReady']
    assert status_from_timeline(row,'推荐人',datetime(2026,9,12))=='适用日期已不在职'
    assert status_from_timeline(row,'推荐人',datetime(2026,10,12))=='适用日期晚于资料采集日期'
    row['edits']=[{'field':'推荐人最后工作日'}]
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31)) is None


def test_ambiguous_histories_and_old_probation_not_used():
    row=person();h=history(regularizationDate='2020-01-01')
    row['employmentTimelines']={'推荐人':[h]}
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31),True) is None
    other=deepcopy(h);other['userID']=2
    row['employmentTimelines']['推荐人'].append(other)
    assert status_from_timeline(row,'推荐人',datetime(2026,7,31)) is None


def test_old_candidate_probation_is_not_carried_to_a_new_employment():
    row=person(工号='zt1',转正日期='2021-12-17');row.update(edits=[],sources=[],provenance={},conflicts=[])
    enrich(row,{'zt1':[history()]},[])
    assert '转正日期' not in row['values']
    assert row['automaticFacts'][0]['before']=='2021-12-17'
    row['values']['转正日期']='2021-12-17';row['edits']=[{'field':'转正日期'}]
    enrich(row,{'zt1':[history()]},[])
    assert row['values']['转正日期']=='2021-12-17'


def test_deterministic_ineligibility_is_not_an_unknown_status():
    row=person(招聘负责人最后工作日='2026-07-10',推荐人工号='')
    result=award(row,'recruitment',202607)
    node=next(n for n in result['nodes'] if n['dueMonth']==202607)
    assert node['status']=='按规则不发' and node['scheduledMonth'] is None
    assert not any('已不在职' in i for i in node['issues'])
    assert result['currentAmount']==0
    row['unresolvedFields']=['招聘负责人最后工作日']
    uncertain=award(row,'recruitment',202607)['nodes'][0]
    assert uncertain['status']=='待核对'
    assert 'scheduledMonth' not in uncertain


def state_with_rows():
    row=person();row['rule'].update(label='2026年7月规则',basis='入职日期');row.update(edits=[],sources=[],provenance={},originKinds=['supplement'])
    state={'month':202607,'revision':1,'rows':[row],'sources':[],'todos':[],'summary':{},'edits':[],'baseline':''}
    calculate_draft(state)
    return state


def test_monthly_export_separates_pending_amounts_and_matches_screen():
    state=state_with_rows();rid=state['rows'][0]['id']
    attention={'current':[{'id':'task','kind':'field:职级','title':'核对职级','rowIds':[rid],'sourceIds':[]}], 'history':[]}
    monthly=project(state,attention)
    assert len(monthly['cases'])==1
    assert all(t['confirmed']==0 for t in monthly['totals'])
    assert any(t['pending']>0 for t in monthly['totals'])
    wb=load_workbook(BytesIO(export(state,monthly)))
    assert wb.sheetnames==['附件一-招聘奖金核算明细','附件二-招聘奖金汇总表','附件三-内推奖金汇总']
    assert wb['附件一-招聘奖金核算明细'].max_row-2==sum(r['current'] for r in monthly['worksheet']['rows'])
    for kind,name in [('recruitment','附件二-招聘奖金汇总表'),('referral','附件三-内推奖金汇总')]:
        summary=monthly['recipientSummaries'][kind]
        headers=[c.value for c in wb[name][2]]
        for cells,row in zip(wb[name].iter_rows(min_row=3,values_only=True),summary['rows']):
            assert cells[headers.index('合计发放' if kind=='recruitment' else '本月发放')]==0
            assert cells[headers.index('待核对金额')]==row['pending']


def test_worksheet_combines_installments_on_employee_row_and_export_matches():
    state=state_with_rows();row=state['rows'][0]
    base=deepcopy(row['calculation']['recruitment']['nodes'][0])
    row['calculation']['recruitment']['nodes']=[
        dict(base,id='owner-one',amount=10.1,dueMonth=202607,issues=[]),
        dict(base,id='owner-two',amount=20.2,dueMonth=202607,issues=[]),
        dict(base,id='assistant-one',role='协助招聘人',recipientId='zt-assist',recipientName='协助人',amount=5.05,dueMonth=202607,issues=[]),
    ]
    row['values'].update(协助招聘人姓名='协助人',协助招聘人工号='zt-assist')
    row['calculation']['referral']['nodes']=[dict(base,id='referral-one',role='推荐人',recipientId='zt-ref',recipientName='推荐人',amount=7.25,dueMonth=202607,issues=[])]
    old=deepcopy(row);old['id']='other-span';old['values']['入职日期']='2026-01-01'
    for kind in ('recruitment','referral'):
        old['calculation'][kind]['nodes']=[]
    state['rows'].append(old)
    monthly=project(state,{'current':[],'history':[]})
    table=monthly['worksheet'];current=[r for r in table['rows'] if r['current']]
    assert len(current)==1 and monthly['currentCount']==4
    assert len(table['rows'])==2  # same employee number, different employment remains separate
    assert current[0]['cells']['recruitmentAmount']==35.35
    assert current[0]['cells']['referralAmount']==7.25
    assert current[0]['cells']['currency']=='欧元'
    assert current[0]['cells']['status']=='已计算'
    wb=load_workbook(BytesIO(export(state,monthly)))
    ws=wb['附件一-招聘奖金核算明细']
    assert list(ws.values)[1]==tuple(c['label'] for c in table['columns'])
    assert ws.freeze_panes=='C3'
    exported=list(ws.values)[2:]
    assert exported==[tuple(None if r['cells'][c['key']]=='' else r['cells'][c['key']] for c in table['columns']) for r in current]
    assert sum(r[-3] for r in list(wb['附件二-招聘奖金汇总表'].values)[2:-1])==35.35
    assert list(wb['附件三-内推奖金汇总'].values)[2][-3]==7.25
    amounts=[i for i,c in enumerate(table['columns'],1) if c['key'] in {'recruitmentAmount','referralAmount'}]
    assert all(ws.cell(3,i).data_type=='n' for i in amounts)


def test_worksheet_keeps_award_currencies_and_never_adds_different_units():
    state=state_with_rows();row=state['rows'][0]
    for kind,currency,amount in [('recruitment','人民币',100),('referral','欧元',25)]:
        award_=row['calculation'][kind];node=deepcopy(award_['nodes'][0])
        award_.update(currency=currency,nodes=[dict(node,currency=currency,amount=amount,dueMonth=202607,issues=[])])
    monthly=project(state,{'current':[],'history':[]});item=monthly['worksheet']['rows'][0]
    assert item['cells']['currency']=='招聘：人民币；内推：欧元'
    assert item['cells']['recruitmentAmount']==100 and item['cells']['referralAmount']==25
    # An individual award with mixed currencies cannot be represented by a single amount.
    row['calculation']['recruitment']['nodes'].append(dict(row['calculation']['recruitment']['nodes'][0],id='mixed',currency='美元',amount=50))
    item=project(state,{'current':[],'history':[]})['worksheet']['rows'][0]
    assert item['cells']['recruitmentAmount'] is None and item['review']
    assert item['amountsByCurrency']['recruitment']==[{'currency':'人民币','amount':100.0},{'currency':'美元','amount':50.0}]
    assert '不同币种' in item['reasons'][0]


def test_worksheet_unknown_is_blank_not_zero_and_current_problem_is_not_lost():
    state=state_with_rows();row=state['rows'][0];row['values']['入职日期']='2026-01-01'
    row['calculation']['recruitment'].update(total=None,currentAmount=None,nodes=[],issues=['招聘启动日期缺失'])
    row['calculation']['referral'].update(total=0,currentAmount=0,nodes=[],issues=[])
    attention={'current':[{'id':'missing','kind':'cycle','title':'核对招聘日期','rowIds':[row['id']],
                            'sourceIds':[],'reasons':['招聘启动日期缺失']}],'history':[]}
    monthly=project(state,attention);item=monthly['worksheet']['rows'][0]
    assert monthly['currentCount']==0 and item['current'] and item['review']
    assert item['cells']['recruitmentAmount'] is None and item['cells']['referralAmount']==0
    assert item['reasons']==['招聘启动日期缺失'] and item['cells']['status']=='待核对'
    wb=load_workbook(BytesIO(export(state,monthly)))
    assert wb['附件一-招聘奖金核算明细'].max_row==3
    indices={c['key']:i for i,c in enumerate(monthly['worksheet']['columns'],1)}
    assert wb['附件一-招聘奖金核算明细'].cell(3,indices['recruitmentAmount']).value is None
    assert wb['附件一-招聘奖金核算明细'].cell(3,indices['referralAmount']).value==0
    assert '招聘启动日期缺失' in wb['附件一-招聘奖金核算明细'].cell(3,indices['note']).value


def test_worksheet_keeps_new_hires_before_first_installment():
    state=state_with_rows();row=state['rows'][0]
    for kind in ('recruitment','referral'):
        for node in row['calculation'][kind]['nodes']:node['dueMonth']=202608
    item=project(state,{'current':[],'history':[]})['worksheet']['rows'][0]
    assert item['current'] and item['newHire'] and not item['review']
    assert item['cells']['recruitmentAmount']==0 and item['cells']['referralAmount']==0


def test_worksheet_future_missing_facts_do_not_become_current_review():
    state=state_with_rows();row=state['rows'][0]
    row['calculation']['recruitment'].update(total=None,currentAmount=None,nodes=[],issues=['招聘启动日期缺失'])
    for node in row['calculation']['referral']['nodes']:node['dueMonth']=202608
    attention={'current':[],'history':[{'id':'later','kind':'cycle','title':'核对招聘日期',
                                       'rowIds':[row['id']],'sourceIds':[]}]}
    monthly=project(state,attention);item=monthly['worksheet']['rows'][0]
    assert item['current'] and not item['review'] and not item['reasons']
    assert item['cells']['recruitmentAmount'] is None
    assert item['cells']['status']=='后续资料待补'
    assert monthly['actionCount']==0 and monthly['reviewPeople']==0


def test_shared_recipient_case_merges_only_shared_cause():
    state=state_with_rows();a=state['rows'][0];a['values']['招聘负责人人员状态']=''
    b=deepcopy(a);b['id']='second';state['rows'].append(b);calculate_draft(state)
    attention={'current':[{'id':'task','kind':'recipient','title':'领取人任职待核','rowIds':[a['id'],b['id']],'sourceIds':[]}], 'history':[]}
    monthly=project(state,attention)
    assert len(monthly['cases'])==1 and len(monthly['cases'][0]['rowIds'])==2
    attention['current'].append({'id':'other','kind':'rehire','title':'上次离职日期','rowIds':[b['id']],'sourceIds':[]})
    assert len(project(state,attention)['cases'])==2


def test_prepare_job_is_single_flight_and_preserves_data_on_failure(tmp_path,monkeypatch):
    from bonus_platform.engine import recruitment_prepare as jobs, recruitment_beisen as adapter
    store=DraftStore(tmp_path)
    state=store.ingest(202607,[('roster.xlsx',book(['姓名','工号','入职日期','人员状态','人员来源'],[['甲','zt1','2026-07-01','正式','EHR入职']]))],0)
    gate=threading.Event();calls=[]
    def fetch(numbers,**kw):
        calls.append(numbers);gate.wait(5);raise ValueError('网络连接未完成')
    monkeypatch.setattr(adapter,'fetch',fetch)
    first=jobs.start(store,202607,state['revision']);second=jobs.start(store,202607,state['revision'])
    assert first['id']==second['id'];gate.set()
    for _ in range(100):
        result=jobs.read(store,202607)
        if result['status']!='running':break
        time.sleep(.01)
    assert result['status']=='failed' and calls==[['zt1']]
    assert store.read(202607)['revision']==state['revision']


def test_prepare_refreshes_new_recipient_with_candidate_already_covered(tmp_path,monkeypatch):
    from bonus_platform.engine import recruitment_prepare as jobs, recruitment_beisen as adapter
    store=DraftStore(tmp_path)
    state=state_with_rows();state['revision']=1
    state['rows'][0]['values'].update({'标签分类':'国内','工号':'zt1','招聘负责人工号':'zt-new'})
    state['sources']=[{'kind':'beisen','forMonth':202607,'scopeJobNumbers':['zt1'],'scopeRecipientJobNumbers':['zt-old'],
        'syncCoverage':{'version':1,'candidates':{'zt1':{'requested':True,'employment':True,'offer':True}},
                        'recipients':{'zt-old':{'requested':True,'employment':True}}}}]
    monkeypatch.setattr(store,'refresh',lambda *args:state)
    monkeypatch.setattr(store,'read',lambda *args:state)
    calls=[]
    monkeypatch.setattr(adapter,'fetch_employment',lambda numbers:calls.append(numbers) or {})
    monkeypatch.setattr(store,'save_beisen_snapshot',lambda *args:state)
    job={'status':'running'}
    jobs.run(store,202607,{'revision':1,**job},False,jobs.key_for(store,202607))
    assert calls and 'zt-new' in calls[0] and 'zt1' not in calls[0]
    assert jobs.read(store,202607)['status']=='complete'


def test_confirmation_and_export_reject_stale_revision(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    store=DraftStore(tmp_path);monkeypatch.setattr(api,'store',store)
    state=state_with_rows();state.update(updatedAt='',parentMonth=None,parentRevision=None,paymentReady=False,manualRows=[])
    store._write(store._folder(202607)/'draft.json',state)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    assert client.get('/api/recruitment/drafts/202607/monthly/export?revision=0').status_code==409
    assert client.get('/api/recruitment/drafts/202607/monthly/export?revision=1').status_code==200
    assert client.post('/api/recruitment/drafts/202607/confirm-month',json={'revision':0,'coverageConfirmed':True}).status_code==409
    assert client.post('/api/recruitment/drafts/202607/confirm-month',json={'revision':1}).status_code==400
    assert client.post('/api/recruitment/drafts/202607/confirm-month',json={'revision':1,'coverageConfirmed':True}).status_code==200
    assert json.loads((store._folder(202607)/'accepted.json').read_text())['revision']==1


def test_upload_and_export_wait_for_background_preparation(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    from bonus_platform.engine import recruitment_prepare as jobs
    store=DraftStore(tmp_path);monkeypatch.setattr(api,'store',store)
    state=state_with_rows();state.update(updatedAt='',parentMonth=None,parentRevision=None,paymentReady=False,manualRows=[])
    store._write(store._folder(202607)/'draft.json',state)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    key=jobs.key_for(store,202607)
    monkeypatch.setitem(jobs.RUNNING,key,'pending-preparation')
    exported=client.get('/api/recruitment/drafts/202607/monthly/export?revision=1')
    uploaded=client.post('/api/recruitment/drafts/202607/sources',data={'revision':'1'},
                         files={'files':('supplement.xlsx',book(['姓名','工号'],[['新增测试','qa-new']]))})
    assert exported.status_code==uploaded.status_code==409
    assert '正在整理' in exported.json()['detail']
    assert store.read(202607)['revision']==1
    assert not (store._folder(202607)/'sources').exists()
    jobs.RUNNING.pop(key)
    assert client.get('/api/recruitment/drafts/202607/monthly/export?revision=1').status_code==200


def test_complete_month_confirm_and_carry_uses_frozen_version(tmp_path,monkeypatch):
    from bonus_platform import recruitment_draft_api as api
    store=DraftStore(tmp_path);monkeypatch.setattr(api,'store',store)
    v=person('2026-10-01',体系外人员编号='test1',招聘启动日期='2026-08-01',推荐人工号='',推荐人姓名='')['values']
    state=store.ingest(202610,[('supplement.xlsx',book(list(v),[list(v.values())],'导入_月度数据')),
        ('roster.xlsx',book(['姓名','工号','入职日期','人员状态','人员来源'],[[v['姓名'],v['工号'],'2026-10-01','正式','OEHR入职']]))],0)
    app=FastAPI();app.include_router(api.router);client=TestClient(app)
    opened=client.post('/api/recruitment/drafts/202610/opening-empty',json={'revision':state['revision'],'confirmed':True})
    assert opened.status_code==200,opened.text
    state=opened.json();assert not state['monthly']['cases']
    confirmed=client.post('/api/recruitment/drafts/202610/confirm-month',json={'revision':state['revision'],'coverageConfirmed':True})
    assert confirmed.status_code==200,confirmed.text
    changed=client.post('/api/recruitment/drafts/202610/changes',json={'revision':state['revision'],'edits':[
        {'rowId':state['rows'][0]['id'],'field':'人工备注','value':'确认之后的新修订','reason':'测试后续修订'}]})
    assert changed.status_code==200,changed.text
    child=store.carry_forward(202611,202610,0)
    assert child['parentRevision']==state['revision']
    assert child['rows'][0]['values'].get('人工备注')!='确认之后的新修订'
    # A month with historical records can never be made an empty opening.
    assert client.post('/api/recruitment/drafts/202611/opening-empty',json={'revision':child['revision'],'confirmed':True}).status_code==400
