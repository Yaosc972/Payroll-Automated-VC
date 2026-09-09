"""Exported detail explanations cover every subject and preserve stable rows."""
import pytest
from openpyxl import load_workbook
from bonus_platform.engine.domestic_labor.exporter import ExcelExporter


@pytest.mark.parametrize('subject', [key for key, _ in ExcelExporter.SUBJECTS])
def test_every_subject_exports_last_process_column_and_equal_heights(tmp_path, subject):
    daily = {'status': 'calculated', 'reason_code': 'calculated', 'amount': 13.8,
             'attendance_hours': 8, 'temperature': 35}
    if subject == 'yeban_butie':
        daily = {'status': 'calculated', 'reason_code': 'generic_rule', 'amount': 3,
                 'night_minutes': 120, 'break_minutes': 60}
    audit = {'formula': '标准÷排班天数×计发天数', 'inputs': {'标准': 600, '排班天数': 20},
             'intermediate_values': {'计发天数': 10}, 'steps': ['600÷20×10=300元']}
    details = {'audit_explanation': audit, 'daily_results': [daily], '小时单价': 1.725,
               '单日封顶': 13.8, '月度封顶': 300, '月度封顶前金额': 13.8}
    records = [{'employee_id': str(i), subject: 300, 'subject_details': {subject: {'details': details}}}
               for i in (1, 2)]
    path = tmp_path / 'result.xlsx'
    ExcelExporter(str(path)).export(records, '202607')
    wb = load_workbook(path)
    ws = next(ws for ws in wb if ws.title in {'计算详情', '每日明细', '岗位补贴核算结果'})
    assert ws.cell(1, ws.max_column).value == '计算过程'
    assert {ws.row_dimensions[i].height for i in (2, 3)} == {24}
    assert all(not cell.alignment.wrap_text for row in ws.iter_rows(min_row=2) for cell in row)
    process = ws.cell(2, ws.max_column).value
    if subject == 'yeban_butie':
        assert '120分钟' in process and process.endswith('=3元')
    elif subject == 'gaowen_butie':
        assert '8小时×1.725元/小时' in process and '月应发=' in process
    else:
        assert '600÷20×10=300元' in process
    assert ws.auto_filter.ref == ws.dimensions
    wb.close()


def test_combined_export_keeps_each_subject_evidence_and_missing_history(tmp_path):
    result = {'employee_id': '1', 'canbu': 50, 'gonglingjiang': 100,
              'subject_details': {'canbu': {'audit_explanation': {'formula': '餐补公式', 'steps': ['餐补步骤']}},
                                  'gonglingjiang': {'audit_explanation': {'formula': '工龄公式', 'steps': ['工龄步骤']}}}}
    path = tmp_path / 'combined.xlsx'
    ExcelExporter(str(path)).export([result, {'employee_id': '2', 'canbu': 10}], '202607')
    wb = load_workbook(path)
    ws = wb['计算详情']
    text = ws.cell(2, ws.max_column).value
    assert '餐补公式' in text and '工龄公式' in text
    assert '未保存计算过程' in ws.cell(3, ws.max_column).value
    wb.close()
