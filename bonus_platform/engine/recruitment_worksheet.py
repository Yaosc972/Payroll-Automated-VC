"""A complete, one-employment worksheet shared by the desktop table and export."""
from collections import defaultdict
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .recruitment_awards import ensure_calculated, money, NODES
from .recruitment_attention import project, named_month
from .recruitment_draft_sources import date_text


BASE_GROUPS = [
    ('人员与组织', ['姓名','工号','证件号','入职日期','人员状态','最后工作日','工作地','职级','二级组织','三级组织','四级组织','五级组织','职位','直线经理','直线经理工号']),
    ('招聘关系', ['标签分类','原始ABC类别','ABC类别','招聘负责人姓名','招聘负责人工号','协助招聘人姓名','协助招聘人工号','原始招聘渠道','招聘渠道','其他渠道','海外offer流程编号']),
    ('招聘周期', ['招聘启动日期','候选人入职时间','转正日期','原计划入职日期','上次离职日期','周期剔除天数','周期剔除依据','周期豁免依据']),
]
TITLES = {'ABC类别':'ABC类别','原始ABC类别':'原始ABC类别','标签分类':'标签分类（国内/海外）','海外offer流程编号':'海外offer流程编号','证件号':'身份证','工号':'工号'}


def columns():
    out=[]
    def add(key,title,group,kind='text',views=('all',)):
        out.append({'key':key,'title':title,'group':group,'type':kind,'views':list(views),'width':190 if '组织' in title or '依据' in title or '备注' in title else 148 if '姓名' in title else 128})
    for group,fields in BASE_GROUPS:
        for field in fields:
            views=['all','facts']
            if field in {'姓名','工号','入职日期','人员状态','最后工作日','职级','ABC类别','招聘负责人姓名','协助招聘人姓名','招聘渠道','招聘启动日期','转正日期'}: views+=['review']
            if field in {'姓名','工号','入职日期','招聘负责人姓名','协助招聘人姓名'}: views+=['payments']
            add(field,TITLES.get(field,field),group,'number' if field=='周期剔除天数' else 'date' if '日期' in field or field=='最后工作日' else 'text',views)
    for key,title in [('standardDays','标准周期'),('actualDays','采用周期'),('holidayDays','法定假期剔除'),('exclusionDays','采用剔除'),('factor','周期系数'),('exempt','周期豁免')]:
        add('cycle.'+key,title,'招聘周期','percent' if key=='factor' else 'text' if key=='exempt' else 'number',('all','review','facts'))
    add('channelFactor','渠道系数','招聘周期','percent',('all','review','facts'))
    add('rule','适用规则','奖金标准',views=('all','review','payments'))
    add('currency','招聘币种','奖金标准',views=('all','review','payments'))
    add('referral.currency','内推币种','奖金标准',views=('all','review','payments'))
    for kind,prefix in [('recruitment','招聘'),('referral','内推')]:
        add(kind+'.standard',prefix+'规则标准','奖金标准','money',('all','review'))
        add(kind+'.total',prefix+'试算总额','奖金标准','money',('all','review'))
        add(kind+'.formulaTotal',prefix+'公式原值','奖金标准','money',('all',))
        add(kind+'.scheduledTotal',prefix+'分期合计','奖金标准','money',('all','review','payments'))
        add(kind+'.roundingDifference',prefix+'分期尾差','奖金标准','money',('all','review','payments'))
        add(kind+'.currentAmount',prefix+'本月安排','本月结果','money',('all','review','payments'))
    for kind,role in [('recruitment','招聘负责人'),('recruitment','协助招聘人'),('referral','推荐人')]:
        for key,title in [('allocatedAmount','原分配总额'),('scheduledAmount','分期合计'),('roundingDifference','分期尾差')]:
            add(f'{kind}.{role}.{key}',title,role+'分期','money',('all','payments'))
        for i,short in enumerate(['满1个月','满3个月','满6个月','转正']):
            add(f'{kind}.{role}.{i}.amount',short+'奖金',role+'分期','money',('all','payments'))
            add(f'{kind}.{role}.{i}.dueDate',short+'到期日',role+'分期','date',('all','payments'))
            add(f'{kind}.{role}.{i}.scheduledMonth',short+'安排月份',role+'分期','text',('all','payments'))
        if kind=='recruitment':add(f'{kind}.{role}.cancelledBalance','本月取消未发余额',role+'分期','money',('all','payments'))
    for prefix in ['招聘负责人','协助招聘人','推荐人']:
        fields=(['姓名','工号','职级','职位','转正日期'] if prefix=='推荐人' else [])+['入职日期','最后工作日','人员状态']
        for suffix in fields:
            add(prefix+suffix,prefix+suffix,'领取人资格',views=('all','facts'))
    for key,title in [('人工备注','备注'),('issues','待核原因'),('history','历史记录'),('payment','实际支付'),('source','采用来源')]:
        add(key,title,'核对与依据',views=('all','review') if key in {'人工备注','issues'} else ('all',))
    return out


def worksheet(state):
    ensure_calculated(state)
    names={s['id']:s['name'] for s in state['sources']}
    sources={s['id']:s for s in state['sources']}
    reference_columns={}
    result=[]
    for row in state['rows']:
        v=row['values'];a=row['calculation'];cells=dict(v)
        cells.update(rule=row['rule']['label'],currency=a['recruitment']['currency'],payment='未核实',history=f"{len(row['history'])} 份台账原记录" if row['history'] else '无历史台账',source='；'.join(dict.fromkeys(names.get(r['sourceId'],'人工补充') for r in row['sources'])))
        for k,value in (a['recruitment']['cycle'] or {}).items(): cells['cycle.'+k]=value
        cells['channelFactor']=a['recruitment']['channelFactor']
        cells['cycle.exempt']='是' if cells.get('cycle.exempt') else '否'
        reasons=[]
        for kind,title in [('recruitment','招聘'),('referral','内推')]:
            award=a[kind]
            for key in ('currency','standard','total','formulaTotal','scheduledTotal','roundingDifference','currentAmount'): cells[kind+'.'+key]=award.get(key)
            for allocation in award.get('recipientTotals',[]):
                for key in ('allocatedAmount','scheduledAmount','roundingDifference'):
                    cells[f'{kind}.{allocation["role"]}.{key}']=allocation[key]
            reasons.extend(title+'：'+s for s in award['issues'])
            for node in award['nodes']:
                if node.get('adjustmentId') or node.get('externalResultId'):continue
                index=NODES.index(node['label'])
                for key in ('amount','dueDate','scheduledMonth'):
                    cells[f'{kind}.{node["role"]}.{index}.{key}']=node.get(key,node['dueMonth'] if key=='scheduledMonth' else '')
            for cancelled in award.get('cancelledBalances',[]):cells[f'{kind}.{cancelled["role"]}.cancelledBalance']=cancelled['amount']
        cells['issues']='；'.join(dict.fromkeys(reasons))
        histories=[h for h in row['history'] if sources.get(h['ref']['sourceId'],{}).get('active',True)]
        if histories:
            chosen=max(histories,key=lambda h:(named_month([names.get(h['ref']['sourceId'],'')]) or 0,h['ref']['row']))
            cells['reference.file']=names.get(chosen['ref']['sourceId'],'')
            cells['reference.row']=chosen['ref']['row']
            for cell in chosen['cells']:
                # Preserve the source's complete financial groups, arrangements,
                # adjustments and notes as reference values, never as formulas.
                if cell['column']<32: continue
                key='reference.'+str(cell['column'])
                parts=cell['label'].split(' / ')
                group=parts[0] if len(parts)>1 else '线下其他记录'
                title=parts[-1].split('\n')[0]
                financial=(any(t in title for t in ('奖金','金额','发放标准','扣减')) or '扣减' in group) and '周期' not in title and '币种' not in title
                reference_columns[key]={'key':key,'title':title,'group':'线下 · '+group,'type':'money' if financial else 'text','views':['reference'],'width':150}
                cells[key]=cell['value']
        result.append({'id':row['id'],'cells':cells})
    cols=columns()+[reference_columns[k] for k in sorted(reference_columns,key=lambda k:int(k.split('.')[1]))]
    return {'columns':cols,'rows':result}


def safe(value):
    if isinstance(value,str) and value.startswith(('=','+','-','@')): return "'"+value
    return value


def style_sheet(ws,header_row=1,freeze='D2'):
    ws.freeze_panes=freeze
    ws.auto_filter.ref=f'A{header_row}:{get_column_letter(ws.max_column)}{ws.max_row}'
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    for cell in ws[header_row]:
        cell.font=Font(name='等线',size=11,bold=True,color='253858')
        cell.fill=PatternFill('solid',fgColor='EAF0F7')
        cell.alignment=Alignment(vertical='center',wrap_text=True)
        cell.border=Border(bottom=Side(style='thin',color='D7DEE8'))
    ws.row_dimensions[header_row].height=34
    body_font=Font(name='等线',size=11,color='172B4D')
    body_alignment=Alignment(vertical='center')
    stripe_fill=PatternFill('solid',fgColor='F7F9FC')
    for row in ws.iter_rows(min_row=header_row+1):
        for cell in row:
            cell.font=body_font
            cell.alignment=body_alignment
            if cell.row%2==0: cell.fill=stripe_fill
            if isinstance(cell.value,(int,float)): cell.number_format='#,##0.00;[Red](#,##0.00);–'
    for index in range(1,ws.max_column+1): ws.column_dimensions[get_column_letter(index)].width=19


def export_results(state):
    data=worksheet(state);wb=Workbook();wb.remove(wb.active)
    cols=data['columns'];ws=wb.create_sheet('核算总明细')
    ws.append([f"{state['month']//100}年{state['month']%100}月 招聘与内推奖金核算明细（草稿）"])
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
    ws.cell(1,1).font=Font(name='等线',size=18,bold=True,color='1E3A8A');ws.row_dimensions[1].height=38
    ws.append([c['group'] for c in cols]);ws.append([c['title'] for c in cols])
    start=1
    for i in range(2,len(cols)+2):
        if i==len(cols)+1 or cols[i-1]['group']!=cols[start-1]['group']:
            if i-start>1: ws.merge_cells(start_row=2,start_column=start,end_row=2,end_column=i-1)
            cell=ws.cell(2,start);cell.font=Font(name='等线',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='385B85');cell.alignment=Alignment(vertical='center')
            start=i
    ws.row_dimensions[2].height=26
    for row in data['rows']: ws.append([safe(row['cells'].get(c['key'])) for c in cols])
    style_sheet(ws,3,'D4')
    for i,col in enumerate(cols,1):
        ws.column_dimensions[get_column_letter(i)].width=max(16,round(col['width']/7))
        for row in ws.iter_rows(min_row=4,min_col=i,max_col=i):
            if col['type']=='percent':row[0].number_format='0%'
            elif col['type']!='money':row[0].number_format='0.########'
    def sheet(name,headers,rows):
        w=wb.create_sheet(name);w.append(headers)
        for r in rows: w.append([safe(v) for v in r])
        style_sheet(w);return w
    nodes=[];summaries=defaultdict(Decimal)
    for row in state['rows']:
        for kind,title in [('recruitment','招聘奖金'),('referral','内推奖金')]:
            for n in row['calculation'][kind]['nodes']:
                if n.get('scheduledMonth',n['dueMonth'])!=state['month']: continue
                nodes.append([row['values'].get('姓名'),row['values'].get('工号'),title,n['role'],n['recipientName'],n['recipientId'],n['label'],n['dueDate'],n['amount'],n['currency'],'待核对' if n['issues'] else '待确认','；'.join(n['issues']),'未确认','未核实',n.get('scheduledMonth',n['dueMonth']),n.get('amountBasis','制度计算'),n.get('openingAmount'),n.get('scheduleText','')])
                summaries[(n['recipientId'],n['recipientName'],n['currency'])]+=Decimal(str(n['amount']))
    sheet('本月到期明细',['候选人','候选人工号','奖金类型','领取角色','领取人','领取人工号','奖金节点','规则到期日','本月安排金额','标准币种','复核状态','待核原因','本次确认','实际支付','安排月份','金额依据','期初分期金额','原发放安排'],nodes)
    sheet('领取人分币种汇总',['领取人','领取人工号','标准币种','本月到期试算','本次确认','实际支付'],[[name,rid,currency,money(value),'未确认','未核实'] for (rid,name,currency),value in sorted(summaries.items())])
    attention=project(state);by_id={r['id']:r for r in state['rows']}
    pending=[]
    for scope,key in [('本月','current'),('历史','history')]:
        for t in attention[key]:
            for rid in t['rowIds'] or ['']:
                row=by_id.get(rid,{});v=row.get('values',{})
                pending.append([scope,v.get('姓名',''),v.get('工号',''),v.get('入职日期',''),t['title'],'；'.join(t['reasons']),'；'.join(names for sid in t['sourceIds'] for names in [next((s['name'] for s in state['sources'] if s['id']==sid),'')]),'','','','','未核实'])
    for row in state['rows']:
        for kind,title in [('recruitment','招聘奖金'),('referral','内推奖金')]:
            for n in row['calculation'][kind]['nodes']:
                if n.get('scheduledMonth',n['dueMonth'])==state['month']:continue
                pending.append([n['timing'],row['values'].get('姓名'),row['values'].get('工号'),row['values'].get('入职日期'),title+' · '+n['label'],'；'.join(n['issues']),'',n['recipientName'] or n['recipientId'],n['dueDate'],n['amount'],n['currency'],'未核实'])
    sheet('未结与异常',['范围','姓名','工号','入职日期','待处理事项','具体原因','相关资料','领取人','规则到期日','节点试算金额','标准币种','实际支付'],pending)
    names={s['id']:s['name'] for s in state['sources']};ops={o['id'] for o in state.get('operations',[]) if o['active']}
    evidence=[]
    for edit in state['edits']:
        location='；'.join(f"{ref.get('sheet','')} 第 {ref.get('row','')} 行"+(f"，第 {ref['column']} 列" if ref.get('column') else '') for ref in edit.get('evidenceRefs',[]))
        evidence.append([by_id.get(edit['rowId'],{}).get('values',{}).get('姓名',''),edit['field'],edit.get('before'),edit['value'],edit['reason'],names.get(edit.get('evidenceId'),''),','.join(map(str,edit.get('evidencePages',[]))),edit['at'],'已采用' if edit.get('active',True) else '已撤销',location])
    for decision in state.get('awardDecisions',[]):
        evidence.append([by_id.get(decision['rowId'],{}).get('values',{}).get('姓名',''),'招聘奖金' if decision['kind']=='recruitment' else '内推奖金','按规则试算','手工设置不发' if decision.get('basisType')=='manual' else '按审批不发',decision['reason'],names.get(decision['sourceId'],''),','.join(map(str,decision['pageNumbers'])),decision.get('at',''),'已采用' if decision['operationId'] in ops else '已撤销'])
    for a in state.get('adjustmentDrafts',[]):
        evidence.append([a['recipientName'],'补发 / 扣回安排','',a['amount'],a['currency']+'；实际执行待核',names.get(a['sourceId'],''),','.join(map(str,a['pageNumbers'])),'','已采用' if a['operationId'] in ops else '已撤销'])
    for plan in state.get('installmentPlans',[]):
        evidence.append([by_id.get(plan['rowId'],{}).get('values',{}).get('姓名',''),'分期安排 · '+plan['label'],plan.get('beforeMonth'),plan.get('month') if plan['action']=='schedule' else '暂缓安排' if plan['action']=='defer' else '恢复规则安排',plan['reason'],names.get(plan['sourceId'],''),'','','已采用' if plan['operationId'] in ops else '已撤销'])
    from .recruitment_special import context, MODES, CONDITIONS
    for record in context(state)['records']:
        typ=record['type'];label={'adjustment':'手工补扣' if record.get('basisType')=='manual' else '审批补扣','exception':'资格例外','cutoff':'核算截点'}[typ]
        value=(str(record['amount'])+' '+record['currency'] if typ=='adjustment' else '、'.join(CONDITIONS[c] for c in record['conditions']) if typ=='exception' else MODES[record['mode']]+(' '+record.get('date','')))
        evidence.append([record.get('name','本月范围'),label,'',value,str(record['month'])+'；'+record['reason']+'；实际支付未核实',names.get(record['sourceId'],''),','.join(map(str,record.get('pageNumbers',[]))),record['at'],'已撤销' if not record['active'] else '依据失效' if not record['sourceValid'] else '已采用'])
    sheet('调整及依据',['姓名','调整内容','原值','采用值','依据说明','依据文件','页码','时间','状态','来源位置'],evidence)
    if state.get('strictInputs'):
        validation=state.get('ruleValidation',{})
        sheet('规则检查',['事项','内容'],[['独立试算批次',state.get('trial',{}).get('name','')],['规则检查','通过' if validation.get('passed') else '未通过'],['核对结果用途','只比较，不参与计算'],['实际支付','未核实']]+[[p['name']+' · '+('招聘奖金' if p['kind']=='recruitment' else '内推奖金'),'；'.join(p['reasons'])] for p in validation.get('problems',[])])
    out=BytesIO();wb.save(out);return out.getvalue()
