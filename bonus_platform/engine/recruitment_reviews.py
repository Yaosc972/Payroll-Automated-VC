"""Review previews, revision-safe application, and editable workbook round trips."""
from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import json
import math
from pathlib import Path
import re
from uuid import uuid4
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .recruitment_time import now as recruitment_now
from .recruitment_awards import calculate_draft, ensure_calculated, money
from .recruitment_draft_sources import DATE_FIELDS, date_text, text
from .recruitment_drafts import EDITABLE_FIELDS
from .recruitment_documents import extract_document, mentioned_attachments

GRID_FIELDS=['姓名','工号','入职日期','职级','ABC类别','招聘渠道','招聘启动日期','招聘负责人工号','招聘负责人姓名','协助招聘人工号','协助招聘人姓名','推荐人工号','推荐人姓名','周期剔除天数','周期剔除依据','原计划入职日期','应届校招','延后入职原因','管培生已批名单资格','周期豁免依据','上次离职日期','人工备注']
FIELD_LABELS={'ABC类别':'岗位类别'}
REVIEW_STATES={'待核实','已认可','待补附件','审批未结束','已驳回'}


def document(store,month,source_id):
    state=store.read(month)
    source=next((s for s in state['sources'] if s['id']==source_id and s.get('attachment')),None)
    if not source: raise ValueError('请选择审批文件或图片')
    folder=store._folder(month)
    data=extract_document(folder/'sources'/source['storedName'],folder/'documents'/source_id)
    full='\n'.join(p['text'] for p in data['pages'])
    compact=re.sub(r'\s+','',full).lower()
    searchable=re.sub(r'\s+',' ',full).lower()
    matches=[]
    for row in state['rows']:
        values=row['values'];keys=[]
        for field in ('工号','海外offer流程编号','海外招聘需求编号'):
            value=text(values.get(field)).lower()
            if len(value)>=5 and re.search(r'(?<![a-z0-9])'+re.escape(value)+r'(?![a-z0-9])',searchable): keys.append(field)
        name=re.sub(r'\s+','',text(values.get('姓名'))).lower()
        if not keys and len(name)>=3 and name in compact: keys.append('姓名，需核对任职')
        if keys: matches.append({'rowId':row['id'],'name':values.get('姓名'),'employeeId':values.get('工号'),'hireDate':values.get('入职日期'),'matchedBy':keys})
    active={o['id'] for o in state.get('operations',[]) if o['active']}
    return {**data,'sourceId':source_id,'name':source['name'],'flowIds':list(dict.fromkeys(re.findall(r'ZTJT[-—]LC[-—]\d+',full,re.I))),
            'mentionedAttachments':mentioned_attachments(data),'matches':matches,
            'reviews':[{**r,'active':not r.get('operationId') or r['operationId'] in active} for r in state.get('reviewEntries',[]) if r['sourceId']==source_id],
            'arrangements':[{**r,'active':r['operationId'] in active} for r in state.get('adjustmentDrafts',[]) if r['sourceId']==source_id],
            'revision':state['revision']}


def evidence_check(store,state,evidence,applying=False):
    if not isinstance(evidence,dict): raise ValueError('审批依据格式无效')
    sid=evidence.get('sourceId')
    source=next((s for s in state['sources'] if s['id']==sid and s.get('active',True)),None)
    if not source or source.get('comparisonOnly') or not source.get('attachment'): raise ValueError('审批原件不存在或已停用')
    info=document(store,state['month'],sid)
    pages=evidence.get('pageNumbers',[])
    if not pages or any(type(p) is not int or p<1 or p>info['pageCount'] for p in pages): raise ValueError('请选择有效的依据页码')
    status=evidence.get('status','待核实')
    if status not in REVIEW_STATES: raise ValueError('审批核对状态无效')
    attachments=evidence.get('attachmentIds',[])
    supplied=[]
    for source_id in attachments:
        found=next((s for s in state['sources'] if s['id']==source_id and s.get('active',True) and not s.get('error')),None)
        if not found or found.get('comparisonOnly'): raise ValueError('关联附件不存在、已停用或仅供结果核对')
        supplied.append(found['name'])
    missing=[name for name in info['mentionedAttachments'] if name not in supplied]
    if applying and (status!='已认可' or not evidence.get('scopeChecked')): raise ValueError('先核对最终认可状态与适用范围')
    if applying and missing: raise ValueError('审批提到的附件尚未关联：'+'、'.join(missing))
    if not text(evidence.get('note')): raise ValueError('请填写审批核对说明')
    active_ops = {o['id'] for o in state.get('operations', []) if o.get('active')}
    previous = next((r for r in reversed(state.get('reviewEntries', [])) if r['sourceId'] == sid
                     and (not r.get('operationId') or r['operationId'] in active_ops)), {})
    scope = evidence.get('personScope', previous.get('personScope'))
    extra = {}
    if scope is not None:
        if not isinstance(scope, list) or not scope or len(scope) > 500:
            raise ValueError('审批人员范围格式无效')
        rows = {r['id']: r for r in state['rows']}
        compact = lambda value: re.sub(r'\s+', '', text(value)).casefold()
        body = compact('\n'.join(p.get('text', '') for p in info['pages'] if p['page'] in pages))
        names = set(); checked = []
        for person in scope:
            if not isinstance(person, dict): raise ValueError('审批人员记录格式无效')
            name = text(person.get('name')); reason = text(person.get('reason')); rid = person.get('rowId')
            if not name or compact(name) in names or compact(name) not in body or not reason:
                raise ValueError('审批人员姓名或匹配说明无效')
            if rid and (rid not in rows or compact(rows[rid]['values'].get('姓名')) != compact(name)):
                raise ValueError('审批人员与本次任职不匹配')
            names.add(compact(name)); checked.append({'name': name, 'rowId': rid, 'reason': reason})
        recruiter = evidence.get('recruiter', previous.get('recruiter', {}))
        name = text(recruiter.get('name')); number = text(recruiter.get('employeeId'))
        known_pair = any(r['values'].get('招聘负责人姓名') == name and
                         r['values'].get('招聘负责人工号') == number for r in state['rows'])
        if not name or compact(name) not in body or not number or not (compact(number) in body or known_pair):
            raise ValueError('审批招聘负责人须有正文姓名及可核实的工号依据')
        extra = {'personScope': checked, 'recruiter': {'name': name, 'employeeId': number}}
    return {**extra,'sourceId':sid,'pageNumbers':sorted(set(pages)),'status':status,'scopeChecked':bool(evidence.get('scopeChecked')),'scopeComplete':bool(evidence.get('scopeComplete')),
            'attachmentIds':attachments,'missingAttachments':missing,'note':text(evidence['note']),'at':recruitment_now().isoformat(timespec='seconds')}


def save_review(store,month,revision,evidence):
    with store._lock(month) as folder:
        state=store.read(month);store._check(state,revision)
        reviewed=evidence_check(store,state,evidence)
        state.setdefault('reviewEntries',[]).append({**reviewed,'id':uuid4().hex,'applied':False})
        return store._save(folder,store._rebuild(folder,state))


def prepare_preview(store,month,revision,edits,*,evidence=None,origin='资料更正',arrangements=None,decisions=None,basis_type=None,approval_reviews=None,reference_selections=None,employment_selections=None):
    with store._lock(month) as folder:
        before=ensure_calculated(store.read(month));store._check(before,revision)
        if not isinstance(edits,list) or len(edits)>5000: raise ValueError('单次最多修改5000个单元格')
        from . import recruitment_reference_facts as reference
        if any(reference.TRUSTED_KEYS.intersection(e) for e in edits if isinstance(e,dict)):
            raise ValueError('可信事实依据只能通过服务端选择生成')
        from .recruitment_field_controls import expand_choices
        edits=expand_choices(deepcopy(edits))
        retained = [edit for edit in edits for review in approval_reviews or []
                    for item in review.get('adoptedItems',[])
                    if edit.get('evidenceId')==review['sourceId'] and edit['rowId']==item.get('rowId')
                    and edit['field'] in item.get('fields',{}) and edit['field'] not in item.get('referenceFields',[])]
        edits.extend(reference.selection_edits(folder,before,reference_selections if reference_selections is not None else [],
                                               retained_edits=retained))
        if employment_selections is not None:
            from .recruitment_employment import employment_edits
            store._prepare_change(folder,deepcopy(before),{'edits':edits},trusted_edits=True,validate_only=True)
            edits.extend(employment_edits(before,employment_selections,proposed_edits=edits))
        cells=[(e.get('rowId'),e.get('field')) for e in edits if isinstance(e,dict)]
        if len(cells)!=len(edits) or len(set(cells))!=len(cells): raise ValueError('更正内容含重复单元格或无效记录')
        manual=basis_type=='manual'
        if manual and evidence:raise ValueError('手工调整不能同时标记为审批调整')
        reviewed=evidence_check(store,before,evidence,True) if evidence else None
        batch_reviews=[{**evidence_check(store,before,r,True),
                        'itemIds':r.get('itemIds',[]),'rowIds':r.get('rowIds',[]),'adoptedItems':r.get('adoptedItems',[])} for r in (approval_reviews or [])]
        if batch_reviews and (evidence or manual):raise ValueError('审批调整的依据类型不一致')
        batch_by_source={r['sourceId']:r for r in batch_reviews}
        edits=deepcopy(edits)
        for edit in edits:
            if reviewed and not edit.get('referenceFact'): edit.update(evidenceId=reviewed['sourceId'],reason=reviewed['note'])
        after=store._prepare_change(folder,deepcopy(before),{'edits':edits},trusted_edits=True)
        submitted={(e['rowId'],e['field']) for e in edits if e.get('referenceFact') or e.get('employmentBasis')}
        if any((e['rowId'],e['field']) in submitted for e in after.get('referenceFacts',{}).get('staleEdits',[])):
            raise ValueError('更正后的任职或来源与所选事实不符，请重新核对')
        decisions=deepcopy(decisions or [])
        if decisions and not reviewed and not manual and not batch_reviews: raise ValueError('不发决定需要关联审批依据')
        for decision in decisions:
            if decision.get('rowId') not in {r['id'] for r in before['rows']} or decision.get('kind') not in {'recruitment','referral'} or decision.get('action')!='exclude': raise ValueError('奖金处理决定无效')
            roles={'recruitment':{'招聘负责人','协助招聘人'},'referral':{'推荐人'}}
            if decision.get('role') and decision['role'] not in roles[decision['kind']]:raise ValueError('奖金处理角色无效')
            proof=reviewed or batch_by_source.get(decision.get('sourceId'))
            if batch_reviews and not proof:raise ValueError('不发决定缺少对应审批原件')
            decision.update(id=uuid4().hex,at=recruitment_now().isoformat(timespec='seconds'),
                basisType='manual' if manual else 'approval',sourceId=proof['sourceId'] if proof else '',
                pageNumbers=proof['pageNumbers'] if proof else [],reason=proof['note'] if proof else '手工设置不发奖金')
            row=next(r for r in after['rows'] if r['id']==decision['rowId'])
            row.setdefault('awardDecisions',[]).append(decision)
        after['revision']+=1;calculate_draft(after)
        people={r['id']:r for r in before['rows']};changed={r['id']:r for r in after['rows']}
        arrangements=deepcopy(arrangements or [])
        if arrangements and not reviewed: raise ValueError('补扣安排需要关联审批依据')
        for item in arrangements:
            amount=item.get('amount');currency=item.get('currency')
            if item.get('rowId') not in people: raise ValueError('补扣人员不在本月底稿中')
            if not isinstance(amount,(int,float)) or isinstance(amount,bool) or not math.isfinite(amount) or amount==0: raise ValueError('请填写非零的有效补扣金额')
            if currency not in {'人民币','欧元','美元','捷克克朗'}: raise ValueError('补扣币种无效')
            if not text(item.get('recipientId')) or not text(item.get('recipientName')): raise ValueError('请补充补扣领取人身份')
            rounded=money(amount)
            if not rounded: raise ValueError('补扣金额四舍五入后不能为零')
            item.update(amount=rounded,id=uuid4().hex,status='待核实际支付',sourceId=reviewed['sourceId'],pageNumbers=reviewed['pageNumbers'],paymentReady=False)
        if not edits and not arrangements and not decisions and not batch_reviews: raise ValueError('没有需要应用的修改')
        changes=[{'rowId':e['rowId'],'name':people[e['rowId']]['values'].get('姓名'),'field':e['field'],'before':people[e['rowId']]['values'].get(e['field'],''),'after':changed[e['rowId']]['values'].get(e['field'],'')} for e in edits]
        changes += [{'rowId':d['rowId'],'name':people[d['rowId']]['values'].get('姓名'),'field':(d.get('role') or ('招聘奖金' if d['kind']=='recruitment' else '内推奖金'))+'处理','before':'按规则试算','after':'手工设置不发' if manual else '按审批不发'} for d in decisions]
        impact=[]; current_impact=[]; node_impact=[]; remaining=[];unresolved_amounts=[]
        edited_ids={e['rowId'] for e in edits+decisions}
        from .recruitment_attention import is_current
        # Assembly can propagate a person's employment facts to other recipients.
        # Compare every affected award, including scheduling-only changes.
        node_fields=('amount','currency','recipientId','recipientName','dueDate','dueMonth','scheduledMonth','status','issues')
        for rid in people:
            if rid not in changed: continue
            for kind,label in [('recruitment','招聘奖金'),('referral','内推奖金')]:
                a=people[rid]['calculation'][kind];b=changed[rid]['calculation'][kind]
                common={'rowId':rid,'name':people[rid]['values'].get('姓名'),'kind':label}
                if rid in edited_ids and (b['total'] is None or b['currentAmount'] is None):
                    unresolved_amounts.append({**common,'reasons':list(dict.fromkeys(b['issues']+[i for n in b['nodes'] for i in n['issues']]))})
                if (a['total'],a['currency'])!=(b['total'],b['currency']):
                    impact.append({**common,'before':a['total'],'after':b['total'],'beforeCurrency':a['currency'],'currency':b['currency']})
                if (a['currentAmount'],a['currency'])!=(b['currentAmount'],b['currency']) and (
                        is_current(people[rid],kind,month,set()) or is_current(changed[rid],kind,month,set())):
                    current_impact.append({**common,'before':a['currentAmount'],'after':b['currentAmount'],'beforeCurrency':a['currency'],'currency':b['currency']})
                old={n['id']:n for n in a['nodes']};new={n['id']:n for n in b['nodes']}
                affected=False
                for nid in dict.fromkeys([*old,*new]):
                    x,y=old.get(nid),new.get(nid)
                    snapshot=lambda n: {f:(n.get('scheduledMonth',n.get('dueMonth')) if f=='scheduledMonth' else n.get(f)) for f in node_fields} if n else None
                    if snapshot(x)!=snapshot(y):
                        affected=True
                        node_impact.append({**common,'nodeId':nid,'label':(y or x)['label'],'role':(y or x)['role'],'before':snapshot(x),'after':snapshot(y)})
                if affected or a['issues']!=b['issues'] or rid in edited_ids:
                    reasons=list(dict.fromkeys(b['issues']+[i for n in b['nodes'] for i in n['issues']]))
                    if reasons:remaining.append({**common,'reasons':reasons})
        token=uuid4().hex
        preview={'id':token,'revision':revision,'month':month,'edits':edits,'evidence':reviewed,'origin':origin,'arrangements':arrangements,'decisions':decisions,'changes':changes,'impact':impact,'currentImpact':current_impact,'nodeImpact':node_impact,'remainingIssues':remaining,'people':len(set(e['rowId'] for e in edits+decisions)|{a['rowId'] for a in arrangements})}
        preview['unresolvedAmounts']=unresolved_amounts
        if batch_reviews:preview['approvalReviews']=batch_reviews
        store._write(folder/'previews'/f'{token}.json',preview)
        return preview


def apply_preview(store,month,revision,token):
    if not re.fullmatch(r'[0-9a-f]{32}',str(token)): raise ValueError('修改预览标识无效')
    with store._lock(month) as folder:
        state=store.read(month)
        if token in state.get('appliedPreviews',{}): return state
        store._check(state,revision)
        path=folder/'previews'/f'{token}.json'
        if not path.exists(): raise ValueError('修改预览不存在，请重新预览')
        preview=json.loads(path.read_text());store._check(state,preview['revision'])
        if preview['evidence']: evidence_check(store,state,preview['evidence'],True)
        for proof in preview.get('approvalReviews',[]):evidence_check(store,state,proof,True)
        signed = [item for proof in preview.get('approvalReviews', []) for item in proof.get('adoptedItems', []) if item.get('sourceSignature')]
        if signed:
            from .recruitment_approval_intake import proposals, resolve_item, conclusion_signature
            current = proposals(store, state)
            items = {item['id']: item for item in current['items']}
            documents = {doc['sourceId']: doc for doc in current['documents']}
            for item in signed:
                actual = items.get(item['id'])
                if not actual or actual['documentIssues'] or documents[actual['sourceId']]['sourceSignature'] != item['sourceSignature']:
                    raise ValueError('审批原件、附件或人员范围已变化，请重新核对')
                if item.get('outcome') == 'no_change':
                    resolved = resolve_item(deepcopy(actual), {row['id']: row for row in state['rows']}, {'rowId': item['rowId']} if item.get('rowId') else {})
                    if conclusion_signature(store, state, resolved, documents[actual['sourceId']]) != item['factSignature']:
                        raise ValueError('审批相关资料已变化，请重新核对')
        from . import recruitment_reference_facts as reference
        valid,stale=reference.effective_edits(folder,{**state,'edits':preview['edits']},proposed_edits=preview['edits'])
        if stale:raise ValueError('参考事实或对应任职已变化，请重新选择并预览')
        if any(e.get('employmentBasis') for e in valid):
            from .recruitment_employment import valid_employment_edit
            if any(e.get('employmentBasis') and not valid_employment_edit(state,e,proposed_edits=valid) for e in valid):
                raise ValueError('任职来源或关联已变化，请重新预览')
        state=store._prepare_change(folder,state,{'edits':valid},trusted_edits=True,validate_only=True)
        if preview['edits']: operation=state['operations'][-1]
        else:
            operation={'id':uuid4().hex,'month':month,'active':True,'at':recruitment_now().isoformat(timespec='seconds')}
            state.setdefault('operations',[]).append(operation);state['redoOperations']=[]
        operation['label']=preview['origin'];operation['previewId']=token
        if preview['evidence']:
            state.setdefault('reviewEntries',[]).append({**preview['evidence'],'id':token,'applied':True,'operationId':operation['id'],'rowIds':sorted({e['rowId'] for e in preview['edits']+preview.get('decisions',[])+preview['arrangements']})})
            for edit in state['edits']:
                if edit.get('operationId')==operation['id']: edit['evidencePages']=preview['evidence']['pageNumbers']
        for index, proof in enumerate(preview.get('approvalReviews',[])):
            state.setdefault('reviewEntries',[]).append({**proof,'id':token+'-'+str(index),'applied':True,'operationId':operation['id']})
            for edit in state['edits']:
                if edit.get('operationId')==operation['id'] and edit.get('evidenceId')==proof['sourceId']:
                    edit['evidencePages']=proof['pageNumbers']
        state.setdefault('awardDecisions',[]).extend({**d,'operationId':operation['id']} for d in preview.get('decisions',[]))
        state.setdefault('adjustmentDrafts',[]).extend({**item,'operationId':operation['id']} for item in preview['arrangements'])
        state.setdefault('appliedPreviews',{})[token]=operation['id']
        return store._save(folder,store._rebuild(folder,state))


def editable_workbook(state):
    wb=Workbook();ws=wb.active;ws.title='可编辑底稿'
    fields=GRID_FIELDS+[f for f in sorted(EDITABLE_FIELDS) if f not in GRID_FIELDS]
    headers=['任职标识','核算月份','底稿版本']+[FIELD_LABELS.get(f,f) for f in fields]
    ws.append(headers)
    for row in state['rows']:
        values=[row['id'],state['month'],state['revision']]+[row['values'].get(f,'') for f in fields]
        ws.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in values])
    ws.freeze_panes='F2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='1E3A8A')
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=22
    out=BytesIO();wb.save(out);return out.getvalue()


def workbook_edits(state,content):
    try: wb=load_workbook(BytesIO(content),data_only=False,read_only=True)
    except (BadZipFile, OSError, KeyError) as exc: raise ValueError('无法读取回传文件，请使用导出的完整表格') from exc
    try:
        if wb.sheetnames!=['可编辑底稿']: raise ValueError('请使用本工作台导出的可编辑底稿回传')
        rows=wb['可编辑底稿'].iter_rows();header=next(rows,());labels=[c.value for c in header]
        reverse={v:k for k,v in FIELD_LABELS.items()}
        fields=[reverse.get(v,v) for v in labels[3:]]
        if labels[:3]!=['任职标识','核算月份','底稿版本'] or len(set(fields))!=len(fields) or any(f not in EDITABLE_FIELDS for f in fields): raise ValueError('表头发生变化，请保留原有列名')
        people={r['id']:r for r in state['rows']};seen=set();edits=[]
        for cells in rows:
            if not any(c.value is not None for c in cells): continue
            if any(c.data_type=='f' for c in cells): raise ValueError('回传含公式，请核对后粘贴为值')
            values=[c.value for c in cells];rid=values[0]
            if rid not in people or rid in seen: raise ValueError('回传含未知或重复任职，请保留唯一任职标识')
            seen.add(rid)
            if values[1]!=state['month'] or values[2]!=state['revision']: raise ValueError('回传月份或版本已过期，请从最新底稿导出后修改')
            for field,value in zip(fields,values[3:]):
                value='' if value is None else value
                if isinstance(value,(datetime,date)): value=date_text(value)
                if isinstance(value,str) and value.startswith("'") and value[1:2] in ('=','+','-','@'): value=value[1:]
                before=people[rid]['values'].get(field,'')
                if text(value)==text(before): continue
                if field in DATE_FIELDS and value: value=date_text(value) or value
                edits.append({'rowId':rid,'field':field,'value':value,'reason':'回传表格核对更正'})
        return edits
    finally: wb.close()
