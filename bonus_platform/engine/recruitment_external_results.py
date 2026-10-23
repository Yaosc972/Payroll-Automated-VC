from __future__ import annotations

from collections import defaultdict
from copy import deepcopy
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import hashlib
import json
from pathlib import Path
import re
from uuid import uuid4

from openpyxl import Workbook, load_workbook

from .recruitment_time import now as recruitment_now
from .recruitment_draft_sources import date_text, text, PARSER_VERSION

VERSION=2
PROJECTS={'招聘负责人奖金':('recruitment','招聘负责人'),'招聘奖金':('recruitment','招聘负责人'),
          '协助招聘人奖金':('recruitment','协助招聘人'),'协助招聘奖金':('recruitment','协助招聘人'),
          '内推奖金':('referral','推荐人'),'推荐人奖金':('referral','推荐人')}
CURRENCIES={'CNY':'人民币','RMB':'人民币','人民币':'人民币','元':'人民币','USD':'美元','美元':'美元',
            'EUR':'欧元','欧元':'欧元','CZK':'捷克克朗','捷克克朗':'捷克克朗','GBP':'英镑','英镑':'英镑',
            'JPY':'日元','日元':'日元','THB':'泰铢','泰铢':'泰铢','VND':'越南盾','越南盾':'越南盾',
            'INR':'印度卢比','印度卢比':'印度卢比','MXN':'墨西哥比索','墨西哥比索':'墨西哥比索'}
IDENTIFIERS=('工号','体系外人员编号','证件号','海外offer流程编号')


def _hash(value):return hashlib.sha256(value.encode()).hexdigest()
def _normal(value):return text(value).casefold().removeprefix('海外-')
def _project(kind,role):return '内推奖金' if kind=='referral' else role+'奖金'
def _slot(record):return (record['month'],record['rowId'],record['kind'],record['role'])
def _now():return recruitment_now().isoformat(timespec='seconds')


def _amount(value):
    if value is None or text(value)=='':return None
    try:
        number=Decimal(str(value).replace(',',''))
        if number.is_finite():return float(number.quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
    except (ValueError,InvalidOperation):pass
    return None


def _month(value):
    raw=text(value).removesuffix('.0')
    if re.fullmatch(r'20\d{4}',raw) and 1<=int(raw[-2:])<=12:return int(raw)
    date=date_text(value)
    return int(date[:7].replace('-','')) if re.fullmatch(r'20\d\d-\d\d-\d\d',date) else None


def _outside(value):return bool(re.search(r'体系外|(?i:\b(?:海外-)?gl\d+\b)',text(value)))


def active_scopes(state):
    active={o['id'] for o in state.get('operations',[]) if o.get('active')}
    return [s for s in state.get('externalScopes',[]) if s.get('operationId') in active and s['month']<=state['month']]


def _scope_key(scope):
    return ('candidate',scope['rowId']) if scope['target']=='candidate' else ('recipient',_normal(scope['recipientId']),scope['kind'],scope['role'])


def _ledger_scopes(state):
    candidates=set();recipients=set()
    if not state:return candidates,recipients
    rows={r['id']:r for r in state['rows']}
    for scope in active_scopes(state):
        if scope['target']=='recipient':
            recipients.add((scope['kind'],scope['role'],_normal(scope['recipientId'])))
        else:
            row=rows.get(scope['rowId'])
            if not row:continue
            values=row['values'];date=date_text(values.get('入职日期') or values.get('候选人入职时间'))
            candidates.update((_normal(values.get(f)),date) for f in IDENTIFIERS if _normal(values.get(f)))
    return candidates,recipients


def _item(source,sheet,rows,values):
    return {'id':_hash(source['id']+'|'+sheet+'|'+','.join(map(str,rows))+'|'+values.get('project',''))[:24],
            'sourceId':source['id'],'sourceName':source['name'],'sheet':sheet,'sourceRows':rows,**values}


def parse(content,source,month,state=None):
    """Interpret either the five-column template or the known global ledger."""
    try:
        workbook=load_workbook(BytesIO(content),data_only=True,read_only=True)
        formulas=load_workbook(BytesIO(content),data_only=False,read_only=True)
    except Exception as exc:raise ValueError('无法读取Excel，请上传有效的xlsx或xlsm表格') from exc
    result=[];recognized=False;scoped_candidates,scoped_recipients=_ledger_scopes(state)
    try:
        for sheet in workbook:
            values=list(sheet.iter_rows(values_only=True))
            if not values:continue
            headers=None;start=None
            for i,line in enumerate(values[:20]):
                names={text(v):j for j,v in enumerate(line) if text(v)}
                if {'被招聘人','奖金项目','领取人','本月核定金额','币种'}<=names.keys():headers=names;start=i+1;break
            if headers is not None:
                recognized=True
                for index,line in enumerate(values[start:],start+1):
                    if not any(v not in (None,'') for v in line):continue
                    def cell(name):
                        col=headers.get(name);return line[col] if col is not None and col<len(line) else None
                    project=text(cell('奖金项目'));kind,role=PROJECTS.get(project,('',''))
                    value=cell('本月核定金额');amount=_amount(value)
                    status=text(cell('发放状态'))
                    statuses={'':'pay','发放':'pay','正常':'pay','正常发放':'pay','应发':'pay','已核定':'pay',
                              '不发':'exclude','不发放':'exclude','取消':'exclude','停发':'exclude','不予发放':'exclude',
                              '暂缓':'hold','暂缓发放':'hold','延后':'hold','待发':'hold','待发放':'hold','延期发放':'hold'}
                    outcome=statuses.get(status,'pay')
                    if amount==0 and not status:outcome='exclude'
                    if outcome=='exclude' and amount is None and status:amount=0.0
                    component={'补发':'supplement','扣回':'deduction','常规':'regular','月度核定':'regular','':'regular'}.get(text(cell('金额类型')),'')
                    if amount is not None and amount<0 and not text(cell('金额类型')):component='deduction'
                    issues=[]
                    if status not in statuses:issues.append('发放状态无法识别，请使用发放、不发或暂缓')
                    for label,value_ in [('被招聘人',text(cell('被招聘人'))),('奖金项目',kind),('领取人',text(cell('领取人')) or outcome in {'exclude','hold'}),('币种',CURRENCIES.get(text(cell('币种')).upper()))]:
                        if not value_:issues.append('请补充'+label if not text(cell(label)) else label+'无法识别')
                    if amount is None:
                        col=headers['本月核定金额']+1
                        formula=formulas[sheet.title].cell(index,col).data_type=='f'
                        issues.append('金额公式没有数值结果，请提供已核定数值' if formula else '请补充本月核定金额（不发请填0）')
                    if not component:issues.append('金额类型请使用常规、补发或扣回')
                    if component=='deduction' and amount is not None and amount>0:issues.append('扣回金额请填负数')
                    if outcome=='exclude' and amount not in (None,0):issues.append('不发放的金额应为0')
                    provided=_month(cell('核算月份'))
                    if text(cell('核算月份')) and provided!=month:issues.append('核算月份与当前活动不一致')
                    result.append(_item(source,sheet.title,[index],{'candidateName':text(cell('被招聘人')),'candidateId':text(cell('人员编号') or cell('被招聘人工号')),
                        'hireDate':date_text(cell('入职日期')),'kind':kind,'role':role,'project':_project(kind,role) if kind else project,
                        'recipientName':text(cell('领取人')),'recipientId':text(cell('领取人编号') or cell('领取人工号')),
                        'amount':amount,'currency':CURRENCIES.get(text(cell('币种')).upper(),text(cell('币种'))),
                        'outcome':outcome,'component':component,'parseIssues':issues,'format':'template'}))
            elif sheet.title=='附件一-招聘奖金核算明细':
                recognized=True
                from .recruitment_opening import COLUMNS, period
                for index,line in enumerate(values,1):
                    def c(col):return line[col-1] if col<=len(line) else None
                    if not text(c(1)) or not text(c(2)) or str(c(1)).strip() in {'序号','1.0'} and not text(c(3)):continue
                    name=text(c(2));number=text(c(3));hire=date_text(c(5))
                    if not name or name in {'姓名','工号'}:continue
                    for kind,roles in COLUMNS.items():
                        for role,columns in roles.items():
                            idcol,namecol={'招聘负责人':(22,23),'协助招聘人':(24,25),'推荐人':(59,58)}[role]
                            recipient_id=text(c(idcol));recipient_name=text(c(namecol))
                            declared=(_normal(number),hire) in scoped_candidates or (_normal(number),'') in scoped_candidates or (kind,role,_normal(recipient_id)) in scoped_recipients
                            if not (_outside(number) or _outside(recipient_id) or declared):continue
                            selected=[(a,p) for a,p,_ in columns if period(c(p))==month]
                            if not selected:continue
                            amounts=[_amount(c(a)) for a,p in selected]
                            issues=[]
                            if any(a is None for a in amounts):issues.append('本月分期金额缺少数值，请核对原表')
                            if not recipient_name:issues.append('请补充领取人')
                            currency=CURRENCIES.get(text(c(56)).upper(),text(c(56)))
                            if not currency:issues.append('请补充币种')
                            total=None if any(a is None for a in amounts) else round(sum(amounts),2)
                            result.append(_item(source,sheet.title,[index],{'candidateName':name,'candidateId':number,'hireDate':hire,
                                'kind':kind,'role':role,'project':_project(kind,role),'recipientName':recipient_name,'recipientId':recipient_id,
                                'amount':total,'currency':currency,'outcome':'exclude' if total==0 else 'pay',
                                'component':'deduction' if total is not None and total<0 else 'regular','parseIssues':issues,'format':'ledger',
                                'scopeExpanded':not (_outside(number) or _outside(recipient_id))}))
    finally:
        workbook.close();formulas.close()
    if not recognized:raise ValueError('未找到体系外核算结果表。请使用模板，或上传含“附件一-招聘奖金核算明细”的原表')
    if not result:raise ValueError('未找到本月体系外核算结果。原台账仅纳入明确GL、体系外编号或已划为体系外的人员及奖金项目')
    return result


def template(month):
    wb=Workbook();sheet=wb.active;sheet.title='体系外核算结果'
    sheet.append(['被招聘人','奖金项目','领取人','本月核定金额','币种','人员编号','入职日期','领取人编号','核算月份','发放状态','金额类型'])
    sheet.freeze_panes='A2';sheet.auto_filter.ref='A1:K1'
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='2563EB')
    for column,width in [('A',24),('B',24),('C',24),('D',20),('E',12),('F',20),('G',18),('H',20),('I',16),('J',16),('K',14)]:sheet.column_dimensions[column].width=width
    output=BytesIO();wb.save(output);return output.getvalue()


def _scope_projection(state):
    from .recruitment_awards import calculate_base, country_currency, present
    scopes=active_scopes(state);expected=[];calculated={}
    for row in state['rows']:
        row.pop('externalCandidateScope',None);row.pop('externalScopeOwnership',None)
        if row.get('externalOnly'):continue
        values=row['values'];owned={}
        for scope in scopes:
            if scope['target']=='candidate':
                if scope['rowId']!=row['id']:continue
                row['externalCandidateScope']=True
                roles=[('recruitment','招聘负责人')]
                if any(present(values.get('协助招聘人'+f)) for f in ('工号','姓名')):roles.append(('recruitment','协助招聘人'))
                if row['id'] not in calculated:calculated[row['id']]=calculate_base(state,row)
                referral_unknown=not present(values.get('招聘渠道')) or bool({'招聘渠道','推荐人工号','推荐人姓名'}&set(row.get('unresolvedFields',[]))) or calculated[row['id']]['referral']['total'] is None
                if any(present(values.get('推荐人'+f)) for f in ('工号','姓名')) or values.get('招聘渠道')=='内推' or referral_unknown:roles.append(('referral','推荐人'))
            else:
                if _normal(values.get(scope['role']+'工号'))!=_normal(scope['recipientId']):continue
                roles=[(scope['kind'],scope['role'])]
            for kind,role in roles:
                owned[(kind,role)]={'kind':kind,'role':role,'scopeId':scope['id']}
                if scope['target']=='recipient':
                    if row['id'] not in calculated:calculated[row['id']]=calculate_base(state,row)
                    award=calculated[row['id']][kind];nodes=[n for n in award['nodes'] if n['role']==role]
                    due=any(n.get('scheduledMonth',n.get('dueMonth'))==state['month'] and n.get('amount')!=0 and n['status']!='按规则不发' for n in nodes)
                    if not due and award['total'] is None and not any(n.get('dueMonth')==state['month'] or n.get('scheduledMonth')==state['month'] for n in nodes):
                        from .recruitment_attention import is_current
                        candidate=deepcopy(row);candidate['calculation']=calculated[row['id']]
                        due=is_current(candidate,kind,state['month'],set())
                    if not due:continue
                currency=country_currency(values)[0]
                number=text(values.get(role+'工号'));name=text(values.get(role+'姓名'))
                component={'component':'regular','recipientId':number,'recipientName':name,'amount':None,'currency':currency,
                           'outcome':'pending','sourceId':'','sheet':'','sourceRows':[],
                           'recipientKey':'number:'+_normal(number) if number else 'scope-payee-'+scope['id']+'-'+row['id']+'-'+role}
                expected.append({'id':'scope-'+scope['id']+'-'+row['id']+'-'+kind+'-'+role,'month':state['month'],
                    'expectedMonth':state['month'],'rowId':row['id'],'candidateName':text(values.get('姓名')),
                    'candidateId':text(values.get('工号') or values.get('体系外人员编号')),'kind':kind,'role':role,
                    'scopeId':scope['id'],'sourceId':'','components':[component]})
        if owned:row['externalScopeOwnership']=list(owned.values())
    state['externalScopeExpected']=expected


def effective(state):
    """Select the latest active version first; invalid proof never revives an older one."""
    active={o['id'] for o in state.get('operations',[]) if o.get('active')}
    sources={s['id']:s for s in state.get('sources',[])};latest={}
    for record in state.get('externalResults',[]):
        if record['operationId'] in active:latest[_slot(record)]=record
    current={key:r for key,r in latest.items() if key[0]==state['month']}
    prior={}
    for key,record in latest.items():
        if key[0]>=state['month'] or not any(c['component']=='regular' for c in record['components']):continue
        anchor=key[1:]
        if anchor not in prior or prior[anchor]['month']<record['month']:prior[anchor]=record
    scopes={s['id'] for s in active_scopes(state)}
    declared={(r['rowId'],r['kind'],r['role']):r for r in state.get('externalScopeExpected',[])
              if r.get('scopeId') in scopes and r.get('expectedMonth')==state['month']}
    scoped_roles={(row['id'],r['kind'],r['role']) for row in state['rows'] for r in row.get('externalScopeOwnership',[])}
    candidates={r['id'] for r in state['rows'] if r.get('externalCandidateScope')}
    expected={key:dict(r,expectedMonth=state['month']) for key,r in prior.items()
              if (state['month'],*key) not in current and (key not in scoped_roles or key in declared or key[0] in candidates)}
    for key,record in declared.items():
        if (state['month'],*key) not in current:expected.setdefault(key,record)
    valid=[];invalid=[]
    for record in current.values():
        proofs=record.get('proofs') or [{'sourceId':record['sourceId'],'sourceDigest':record['sourceDigest']}]
        verified=all((source:=sources.get(proof['sourceId'],{})) and source.get('active',True) and not source.get('error')
                     and source.get('contentDigest',source.get('id'))==proof['sourceDigest'] for proof in proofs)
        if verified:
            valid.append(record)
        else:invalid.append(record)
    return {'valid':valid,'invalid':invalid,'expected':list(expected.values()),'latest':list(current.values())}


def _subject_for(state,item,resolutions):
    chosen=resolutions.get(item['id']);rows=[r for r in state['rows'] if not r.get('externalOnly')]
    subjects=state.get('externalSubjects',[])
    def choice(row):
        values=row['values'];return {'rowId':row['id'],'name':text(values.get('姓名')),'number':text(values.get('工号') or values.get('体系外人员编号')),
                                    'hireDate':date_text(values.get('入职日期') or values.get('候选人入职时间'))}
    if chosen:
        _,offered,issues=_subject_for(state,item,{})
        if not offered or chosen not in {c['rowId'] for c in offered}:
            return None,offered,['请选择该人员对应的任职记录']
        matches=[r for r in rows if r['id']==chosen]
        known=next((v for v in subjects if v['rowId']==chosen),None)
        if not matches:
            return (deepcopy(known),offered,[]) if known else (None,offered,['所选人员已变化，请重新选择'])
        return {'id':'subject-'+matches[0]['id'],'rowId':matches[0]['id'],'candidateName':text(matches[0]['values'].get('姓名')),
                'candidateId':item['candidateId'],'hireDate':item['hireDate'],'externalOnly':False},offered,[]
    number=_normal(item.get('candidateId'));name=_normal(item['candidateName']);hire=item['hireDate']
    matches=[r for r in rows if (number and number in {_normal(r['values'].get(f)) for f in IDENTIFIERS}) or
             (not number and _normal(r['values'].get('姓名'))==name)]
    if hire:matches=[r for r in matches if date_text(r['values'].get('入职日期') or r['values'].get('候选人入职时间'))==hire]
    if number and not matches:
        named=[r for r in rows if _normal(r['values'].get('姓名'))==name and (not hire or not date_text(r['values'].get('入职日期')) or date_text(r['values'].get('入职日期'))==hire)]
        if named:
            if any(any(text(r['values'].get(f)) for f in IDENTIFIERS) for r in named):
                return None,[],['人员编号未匹配，但相同姓名已有任职记录，请核对人员编号后重新上传']
            matches=named
    if number and matches and any(_normal(r['values'].get('姓名'))!=name for r in matches):
        return None,[],['被招聘人姓名与人员编号对应的人员不一致，请更正表格']
    if len(matches)>1:return None,[choice(r) for r in matches],['同一人员有多条任职记录，请选择本次任职']
    if len(matches)==1:
        row=matches[0]
        return {'id':'subject-'+row['id'],'rowId':row['id'],'candidateName':text(row['values'].get('姓名')),
                'candidateId':item['candidateId'],'hireDate':item['hireDate'],'externalOnly':False},[choice(row)],[]
    known=[s for s in subjects if (number and _normal(s.get('candidateId'))==number or not number and _normal(s['candidateName'])==name)
           and (not hire or not s.get('hireDate') or s.get('hireDate')==hire)]
    if not known and number:
        named=[s for s in subjects if _normal(s['candidateName'])==name and (not hire or not s.get('hireDate') or s.get('hireDate')==hire)]
        if any(s.get('candidateId') for s in named):return None,[],['相同姓名已有不同体系外编号，请核对人员编号后重新上传']
        known=named
    if len(known)>1:return None,[{'rowId':s['rowId'],'name':s['candidateName'],'number':s.get('candidateId',''),'hireDate':s.get('hireDate','')} for s in known],['人员存在多个体系外身份，请提供人员编号或入职日期']
    if known and _normal(known[0]['candidateName'])!=name:return None,[],['被招聘人姓名与体系外人员编号不一致，请更正表格']
    if known:
        subject=deepcopy(known[0])
        if not subject.get('candidateId') and item['candidateId']:subject['candidateId']=item['candidateId']
        if not subject.get('hireDate') and hire:subject['hireDate']=hire
        return subject,[],[]
    if not number:
        competing=[r for r in rows if _normal(r['values'].get('姓名'))==name]
        if competing:return None,[choice(r) for r in competing],['姓名对应任职与入职日期不一致，请选择人员或更正表格']
    subject_id=item.get('newSubjectId') or uuid4().hex
    return {'id':subject_id,'rowId':'external-'+subject_id,'candidateName':item['candidateName'],
            'candidateId':item['candidateId'],'hireDate':hire,'externalOnly':True},[],[]


def _recipient_identity(state,item):
    identities=defaultdict(set)
    for row in state['rows']:
        values=row['values']
        for prefix in ('','招聘负责人','协助招聘人','推荐人'):
            number=_normal(values.get(prefix+'工号'));name=text(values.get(prefix+'姓名'))
            if number and name:identities[number].add(name)
    for record in state.get('externalResults',[]):
        for component in record['components']:
            number=_normal(component.get('recipientId'));name=component.get('recipientName','')
            if number and name:identities[number].add(name)
    number=_normal(item['recipientId']);name=_normal(item['recipientName'])
    if number and number in identities:
        names=identities[number]
        if len(names)!=1 or name and name not in {_normal(n) for n in names}:
            return ['领取人姓名与编号对应的人员不一致，请核对领取人']
        if not name:item['recipientName']=next(iter(names))
    elif not number and name:
        known=[number for number,names in identities.items() if len(names)==1 and name==_normal(next(iter(names)))]
        if len(known)==1:item['recipientId']=known[0]
    return []


def _special_conflict(state,row_id,kind,role):
    from .recruitment_special import active
    return [r for r in active(state,'specialAdjustments') if r.get('rowId')==row_id and r.get('kind')==kind
            and r.get('role')==role and r.get('accountingMonth',r.get('month'))==state['month']]


def _signature(components):
    fields=('component','recipientId','recipientName','amount','currency','outcome')
    return sorted((tuple(c.get(k) for k in fields) for c in components),key=lambda v:json.dumps(v,ensure_ascii=False))


def _batch_identities(items):
    normalized=deepcopy(items);issues=defaultdict(list)
    by_name=defaultdict(list)
    for item in normalized:by_name[_normal(item['candidateName'])].append(item)
    for name,people in by_name.items():
        if not name:continue
        for item in people:
            if item['candidateId']:continue
            known={_normal(other['candidateId']):other['candidateId'] for other in people if other['candidateId'] and
                   (not item['hireDate'] or not other['hireDate'] or item['hireDate']==other['hireDate'])}
            if len(known)==1:item['candidateId']=next(iter(known.values()))
            elif len(known)>1:issues[item['id']].append('同名人员有多个编号，请补充本行人员编号')
        for item in people:
            if item['hireDate'] or issues[item['id']]:continue
            dates={other['hireDate'] for other in people if other['hireDate'] and not issues[other['id']] and
                   _normal(other['candidateId'])==_normal(item['candidateId'])}
            if len(dates)==1:item['hireDate']=next(iter(dates))
            elif len(dates)>1:issues[item['id']].append('同一人员有多段任职，请补充本行入职日期')
    by_number=defaultdict(set)
    for item in normalized:
        if item['candidateId']:by_number[_normal(item['candidateId'])].add(_normal(item['candidateName']))
    for item in normalized:
        if item['candidateId'] and len(by_number[_normal(item['candidateId'])])>1:
            issues[item['id']].append('同一人员编号对应不同姓名，请核对后重新上传')
    return {item['id']:(item,issues[item['id']]) for item in normalized}


def _preview(state,batch):
    rows=[];groups=defaultdict(list);resolutions=batch.get('resolutions',{});subjects={}
    selection=effective(state)
    selected={_slot(r):r for r in selection['latest']}
    valid_ids={r['id'] for r in selection['valid']}
    adopted_ids={item_id for r in selection['latest'] for item_id in r.get('itemIds',[])}
    identities=_batch_identities(batch['items'])
    for original in batch['items']:
        item,identity_issues=identities[original['id']];issues=list(item['parseIssues'])+identity_issues
        issues+=_recipient_identity(state,item)
        source=next((s for s in state['sources'] if s['id']==item['sourceId']),None)
        if source and (not source.get('active',True) or source.get('error')):issues.append('来源已停用或原件已变化，请重新上传该核算结果表')
        subject,candidates,matching=(None,[],[]) if identity_issues else _subject_for(state,item,resolutions)
        issues+=matching
        if subject:
            key=(_normal(item.get('candidateId')) or _normal(item['candidateName']),item.get('hireDate',''))
            if subject.get('externalOnly') and key in subjects:subject=subjects[key]
            subjects[key]=subject
            original.setdefault('newSubjectId',subject['id'])
            item.update(rowId=subject['rowId'],subject=subject)
            if item.get('scopeExpanded'):
                row=next((r for r in state['rows'] if r['id']==subject['rowId']),None)
                current_id=_normal(row['values'].get(item['role']+'工号')) if row else ''
                if current_id and current_id!=_normal(item['recipientId']):
                    issues.append('领取人与当前安排不同，需先核对原领取人的已安排金额及补发、扣回')
            if _special_conflict(state,subject['rowId'],item['kind'],item['role']):issues.append('同月同角色已有补发或扣回安排，请先在调整记录中核对并撤销重复安排')
        else:item['rowId']=''
        if item['id'] in batch.get('adoptedItemIds',[]) and item['id'] in adopted_ids:status='adopted'
        else:status='blocked' if issues else 'ready'
        item.update(status=status,issues=issues,candidates=candidates,change='new')
        rows.append(item)
        if subject and status!='adopted':groups[(state['month'],subject['rowId'],item['kind'],item['role'])].append(item)
    for key,items in groups.items():
        duplicates=defaultdict(list)
        for item in items:duplicates[item['component']].append(item)
        if any(len(group)>1 for group in duplicates.values()):
            for item in items:item['issues'].append('同一人员及奖金项目存在重复金额行，请合并后重新上传');item['status']='blocked'
        if any(item['status']=='blocked' for item in items):
            for item in items:
                if item['status']=='ready':item['status']='blocked';item['issues'].append('同一奖金项目另有异常行，请一并更正后采用')
        previous=selected.get(key)
        if previous:
            replaced=previous['revision']>batch.get('createdRevision',state['revision']) and previous['id'] in valid_ids and bool({i['sourceId'] for i in items}-{p['sourceId'] for p in previous.get('proofs',[{'sourceId':previous['sourceId']}])})
            if replaced:
                for item in items:item.update(status='superseded',change='replacement')
                continue
            unchanged=previous['id'] in valid_ids and not any(i['issues'] for i in items) and _signature(items)==_signature(previous['components'])
            for item in items:
                item['change']='unchanged' if unchanged else 'replacement'
                if unchanged and item['status']=='ready':item['status']='unchanged'
    ready=[r for r in rows if r['status']=='ready'];totals=defaultdict(Decimal)
    for item in ready:
        if item['outcome']!='hold':totals[item['currency']]+=Decimal(str(item['amount']))
    return {'token':batch['token'],'revision':state['revision'],'month':state['month'],'files':batch['files'],
            'rows':rows,'readyCount':len(ready),'blockedCount':sum(r['status']=='blocked' for r in rows),
            'unchangedCount':sum(r['status']=='unchanged' for r in rows),'appliedCount':sum(r['status']=='adopted' for r in rows),
            'peopleCount':len({r['rowId'] or (r['candidateId'],r['candidateName'],r['hireDate']) for r in rows}),'totalsByCurrency':{c:float(v) for c,v in totals.items()},'canApply':bool(ready)}


def _batch_path(store,month,token):
    if not isinstance(token,str) or not re.fullmatch(r'[a-f0-9]{32}',token):raise ValueError('导入预览不存在，请重新上传')
    return store._folder(month)/'external-previews'/f'{token}.json'


def apply_scope(store,month,revision,scopes):
    from .recruitment_trials import _require_idle
    _require_idle(month,store)
    if not isinstance(scopes,list) or not scopes or len(scopes)>1000:raise ValueError('请选择要按体系外核算的人员或领取人')
    if any(not isinstance(s,dict) for s in scopes):raise ValueError('体系外人员范围格式无效')
    request=_hash(json.dumps(sorted(scopes,key=lambda s:json.dumps(s,sort_keys=True,ensure_ascii=False)),sort_keys=True,ensure_ascii=False))
    with store._lock(month) as folder:
        state=store.read(month)
        if any(o.get('externalScopeRequest')==request and o.get('requestRevision')==revision for o in state.get('operations',[])):return state
        store._check(state,revision)
        if (folder/'accepted.json').exists():raise ValueError('本月核算已确认，请新建核算记录后调整体系外范围')
        rows={r['id']:r for r in state['rows'] if not r.get('externalOnly')};normalized=[]
        for scope in scopes:
            target=scope.get('target')
            if target=='candidate':
                if set(scope)-{'target','rowId'}:raise ValueError('被招聘人体系外范围仅接受任职记录编号')
                row=rows.get(scope.get('rowId'))
                if not row:raise ValueError('被招聘人的任职记录不存在，请刷新后重试')
                values=row['values']
                normalized.append({'target':target,'rowId':row['id'],'candidateName':text(values.get('姓名')),
                    'candidateId':text(values.get('工号') or values.get('体系外人员编号')),
                    'hireDate':date_text(values.get('入职日期') or values.get('候选人入职时间'))})
            elif target=='recipient':
                if set(scope)-{'target','recipientId','recipientName','kind','role'}:raise ValueError('领取人体系外范围字段无效')
                number=_normal(scope.get('recipientId'));kind=scope.get('kind');role=scope.get('role')
                if not number or number in {'-','无','无工号'} or len(number)>100:raise ValueError('请提供领取人编号，不能仅按姓名划定体系外范围')
                if (kind,role) not in set(PROJECTS.values()):raise ValueError('体系外奖金项目或领取人角色无效')
                matches=[r for r in rows.values() if _normal(r['values'].get(role+'工号'))==number]
                if not matches:raise ValueError('未找到该编号对应的奖金领取人，请核对编号和奖金项目')
                names={text(r['values'].get(role+'姓名')) for r in matches if text(r['values'].get(role+'姓名'))}
                if len({_normal(n) for n in names})>1:raise ValueError('该领取人编号对应多个姓名，请先核对人员身份')
                name=text(scope.get('recipientName'))
                if name and names and _normal(name) not in {_normal(n) for n in names}:raise ValueError('领取人姓名与编号不一致，请核对后重试')
                normalized.append({'target':target,'recipientId':number,'recipientName':name or next(iter(names),''),'kind':kind,'role':role})
            else:raise ValueError('请选择被招聘人或奖金领取人的体系外范围')
        existing={_scope_key(s) for s in active_scopes(state)};additions={_scope_key(s):s for s in normalized if _scope_key(s) not in existing}
        if not additions:return state
        operation_id=uuid4().hex
        for scope in additions.values():state.setdefault('externalScopes',[]).append({**scope,'id':uuid4().hex,'month':month,'operationId':operation_id})
        state.setdefault('operations',[]).append({'id':operation_id,'month':month,'active':True,'label':'调整体系外核算范围','at':_now(),
                                                'externalScopeRequest':request,'requestRevision':revision})
        state['redoOperations']=[]
        return store._save(folder,store._rebuild(folder,state))


def prepare_import(store,month,revision,files):
    from .recruitment_trials import _require_idle
    _require_idle(month,store)
    with store._lock(month) as folder:
        state=store.read(month);store._check(state,revision)
        if (folder/'accepted.json').exists():raise ValueError('本月核算已确认，请新建核算记录后导入更正资料')
        payload=[]
        if not files:raise ValueError('请选择体系外核算结果表')
        for name,content in files:
            name=Path(name).name
            if Path(name).suffix.lower() not in {'.xlsx','.xlsm'}:raise ValueError('请上传xlsx或xlsm核算结果表')
            if len(content)>80*1024*1024:raise ValueError('单份文件不能超过80MB')
            digest=hashlib.sha256(content).hexdigest();sid=_hash('external-results|'+str(month)+'|'+digest)
            if any((s['id']==digest or s.get('contentDigest')==digest) and s.get('purpose','')!='external-results' for s in state['sources']):
                raise ValueError('相同文件已按其他用途导入，请使用对应入口核对，不能重复改变资料用途')
            source=next((s for s in state['sources'] if s['id']==sid),None)
            source=deepcopy(source) if source else {'id':sid,'name':name,'storedName':sid+Path(name).suffix.lower(),'bytes':len(content),
                'addedAt':_now(),'active':True,'forMonth':month,'purpose':'external-results','roles':['external_results'],
                'roleCounts':{},'recordCount':0,'attachment':False,'sheets':[],'parserVersion':PARSER_VERSION,
                'contentDigest':digest,'parsedContentDigest':digest}
            parsed=parse(content,source,month,state)
            source.update(recordCount=len(parsed),roleCounts={'external_results':len(parsed)},sheets=list(dict.fromkeys(r['sheet'] for r in parsed)))
            payload.append((source,content,parsed))
        restored=False
        for source,content,_ in payload:
            original=next((s for s in state['sources'] if s['id']==source['id']),None)
            if not original:continue
            path=folder/'sources'/original['storedName'];incoming=hashlib.sha256(content).hexdigest()
            if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest()!=incoming or not original.get('active',True) or original.get('error'):
                path.parent.mkdir(parents=True,exist_ok=True)
                temporary=path.with_name(path.name+'.'+uuid4().hex+'.tmp');temporary.write_bytes(content);temporary.replace(path)
                original.update(active=True,contentDigest=incoming)
                original.pop('error',None);original.pop('sourceCacheError',None);restored=True
        digest_key=_hash('|'.join(sorted(s['id'] for s,_,_ in payload)))
        existing=next((b for b in state.get('externalImports',[]) if b['digestKey']==digest_key),None)
        if existing:
            batch=json.loads(_batch_path(store,month,existing['token']).read_text())
            original_items={i['id']:i for i in batch['items']};items=[i for _,_,parsed in payload for i in parsed]
            for item in items:
                if original_items.get(item['id'],{}).get('newSubjectId'):item['newSubjectId']=original_items[item['id']]['newSubjectId']
            reparsed=items!=batch['items'];batch['items']=items
            if reparsed:
                for source,_,_ in payload:
                    original=next(s for s in state['sources'] if s['id']==source['id'])
                    original.update(recordCount=source['recordCount'],roleCounts=source['roleCounts'],sheets=source['sheets'])
            previous_pending=deepcopy(existing)
            _pending(state,batch,_preview(state,batch))
            if restored or reparsed or existing!=previous_pending:state=store._save(folder,store._rebuild(folder,state))
            preview=_preview(state,{**batch,'adoptedItemIds':[]})
            store._write(_batch_path(store,month,batch['token']),batch)
            return preview
        batch={'token':uuid4().hex,'createdRevision':state['revision'],'digestKey':digest_key,'files':[{'id':s['id'],'name':s['name']} for s,_,_ in payload],
               'items':[item for _,_,items in payload for item in items],'resolutions':{},'adoptedItemIds':[]}
        for source,content,_ in payload:
            if not any(s['id']==source['id'] for s in state['sources']):
                path=folder/'sources'/source['storedName'];path.parent.mkdir(parents=True,exist_ok=True)
                temporary=path.with_suffix(path.suffix+'.'+uuid4().hex+'.tmp');temporary.write_bytes(content);temporary.replace(path)
                state['sources'].append(source)
                store._write(folder/'parsed'/f"{source['id']}.json",{'records':[],'attachment':False,'sheets':source['sheets'],'contentDigest':source['contentDigest'],'parserVersion':PARSER_VERSION})
        state.setdefault('externalImports',[]).append({'token':batch['token'],'digestKey':digest_key,'name':'、'.join(s['name'] for s,_,_ in payload)})
        preview=_preview(state,batch)
        _pending(state,batch,preview)
        store._write(_batch_path(store,month,batch['token']),batch)
        state=store._save(folder,store._rebuild(folder,state))
        preview=_preview(state,batch);store._write(_batch_path(store,month,batch['token']),batch)
        return preview


def repreview(store,month,revision,token,resolutions):
    from .recruitment_trials import _require_idle
    _require_idle(month,store)
    with store._lock(month):
        state=store.read(month);store._check(state,revision)
        path=_batch_path(store,month,token)
        if not path.is_file():raise ValueError('导入预览不存在，请重新上传')
        batch=json.loads(path.read_text());ids={i['id'] for i in batch['items']}
        if not isinstance(resolutions,list):raise ValueError('人员选择格式无效')
        for resolution in resolutions:
            if not isinstance(resolution,dict) or resolution.get('itemId') not in ids or not text(resolution.get('rowId')):raise ValueError('人员选择无效')
            batch['resolutions'][resolution['itemId']]=resolution['rowId']
        result=_preview(state,batch);store._write(path,batch);return result


def apply_import(store,month,revision,token):
    from .recruitment_trials import _require_idle
    _require_idle(month,store)
    with store._lock(month) as folder:
        state=store.read(month);path=_batch_path(store,month,token)
        if not path.is_file():raise ValueError('导入预览不存在，请重新上传')
        batch=json.loads(path.read_text())
        completed=next((o for o in state.get('operations',[]) if o.get('externalPreview')==token and o.get('requestRevision')==revision),None)
        if completed:return state,0
        store._check(state,revision)
        if (folder/'accepted.json').exists():raise ValueError('本月核算已确认，不能覆盖已确认的体系外结果')
        preview=_preview(state,batch);ready=[r for r in preview['rows'] if r['status']=='ready']
        if not ready:return state,0
        operation_id=uuid4().hex;grouped=defaultdict(list)
        for item in ready:
            source=next(s for s in state['sources'] if s['id']==item['sourceId'])
            if not source.get('active',True) or source.get('error') or source.get('contentDigest')!=hashlib.sha256((folder/'sources'/source['storedName']).read_bytes()).hexdigest():raise ValueError('原件已变化，请重新上传核算结果')
            grouped[(item['rowId'],item['kind'],item['role'])].append(item)
        active=effective(state);previous={_slot(r):r for r in active['latest']}
        for (row_id,kind,role),items in grouped.items():
            subject=items[0]['subject']
            existing_subject=next((s for s in state.get('externalSubjects',[]) if s['id']==subject['id']),None)
            if existing_subject is None:state.setdefault('externalSubjects',[]).append(subject)
            else:
                for field in ('candidateId','hireDate'):
                    if not existing_subject.get(field) and subject.get(field):existing_subject[field]=subject[field]
            source=next(s for s in state['sources'] if s['id']==items[0]['sourceId'])
            components=[];recipient_keys={}
            for item in items:
                old=previous.get((month,row_id,kind,role),{})
                same=next((c for c in old.get('components',[]) if _normal(c.get('recipientId'))==_normal(item['recipientId']) and c.get('recipientName')==item['recipientName']),{})
                payee=(_normal(item['recipientId']),_normal(item['recipientName']))
                recipient_key=recipient_keys.get(payee) or (('number:'+payee[0]) if item['recipientId'] else same.get('recipientKey') or 'external-payee-'+uuid4().hex)
                recipient_keys[payee]=recipient_key
                components.append({k:item[k] for k in ('component','recipientId','recipientName','amount','currency','outcome','sourceId','sheet','sourceRows')}|{'recipientKey':recipient_key})
            old=previous.get((month,row_id,kind,role))
            record={'id':uuid4().hex,'month':month,'rowId':row_id,'subjectId':subject['id'],'candidateName':subject['candidateName'],
                    'kind':kind,'role':role,'itemIds':[i['id'] for i in items],'components':components,'sourceId':source['id'],'sourceDigest':source['contentDigest'],
                    'proofs':[{'sourceId':sid,'sourceDigest':next(s['contentDigest'] for s in state['sources'] if s['id']==sid)} for sid in dict.fromkeys(i['sourceId'] for i in items)],
                    'parserVersion':VERSION,'operationId':operation_id,'revision':state['revision']+1,'supersedes':old['id'] if old else None}
            state.setdefault('externalResults',[]).append(record)
        state.setdefault('operations',[]).append({'id':operation_id,'month':month,'active':True,'label':'采用体系外核算结果','at':_now(),
                                                'externalPreview':token,'requestRevision':revision})
        state['redoOperations']=[]
        batch['adoptedItemIds']=list(dict.fromkeys(batch['adoptedItemIds']+[r['id'] for r in ready]))
        _pending(state,batch,_preview(state,batch))
        for entry in state.get('externalImports',[]):
            old_path=_batch_path(store,month,entry['token'])
            if entry['token']!=token and old_path.is_file():
                old_batch=json.loads(old_path.read_text());_pending(state,old_batch,_preview(state,old_batch))
        state=store._save(folder,store._rebuild(folder,state))
        store._write(path,batch)
        return state,len(ready)


def _pending(state,batch,preview):
    entry=next(e for e in state['externalImports'] if e['token']==batch['token'])
    entry.update({k:preview[k] for k in ('readyCount','blockedCount','unchangedCount')})
    entry['sourceIds']=[s['id'] for s in batch['files']]
    slots={(r['rowId'],r['kind'],r['role']) for r in preview['rows'] if r['rowId'] and r['status'] in {'ready','blocked','unchanged','adopted'}}
    entry['coveredSlots']=[{'rowId':row_id,'kind':kind,'role':role} for row_id,kind,role in sorted(slots)]


def _covered_slots(entry):
    return {(r['rowId'],r['kind'],r['role']) for r in entry.get('coveredSlots',[])}


def context(store,state):
    selected=effective(state);records=[];totals=defaultdict(Decimal);people=set()
    for record in selected['latest']:
        status='adopted' if record in selected['valid'] else 'invalid'
        source=next((s for s in state['sources'] if s['id']==record['sourceId']),{})
        for component in record['components']:
            records.append({**component,'id':record['id']+'-'+component['component'],'rowId':record['rowId'],
                'candidateName':record['candidateName'],'project':_project(record['kind'],record['role']),
                'sourceName':next((s['name'] for s in state['sources'] if s['id']==component['sourceId']),source.get('name','')),'status':status})
            if status=='adopted' and component['outcome']!='hold':totals[component['currency']]+=Decimal(str(component['amount']))
        people.add(record['rowId'])
    pending=[];covered_slots=set()
    for entry in state.get('externalImports',[]):
        path=_batch_path(store,state['month'],entry['token'])
        if not path.is_file():continue
        preview=_preview(state,json.loads(path.read_text()))
        if preview['readyCount'] or preview['blockedCount']:
            pending.append({**entry,'sourceIds':[f['id'] for f in preview['files']],**{k:preview[k] for k in ('readyCount','blockedCount','unchangedCount')}})
            covered_slots.update((r['rowId'],r['kind'],r['role']) for r in preview['rows'] if r['rowId'] and r['status'] in {'ready','blocked','unchanged'})
    latest_ids={r['id'] for r in selected['latest']};active_ops={o['id'] for o in state.get('operations',[]) if o.get('active')}
    source_statuses={}
    for source in state['sources']:
        if source.get('purpose')!='external-results':continue
        sid=source['id']
        history=[r for r in state.get('externalResults',[]) if any(p['sourceId']==sid for p in r.get('proofs',[{'sourceId':r['sourceId']}]))]
        current=[r for r in history if r['id'] in latest_ids]
        source_statuses[sid]='invalid' if any(r in selected['invalid'] for r in current) else 'adopted' if current else 'superseded' if any(r['operationId'] in active_ops for r in history) else 'revoked' if history else 'pending'
    expected=[];rows={r['id']:r for r in state['rows']}
    for record in selected['expected']:
        if (record['rowId'],record['kind'],record['role']) in covered_slots:continue
        component=next((c for c in record['components'] if c['component']=='regular'),record['components'][0])
        values=rows.get(record['rowId'],{}).get('values',{})
        expected.append({'id':record['id']+'-expected','rowId':record['rowId'],'candidateName':record['candidateName'],
            'candidateId':record.get('candidateId') or text(values.get('工号') or values.get('体系外人员编号')),
            'project':_project(record['kind'],record['role']),'recipientName':component['recipientName'],'recipientId':component['recipientId'],
            'status':'missing','amount':None,'currency':component['currency'],'outcome':'pending'})
    return {'revision':state['revision'],'month':state['month'],'sourceStatuses':source_statuses,'records':records,'pendingImports':pending,
            'expected':expected,
            'summary':{'people':len(people),'records':len(records),'totalsByCurrency':{c:float(v) for c,v in totals.items()}}}


def attach_rows(state):
    _scope_projection(state)
    selected=effective(state);owned=selected['latest']+selected['expected'];by_row=defaultdict(list)
    for record in owned:by_row[record['rowId']].append(record)
    subjects={s['rowId']:s for s in state.get('externalSubjects',[])}
    state['todos']=[t for t in state.get('todos',[]) if not t['kind'].startswith('external_results_')]
    state['rows']=[r for r in state['rows'] if not r.get('externalOnly')]
    existing={r['id']:r for r in state['rows']}
    for row_id,records in by_row.items():
        if row_id not in existing:
            subject=subjects.get(row_id)
            if not subject or not subject.get('externalOnly'):continue
            values={'姓名':subject['candidateName'],'体系外人员编号':subject.get('candidateId',''),'入职日期':subject.get('hireDate','')}
            for record in records:
                component=next(iter(record['components']),{})
                values[record['role']+'姓名']=component.get('recipientName','');values[record['role']+'工号']=component.get('recipientId','')
            row={'id':row_id,'key':row_id,'values':values,'original':deepcopy(values),'provenance':{},'sources':[],
                 'history':[],'conflicts':[],'edits':[],'originKinds':['external_results'],'offerMatched':False,'scopeStatus':'已纳入底稿',
                 'needsReview':False,'unresolvedFields':[],'rule':{'version':'external','label':'体系外线下核定','basis':'月度核定结果','specialScope':''},'externalOnly':True}
            state['rows'].append(row);existing[row_id]=row
        row=existing[row_id];row['externalOwnership']=[{'kind':r['kind'],'role':r['role']} for r in records]
        for record in records:
            for c in record['components']:
                if not c.get('sourceId') or not c.get('sourceRows'):continue
                ref={'sourceId':c['sourceId'],'sheet':c['sheet'],'row':c['sourceRows'][0],'role':'external_results'}
                if ref not in row['sources']:row['sources'].append(ref)
    expected_slots={(r['rowId'],r['kind'],r['role']) for r in selected['expected']}
    covered_slots=set()
    for entry in state.get('externalImports',[]):
        coverage=_covered_slots(entry);covered_slots.update(coverage)
        if entry.get('readyCount') or entry.get('blockedCount') or coverage&expected_slots:
            state['todos'].append({'id':_hash('external_results_pending'+entry['token']),'kind':'external_results_pending',
                'title':'体系外结果有待采用或待补充内容','detail':'','rowIds':[],'sourceIds':entry.get('sourceIds',[])})
    for row in state['rows']:
        if row['id'] not in by_row:row.pop('externalOwnership',None)
    missing=[r for r in selected['expected'] if (r['rowId'],r['kind'],r['role']) not in covered_slots]
    for kind,records,title in [('external_results_missing',missing,'体系外本月核算结果待导入'),('external_results_invalid',selected['invalid'],'体系外核定资料已失效，请恢复原件或重新上传')]:
        if records:state['todos'].append({'id':_hash(kind),'kind':kind,'title':title,'detail':'','rowIds':list(dict.fromkeys(r['rowId'] for r in records)),
                                         'sourceIds':list(dict.fromkeys(r['sourceId'] for r in records if r.get('sourceId')))})
    if state['rows'] and all(r.get('externalOnly') for r in state['rows']):
        state['todos']=[t for t in state['todos'] if t['kind'] not in {'baseline','coverage'}]
    state['summary']['people']=len(state['rows'])


def automatic_roles(row,kind,month=None):
    if row.get('externalOnly') or row.get('externalCandidateScope'):return set()
    values=row['values'];roles={'推荐人'} if kind=='referral' else {'招聘负责人'}
    if kind=='referral' and not any(text(values.get(f)) for f in ('推荐人工号','推荐人姓名')) and values.get('招聘渠道')!='内推' and '推荐人工号' not in row.get('unresolvedFields',[]):roles=set()
    if kind=='recruitment' and (values.get('协助招聘人工号') or values.get('协助招聘人姓名')):roles.add('协助招聘人')
    owned={r['role'] for r in row.get('externalOwnership',[])+row.get('externalScopeOwnership',[]) if r['kind']==kind}
    return roles-owned


def automatic_dependency(row,kind,field=''):
    roles=automatic_roles(row,kind)
    for prefix in ('招聘负责人','协助招聘人','推荐人'):
        if field.startswith(prefix):return prefix in roles
    return bool(roles)


def empty_award(kind):
    return {'kind':kind,'standard':None,'total':0,'currency':'','formulaTotal':None,'countryBasis':'','source':'体系外线下核定',
            'cycle':None,'channelFactor':None,'issues':[],'reasons':[],'nodes':[],'currentAmount':0,'historyStatus':'未核实支付',
            'roundingPolicy':'线下核定金额','scheduledTotal':0,'roundingDifference':0,'recipientTotals':[]}


def overlay(state,row,result):
    if not row.get('externalOwnership') and not row.get('externalCandidateScope'):return result
    selected=effective(state);owned=[r for r in selected['latest']+selected['expected'] if r['rowId']==row['id']]
    candidate_scope=row.get('externalCandidateScope')
    if not owned and not candidate_scope:return result
    month=state['month']
    for kind in ('recruitment','referral'):
        records=[r for r in owned if r['kind']==kind]
        award=result[kind];roles={r['role'] for r in records};auto=automatic_roles(row,kind)
        if not records and not any(r['kind']==kind for r in row.get('externalScopeOwnership',[])):continue
        award['monthlyResult']={'source':'external','roles':sorted(roles),'automaticRoles':sorted(auto),'fullyExternal':not auto,
                                'unknown':bool(candidate_scope and not records)}
        award['nodes']=[n for n in award['nodes'] if not ((candidate_scope or n['role'] in roles) and not n.get('openingLocked') and (n.get('scheduledMonth',n.get('dueMonth'))==month or n.get('dueMonth')==month) and not n.get('adjustmentId'))]
        if not auto:
            award['diagnosticIssues']=award['issues'];award['issues']=[]
            award['accountingIssues']=[i for i in award.get('accountingIssues',[]) if i.get('accountingMonth')!=month]
        else:
            award['issues']=[issue for issue in award['issues'] if not any(issue.startswith(p) for p in roles)]
            for node in award['nodes']:
                if node['role'] in auto:node['issues']=[issue for issue in node['issues'] if not any(issue.startswith(p) for p in roles)]
        for record in records:
            valid=record in selected['valid'];expected=record in selected['expected']
            for component in record['components']:
                if expected and component['component']!='regular':continue
                issue='体系外本月核算结果待导入' if expected else '体系外核定资料已失效，请恢复原件或重新上传'
                status='暂缓安排' if component['outcome']=='hold' and valid else '按规则不发' if component['outcome']=='exclude' and valid else '本月到期' if valid else '待导入核定结果' if expected else '依据已失效'
                value=component['amount'] if valid else None
                award['nodes'].append({'id':record['id']+'-'+component['component']+'-'+str(month),'externalResultId':record['id'],
                    'rowId':row['id'],'role':record['role'],'recipientId':component['recipientId'],'recipientName':component['recipientName'],
                    'recipientKey':component['recipientKey'],'label':'体系外月度核定' if component['component']=='regular' else '体系外补发' if component['component']=='supplement' else '体系外扣回',
                    'component':component['component'],'amount':value,'currency':component['currency'],'dueDate':'','dueMonth':month,
                    'scheduledMonth':month if valid and component['outcome']!='hold' else None,'status':status,'timing':'本月核定',
                    'issues':[] if valid else [issue],'paidAmount':None,'paymentStatus':'未核实','paymentReady':False,
                    'scheduleText':'体系外线下核定','sourceId':component['sourceId'],'outcome':component['outcome']})
        current=[n for n in award['nodes'] if n.get('scheduledMonth',n.get('dueMonth'))==month]
        units={n['currency'] for n in current};known=all(n.get('amount') is not None for n in current)
        award['currentAmount']=float(sum((Decimal(str(n['amount'])) for n in current),Decimal(0))) if known and len(units)<=1 and records and all(r in selected['valid'] for r in records) else None
    return result
