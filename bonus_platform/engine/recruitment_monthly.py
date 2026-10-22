"""One business projection shared by the monthly screen and its export."""
from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .recruitment_detail_table import columns as detail_columns, cells as detail_cells


def worksheet(state, nodes, row_tasks):
    """Keep the familiar employee row while summing only this month's installments."""
    by_row=defaultdict(list)
    for node in nodes:by_row[node['rowId']].append(node)
    result=[]
    has_adjustments=any(n.get('adjustmentId') and n.get('accountingMonth')==state['month'] for n in nodes)
    for row in state['rows']:
        rid=row['id'];v=row['values'];all_nodes=by_row[rid]
        current_nodes=[n for n in all_nodes if n['scope']=='current']
        hire=str(v.get('入职日期') or v.get('候选人入职时间') or '')
        new_hire=hire[:7].replace('-','')==str(state['month'])
        scheduled=any(n['month']==state['month'] or n['scope']=='held' and n['dueMonth']==state['month'] for n in all_nodes)
        reasons=list(dict.fromkeys([reason for n in current_nodes for reason in n['reasons']]+
                    [reason for task in row_tasks.get(rid,[]) for reason in task.get('rowReasons',{}).get(rid,task.get('reasons',[])) or [task['title']]]))
        current=bool(new_hire or scheduled or row_tasks.get(rid))
        amounts={};currencies={};breakdown={}
        for kind,label in [('recruitment','招聘奖金'),('referral','内推奖金')]:
            award=row['calculation'][kind];grouped=defaultdict(Decimal)
            selected=[n for n in current_nodes if n['kind']==kind]
            external=award.get('monthlyResult',{})
            unknown=external.get('unknown',False) or (award['total'] is None and not external.get('fullyExternal')) or any(n.get('amount') is None for n in selected)
            if external and any(n.get('externalResultId') and n.get('amount') is None for n in all_nodes if n['kind']==kind):unknown=True
            for node in selected:
                if node.get('amount') is not None:grouped[node['currency']]+=Decimal(str(node['amount']))
            breakdown[kind]=[{'currency':currency,'amount':float(amount)} for currency,amount in grouped.items()]
            currencies[kind]=list(grouped) or ([award['currency']] if award.get('currency') else [])
            amounts[kind]=None if unknown or len(grouped)>1 else float(sum(grouped.values(),Decimal(0)))
            if current and len(grouped)>1:
                reasons.append(label+'包含不同币种，金额请查看分期明细')
        currencies_all=list(dict.fromkeys(currencies['recruitment']+currencies['referral']))
        currency=currencies_all[0] if len(currencies_all)==1 else '；'.join(
            label+'：'+'、'.join(currencies[kind]) for kind,label in [('recruitment','招聘'),('referral','内推')] if currencies[kind])
        reasons=list(dict.fromkeys(reasons));review=bool(reasons)
        note='；'.join(dict.fromkeys(str(v[k]) for k in ('人工备注','备注') if v.get(k)))
        eligibility_notes=list(dict.fromkeys(('内推奖金：' if n['kind']=='referral' else '招聘奖金：')+
                    '；'.join(n.get('eligibilityReasons',[])) for n in all_nodes
                    if n['scope']=='held' and n['dueMonth']==state['month'] and n.get('eligibilityReasons')))
        if eligibility_notes:note='；'.join(filter(None,[note,*eligibility_notes]))
        if reasons:note='；'.join(filter(None,[note,'；'.join(reasons)]))
        cells,cell_notes=detail_cells(row)
        if has_adjustments:
            for kind in ('recruitment','referral'):
                for direction in ('supplement','deduction'):
                    selected=[n for n in current_nodes if n['kind']==kind and n.get('adjustmentId') and (n['amount']>0)==(direction=='supplement')]
                    units={n['currency'] for n in selected}
                    cells[kind+'.'+direction]=float(sum((Decimal(str(n['amount'])) for n in selected),Decimal(0))) if len(units)<=1 else None
                    cell_notes[kind+'.'+direction]='；'.join(f"{n['recipientName']}：{n['amount']:,.2f} {n['currency']} · {n['scheduleText']}" for n in selected)
            adjustment_notes=[f"{n['label']} {n['recipientName']} {n['amount']:,.2f} {n['currency']}：{n['scheduleText']}" for n in current_nodes if n.get('adjustmentId')]
            note='；'.join(filter(None,[note,*adjustment_notes]))
        cells.update({'name':v.get('姓名',''),'employeeId':v.get('工号',''),'hireDate':hire,
               'channel':v.get('招聘渠道',''),'recruiter':v.get('招聘负责人姓名') or v.get('招聘负责人工号',''),
               'assistant':v.get('协助招聘人姓名') or v.get('协助招聘人工号',''),
               'recruitmentAmount':amounts['recruitment'],'referrer':v.get('推荐人姓名') or v.get('推荐人工号',''),
               'referralAmount':amounts['referral'],'currency':currency,
               'status':'待核对' if review else '后续资料待补' if None in amounts.values() else '体系外线下核定' if row.get('externalOwnership') else '已计算','note':note})
        result.append({'id':rid,'cells':cells,'current':current,'review':review,'reasons':reasons,
                       'newHire':new_hire,'amountsByCurrency':breakdown,'cellNotes':cell_notes})
    cols=detail_columns()
    if has_adjustments:
        cols[-3:-3]=[{'key':kind+'.'+direction,'label':label+title,'group':'本月核算','type':'money','width':136}
                     for kind,label in [('recruitment','招聘奖金'),('referral','内推奖金')]
                     for direction,title in [('supplement','补发'),('deduction','扣回')]]
    return {'columns':cols,'rows':result}


def recipient_summaries(state, nodes):
    """Match offline summary groups without changing source installment currencies."""
    people={r['id']:r['values'] for r in state['rows']}
    result={}
    stages={'入职满一个月':'month1','入职满三个月':'month3','入职满六个月':'month6','转正':'probation'}
    for kind in ('recruitment','referral'):
        columns=[('employeeId','工号'),('name','姓名'),('position','职位')]
        columns+= [('grade','职级')] if kind=='referral' else []
        columns += [('currency','币种'),('employment','人员状态')]
        if kind=='recruitment':columns += [('lastDay','最后工作日')]
        money=[('month1','入职1个月奖金'),('month3','入职3个月奖金'),('month6','入职6个月奖金'),('probation','转正奖金')] if kind=='recruitment' else []
        has_adjustments=any(n.get('adjustmentId') and n['kind']==kind and n.get('accountingMonth')==state['month'] for n in nodes)
        if has_adjustments:money += [('supplement','补发金额'),('deduction','扣回金额')]
        money += [('confirmed','本月已核对合计'),('pending','待核对金额')]
        grouped={}
        for n in nodes:
            if n['kind']!=kind or n['scope']!='current':continue
            # An unidentified payee must not merge with another unidentified person.
            number=str(n.get('recipientId') or '').strip()
            group_currency=n['currency']
            key=(n.get('recipientKey') or number.lower() or n['rowId']+':'+n['role'],group_currency)
            r=grouped.setdefault(key,{'employeeId':number,'name':n.get('recipientName',''),'currency':group_currency,'sourceCurrencies':[],
                **{k:Decimal(0) for k,_ in money},'unknown':0,'reasons':[], 'rowIds':[], '_meta':defaultdict(set)})
            if n['currency'] not in r['sourceCurrencies']:r['sourceCurrencies'].append(n['currency'])
            v=people.get(n['rowId'],{});prefix=n['role']
            for k,f in [('position','职位'),('grade','职级'),('employment','人员状态'),('lastDay','最后工作日')]:
                if not n.get('adjustmentId') and v.get(prefix+f):r['_meta'][k].add(str(v[prefix+f]))
            if n['rowId'] not in r['rowIds']:r['rowIds'].append(n['rowId'])
            r['reasons']=list(dict.fromkeys(r['reasons']+n.get('reasons',[])))
            if n.get('amount') is None:r['unknown']+=1;continue
            value=Decimal(str(n['amount']))
            r['pending' if n['review'] else 'confirmed']+=value
            if n.get('adjustmentId') and not n['review']:
                r['supplement' if value>0 else 'deduction']+=value
            if kind=='recruitment' and not n['review']:
                stage=stages.get(n['label'])
                if stage:r[stage]+=value
        rows=[]
        for r in grouped.values():
            for k,_ in money:r[k]=float(r[k])
            if kind=='recruitment' and r['currency']!='人民币' and r['employeeId']:
                while r['employeeId'].startswith('海外-'):r['employeeId']=r['employeeId'].removeprefix('海外-')
                r['employeeId']='海外-'+r['employeeId']
            meta=r.pop('_meta')
            for k in ('position','grade','employment','lastDay'):
                r[k]=' / '.join(sorted(meta[k]))
            r['note']='；'.join(([f"{r['unknown']} 笔金额待确认"] if r['unknown'] else [])+r['reasons'])
            rows.append(r)
        rows.sort(key=lambda r:(r['currency'],r['employeeId'].lower(),r['name']))
        totals=[]
        for currency in sorted({r['currency'] for r in rows}):
            selected=[r for r in rows if r['currency']==currency]
            totals.append({'currency':currency,**{k:float(sum((Decimal(str(r[k])) for r in selected),Decimal(0))) for k,_ in money},
                           'unknown':sum(r['unknown'] for r in selected)})
        result[kind]={'columns':[{'key':k,'label':v,'type':'money' if (k,v) in money else 'text'} for k,v in columns+money+[('note','备注')]],'rows':rows,'totals':totals,'grouping':'offline'}
    return result


def project(state, attention):
    from .recruitment_work_items import project as work_items
    month=state['month'];rows={r['id']:r for r in state['rows']}
    row_tasks=defaultdict(list);cases=[]
    for task in attention['current']:
        for rid in task['rowIds']:row_tasks[rid].append(task)
        if not task['rowIds']:
            cases.append({'id':task['id'],'title':task['title'],'rowIds':[],
                          'taskIds':[task['id']],'sourceIds':task['sourceIds'],'reasons':task.get('reasons',[])})
    grouped={}
    for rid,tasks in row_tasks.items():
        r=rows[rid];v=r['values'];key=rid;name=v.get('姓名') or v.get('工号') or '人员资料'
        if all(t['kind']=='rehire' for t in tasks) and v.get('上次离职日期') and not v.get('二次入职专项口径依据'):
            key='rule:rehire:'+r['rule']['version'];name='确认二次入职专项口径'
        # Merge only a proven shared-recipient root cause, not every case with
        # the same category. Independent candidate decisions remain separate.
        if all(t['kind']=='recipient' for t in tasks):
            issues=[i for k in ('recruitment','referral') for n in r['calculation'][k]['nodes']
                    if n.get('scheduledMonth',n['dueMonth'])==month for i in n['issues']]
            prefixes=[p for p in ('招聘负责人','协助招聘人','推荐人') if any(i.startswith(p+'：') for i in issues)]
            ids={v.get(p+'工号') for p in prefixes if v.get(p+'工号')}
            if len(ids)==1 and issues and all(any(i.startswith(p+'：') for p in prefixes) for i in issues):
                recipient=next(iter(ids));key='recipient:'+recipient;name=' / '.join(prefixes)+' '+recipient
        case=grouped.setdefault(key,{'id':key,'title':name,'rowIds':[],'taskIds':[], 'sourceIds':[], 'reasons':[]})
        case['rowIds'].append(rid)
        for task in tasks:
            for target,values in [('taskIds',[task['id']]),('sourceIds',task['sourceIds']),('reasons',[task['title']])]:
                case[target]=list(dict.fromkeys(case[target]+values))
    cases+=list(grouped.values())
    nodes=[];totals=defaultdict(lambda:{'confirmed':Decimal(0),'pending':Decimal(0),'count':0})
    for rid,row in rows.items():
        v=row['values']
        for kind,label in [('recruitment','招聘奖金'),('referral','内推奖金')]:
            award=row['calculation'][kind]
            for n in award['nodes']:
                schedule=n.get('scheduledMonth',n['dueMonth'])
                excluded=n['status'] in {'按规则不发','暂缓安排','依据已失效','安排需重新核对','待导入核定结果'}
                scope='held' if excluded else 'current' if schedule==month else 'future' if schedule and schedule>month else 'history'
                if n.get('adjustmentId') and n.get('accountingMonth')==month:scope='current'
                tasks=[] if n.get('adjustmentId') or n.get('externalResultId') else [t for t in row_tasks.get(rid,[]) if kind in t.get('rowAwardKinds',{}).get(rid,('recruitment','referral'))]
                reasons=list(dict.fromkeys(n['issues']+([t['title'] for t in tasks] if scope=='current' else [])))
                status='待确认' if reasons else '本月可核对' if scope=='current' else n['status']
                item={**n,'candidate':v.get('姓名',''),'employeeId':v.get('工号',''),'kind':kind,'project':label+' · '+n['label'],
                      'month':schedule,'scope':scope,'review':bool(reasons),'reasons':reasons,'displayStatus':status}
                nodes.append(item)
                if scope=='current':
                    total=totals[n['currency']]
                    if n.get('amount') is not None:total['pending' if reasons else 'confirmed']+=Decimal(str(n['amount']))
                    total['count']+=1
    active=[s for s in state['sources'] if s.get('active',True) and not s.get('comparisonOnly')]
    synced=[s for s in active if s.get('kind')=='beisen' and not s.get('employmentOnly') and not s.get('channelDetailsOnly') and not s.get('error') and s.get('forMonth')==month]
    from .recruitment_external_results import automatic_roles
    domestic={str(r['values'].get('工号','')).lower() for r in rows.values() if r['values'].get('标签分类')=='国内' and r['values'].get('工号') and any(automatic_roles(r,k) for k in ('recruitment','referral'))}
    from .recruitment_beisen import coverage_sets
    coverage=coverage_sets([s for s in active if s.get('kind')=='beisen' and s.get('forMonth')==month])
    queried=coverage['candidateEmployment']
    from .recruitment_attention import recipient_requirements
    missing_recipients=recipient_requirements(state,month)['domesticFetchIds']
    work=work_items(state,attention)
    return {'nodes':nodes,'recipientSummaries':recipient_summaries(state,nodes),'worksheet':worksheet(state,nodes,row_tasks),'cases':cases,'reviewPeople':len(row_tasks),'workItems':work,
            'automaticResults':automatic_results(state,nodes,attention.get('notices',[])),
            'actionCount':sum(w['current'] for w in work),'preparationCount':sum(not w['current'] for w in work),
            'operatorActionCount':sum(w['current'] and not w.get('ruleMaintenance') for w in work),
            'ruleActionCount':sum(w['current'] and w.get('ruleMaintenance',False) for w in work),
            'preparationPeople':len({rid for w in work for rid in w['rowIds']} - set(row_tasks)),
            'totals':[{'currency':currency,**{k:float(v) if isinstance(v,Decimal) else v for k,v in total.items()}} for currency,total in sorted(totals.items())],
            'currentCount':sum(n['scope']=='current' for n in nodes),
            'recipientCount':len({n['recipientId'] for n in nodes if n['scope']=='current' and n['recipientId']}),
            'coverage':{'domesticPeople':len(domestic),'queriedPeople':len(domestic&queried),
                        'recipientSyncRequired':bool(missing_recipients),'recipientQueryCount':len(missing_recipients),
                        'candidateOfferPeople':len(domestic&coverage['candidateOffer']),
                        'attemptedPeople':len(domestic&coverage['requestedCandidate']),
                        'syncRequired':bool(domestic-queried),'syncedAt':max((s.get('addedAt','') for s in synced),default=''),
                        'discoveryComplete':False,'sourceErrors':sum(bool(s.get('error')) for s in active if s.get('purpose')!='external-results'),
                        'hasOpening':bool(state.get('baseline') or state.get('parentMonth') or state.get('openingEmptyConfirmed')),
                        'notice':'国内同步覆盖已加入本月的人员；新增人员仍需花名册，海外和体系外资料请补充。' if domestic else '本月按原始资料核算；请核对海外和体系外人员是否完整纳入。'}}


def automatic_results(state,nodes,notices=()):
    month=state['month'];items=[]
    catchup=[n for n in nodes if n.get('scheduleBasis')=='上月待转正余额承接' and n['scope']=='current']
    if catchup:items.append({'title':f'上月待转正奖金已接续 · {len(catchup)} 笔','detail':'保留上月已安排金额，按本次实际转正资料计算剩余奖金。'})
    if state.get('openingSelection',{}).get('mode')=='uniquePreviousMonth':
        items.append({'title':'上月核算表已承接','detail':'已识别唯一的上月表，继续计算未结束的奖金，无需重复确认。'})
    external={n['recipientId'] for n in nodes if n['scope']=='current' and n.get('employmentBasis')}
    if external:items.append({'title':f'体系外／GL 人员 {len(external)} 人','detail':'使用已导入的线下资料，不要求内部任职记录。'})
    probation={n['recipientId'] for n in nodes if n['dueMonth']==month and n['scope']=='held' and any('推荐人：适用日期尚未转正'==r for r in n.get('eligibilityReasons',[]))}
    if probation:items.append({'title':f'推荐人试用状态 {len(probation)} 人','detail':'按花名册状态判断，本月相应内推分期不安排；原因已列入核算明细备注。'})
    special={n['rowId'] for n in nodes if n['kind']=='referral' and n['dueMonth']==month}
    special={r['id'] for r in state['rows'] if r['id'] in special and r['rule']['version']=='legacy' and r['rule'].get('specialScope')}
    if special:items.append({'title':f'德国／捷克旧规则 {len(special)} 人','detail':'按 7 月前入职及已确认的部门边界自动适用，保留推荐人资格条件。'})
    due={n['rowId'] for n in nodes if n['kind']=='recruitment' and n['dueMonth']==month}
    rehires=[r for r in state['rows'] if r['calculation']['recruitment'].get('rehireBasis')
             and (r['id'] in due or r['calculation']['recruitment']['rehireBasis']['hireDate'][:7].replace('-','')==str(month))]
    if rehires:items.append({'title':f'二次入职渠道比例已确定 · {len(rehires)} 人',
        'detail':'离职超过12个月，专项条款与通用说明均按40%计算；具体日期及依据可在个人奖金明细查看。'})
    for notice in notices:
        dates='、'.join(dict.fromkeys(f"{int(r['lastDay'][5:7])}月{int(r['lastDay'][8:10])}日" for r in notice['records']))
        items.append({'title':str(notice['name'])+' · 离职流程不影响本月核算',
            'detail':f'流程尚未结束，预计最后工作日为{dates}；按现有日期对照，本月奖金不变，无需选择人员状态或补依据。后续导入最终结果时会重新判断。'})
    return items


def export_filename(month):
    return f'全球招聘奖金核算明细表（非操作类）{month}.xlsx'


def export(state, monthly):
    """Three offline-style worksheets, using the same amounts as the screen."""
    from openpyxl.styles import Border, Side
    wb=Workbook();wb.remove(wb.active)
    line=Side(style='thin',color='D9D9D9')
    def safe(value):
        return "'"+value if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')) else value
    def color(group):
        if '协助' in group:return 'C6E0B4'
        if '内推' in group:return 'B4C6E7'
        if '奖金' in group:return 'A9D08E'
        if group=='本月核算':return 'FFF2CC'
        if group=='备注':return 'D9E1F2'
        return 'FCE4D6'
    def sheet(name,cols,values,groups,total_count=0):
        ws=wb.create_sheet(name)
        ws.append(groups);ws.append([c['label'] for c in cols])
        for values_row in values:ws.append([safe(v) for v in values_row])
        start=1
        for i in range(2,len(cols)+2):
            if i==len(cols)+1 or groups[i-1]!=groups[start-1]:
                if i-start>1:ws.merge_cells(start_row=1,start_column=start,end_row=1,end_column=i-1)
                start=i
        last_row=ws.max_row
        ws.freeze_panes='C3';ws.auto_filter.ref=f'A2:{get_column_letter(len(cols))}{max(2,last_row-total_count)}'
        ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85
        ws.row_dimensions[1].height=25;ws.row_dimensions[2].height=40
        for i,c in enumerate(cols,1):
            ws.column_dimensions[get_column_letter(i)].width=max(12,min(38,c.get('width',126)/7))
            for row in range(1,last_row+1):
                cell=ws.cell(row,i);cell.font=Font(name='微软雅黑',size=10,bold=row<=2 or row>last_row-total_count,color='000000')
                cell.border=Border(left=line,right=line,top=line,bottom=line)
                cell.alignment=Alignment(vertical='center',horizontal='center' if row<=2 else 'right' if c['type']=='money' else 'left',wrap_text=True)
                if row<=2:cell.fill=PatternFill('solid',fgColor=color(str(groups[i-1])))
                elif row>last_row-total_count:cell.fill=PatternFill('solid',fgColor='FFF2CC')
                elif c['key']=='pending':cell.fill=PatternFill('solid',fgColor='FFF8EB')
                if row>=3 and c['type'] in {'money','percent','number'}:
                    cell.number_format={'money':'#,##0.00','percent':'0.##%','number':'0.##'}[c['type']]
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.print_title_rows='1:2';ws.print_options.horizontalCentered=True
        ws.oddFooter.center.text=f"{state['month']} · 第 &P 页 / 共 &N 页"
        return ws
    table=monthly['worksheet'];columns=table['columns']
    sheet('附件一-招聘奖金核算明细',columns,
          [[r['cells'].get(c['key'],'') for c in columns] for r in table['rows'] if r['current']],
          [c['group'] for c in columns])
    summaries=monthly.get('recipientSummaries') or recipient_summaries(state,monthly['nodes'])
    for kind,name in [('recruitment','附件二-招聘奖金汇总表'),('referral','附件三-内推奖金汇总')]:
        data=summaries[kind];cols=[dict(c) for c in data['columns']]
        labels={'month1':'入职1个月发放奖金','month3':'入职3个月发放奖金','month6':'入职6个月发放奖金','confirmed':'合计发放' if kind=='recruitment' else '本月发放'}
        for c in cols:
            c['label']=labels.get(c['key'],c['label'])
            if kind=='referral' and c['key'] in {'employeeId','name'}:c['label']='内推人员'+c['label']
            c['width']=224 if c['key']=='note' else 154 if c['type']=='money' else 126
        values=[[r.get(c['key'],'') for c in cols] for r in data['rows']]
        values += [[t.get(c['key'],0) if c['type']=='money' else '合计' if c['key']=='name' else t['currency'] if c['key']=='currency' else '' for c in cols] for t in data['totals']]
        ws=sheet(name,cols,values,[str(state['month']) if c['type']=='money' else '人员资料' if c['key']!='note' else '备注' for c in cols],len(data['totals']))
        for i,c in enumerate(cols,1):
            for row in (1,2):ws.cell(row,i).fill=PatternFill('solid',fgColor='FFF2CC' if kind=='recruitment' and c['type']=='money' else 'FCE4D6')
    out=BytesIO();wb.save(out);return out.getvalue()
