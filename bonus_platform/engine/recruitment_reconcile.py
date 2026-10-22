"""Independent current-ledger comparison; the reference is excluded from calculation inputs."""
from collections import Counter, defaultdict
from copy import deepcopy
from decimal import Decimal
from io import BytesIO
import json
from hashlib import sha256
import re

from openpyxl import Workbook
from .recruitment_awards import calculate_draft, ensure_calculated, lookup, number, money, NODES
from .recruitment_draft_sources import date_text, text
from .recruitment_drafts import staff_id

FIELDS=('ABC类别','职级','招聘渠道','招聘启动日期','招聘负责人工号','协助招聘人工号','推荐人工号','周期剔除天数')


def comparison_month(source):
    periods = {int(year) * 100 + int(month) for year, month in re.findall(
        r'(?<!\d)(20\d{2})[年._\s-]*(0?[1-9]|1[0-2])月?(?!\d)', source['name'])}
    if periods:
        return next(iter(periods)) if len(periods) == 1 else None
    months = {int(month) for month in re.findall(r'(?<!\d)(0?[1-9]|1[0-2])\s*月', source['name'])}
    if months:
        if len(months) == 1 and source.get('forMonth'):
            return source['forMonth'] // 100 * 100 + next(iter(months))
        return None
    if source.get('comparisonOnly'):
        return source.get('forMonth')
    return None


def unavailable_comparison(source=None):
    return {'comparisonScope': 'history' if source else 'not_provided', 'comparisonReady': False,
            'summary': {'differenceItems': 0}, 'differences': [], 'totalsByCurrency': [],
            'basis': '所选台账用于全历史与字段核对，尚未进行本月线下比较' if source else '尚未导入本月线下核对表'}


def identity(values):
    number=staff_id(values.get('工号'))
    return (number,date_text(values.get('入职日期') or values.get('候选人入职时间'))) if number and number not in {'-','无工号'} else None


def comparable(field,value):
    value=text(value)
    if value in {'-','无','N/A'}: return ''
    if field.endswith('工号'): return staff_id(value)
    if field=='ABC类别': return value.replace('类','').replace('C1','C')
    return value


def reference_matches(values, rows):
    key=identity(values)
    if key:
        return [r for r in rows if identity(r['values'])==key]
    fields=('职位','二级组织','三级组织','工作地','招聘启动日期','海外offer流程编号')
    matches=[]
    for row in rows:
        v=row['values']
        if text(values.get('姓名')).casefold()!=text(v.get('姓名')).casefold() or not date_text(values.get('入职日期')) or date_text(values.get('入职日期'))!=date_text(v.get('入职日期')):continue
        comparisons=[comparable(f,values[f])==comparable(f,v[f]) for f in fields if values.get(f) and v.get(f)]
        if len(comparisons)>=2 and all(comparisons):matches.append(row)
    return matches


def match_reference(values,rows):
    matches=reference_matches(values,rows)
    return matches[0] if len(matches)==1 else None


def comparison_signature(value):
    return sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,default=str).encode()).hexdigest()


def roster_comparison(saved,reference,source):
    by_ref=defaultdict(list)
    for row in saved['rows']:
        for h in row.get('history',[]):
            if h['ref']['sourceId']==source['id']:by_ref[(h['ref']['sheet'],h['ref']['row'])].append(row)
    by_identity=defaultdict(list)
    for row in saved['rows']:
        key=identity(row['values'])
        if key:by_identity[key].append(row)
    candidates=[]
    for rec in reference:
        key=identity(rec['values'])
        candidates.append(by_ref.get((rec['ref']['sheet'],rec['ref']['row'])) or (by_identity.get(key,[]) if key else reference_matches(rec['values'],saved['rows'])))
    uses=Counter(row['id'] for matches in candidates for row in matches)
    assignments=[]; differences=[]; covered=set()
    for rec,matches in zip(reference,candidates):
        covered.update(r['id'] for r in matches)
        row=matches[0] if len(matches)==1 and uses[matches[0]['id']]==1 else None
        status='matched' if row else 'ambiguous' if matches else 'unmatched'
        assignments.append((rec,row,status))
        if row:continue
        ident={k:rec['values'].get(k,'') for k in ('姓名','工号','入职日期')}
        ref={**rec['ref'],'sourceId':source['id']}
        subject={'side':'reference','sourceRef':ref,'identity':ident,'matchStatus':status,
                 'taskIds':[],'itemIds':[],'monthImpacts':['scheduled_current'] if any(n['period']==str(saved['month']) and n['amount'] for n in reference_nodes(rec)) else ['future']}
        subject['fingerprint']=comparison_signature([subject,rec,source.get('contentDigest',source['id']),[r['id'] for r in matches]])
        differences.append({'subject':subject,'side':subject['side'],'sourceRef':ref,'rowId':'','identity':ident,'matchStatus':status,
            'name':ident['姓名'],'employeeId':ident['工号'],'hireDate':date_text(ident['入职日期']),
            'sourceId':source['id'],'sheet':ref['sheet'],'sourceRow':ref['row'],'kind':'人员范围','item':'线下任职未匹配' if status=='unmatched' else '任职匹配不唯一',
            'category':'线下未匹配' if status=='unmatched' else '匹配待核','reason':'原始资料未对应本次任职' if status=='unmatched' else '存在多个候选匹配或多条线下记录对应同一任职，不能判为确定遗漏',
            'reference':None,'calculated':None,'currency':'','calculatedCurrency':'','referencePeriod':'','calculatedPeriod':None,
            'note':rec['values'].get('人工备注',''),'taskIds':[],'itemIds':[],'monthImpacts':subject['monthImpacts'],'fingerprint':subject['fingerprint']})
    for row in saved['rows']:
        if row['id'] in covered:continue
        ident={k:row['values'].get(k,'') for k in ('姓名','工号','入职日期')}
        current=any(n.get('scheduledMonth',n.get('dueMonth'))==saved['month'] and n.get('amount') for a in row.get('calculation',{}).values() if isinstance(a,dict) for n in a.get('nodes',[]))
        subject={'side':'draft','rowId':row['id'],'identity':ident,'matchStatus':'draft_only','taskIds':[],'itemIds':[],'monthImpacts':['scheduled_current','future'] if current else ['future']}
        subject['fingerprint']=comparison_signature([subject,row['values'],row.get('unresolvedFields'),source.get('contentDigest',source['id'])])
        differences.append({'subject':subject,'side':'draft','rowId':row['id'],'identity':ident,'matchStatus':'draft_only',
            'name':ident['姓名'],'employeeId':ident['工号'],'hireDate':date_text(ident['入职日期']),
            'sourceId':source['id'],'sheet':'','sourceRow':'','kind':'人员范围','item':'平台任职独有','category':'平台独有',
            'reason':'当前底稿有本次任职，线下未列；先核对身份及适用范围，关联现有资料事项继续处理',
            'reference':None,'calculated':None,'currency':'','calculatedCurrency':'','referencePeriod':'','calculatedPeriod':None,'note':'',
            'taskIds':[],'itemIds':[],'monthImpacts':subject['monthImpacts'],'fingerprint':subject['fingerprint']})
    return assignments,differences


def reference_nodes(rec):
    """Read amounts and their adjacent period from the original named groups."""
    cells=sorted(rec.get('historyCells',[]),key=lambda c:c['column'])
    by_column={c['column']:c for c in cells}
    for cell in cells:
        label=cell['label']
        if not any(word in label for word in ('发放奖金','发放金额')): continue
        amount=number(cell['value'])
        if amount is None: continue
        role='推荐人' if '内推' in label else '协助招聘人' if '协助' in label else '招聘负责人' if '招聘负责人' in label or '招聘人转正' in label else ''
        if not role: continue
        index=next((i for i,word in enumerate(('入职1个月','入职3个月','入职6个月','转正')) if word in label),None)
        if index is None: continue
        following=by_column.get(cell['column']+1,{})
        period=text(following.get('value','')) if '周期' in following.get('label','') else ''
        yield {'kind':'referral' if role=='推荐人' else 'recruitment','role':role,'label':NODES[index],
               'amount':float(amount),'period':period,'column':cell['column']}


def reference_currency(rec,kind):
    # A legacy Czech referral can have a different currency from recruitment.
    # Read that explicit annotation before the ledger's general currency column.
    if kind=='referral':
        notes=' '.join(text(c['value']) for c in rec.get('historyCells',[]) if '内推' in c['label'])
        if '内推奖币种' in notes and any(s in notes for s in ('捷克郎','捷克克朗','CZK')):return '捷克克朗'
    raw=text(rec['values'].get('标准币种'))
    aliases={'CNY':'人民币','RMB':'人民币','人民币':'人民币','USD':'美元','美元':'美元','EUR':'欧元','欧元':'欧元','CZK':'捷克克朗','捷克郎':'捷克克朗','捷克克朗':'捷克克朗'}
    return aliases.get(raw.upper(),aliases.get(raw,'币种待核'))


def monthly_comparison(saved,reference,source):
    """Compare independent dimensions without adopting any reference money."""
    saved=ensure_calculated(saved,copy_if_stale=True)
    month=saved['month'];counts=Counter();items=[]
    totals=defaultdict(lambda:{'reference':Decimal(0),'calculated':Decimal(0)})
    assignments,roster=roster_comparison(saved,reference,source)
    matched=set();seen=set();ref_by_row={}
    labels={'amount':'金额','period':'月份','currency':'币种','recipient':'领取人','role':'角色分配','eligibility':'资格'}
    def dim(status,reason):return {'status':status,'reason':reason}
    def eligibility(a,node):
        reasons=list(dict.fromkeys((a or {}).get('issues',[])+(a or {}).get('reasons',[])+(node or {}).get('eligibilityReasons',[])+(node or {}).get('issues',[])))
        if node and node.get('status') in {'按规则不发','安排需重新核对','依据已失效'} and node.get('scheduleText'):reasons.append(node['scheduleText'])
        if (a or {}).get('approvalDecision'):reasons.append('已采用审批不发决定')
        return dim('blocked','；'.join(dict.fromkeys(reasons))) if reasons else dim('unknown','金额及安排一致不代表资格或实际支付已经核实')
    for rec,row,status in assignments:
        refs=list(reference_nodes(rec))
        if row:matched.add(row['id']);ref_by_row[row['id']]=(rec,refs)
        common={'name':rec['values'].get('姓名'),'employeeId':rec['values'].get('工号'),'hireDate':date_text(rec['values'].get('入职日期')),
                'rowId':row['id'] if row else '', 'sourceRow':rec['ref']['row'],'sourceId':source['id'],'sheet':rec['ref']['sheet']}
        for ref in refs:
            if ref['period']!=str(month) or not ref['amount']:continue
            counts['referenceNodes']+=1
            a=row['calculation'][ref['kind']] if row else None
            node=next((n for n in a['nodes'] if n['role']==ref['role'] and n['label']==ref['label'] and not n.get('adjustmentId')),None) if a else None
            currency=reference_currency(rec,ref['kind']);actual_currency=(a or {}).get('currency','')
            totals[currency]['reference']+=Decimal(str(ref['amount']))
            if row:seen.add((row['id'],ref['kind'],ref['role'],ref['label']))
            actual=node['amount'] if node else 0 if a and a['total'] is not None else None
            actual_month=node.get('scheduledMonth',node.get('dueMonth')) if node else None
            same=actual is not None and money(actual)==money(ref['amount'])
            counts['equalAmounts' if same else 'unresolvedAmounts' if actual is None else 'differentAmounts']+=1
            if actual_month==month:counts['equalPeriods']+=1
            expected_id=staff_id(rec['values'].get(ref['role']+'工号'));expected_name=text(rec['values'].get(ref['role']+'姓名')).casefold().replace(' ','')
            actual_id=staff_id((node or {}).get('recipientId'));actual_name=text((node or {}).get('recipientName')).casefold().replace(' ','')
            same_recipient=bool(expected_id and actual_id and expected_id==actual_id) if expected_id and actual_id else bool(expected_name and actual_name and expected_name==actual_name)
            same_currency=currency==actual_currency and currency!='币种待核'
            if not same_currency:counts['differentCurrencies']+=1
            if node and not same_recipient:counts['differentRecipients']+=1
            ref_roles={n['role']:money(n['amount']) for n in refs if n['kind']==ref['kind'] and n['label']==ref['label'] and n['amount']}
            actual_roles={n['role']:money(n['amount']) for n in (a or {}).get('nodes',[]) if n['label']==ref['label'] and n['amount'] and not n.get('adjustmentId')}
            same_roles=ref_roles==actual_roles
            amount_reasons=list((a or {}).get('issues',[]))
            if not same and actual is not None and a:
                amount_reasons=list(a.get('reasons',[]))+amount_reasons
                cycle=a.get('cycle')
                if cycle and cycle.get('exempt'):amount_reasons.append('已采用认可的周期豁免，周期系数100%')
                elif cycle and cycle.get('factor') is not None:
                    amount_reasons.append(f'采用周期 {cycle["actualDays"]:g} 天（剔除 {cycle["exclusionDays"]:g} 天），周期系数 {cycle["factor"]:.0%}，渠道系数 {a["channelFactor"]:.0%}')
            dimensions={
                'amount':dim('equal' if same else 'unknown' if actual is None else 'different','金额一致' if same else '；'.join(dict.fromkeys(amount_reasons)) or '规则计算金额与线下不同，请核对适用标准、周期及分配依据'),
                'period':dim('equal' if actual_month==month else 'different','均安排在本月' if actual_month==month else (node or {}).get('scheduleText') or f'线下安排本月，当前安排为 {actual_month or "尚未安排"}'),
                'currency':dim('equal' if same_currency else 'unknown' if currency=='币种待核' else 'different',f'线下 {currency}，当前 {actual_currency or "待核"}'),
                'recipient':dim('equal' if same_recipient else 'different' if expected_id or expected_name else 'unknown',f'线下领取人 {rec["values"].get(ref["role"]+"姓名") or expected_id or "待核"}，当前 {(node or {}).get("recipientName") or actual_id or "待核"}'),
                'role':dim('equal' if same_roles else 'different','角色及分配一致' if same_roles else '线下与平台各角色分配不同；合计一致不能替代逐角色核对'),
                'eligibility':eligibility(a,node)}
            if not row:dimensions['amount']=dim('unknown','台账原行尚未唯一对应底稿任职')
            four_equal=same and actual_month==month and same_currency and same_recipient
            if four_equal:counts['equalNodes']+=1
            qualifying=[k for k in ('role','eligibility') if dimensions[k]['status'] in {'different','blocked'}]
            if qualifying:
                counts['qualificationDifferenceNodes']+=1
                if four_equal:counts['additionalQualificationNodes']+=1
            if not four_equal or qualifying:
                changed=[k for k,v in dimensions.items() if v['status'] in {'different','blocked'} or k=='amount' and v['status']=='unknown']
                category='、'.join(labels[k] for k in changed)+'待核' if len(changed)>1 else {'period':'发放月份不同','amount':'金额不同' if actual is not None else '无法试算','currency':'币种不同','recipient':'领取人不同或待核','role':'角色分配不同','eligibility':'资格待核'}.get(next(iter(changed),''),'任职待核')
                items.append({**common,'kind':'内推奖金' if ref['kind']=='referral' else '招聘奖金','item':ref['role']+' · '+ref['label'],
                    'reference':ref['amount'],'calculated':actual,'currency':currency,'calculatedCurrency':actual_currency,'referencePeriod':ref['period'],'calculatedPeriod':actual_month,
                    'category':category,'reason':'；'.join(dict.fromkeys(dimensions[k]['reason'] for k in changed)),'dimensions':dimensions,'fourDimensionsEqual':four_equal,'note':rec['values'].get('人工备注','')})
    for row in saved['rows']:
        rec,refs=ref_by_row.get(row['id'],(None,[]))
        for kind,title in [('recruitment','招聘奖金'),('referral','内推奖金')]:
            a=row['calculation'][kind]
            for node in a['nodes']:
                if node.get('scheduledMonth',node.get('dueMonth'))!=month or not node['amount']:continue
                totals[node['currency']]['calculated']+=Decimal(str(node['amount']))
                if (row['id'],kind,node['role'],node['label']) in seen:continue
                counts['draftOnlyNodes']+=1
                ref=next((n for n in refs if n['kind']==kind and n['role']==node['role'] and n['label']==node['label']),None)
                refusal=bool(ref and any(t in ref['period'] for t in ('不发','不享受','取消')))
                reference_status='refused' if refusal else 'blank_period' if ref and not ref['period'] else 'other_period' if ref else 'absent_node' if rec else 'absent_row'
                period_reason='线下明确不发或不享受：'+ref['period'] if refusal else '线下原格月份为空' if reference_status=='blank_period' else '线下对应节点不在本月安排' if ref else '线下未列对应角色节点' if rec else '线下未匹配本次任职'
                dimensions={'amount':dim('different' if ref else 'unknown','当前有本月非零安排，参考金额仅作核对'),
                    'period':dim('different',period_reason),'currency':dim('equal' if rec and reference_currency(rec,kind)==node['currency'] else 'unknown','按原币种分别核对'),
                    'recipient':dim('unknown','核对对应领取人来源'),'role':dim('different' if rec and not ref else 'unknown','线下未列该角色分配' if rec and not ref else '逐角色核对'),
                    'eligibility':dim('blocked',period_reason) if refusal else eligibility(a,node)}
                items.append({'rowId':row['id'],'name':row['values'].get('姓名'),'employeeId':row['values'].get('工号'),'hireDate':row['values'].get('入职日期'),
                    'sourceRow':rec['ref']['row'] if rec else '','sheet':rec['ref']['sheet'] if rec else '','sourceId':source['id'],'kind':title,'item':node['role']+' · '+node['label'],'reference':ref['amount'] if ref else None,'calculated':node['amount'],
                    'currency':reference_currency(rec,kind) if rec else '','calculatedCurrency':node['currency'],'referencePeriod':ref['period'] if ref else '','calculatedPeriod':month,'category':'线下明确不发' if refusal else '台账本月未列',
                    'referenceStatus':reference_status,'dimensions':dimensions,'fourDimensionsEqual':False,'reason':period_reason+'；核对资格、角色或延期依据','note':rec['values'].get('人工备注','') if rec else row['values'].get('人工备注','')})
    return {'comparisonScope':'current','comparisonReady':True,
            'summary':{**dict(counts),'matchedPeople':len(matched),'differenceItems':len(items),'unmatchedPeople':sum(d['matchStatus']=='unmatched' for d in roster),'draftOnlyPeople':sum(d['matchStatus']=='draft_only' for d in roster),'ambiguousPeople':sum(d['matchStatus']=='ambiguous' for d in roster),'rosterDifferenceItems':len(roster)},
            'differences':items,'rosterDifferences':roster,
            'matchedRosterSubjects':[{'rowId':row['id'],'identity':{k:rec['values'].get(k,'') for k in ('姓名','工号','入职日期')}} for rec,row,status in assignments if row],
            'totalsByCurrency':[{'currency':c,**{k:money(v) for k,v in values.items()}} for c,values in sorted(totals.items())],
            'basis':saved.get('offlineReplay',{}).get('basis') or '四维一致仅指金额、月份、币种、领取人；角色与资格另列待核，均不证明实际支付。参考金额不进入计算。人员范围与金额节点分别统计。'}


def reconcile(store,month,source_id):
    saved=store.read(month)
    source=next((s for s in saved['sources'] if s['id']==source_id and s.get('active',True) and 'ledger' in s.get('roles',[])),None)
    if not source: raise ValueError('请选择使用中的核对台账')
    parsed=json.loads((store._folder(month)/'parsed'/f'{source_id}.json').read_text())
    reference=[r for r in parsed['records'] if r['role']=='ledger']
    working=deepcopy(saved);working['sources']=[s for s in working['sources'] if s['id']!=source_id]
    working['rows']=[]
    if working.get('baseline')==source_id: working['baseline']=''
    working=store._rebuild(store._folder(month),working);calculate_draft(working)
    assignments,roster=roster_comparison(working,reference,source)
    money_diffs=[];field_diffs=[];missing=[];counts=Counter()
    for rec,row,match_status in assignments:
        common={'name':rec['values'].get('姓名'),'employeeId':rec['values'].get('工号'),'hireDate':date_text(rec['values'].get('入职日期')),'sourceId':source_id,'sheet':rec['ref']['sheet'],'sourceRow':rec['ref']['row']}
        if not row:
            missing.append({**common,'reason':'任职匹配不唯一，需核对' if match_status=='ambiguous' else '原始资料未对应本次任职'});continue
        counts['matchedPeople']+=1;common['rowId']=row['id']
        changed_fields=[]
        for field in FIELDS:
            actual=row['values'].get(field,'');expected=rec['values'].get(field,'')
            if comparable(field,actual)!=comparable(field,expected):
                changed_fields.append(field);field_diffs.append({**common,'field':field,'reference':expected,'calculated':actual})
        for kind,title,keyword in [('recruitment','招聘奖金','奖金标准'),('referral','内推奖金','推荐奖金标准')]:
            a=row['calculation'][kind];rule=lookup(row,kind)
            candidates=[c for c in rec['historyCells'] if c['label'].split(' / ')[-1]==keyword]
            expected=number(candidates[0]['value']) if candidates else None
            actual=number(rule['amount']) if rule else None
            if expected is not None:
                same=(expected==0 and a['total']==0) or (a['total'] is not None and actual==expected)
                if not same and kind=='recruitment' and actual is not None and a['channelFactor'] is not None and a['total'] is not None:
                    same=money(actual*Decimal(str(a['channelFactor'])))==money(expected)
                if same: counts['equalStandards']+=1
                else:
                    reason='；'.join(a['issues']) if a['total'] is None else '规则标准与台账不同'
                    money_diffs.append({**common,'kind':title,'item':'奖金标准','reference':float(expected),'calculated':float(actual) if actual is not None and a['total'] is not None else None,'currency':a['currency'],'reason':reason,'differentFields':changed_fields})
            labels={'入职满一个月':'入职1个月','入职满三个月':'入职3个月','入职满六个月':'入职6个月','转正':'转正'}
            nodes={(n['role'],n['label']):n for n in a['nodes']}
            # Include reference-only nodes: a cancellation or removal of a helper
            # must create a difference rather than hide a historical allocation.
            for role in (['推荐人'] if kind=='referral' else ['招聘负责人','协助招聘人']):
                for node_label,raw_label in labels.items():
                    cells=[c for c in rec['historyCells'] if raw_label in c['label'] and ('发放奖金' in c['label'] or '发放金额' in c['label']) and (('内推' in c['label']) if kind=='referral' else ('协助' in c['label']) if role=='协助招聘人' else ('招聘负责人' in c['label'] or '招聘人转正' in c['label']))]
                    expected=number(cells[0]['value']) if cells else None
                    node=nodes.get((role,node_label))
                    if expected is None:
                        if node: counts['blankReferenceNodes']+=1
                        continue
                    actual=node['amount'] if node else 0 if a['total'] is not None else None
                    if actual is not None and money(expected)==money(actual):
                        counts['equalNodes']+=1;continue
                    reasons=[]
                    if changed_fields: reasons.append('采用字段不同：'+'、'.join(changed_fields))
                    if a.get('approvalDecision'): reasons.append('已采用审批不发决定')
                    if actual is None: reasons += a['issues']
                    if not reasons: reasons.append('采用字段一致；台账节点金额与规则试算不同，原调整依据待核')
                    money_diffs.append({**common,'kind':title,'item':role+' · '+node_label,'reference':float(expected),'calculated':actual,'currency':a['currency'],'reason':'；'.join(reasons),'differentFields':changed_fields})
    missing=[d for d in roster if d['matchStatus']=='unmatched']
    unresolved=sum(d['calculated'] is None for d in money_diffs)
    return {'month':month,'revision':saved['revision'],'referenceId':source_id,'referenceName':source['name'],'referenceRows':len(reference),'calculationPeople':len(working['rows']),
            'summary':{**dict(counts),'moneyDifferences':len(money_diffs),'unresolvedMoneyItems':unresolved,'numericMoneyDifferences':len(money_diffs)-unresolved,'fieldDifferences':len(field_diffs),'unmatchedPeople':len(missing),'draftOnlyPeople':sum(d['matchStatus']=='draft_only' for d in roster),'ambiguousPeople':sum(d['matchStatus']=='ambiguous' for d in roster)},'rosterDifferences':roster,'draftOnlyPeople':[d for d in roster if d['matchStatus']=='draft_only'],'moneyDifferences':money_diffs,'fieldDifferences':field_diffs,'unmatchedPeople':missing,
            'monthly':monthly_comparison(saved,reference,source) if comparison_month(source)==month else unavailable_comparison(source),
            'ruleValidation':saved.get('ruleValidation',{}),'strictInputs':saved.get('strictInputs',False),
            'basis':'本页独立核对排除整份核对台账，使用其他原始资料、历史底稿和已应用更正；含历史分期。历史支付未核实。'}


def export_reconciliation(result):
    wb=Workbook();wb.remove(wb.active)
    def sheet(name,fields,rows):
        ws=wb.create_sheet(name);ws.append([label for label,key in fields])
        for row in rows:
            values=[row.get(key,'') for label,key in fields]
            ws.append(["'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in values])
        ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=24
    common=[('任职标识','rowId'),('姓名','name'),('工号','employeeId'),('入职日期','hireDate'),('台账行','sourceRow')]
    monthly=result.get('monthly',{})
    sheet('本月安排差异',common+[('奖金类型','kind'),('项目','item'),('线下金额','reference'),('当前试算','calculated'),('线下币种','currency'),('试算币种','calculatedCurrency'),('线下发放月','referencePeriod'),('规则到期月','calculatedPeriod'),('差异类型','category'),('具体原因','reason'),('线下备注','note')],monthly.get('differences',[]))
    sheet('本月分币种合计',[('币种','currency'),('线下本月安排','reference'),('当前本月到期试算','calculated')],monthly.get('totalsByCurrency',[]))
    sheet('金额差异',common+[('奖金类型','kind'),('项目','item'),('台账值','reference'),('独立试算值','calculated'),('币种','currency'),('待核原因','reason')],result['moneyDifferences'])
    sheet('采用字段差异',common+[('字段','field'),('台账值','reference'),('原始资料与更正采用值','calculated')],result['fieldDifferences'])
    sheet('身份待核',common+[('差集侧','sideLabel'),('匹配状态','matchStatus'),('原因','reason')],
          [{**d,'sideLabel':'线下' if d['side']=='reference' else '平台'} for d in result.get('rosterDifferences',[])] if 'rosterDifferences' in result else result['unmatchedPeople'])
    sheet('核对范围',[('事项','label'),('内容','value')],[{'label':'核算月份','value':result['month']},{'label':'底稿版本','value':result['revision']},{'label':'核对台账','value':result['referenceName']},{'label':'本月核对方法','value':monthly.get('basis','')},{'label':'全历史独立核对方法','value':result['basis']}])
    from .recruitment_worksheet import style_sheet
    for ws in wb:
        style_sheet(ws)
        for col in ws.columns:
            if col[0].value in {'台账行','线下发放月','规则到期月'}:
                for cell in col[1:]: cell.number_format='0'
    for row in wb['核对范围'].iter_rows(min_row=2):
        if row[0].value in {'核算月份','底稿版本'}: row[1].number_format='0'
    out=BytesIO();wb.save(out);return out.getvalue()
