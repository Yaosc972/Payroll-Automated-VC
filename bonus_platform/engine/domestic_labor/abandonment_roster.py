"""Batch-scoped employee abandonment roster for housing allowance."""
from io import BytesIO

from openpyxl import load_workbook


def parse_abandonment_roster(content: bytes) -> list[dict]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise ValueError("自离名单无法读取，请上传有效的 .xlsx 文件") from exc
    roster = {}
    try:
        for sheet in workbook:
            rows = iter(sheet.iter_rows(values_only=True))
            header = next(rows, ())
            if not any(value is not None for value in header):
                continue
            columns = [str(value or "").strip() for value in header]
            if columns.count("工号") != 1 or columns.count("姓名") != 1:
                raise ValueError(f"自离名单工作表“{sheet.title}”首行必须包含唯一的“工号”和“姓名”列")
            for index, row in enumerate(rows, 2):
                if not any(value is not None and str(value).strip() for value in row):
                    continue
                employee_id = str(row[columns.index("工号")] or "").strip()
                name = str(row[columns.index("姓名")] or "").strip()
                if not employee_id or not name or employee_id.startswith("=") or name.startswith("="):
                    raise ValueError(f"自离名单第{index}行工号和姓名必须填写为实际值")
                if employee_id in roster and roster[employee_id] != name:
                    raise ValueError(f"自离名单工号 {employee_id} 对应多个姓名，请核对")
                roster[employee_id] = name
    finally:
        workbook.close()
    if not roster:
        raise ValueError("自离名单为空，请补充人员，或选择本月没有自离员工")
    return [{"employee_id": employee_id, "employee_name": name} for employee_id, name in roster.items()]


def validate_abandonment_roster(roster: list[dict], monthly: list[dict]) -> None:
    employees = {}
    for row in monthly:
        employees.setdefault(str(row.get("工号", "")).strip(), set()).add(str(row.get("姓名", "")).strip())
    for item in roster:
        employee_id = item["employee_id"]
        if employee_id not in employees:
            raise ValueError(f"自离名单工号 {employee_id} 未匹配本次月考勤，请核对核算月份和工号")
        if employees[employee_id] != {item["employee_name"]}:
            raise ValueError(f"自离名单工号 {employee_id} 的姓名与月考勤不一致，请核对")
