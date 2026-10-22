"""Executable work packages, not one task per affected employee.

Grouping never clears a diagnostic or changes money. Uniform edits are allowed
only for a shared policy or the same identified recipient fact.
"""
from .recruitment_draft_sources import digest


def approval_people(state, source_id):
    """Show reviewed person matching separately from missing attachments.

    Processing status is derived from current fields so undoing a correction
    cannot leave a stale 'processed' label. This projection never changes facts.
    """
    active = {o['id'] for o in state.get('operations', []) if o.get('active')}
    review = next((r for r in reversed(state.get('reviewEntries', []))
                   if r['sourceId'] == source_id and
                   (not r.get('operationId') or r['operationId'] in active)), {})
    if review.get('status') != '已认可' or not review.get('scopeChecked'):
        return []
    rows = {r['id']: r for r in state['rows']}
    target = review.get('recruiter', {})
    result = []
    for person in review.get('personScope', []):
        row = rows.get(person.get('rowId')); values = row['values'] if row else {}
        processed = bool(row and target.get('name') and target.get('employeeId') and
                         values.get('招聘负责人姓名') == target['name'] and
                         values.get('招聘负责人工号') == target['employeeId'])
        status = '已处理' if processed else '待更正' if row else '待匹配'
        reason = ('招聘负责人已更正为' + target['name']) if processed else person['reason']
        result.append({**person, 'status': status, 'reason': reason})
    return result


def opening_candidates(state):
    from .recruitment_attention import named_month
    previous=(state['month']//100-1)*100+12 if state['month']%100==1 else state['month']-1
    return [{'id':s['id'],'name':s['name']} for s in state['sources']
            if s.get('active',True) and not s.get('error') and not s.get('comparisonOnly') and 'ledger' in s.get('roles',[])
            and (named_month([s['name']]) or s.get('forMonth')) in (None,previous)
            and (not state.get('strictInputs') or named_month([s['name']])==previous)]


def people_choices(folder,state,query):
    """Search explicit name/ID pairs in this activity's existing materials.

    This is a directory lookup, not evidence of a referral relationship or
    employment eligibility. It never queries HR or reparses/mutates a file.
    """
    import json
    from .recruitment_awards import present
    query=str(query).strip().casefold()
    if not query:return []
    matches={}
    for source in state['sources']:
        if not source.get('active',True) or source.get('error') or source.get('comparisonOnly'):continue
        path=folder/'parsed'/f"{source['id']}.json"
        if not path.is_file():continue
        records=json.loads(path.read_text())['records']
        for record in records:
            v=record.get('values',{})
            for prefix in ('','招聘负责人','协助招聘人','推荐人'):
                name=v.get(prefix+'姓名','');number=v.get(prefix+'工号','')
                if not present(name) or not present(number):continue
                name=str(name).strip();number=str(number).strip()
                if query not in (name+' '+number).casefold():continue
                key=(number.casefold(),name)
                option=matches.setdefault(key,{'employeeId':number,'name':name,'sources':[]})
                ref={'sourceId':source['id'],'file':source['name'],'sheet':record.get('ref',{}).get('sheet',''),'row':record.get('ref',{}).get('row')}
                if ref not in option['sources']:option['sources'].append(ref)
    return sorted(matches.values(),key=lambda p:(p['name'].casefold()!=query and p['employeeId'].casefold()!=query,p['name'].casefold(),p['employeeId']))[:20]


def project(state, attention):
    from .recruitment_awards import CORE_FIELDS, present
    from .models import as_date
    rows={r['id']:r for r in state['rows']};groups={}
    employment_ids={rid for scope in ('current','history') for t in attention[scope] if t['kind']=='employment' for rid in t['rowIds']}
    source_map = {s['id']: s for s in state['sources']}
    from .recruitment_assignments import missing_attachments
    for scope in ('current','history'):
        # Attention can aggregate multiple approvals. Split by source before
        # deciding whether this is missing a list or matching named people.
        tasks = [part for task in attention[scope] for part in (
            [{**task, 'sourceIds': [sid]} for sid in task['sourceIds']]
            if task['kind'] == 'attachment_review' else [task])]
        for task in tasks:
            for rid in task['rowIds'] or [None]:
                row=rows.get(rid);v=row['values'] if row else {};kind=task['kind']
                # Future preparation is explicit and bounded to this month's
                # new hires, not every diagnostic in the historical ledger.
                if scope=='history' and (not row or str(v.get('入职日期',''))[:7].replace('-','')!=str(state['month'])):continue
                # Preserve financial diagnostics in attention. A rule version
                # waiting for an actual hire date is not another user decision.
                if kind=='standard' and row and row['rule']['version']=='unresolved' and rid in employment_ids and all(present(v.get(f)) and f not in row.get('unresolvedFields',[]) for f in ('标签分类','奖金地区类型','ABC类别','职级')):continue
                if kind.startswith('external_results_'):
                    key='external_results';title='处理体系外核算结果';mode='externalResults';field=None
                    detail=''
                elif kind == 'attachment_review' and task['sourceIds'] and not missing_attachments(state, source_map.get(task['sourceIds'][0], {})):
                    sid = task['sourceIds'][0]
                    key = 'approval:' + sid; mode = 'sources'; field = None
                    people = approval_people(state, sid)
                    title = '核对审批人员匹配' if people else '核对审批适用范围'
                    detail = source_map.get(sid, {}).get('name', '') + '：' + (
                        '；'.join(p['name'] + ' · ' + p['status'] for p in people)
                        if people else '核对审批内容及适用人员，不要求重复补填已有资料。')
                    if source_map.get(sid,{}).get('approvalSummary'):
                        summary=source_map[sid]['approvalSummary']
                        key='approval:blocked' if summary.get('blockedCount') else 'approval:ready'
                        title='核对审批调整内容' if summary.get('blockedCount') else '核对并采用审批调整'
                        detail=summary.get('detail') or '；'.join(summary.get('issues',[]))
                elif kind=='intake_scope' or (row and not row.get('offerMatched') and (
                        kind.startswith('field:') and kind.split(':',1)[1] in {'ABC类别','职级','标签分类','奖金地区类型','招聘负责人工号','招聘渠道','推荐人工号'}
                        and not present(v.get(kind.split(':',1)[1]))
                        or kind=='cycle' and not as_date(v.get('招聘启动日期'))
                        or kind=='standard' and any(not present(v.get(f)) for f in CORE_FIELDS))):
                    region=v.get('标签分类') or '待定';key='intake:'+region;title='补齐'+region+'录用资料';mode='intake';field=None
                    detail='核对下列缺项，上传对应的录用资料或一张补充表，系统重新匹配并更新核算；已知信息不需要重填。'
                elif kind in {'field:招聘负责人工号','field:招聘负责人姓名','field:协助招聘人工号','field:协助招聘人姓名'}:
                    key='recruitment-assignment';title='核对招聘分工';mode='assignments';field=None
                    detail='原始资料中的招聘分工不一致或缺少负责人。对照依据，集中更正负责人和协助招聘人。'
                elif kind=='recipient':
                    key='recipient';title='补齐招聘负责人、协助招聘人和推荐人的任职资料';mode='recipient';field=None
                    detail='国内工号自动补查任职；同一人的资料补充一次，更新全部关联奖金。'
                elif kind=='channel_policy':
                    key=kind;title='核定招聘渠道归类及系数';mode='facts';field=None
                    detail='原始渠道已经提供，需核定适用归类及系数；保留原文和当前不可计算结果。'
                elif kind.startswith('calculation:') and row and any(
                        r['reason']==task['title'] for award in row['calculation'].values() if isinstance(award,dict)
                        for r in award.get('recoveries',[])):
                    key=kind;title=task['title'];mode='settlement';field=None
                    detail='核对原历史安排与当前规则差异，预览处理影响后采用。'
                elif kind.startswith('calculation:') and row and any(
                        i.get('recordId') and i['reason']==task['title']
                        for award in row['calculation'].values() if isinstance(award,dict)
                        for i in award.get('accountingIssues',[])):
                    key=kind;title=task['title'];mode='specialRecords';field=None
                    detail='打开奖金调整记录，核对需重新处理的原安排。'
                elif kind=='rehire' and rid in employment_ids:
                    key='employment';title='核对任职资料与二次入职';mode='facts';field=None
                    detail='在同一次核对中确认任职关系、上段离职日及适用规则。'
                elif kind=='rehire' and v.get('上次离职日期') and '二次入职专项条款执行依据待核' in row['calculation']['recruitment']['issues'] and not set(row.get('unresolvedFields',[])) & {'上次离职日期','入职日期','候选人入职时间','招聘渠道'}:
                    key='rehire:'+row['rule']['version'];title='选择二次入职招聘奖金口径';mode='policy';field='二次入职执行口径'
                    detail='离职间隔不超过12个月，选择不发或按40%核算。'
                else:
                    key='employment' if kind=='rehire' else kind;mode='sources' if not rid else 'facts';field=None
                    title={'cycle':'核对招聘日期与剔除天数','employment':'核对任职资料与二次入职','rehire':'核对任职资料与二次入职','field:招聘渠道':'确认实际招聘渠道','field:推荐人工号':'补充推荐人','standard':'核对适用奖金标准','resignation_review':'核对离职流程结果'}.get(kind,task['title'])
                    detail='在一张清单中批量更正，仅保存实际修改的字段。' if rid else '查看实际审批与引用附件，按适用范围采用，不把审批数量拆成人员待办。'
                    if not rid and kind=='baseline':
                        candidates=opening_candidates(state)
                        title='确认采用上月核算表' if candidates else '补充上月核算表'
                        detail='上月表已经上传，确认采用后继续计算未发完的奖金。' if candidates else '导入上月核算表；确无历史未结奖金时，可直接确认，无须补空表。'
                    elif not rid and kind!='attachment_review':detail='在资料页查看读取结果并补充原始文件，系统会自动更新关联人员。'
                g=groups.setdefault(key,{'id':digest('work:'+key),'key':key,'title':title,'mode':mode,'field':field,'detail':detail,
                    'rowIds':[],'currentRowIds':[],'taskIds':[],'sourceIds':[],'reasons':[],'current':False})
                if rid and rid not in g['rowIds']:g['rowIds'].append(rid)
                if scope=='current':
                    g['current']=True
                    if rid and rid not in g['currentRowIds']:g['currentRowIds'].append(rid)
                for k,values in [('taskIds',[task['id']]),('sourceIds',task['sourceIds']),('reasons',[task['title']])]:
                    g[k]=list(dict.fromkeys(g[k]+values))
    from .recruitment_actions import task_fields
    for g in groups.values():
        ids=g['currentRowIds'] if g['current'] else g['rowIds']
        g['affectedCount']=len(ids)
        g['previews']=[{'rowId':rid,'name':rows[rid]['values'].get('姓名') or rows[rid]['values'].get('工号'),
                        'reason':operator_reason(rows[rid],g['key'],state['month'],state)} for rid in ids[:3]]
        g['ruleMaintenance']=g['key'].startswith(('special:','rehire:')) or g['key']=='channel_policy'
        if g['key']=='standard' and all(all(present(rows[rid]['values'].get(f)) and f not in rows[rid].get('unresolvedFields',[]) for f in CORE_FIELDS) for rid in g['rowIds']):
            g.update(ruleMaintenance=True,title='补充适用奖金标准',detail='人员资料已齐全，制度标准尚未对应；无需修改真实职级或类别。')
        if g['key']=='baseline':g['openingCandidates']=opening_candidates(state)
        g['actionClass']=action_class(state,g,attention)
        if g['mode']=='externalResults':g['action']='查看体系外结果'
        if g['mode']=='intake' and g['actionClass']=='review_available':
            g['title']='核对'+g['key'].split(':',1)[1]+'录用资料'
        g['monthImpacts']=month_impacts(state,g)
        g['taskCount']=len(g['taskIds'])
        g['sourceCount']=len(g['sourceIds'])
        if g['key'].startswith('approval:'):
            g['approvalPeople'] = [p for sid in g['sourceIds'] for p in approval_people(state,sid)]
        if g['key']=='attachment_review':
            from .recruitment_assignments import missing_attachments
            missing=[s['name']+'：缺少 '+'、'.join(missing_attachments(state,s)) for s in state['sources']
                     if s['id'] in g['sourceIds'] and missing_attachments(state,s)]
            if missing:g.update(title='补充审批引用的附件',detail='；'.join(missing))
        if g['mode']=='recipient':
            subjects=set()
            for rid in ids:
                v=rows[rid]['values']
                for f in task_fields(rows[rid],{'kind':'recipient'},state['month']):
                    prefix=next((p for p in ('招聘负责人','协助招聘人','推荐人') if f.startswith(p)),None)
                    subjects.add((str(v.get(prefix+'工号') or v.get(prefix+'姓名') or rid).lower(),v.get(prefix+'入职日期','')) if prefix else (rid,''))
            g['subjectCount']=len(subjects)
    return sorted(groups.values(),key=lambda g:(not g['current'],{'policy':0,'recipient':1,'intake':2,'sources':3,'facts':4,'assignments':4,'settlement':4,'specialRecords':4,'externalResults':2}[g['mode']],g['key']))


def operator_reason(row,key,month,state=None):
    from .models import as_date
    from .recruitment_attention import intake_fields, LABELS
    v=row['values']
    if key.startswith('intake:'):
        fields=intake_fields(row,month)
        names={'招聘负责人工号':'招聘负责人','推荐人工号':'推荐人',**LABELS}
        suggestions=(state or {}).get('referenceFacts',{}).get('byRowId',{}).get(row['id'],[])
        available={field for p in suggestions if p.get('canAdopt') and p.get('status')=='available' for field in p.get('fields',[])}
        supplied=[field for field in fields if field in available]
        if supplied:
            missing=[field for field in fields if field not in available]
            return '已有'+ '、'.join(names.get(f,f) for f in supplied)+'可核对采用'+('；仍缺少'+'、'.join(names.get(f,f) for f in missing) if missing else '，无需重复上传')+'。'
        return '缺少'+ '、'.join(names.get(f,f) for f in fields)+'；上传录用资料或补充表。' if fields else '录用资料存在差异，请对照来源核对。'
    if key=='cycle':
        start=as_date(v.get('招聘启动日期'));hire=as_date(v.get('入职日期') or v.get('候选人入职时间'))
        if start and hire and start>hire:return f'招聘启动日 {start:%Y-%m-%d} 晚于入职日 {hire:%Y-%m-%d}，请核对日期。'
        return f"招聘启动日 {v.get('招聘启动日期') or '未提供'}，剔除 {v.get('周期剔除天数',0)} 天；核对日期或剔除依据。"
    if key=='field:招聘渠道':
        resolution=row.get('channelResolution',{})
        if resolution.get('description')=='渠道说明及备注不一致':return '现有资料的渠道说明不一致，请对照来源确认实际渠道。'
        if row.get('channelChecks') and all(c.get('status')=='备注为空' for c in row['channelChecks']):
            return '已检查导入资料并补查国内 Offer；渠道为「其他」，接口备注为空。仍缺少具体渠道名称，可补充原始资料。'
        if resolution.get('sources'):
            return '已检查导入资料的渠道说明及备注，仍无法唯一识别；请补充明确的实际渠道。'
        return f"原始渠道为「{v.get('招聘渠道') or '未填写'}」，缺少具体渠道名称；填写后自动核算。"
    if key=='channel_policy':return f"原始渠道已写明「{v.get('其他渠道') or v.get('招聘渠道')}」，适用归类及系数待核定。"
    if key=='employment':
        dates=list(dict.fromkeys(d for d in [v.get('入职日期'),v.get('候选人入职时间')]+[o.get('value') for o in row.get('employmentObservations',[])] if d))
        if len(dates)>1:return '入职日期存在 '+ ' / '.join(dates)+' 不同记录，请对照原始资料。'
        return '本次入职日期 '+(dates[0] if dates else '未提供')+'，请核对是否重新入职及此前离职记录。'
    if key=='field:推荐人工号':return '已记录内推渠道，但没有推荐人姓名；核对实际推荐关系。'
    if key=='standard':return f"{v.get('标签分类','')} · {v.get('ABC类别','')} · {v.get('职级','')} 尚未对应奖金标准，不需要更改人员资料。"
    return ''


def original_cycle_option(row):
    from math import isfinite
    from .models import as_date
    from .recruitment_awards import number, present
    v=row['values'];start=as_date(v.get('招聘启动日期'));hire=as_date(v.get('入职日期') or v.get('候选人入职时间'))
    cycle=row.get('calculation',{}).get('recruitment',{}).get('cycle') or {}
    exclusion=number(v.get('周期剔除天数')) if present(v.get('周期剔除天数')) else number(0)
    conflicts=set(row.get('unresolvedFields',[])) & {'招聘启动日期','入职日期','候选人入职时间','周期剔除天数'}
    if not start or not hire or start<=hire or exclusion is None or exclusion<0 or not isfinite(float(exclusion)) or conflicts or cycle.get('standardDays') is None:
        return None
    return {'startDate':start.strftime('%Y-%m-%d'),'hireDate':hire.strftime('%Y-%m-%d'),
            'exclusionDays':float(exclusion),'actualDays':(hire-start).days-float(exclusion),
            'standardDays':cycle['standardDays'],'confirmed':cycle.get('factor') is not None}


def confirm_original_cycle(store,month,revision,row_ids):
    from .recruitment_awards import ensure_calculated
    if not isinstance(row_ids,list) or not row_ids or len(row_ids)>500 or any(not isinstance(rid,str) for rid in row_ids):
        raise ValueError('请选择需要按原日期核算的人员')
    with store._lock(month) as folder:
        state=ensure_calculated(store.read(month));store._check(state,revision)
        rows={r['id']:r for r in state['rows']};edits=[]
        for rid in dict.fromkeys(row_ids):
            row=rows.get(rid)
            if row is None:raise ValueError('人员记录不存在，请刷新待办')
            choice=original_cycle_option(row)
            if choice is None:raise ValueError('日期或剔除天数尚未核清，请先修改日期或核对来源')
            if choice['confirmed']:continue
            values={'负周期确认启动日期':choice['startDate'],'负周期确认入职日期':choice['hireDate'],
                    '负周期确认剔除天数':choice['exclusionDays'],'负周期确认依据':'用户选择按原日期核算招聘周期'}
            edits.extend({'rowId':rid,'field':field,'value':value,'reason':'按原日期核算招聘周期'} for field,value in values.items())
        if not edits:return state
        state=store._prepare_change(folder,state,{'edits':edits})
        state['operations'][-1]['label']='按原日期核算招聘周期'
        return store._save(folder,state)


def detail(state,item,attention=None):
    from .recruitment_actions import task_fields,task_detail
    from .recruitment_attention import project as attention_project, intake_fields
    from .recruitment_awards import present
    if attention is None:attention=attention_project(state)
    tasks={t['id']:t for scope in ('current','history') for t in attention[scope]}
    rows={r['id']:r for r in state['rows']};fields={};employment={}
    if item['key']=='employment':
        for tid in item['taskIds']:
            if tasks[tid]['kind'] in {'employment','rehire'}:
                employment.update({p['rowId']:p['employmentReview'] for p in task_detail(state,tid)['rows']})
    for rid in item['rowIds']:
        row=rows[rid];v=row['values'];required=[]
        for tid in item['taskIds']:
            t=tasks[tid]
            if rid in t['rowIds']:required+=task_fields(row,t,state['month'])
        if item['key']=='employment':
            from .recruitment_employment import employment_review
            review=employment_review(state,rid)
            if review['selection']:
                required=[f for f in required if f not in {'任职说明','核算范围'}]
            if review['priorExit']:
                required=[f for f in required if f!='上次离职日期']
            employment.setdefault(rid,review)
        if item['mode']=='intake':
            required+=intake_fields(row,state['month'])
            # A cycle review may legitimately need known dates. Do not turn
            # its optional exclusion inputs into mandatory missing Offer facts.
            cycle_issues=row['calculation']['recruitment']['issues']
            required=[f for f in required if not (
                f in {'入职日期','候选人入职时间'} and present(v.get(f)) and f not in row.get('unresolvedFields',[]) or
                f=='招聘负责人工号' and not present(v.get(f)) and present(v.get('招聘负责人姓名')) and f not in row.get('unresolvedFields',[]) or
                f in {'周期剔除天数','周期剔除依据'} and not any('剔除' in issue or '免周期' in issue for issue in cycle_issues))]
        if item['mode']=='policy':required=[item['field']]
        if item['mode'] in {'settlement','specialRecords'}:required=[]
        if item.get('ruleMaintenance') and item['key']=='standard':required=[]
        # A missing source groups the work, but does not invalidate facts already
        # present. Only expose fields relevant to the actual remaining issues.
        for field in dict.fromkeys(required):
            person=v.get('姓名') or v.get('工号') or rid
            identity=rid;fact=field
            if item['mode']=='policy':identity=item['key'];person='本批统一口径'
            elif item['mode']=='recipient':
                for prefix in ('招聘负责人','协助招聘人','推荐人'):
                    if field.startswith(prefix) and not field.endswith('工号') and v.get(prefix+'工号'):
                        identity='recipient:'+str(v[prefix+'工号']).lower()+'|'+str(v.get(prefix+'入职日期',''));fact=field[len(prefix):]
                        person=v.get(prefix+'姓名') or v[prefix+'工号'];break
            key=(identity,fact);f=fields.setdefault(key,{'field':field,'person':person,'identity':identity,'value':v.get(field,''),'options':[],'targets':[],'bindings':[]})
            if v.get(field,'')!=f['value']:f['mixed']=True
            if rid not in f['targets']:f['targets'].append(rid)
            binding={'rowId':rid,'field':field}
            if binding not in f['bindings']:f['bindings'].append(binding)
            for ref in row.get('provenance',{}).get(field,[]):
                value=ref.get('value')
                if present(value) and not any(o['value']==value for o in f['options']):
                    f['options'].append({'value':value,'label':str(value),'sources':[ref]})
            if field=='入职日期' and rid in employment:
                for observation in employment[rid]['observations']:
                    value=observation['value']
                    if observation['label'].endswith('入职日') and value!='未提供' and not any(o['value']==value for o in f['options']):
                        f['options'].append({'value':value,'label':value+'（'+observation['label']+'）','sources':observation['sources']})
    checks={}
    if item['mode']=='recipient':
        from .recruitment_beisen import coverage_sets
        coverage=coverage_sets([s for s in state['sources'] if s.get('kind')=='beisen' and s.get('forMonth')==state['month']])
        attempted=coverage['requestedRecipient']|coverage['requestedCandidate']
        for f in fields.values():
            for binding in f['bindings']:
                row=rows[binding['rowId']];v=row['values'];field=binding['field']
                for prefix in ('招聘负责人','协助招聘人','推荐人'):
                    if not field.startswith(prefix) or not v.get(prefix+'工号'):continue
                    number=str(v[prefix+'工号']).lower();facts=row.get('employmentTimelines',{}).get(prefix,[])
                    message=('已返回任职记录，但仍未提供可确定本次领取资格的完整事实' if facts else
                             '已查询，未返回可用任职记录' if number in attempted else
                             '可通过接口补查' if number.startswith('zt') else '体系外编号，需要补充原始任职资料')
                    checks[number]={'employeeId':number,'name':v.get(prefix+'姓名') or number,'message':message}
    current_fields=[]
    for f in fields.values():
        bindings=[b for b in f['bindings'] if b['rowId'] in item['currentRowIds']]
        if bindings:
            values=[rows[b['rowId']]['values'].get(b['field'],'') for b in bindings]
            current_fields.append({**f,'value':values[0],'mixed':any(v!=values[0] for v in values),
                                   'bindings':bindings,'targets':list(dict.fromkeys(b['rowId'] for b in bindings))})
    from .recruitment_resignations import review_evidence
    from .recruitment_field_controls import controls
    for f in fields.values():
        definitions=[controls(state,rows[rid]).get(f['field']) for rid in f['targets']]
        if definitions and definitions[0]:
            shared=[c for c in definitions[0]['choices'] if all(d and c in d['choices'] for d in definitions)]
            f.update({**definitions[0],'choices':shared})
    for f in current_fields:
        if definition:=controls(state,rows[f['targets'][0]]).get(f['field']):
            shared=[c for c in definition['choices'] if all(c in controls(state,rows[rid])[f['field']]['choices'] for rid in f['targets'])]
            f.update({**definition,'choices':shared})
    return {**item,'revision':state['revision'],'fields':list(fields.values()),'currentFields':current_fields,'checks':list(checks.values()),
        'checkedSources':[{'id':s['id'],'name':s['name']} for s in state['sources']
                          if item['mode']=='intake' and s.get('active',True) and not s.get('error') and not s.get('comparisonOnly')
                          and ('domestic_offer' if item['key']=='intake:国内' else 'overseas_offer') in s.get('roles',[])],
        'people':[{'id':rid,'name':rows[rid]['values'].get('姓名',''),'employeeId':rows[rid]['values'].get('工号',''),
                   'previousExitDate':rows[rid]['values'].get('上次离职日期',''),
                   'hireDate':rows[rid]['values'].get('入职日期',''),'jobType':rows[rid]['values'].get('职类',''),
                   'offerStatus':'已匹配录用记录，字段仍需补充或核对' if rows[rid].get('offerMatched') else '未匹配到本次入职的录用记录',
                   'missingFields':[f['field'] for f in fields.values() if rid in f['targets'] and not present(rows[rid]['values'].get(f['field']))],
                   'reviewFields':[f['field'] for f in fields.values() if rid in f['targets'] and present(rows[rid]['values'].get(f['field']))],
                   'values':{f:rows[rid]['values'].get(f,'') for f in ('标签分类','奖金地区类型','ABC类别','职级','招聘渠道','其他渠道')},
                   'reasons':list(dict.fromkeys(reason for tid in item['taskIds'] if rid in tasks[tid]['rowIds'] for reason in tasks[tid].get('rowReasons',{}).get(rid) or [tasks[tid]['title']])),
                   'explanation':operator_reason(rows[rid],item['key'],state['month'],state),
                   'employmentReview':employment.get(rid),
                   'resignationReview':review_evidence(rows[rid]) if item['key']=='resignation_review' else None,
                   'originalCycle':original_cycle_option(rows[rid]) if item['key']=='cycle' else None,
                   'current':rid in item['currentRowIds']} for rid in item['rowIds']]}


def supplement_workbook(state,data,include_future=False):
    """Use the existing supplement importer and stable employment identity."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from .recruitment_awards import present
    people=data['people'] if include_future or not data['currentRowIds'] else [p for p in data['people'] if p['current']]
    ids={p['id'] for p in people};rows={r['id']:r for r in state['rows']}
    # Never offer a name-only reimport that would create a second person. Such
    # a record can still be corrected directly using its existing row ID.
    identity_fields=('工号','体系外人员编号','海外offer流程编号')
    if any(not any(present(rows[p['id']]['values'].get(f)) for f in identity_fields) for p in people):
        raise ValueError('清单中有人员缺少可用于回传匹配的编号，请使用「手工补充」直接更正原记录；无需新建人员。')
    anchors=[f for f in identity_fields if f=='工号' or any(present(rows[p['id']]['values'].get(f)) for p in people)]
    headers=list(dict.fromkeys(['姓名',*anchors,'入职日期']+[f['field'] for f in data['fields'] if ids.intersection(f['targets'])]))
    book=Workbook();sheet=book.active;sheet.title='导入_月度数据';sheet.append(headers)
    for p in people:sheet.append([rows[p['id']]['values'].get(f,'') for f in headers])
    sheet.freeze_panes='D2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True,color='1E3A8A');cell.fill=PatternFill('solid',fgColor='EEF3F7')
    for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=22
    buffer=BytesIO();book.save(buffer);return buffer.getvalue()


def action_class(state, item, attention=None):
    if item['mode']=='externalResults':return 'missing_fact'
    if item.get('ruleMaintenance'):
        return 'policy'
    if item['mode'] in {'settlement','specialRecords'}:
        return 'review_history'
    proposals=state.get('referenceFacts',{}).get('byRowId',{})
    if item['key'].startswith('approval:'):
        return 'review_available'
    from .recruitment_actions import task_fields
    from .recruitment_attention import intake_fields
    rows={row['id']:row for row in state['rows']}
    tasks=[task for scope in ('current','history') for task in (attention or {}).get(scope,[])
           if task['id'] in item.get('taskIds',[])]
    for rid in item['rowIds']:
        row=rows[rid]
        if item['mode']=='intake':
            required=set(intake_fields(row,state['month']))
        else:
            obligations=[task for task in tasks if rid in task['rowIds']] or [{'kind':item['key']}]
            required={field for task in obligations for field in task_fields(row,task,state['month'])}
        if any(p.get('canAdopt') and p.get('status')=='available' and required.intersection(p.get('fields',[]))
               for p in proposals.get(rid,[])):
            return 'review_available'
    if item['mode'] in {'intake','recipient'} or item['key']=='attachment_review' or item['key']=='baseline' and not item.get('openingCandidates'):
        return 'missing_fact'
    return 'resolve_fact'


def month_impacts(state, item):
    impacts=set()
    month=state['month']
    rows={r['id']:r for r in state['rows']}
    if item['mode'] in {'settlement','specialRecords'}:
        impacts.add('historical')
    for rid in item['rowIds']:
        row=rows[rid]
        for kind in ('recruitment','referral'):
            award=row['calculation'][kind]
            for node in award.get('nodes',[]):
                scheduled=node.get('scheduledMonth',node.get('dueMonth'))
                if scheduled==month:
                    impacts.add('scheduled_current' if node.get('amount') is not None else 'potential_current_unknown')
                elif scheduled and scheduled>month:
                    impacts.add('future')
            if any(award.get(k) for k in ('recoveries','accountingIssues')):
                impacts.add('historical')
    if item['current'] and not impacts & {'scheduled_current','historical'}:
        impacts.add('potential_current_unknown')
    if not item['current'] and not impacts:
        impacts.add('future')
    return [impact for impact in ('scheduled_current','potential_current_unknown','historical','future') if impact in impacts]
