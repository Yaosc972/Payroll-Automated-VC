"""Styled result workbook and immutable attendance snapshot export."""
from datetime import datetime
from io import BytesIO
from decimal import Decimal
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FIELDS = {
    "employeeId": "工号", "employeeName": "姓名", "secondOrg": "二级组织", "thirdOrg": "三级组织",
    "fourthOrg": "四级组织", "grade": "职级", "employmentType": "人员类型", "entryDate": "入职日期",
    "departureDate": "离职日期", "targetDays": "应出勤天数", "sickDays": "病假天数",
    "casualDays": "事假天数", "absentDays": "旷工天数", "maternityDays": "工作日产假天数",
    "normalDays": "正常出勤天数", "amount": "职场搬迁补贴", "status": "核算状态", "reason": "判定说明",
}
RAW_FIELDS = {
    "TargetDays": "应出勤天数",
    "AbsenteeismDuration": "旷工天数",
    "AttendanceItem5": "事假小时数",
    "AttendanceItem6": "病假小时数",
    "AttendanceItem16": "工作日产假天数",
}


def write_cell(ws, row, col, value):
    cell = ws.cell(row, col, value)
    if isinstance(value, str):
        cell.data_type = "s"  # API-provided strings must never become Excel formulas.
    return cell


def style_sheet(ws, headers, records, widths, amount_col=None):
    end = get_column_letter(len(headers))
    ws.sheet_view.showGridLines = False
    for col, header in enumerate(headers, 1):
        c = ws.cell(1, col, header)
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.row_dimensions[1].height = 36
    for row_index, record in enumerate(records, 2):
        ws.row_dimensions[row_index].height = 32
        for col, value in enumerate(record, 1):
            c = write_cell(ws, row_index, col, value)
            c.font = Font(name="微软雅黑", size=11, color="1E293B")
            c.fill = PatternFill("solid", fgColor="F1F5F9" if row_index % 2 == 0 else "FFFFFF")
            c.alignment = Alignment(vertical="center", horizontal="right" if isinstance(value, (int, float)) else "left", wrap_text=True)
            c.border = Border(bottom=Side(style="hair", color="E2E8F0"))
            if isinstance(value, (int, float)):
                places = max(0, -Decimal(str(value)).normalize().as_tuple().exponent)
                c.number_format = "0" if not places else "0." + "0" * places
            if col == amount_col:
                c.number_format = '#,##0.00'
                c.fill = PatternFill("solid", fgColor="DBEAFE")
                c.font = Font(name="微软雅黑", size=11, bold=True, color="1E3A8A")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{end}{max(1, ws.max_row)}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


def build_workbook(run):
    year, month = map(int, run["month"].split("-"))
    month_text = f"{year}年{month}月"
    calculated = datetime.fromisoformat(run.get("calculatedAt") or run["createdAt"])
    if calculated.tzinfo:
        calculated = calculated.astimezone(ZoneInfo("Asia/Shanghai"))
    filename = f"{year} 年 {month} 月 ABU职场搬迁补贴{calculated:%Y%m%d%H%M}.xlsx"
    workbook = Workbook()
    ws = workbook.active
    ws.title = "职场搬迁补贴"
    result = run["result"]
    headers = ["核算月份", *FIELDS.values()]
    records = [[month_text, *[r.get(k) for k in FIELDS]] for r in result["results"]]
    widths = [16, 16, 16, 30, 30, 30, 12, 16, 16, 16, 15, 15, 15, 15, 18, 18, 21, 14, 36]
    style_sheet(ws, headers, records, widths, amount_col=17)
    ws.sheet_properties.tabColor = "2563EB"
    source_ws = workbook.create_sheet("考勤源数据")
    snapshot = run.get("attendanceSource")
    if snapshot is None:
        style_sheet(source_ws, ["源数据状态"], [["本批次未保存原始考勤快照，请重新核算后导出。"]], [80])
    else:
        identity = ["employeeId", "employeeName", "secondOrg", "thirdOrg", "fourthOrg"]
        raw_headers = ["考勤月份", *[FIELDS[k] for k in identity], *RAW_FIELDS.values()]
        raw_rows = []
        for row in snapshot:
            raw = row["raw"]
            props = raw.get("Properties") or {}
            values = [raw.get(k) if k in {"TargetDays", "AbsenteeismDuration"} else props.get(k) for k in RAW_FIELDS]
            raw_rows.append([month_text, *[row.get(k) for k in identity], *values])
        style_sheet(source_ws, raw_headers, raw_rows, [16, 16, 16, 30, 30, 30, 18, 16, 18, 18, 22])
    source_ws.sheet_properties.tabColor = "0F766E"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue(), filename
